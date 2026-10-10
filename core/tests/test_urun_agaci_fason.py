"""ÜRÜN AĞACI — Maliyet sekmesinde İŞLEM / FASON BEDELİ grubu (Ürün/Maliyet, Excel, Karşılaştır): ağaçtaki her üretilen çıktı (PARÇALA'da tüm
çıktılar, ürünün kendisi dahil) için birim ihtiyaç adedi × fasoncunun seçilen tarihte geçerli FasonFiyat'ı; döviz fiyat TCMB kuruyla TL; fasoncu ve
tarih seçimi; fiyatsız üretilen parça sayısı; Malzeme / İşlem-fason / GENEL toplamları. Malzeme hesabı değişmez.
Fixture canlı 152-22-0003 (4+4) zincirinin kopyası: malzeme 784,15 TL; Gürbüz fason 151-10-0017 ×2 ×16,06 + 151-10-0018 ×2 ×16,06 +
151-20-0014 ×2 ×5,93 + 151-20-0015 ×2 ×5,93 = 87,96 TL → genel 872,11 TL."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.models import Cari, FasonFiyat, IsIstasyonu, Kur, Operasyon, Stok
from core.services import fason as fs
from core.services import urun_agaci as ua
from core.services.uretim import operasyon_olustur
from core.services.urun_agaci_xlsx import karsilastir_xlsx, maliyet_xlsx
from core.tests.test_uretim_tam_boy import TamBoyBase

D = Decimal
BUGUN = date(2026, 10, 9)


class FasonMaliyetBase(TamBoyBase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.lazer = IsIstasyonu.objects.create(kod="10", ad="BORU LAZER")
        s = lambda kod, birim=None, **kw: Stok.objects.create(kod=kod, ad=kod.lower(), kategori=cls.kat, uretim_birimi=birim or cls.adet, fatura_birimi=birim or cls.adet,
                                                              uretim_urunu=kw.pop("uretim", True), satinalma_urunu=kw.pop("satinalma", False), satis_urunu=kw.pop("satis", False))
        cls.p_mentese = s("150-20-0002", cls.boy, uretim=False, satinalma=True)
        cls.p_on = s("150-10-0007", cls.boy, uretim=False, satinalma=True)
        cls.p_arka = s("150-10-0008", cls.boy, uretim=False, satinalma=True)
        Stok.objects.filter(pk=cls.p_mentese.pk).update(ort_maliyet_try=D("453.12"), ort_maliyet_usd=D("11.328"))      # 0,03125 BOY → 14,16 TL
        for p in (cls.p_on, cls.p_arka):
            Stok.objects.filter(pk=p.pk).update(ort_maliyet_try=D("384.995"), ort_maliyet_usd=D("9.624875"))          # 1 BOY → 384,995 TL
        cls.on, cls.arka = s("151-10-0017"), s("151-10-0018")
        cls.sag, cls.sol = s("151-20-0014"), s("151-20-0015")
        cls.urun = s("152-22-0003", satis=True)
        cls.a41 = s("152-10-0003", satis=True)
        operasyon_olustur(istasyon_id=cls.lazer.pk, cikti_id=cls.on.pk, cikti_miktar=D("2"), satirlar=[(cls.p_on, D("1"))])
        operasyon_olustur(istasyon_id=cls.lazer.pk, cikti_id=cls.arka.pk, cikti_miktar=D("2"), satirlar=[(cls.p_arka, D("1"))])
        operasyon_olustur(istasyon_id=cls.lazer.pk, satirlar=[(cls.p_mentese, D("1"))], tur=Operasyon.Tur.PARCALA, pay_anahtari=Operasyon.PayAnahtari.ESIT,
                          tam_boy=False, ciktilar=[(cls.sag, D("64"), None, None), (cls.sol, D("64"), None, None)])
        operasyon_olustur(istasyon_id=cls.montaj.pk, cikti_id=cls.urun.pk, cikti_miktar=D("1"),
                          satirlar=[(cls.on, D("2")), (cls.arka, D("2")), (cls.sag, D("2")), (cls.sol, D("2"))])
        operasyon_olustur(istasyon_id=cls.montaj.pk, cikti_id=cls.a41.pk, cikti_miktar=D("1"), satirlar=[(cls.on, D("2"))])     # A41: yalnız ön ayak
        cls.gurbuz = Cari.objects.create(kod="320-10-0001", unvan="GÜRBÜZ FASON", muhasebe_kodu="320.10.0001")
        cls.salim = Cari.objects.create(kod="320-10-0002", unvan="SALİM FASON", muhasebe_kodu="320.10.0002")
        for st, f in ((cls.on, "16.06"), (cls.arka, "16.06"), (cls.sag, "5.93"), (cls.sol, "5.93")):
            fs.fiyat_olustur(cari_id=cls.gurbuz.pk, stok_id=st.pk, birim_fiyat=D(f), gecerlilik_baslangic=date(2026, 1, 1), fasoncu_kodu="GZ-" + st.kod[-4:])
        fs.fiyat_olustur(cari_id=cls.salim.pk, stok_id=cls.sag.pk, birim_fiyat=D("0.2"), para_birimi="USD", gecerlilik_baslangic=date(2026, 1, 1))
        Kur.objects.create(tarih=BUGUN, usd_alis=D("40"), eur_alis=D("44"))

    def graf(self):
        return ua.graf_yukle(BUGUN)


class ServisTest(FasonMaliyetBase):
    def test_canli_rakamlar_fason_grubu_ve_genel_toplam(self):
        m = ua.urun_maliyet(self.graf(), self.urun, D("1"))
        self.assertEqual(m["toplam_try"], D("784.15"))                                                       # malzeme aynen
        f = m["fason"]
        self.assertEqual(sorted((s["stok"].kod, s["tuketim"], s["birim_try"], s["tutar_try"], s["fasoncu_kodu"]) for s in f["satirlar"]),
                         [("151-10-0017", D("2"), D("16.06"), D("32.12"), "GZ-0017"), ("151-10-0018", D("2"), D("16.06"), D("32.12"), "GZ-0018"),
                          ("151-20-0014", D("2"), D("5.93"), D("11.86"), "GZ-0014"), ("151-20-0015", D("2"), D("5.93"), D("11.86"), "GZ-0015")])
        self.assertEqual((f["toplam_try"], m["genel_try"]), (D("87.96"), D("872.11")))
        self.assertEqual((f["fiyatli"], f["fiyatsiz"], f["uretilen"]), (4, 1, 5))                             # fiyatsız: 152-22-0003 (montaj)
        self.assertEqual(f["fasoncu"], self.gurbuz)                                                          # varsayılan: en çok parçayı fiyatlayan
        self.assertEqual([(s["cari"].pk, s["sayi"]) for s in f["fasoncular"]], [(self.gurbuz.pk, 4), (self.salim.pk, 1)])
        self.assertEqual(f["toplam_usd"], D("2.199"))                                                        # 87,96 / 40
        self.assertEqual((m["genel_usd"], m["toplam_usd"]), (m["toplam_usd"] + f["toplam_usd"], D("19.60375")))
        self.assertEqual((f["tarih"], f["kursuz"], f["usd_eksik"]), (BUGUN, 0, 0))
        self.assertEqual(f["kur_notu"], "Fason fiyatı çevrimi: TCMB 09.10.2026, USD 40,0000")               # TRY fiyat: USD sütunu için kur
        m3 = ua.urun_maliyet(self.graf(), self.urun, D("3"))
        self.assertEqual((m3["fason"]["toplam_try_n"], m3["genel_try_n"], m3["fason"]["satirlar"][0]["tutar_try_n"]), (D("263.88"), D("2616.33"), D("96.36")))

    def test_parcala_ciktilari_ayri_satir_ve_urunun_kendisi(self):
        m = ua.urun_maliyet(self.graf(), self.sag, D("1"))                                                   # ara parça: kendisi 1 adet
        self.assertEqual([(s["stok"].kod, s["tuketim"]) for s in m["fason"]["satirlar"]], [("151-20-0014", D("1"))])
        self.assertEqual((m["fason"]["toplam_try"], m["toplam_try"], m["genel_try"]), (D("5.93"), D("3.54"), D("9.47")))   # 1/64 × ½ × 453,12
        m = ua.urun_maliyet(self.graf(), self.urun, D("1"))
        self.assertEqual({s["stok"].kod for s in m["fason"]["satirlar"]} & {"151-20-0014", "151-20-0015"}, {"151-20-0014", "151-20-0015"})

    def test_fasoncu_secimi_ve_doviz_fiyat(self):
        m = ua.urun_maliyet(self.graf(), self.urun, D("1"), fasoncu_id=self.salim.pk)
        f = m["fason"]
        self.assertEqual(f["fasoncu"], self.salim)
        s = f["satirlar"][0]
        self.assertEqual((s["stok"].kod, s["pb"], s["fiyat"], s["birim_try"], s["birim_usd"], s["tutar_try"], s["tutar_usd"]),
                         ("151-20-0014", "USD", D("0.2"), D("8"), D("0.2"), D("16"), D("0.4")))                 # 0,2 USD × 40
        self.assertEqual((f["toplam_try"], f["fiyatli"], f["fiyatsiz"], m["genel_try"]), (D("16"), 1, 4, D("800.15")))
        self.assertEqual(f["kur_notu"], "Fason fiyatı çevrimi: TCMB 09.10.2026, USD 40,0000")
        self.assertEqual(ua.urun_maliyet(self.graf(), self.urun, D("1"), fasoncu_id=999999)["fason"]["fasoncu"], self.gurbuz)   # listede yok → varsayılan

    def test_tarih_secimi(self):
        fs.fiyat_olustur(cari_id=self.gurbuz.pk, stok_id=self.sag.pk, birim_fiyat=D("7"), gecerlilik_baslangic=date(2026, 7, 1))
        Kur.objects.create(tarih=date(2026, 6, 30), usd_alis=D("38"))
        g = self.graf()
        self.assertEqual(ua.urun_maliyet(g, self.urun, D("1"))["fason"]["toplam_try"], D("90.10"))            # bugün: 7,00 geçerli
        eski = ua.urun_maliyet(g, self.urun, D("1"), tarih=date(2026, 6, 30))
        self.assertEqual((eski["fason"]["toplam_try"], eski["fason"]["tarih"]), (D("87.96"), date(2026, 6, 30)))   # başlangıcı geçmeyen en son
        on = next(s for s in eski["fason"]["satirlar"] if s["stok"].kod == "151-10-0017")
        self.assertEqual((on["birim_usd"], eski["fason"]["kur_notu"]), (D("0.422632"), "Fason fiyatı çevrimi: TCMB 30.06.2026, USD 38,0000"))   # 16,06 / 38, 6 ondalık
        self.assertEqual(eski["fason"]["toplam_usd"], sum(s["tutar_usd"] for s in eski["fason"]["satirlar"]))
        hic = ua.urun_maliyet(g, self.urun, D("1"), tarih=date(2025, 12, 31))
        self.assertEqual((hic["fason"]["satirlar"], hic["fason"]["fasoncular"], hic["fason"]["fasoncu"], hic["fason"]["fiyatsiz"]), ([], [], None, 5))
        self.assertEqual(hic["genel_try"], hic["toplam_try"])

    def test_kur_yoksa_doviz_fiyat_toplama_girmez(self):
        m = ua.urun_maliyet(self.graf(), self.urun, D("1"), fasoncu_id=self.salim.pk, tarih=date(2026, 9, 1))    # 7 gün içinde kur yok
        f = m["fason"]
        self.assertTrue(f["satirlar"][0]["maliyet_yok"])
        self.assertEqual((f["toplam_try"], f["kursuz"], f["fiyatli"]), (D("0"), 1, 1))
        self.assertEqual(f["kur_uyarilari"], ["151-20-0014: fason fiyatı (USD) çevrilemedi — USD için TCMB kuru bulunamadı"])
        self.assertEqual(m["genel_try"], m["toplam_try"])

    def test_pasif_silinmis_fiyat_ve_fasoncu_sayilmaz(self):
        f = FasonFiyat.objects.get(cari=self.salim, stok=self.sag)
        f.aktif = False
        f.save()
        self.assertEqual([s["cari"].pk for s in ua.urun_maliyet(self.graf(), self.urun, D("1"))["fason"]["fasoncular"]], [self.gurbuz.pk])
        f.aktif = True
        f.save()
        fs.fiyat_sil(f)
        self.assertEqual([s["cari"].pk for s in ua.urun_maliyet(self.graf(), self.urun, D("1"))["fason"]["fasoncular"]], [self.gurbuz.pk])

    def test_malzeme_hesabi_degismedi(self):
        m = ua.urun_maliyet(self.graf(), self.urun, D("2"))
        satir = {s["stok"].kod: s for gr in m["gruplar"] for s in gr["satirlar"]}
        self.assertEqual(set(satir), {"150-20-0002", "150-10-0007", "150-10-0008"})                        # fason satırları malzeme gruplarında değil
        self.assertEqual(satir["150-20-0002"]["tuketim"], D("0.03125"))
        self.assertEqual((m["miktar_toplam_try"], m["satir_sayisi"], m["maliyetsiz"]), (D("1568.30"), 3, 0))
        self.assertEqual(sum(s["pay"] for s in satir.values()).quantize(D("0.000001")), D("100"))         # pay malzeme toplamına göre
        self.assertEqual(sum(gr["toplam_try"] for gr in m["gruplar"]), m["toplam_try"])

    def test_karsilastir_genel_toplam(self):
        g = self.graf()
        k = ua.karsilastir(g, [self.urun, self.a41], "maliyet", "TL")
        self.assertEqual([s["stok"].kod for s in k["fason"]["satirlar"]], ["151-10-0017", "151-10-0018", "151-20-0014", "151-20-0015"])
        t0, t1 = k["toplamlar"]
        self.assertEqual((t0["toplam"], t0["fason"], t0["genel"], t0["fiyatsiz"]), (D("784.15"), D("87.96"), D("872.11"), 1))
        self.assertEqual((t1["toplam"], t1["fason"], t1["genel"], t1["fiyatsiz"]), (D("384.995"), D("32.12"), D("417.115"), 1))
        on = next(s for s in k["fason"]["satirlar"] if s["stok"].kod == "151-10-0017")
        self.assertEqual([(h["deger"], h["miktar"], h["pb"]) for h in on["hucreler"]], [(D("32.12"), D("2"), "TRY"), (D("32.12"), D("2"), "TRY")])
        sag = next(s for s in k["fason"]["satirlar"] if s["stok"].kod == "151-20-0014")
        self.assertIsNone(sag["hucreler"][1])                                                             # A41 menteşe kullanmıyor
        self.assertEqual((k["fason"]["fasoncu"], k["fason"]["tarih"], [s["cari"].pk for s in k["fason"]["fasoncular"]]), (self.gurbuz, BUGUN, [self.gurbuz.pk, self.salim.pk]))
        usd = ua.karsilastir(g, [self.urun], "maliyet", "USD")
        self.assertEqual(usd["toplamlar"][0]["fason"], D("2.199"))
        self.assertEqual(usd["toplamlar"][0]["genel"], usd["toplamlar"][0]["toplam"] + usd["toplamlar"][0]["fason"])
        self.assertIsNone(ua.karsilastir(g, [self.urun], "miktar")["fason"])
        ks = ua.karsilastir(g, [self.urun], "maliyet", "TL", fasoncu_id=self.salim.pk, tarih=date(2026, 9, 1))        # seçim Ürün görünümüyle aynı
        self.assertEqual((ks["fason"]["fasoncu"], ks["toplamlar"][0]["fason"], ks["toplamlar"][0]["fason_eksik"]), (self.salim, D("0"), 1))
        self.assertTrue(ks["fason"]["satirlar"][0]["hucreler"][0]["maliyet_yok"])

    def test_excel_maliyet_ve_karsilastir(self):
        g = self.graf()
        ws = load_workbook(BytesIO(maliyet_xlsx(self.urun, D("1"), ua.urun_maliyet(g, self.urun, D("1"))))).active
        satirlar = list(ws.iter_rows(min_row=5, values_only=True))
        etiketler = [str(r[0]) for r in satirlar]
        self.assertIn("MALZEME TOPLAMI", etiketler)
        self.assertTrue(any(e.startswith("İŞLEM / FASON BEDELİ — GÜRBÜZ FASON · 09.10.2026") for e in etiketler))
        kod = {r[1]: r for r in satirlar if r[1]}
        self.assertEqual((kod["151-20-0014"][0], kod["151-20-0014"][4], D(str(kod["151-20-0014"][5])), D(str(kod["151-20-0014"][8]))),
                         ("İşlem / fason bedeli", "Fason fiyatı (TRY) · GZ-0014", D("2"), D("11.86")))
        fason = next(r for r in satirlar if str(r[0]).startswith("İŞLEM/FASON TOPLAMI · fason fiyatı olmayan üretilen parça: 1"))
        genel = next(r for r in satirlar if str(r[0]).startswith("GENEL TOPLAM"))
        malzeme = next(r for r in satirlar if r[0] == "MALZEME TOPLAMI")
        self.assertEqual((D(str(malzeme[8])), D(str(fason[8])), D(str(genel[8]))), (D("784.15"), D("87.96"), D("872.11")))
        self.assertIn("4 parça fason fiyatlı (GÜRBÜZ FASON), fason fiyatı olmayan üretilen parça: 1", ws["A2"].value)
        wb = load_workbook(BytesIO(karsilastir_xlsx([ua.karsilastir(g, [self.urun, self.a41], "maliyet", pb) for pb in ("TL", "USD")])))
        satirlar = list(wb["Maliyet TL"].iter_rows(min_row=5, values_only=True))
        genel = next(r for r in satirlar if str(r[0]).startswith("GENEL TOPLAM (TL)"))
        self.assertEqual((D(str(genel[5])), D(str(genel[6]))), (D("872.11"), D("417.115")))
        fason = next(r for r in satirlar if str(r[0]).startswith("İŞLEM/FASON TOPLAMI (TL)"))
        self.assertEqual((D(str(fason[5])), D(str(fason[6]))), (D("87.96"), D("32.12")))
        self.assertTrue(any(r[1] == "151-20-0015" and r[4] == "Fason fiyatı" and r[6] is None for r in satirlar))      # A41 hücresi boş
        self.assertTrue(any(str(r[0]).startswith("fason fiyatı olmayan üretilen parça sayısı") for r in satirlar))
        self.assertIn("GENEL TOPLAM (USD)", [str(r[0]) for r in wb["Maliyet USD"].iter_rows(min_row=5, values_only=True)])


class EkranTest(FasonMaliyetBase):
    def setUp(self):
        self.client.force_login(User.objects.create_superuser("uaf_y", password="x"))
        Kur.objects.get_or_create(tarih=timezone.localdate(), defaults={"usd_alis": D("40")})
        self.url = reverse("core:urun_agaci")

    def test_maliyet_sekmesi_fason_grubu(self):
        r = self.client.get(self.url, {"urun": self.urun.pk, "sekme": "maliyet", "tarih": "2026-10-09"})
        h = r.content.decode()
        for parca in ("İşlem / fason bedeli · GÜRBÜZ FASON · 09.10.2026", "Malzeme toplamı", "İşlem/fason toplamı", "GENEL TOPLAM", "87,96", "872,11", "784,15",
                      "fason fiyatı olmayan üretilen parça: <strong>1</strong>", 'name="fasoncu"', 'name="tarih"', 'value="2026-10-09"', "GZ-0014", "kaynak-fason",
                      "<strong>4</strong> parça fason fiyatlı (GÜRBÜZ FASON)"):
            self.assertIn(parca, h)
        self.assertEqual(r.context["maliyet"]["fason"]["fasoncu"], self.gurbuz)
        r = self.client.get(self.url, {"urun": self.urun.pk, "sekme": "maliyet", "tarih": "2026-10-09", "fasoncu": self.salim.pk})
        h = r.content.decode()
        for parca in ("SALİM FASON", "16,00", "Fason fiyatı (USD)", f'<option value="{self.salim.pk}" selected>', "Fason fiyatı çevrimi: TCMB 09.10.2026, USD 40,0000"):
            self.assertIn(parca, h)
        r = self.client.get(self.url, {"urun": self.urun.pk, "sekme": "maliyet", "tarih": "2025-12-31"})
        self.assertContains(r, "5 üretilen parça için 31.12.2025 tarihinde geçerli fason fiyatı yok")
        self.assertNotContains(r, 'name="fasoncu"')
        r = self.client.get(self.url, {"urun": self.urun.pk, "sekme": "maliyet", "tarih": "bozuk"})         # geçersiz tarih → bugün
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context["fason_tarih"], timezone.localdate().isoformat())
        r = self.client.get(self.url, {"urun": self.urun.pk, "sekme": "maliyet", "tarih": "2026-10-09", "miktar": "3", "xlsx": "1"})
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn("fason fiyatı olmayan üretilen parça: 1", load_workbook(BytesIO(r.content)).active["A2"].value)

    def test_karsilastir_ekrani(self):
        r = self.client.get(self.url, {"gorunum": "karsilastir", "mod": "maliyet", "urunler": [self.urun.pk, self.a41.pk], "tarih": "2026-10-09"})
        h = r.content.decode()
        for parca in ("MALZEME TOPLAMI (TL)", "İŞLEM/FASON TOPLAMI (TL)", "GENEL TOPLAM (TL)", "872,11", "417,12", "fiyatsız üretilen parça: 1", 'name="fasoncu"',
                      "İşlem / fason bedeli · GÜRBÜZ FASON · 09.10.2026"):
            self.assertIn(parca, h)
        r = self.client.get(self.url, {"gorunum": "karsilastir", "mod": "maliyet", "urunler": [self.urun.pk], "tarih": "2026-10-09", "fasoncu": self.salim.pk, "pb": "USD"})
        self.assertContains(r, "GENEL TOPLAM (USD)")
        self.assertEqual(r.context["sonuc"]["toplamlar"][0]["fason"], D("0.4"))
        r = self.client.get(self.url, {"gorunum": "karsilastir", "mod": "miktar", "urunler": [self.urun.pk]})
        self.assertNotContains(r, "GENEL TOPLAM")
        self.assertNotContains(r, 'name="tarih"')
