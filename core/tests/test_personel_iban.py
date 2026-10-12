"""Personel kartı BANKA BİLGİLERİ (2026-10-12): IBAN kanonikleştirme (boşluk temizleme, büyük harf), TR + 26 karakter + mod-97 doğrulaması ve
Türkçe hata, 4'lü gruplu gösterim, banka adının IBAN banka kodundan önerilmesi (elle değiştirilebilir), hesap sahibi varsayılanı (ad soyad),
kart formu/detayında Banka Bilgileri bölümü, personel listesinde yok, mesai/atölye hesabı göremez, Excel dökümünde IBAN sütunu."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.dogrulama import iban_banka_adi, iban_bicimli, iban_dogrula, iban_gecerli_tr, iban_kanonik
from core.models import EkranYetki, Personel
from core.services import personel as personel_servis
from core.services.mesai_hesap import hesap_olustur
from core.services.personel_dokum import dokum_xlsx
from core.tests.ik_yardimci import personel_kur

GECERLI = "TR330006100519786457841326"          # ISO örneği (banka kodu 00061: listede yok → öneri boş)


def _tr_iban(banka_kodu="00010", hesap="0051978645784132"):
    """Geçerli TR IBAN üretir: TR00 + banka kodu + 0 + hesap → kontrol haneleri mod-97."""
    bban = banka_kodu + "0" + hesap
    donmus = bban + "TR00"
    sayi = "".join(str(ord(c) - 55) if c.isalpha() else c for c in donmus)
    kontrol = 98 - int(sayi) % 97
    return f"TR{kontrol:02d}{bban}"


class DogrulamaTest(TestCase):
    def test_kanonik_bicim_mod97(self):
        self.assertEqual(iban_kanonik(" tr33 0006 1005-1978 6457 8413 26 "), GECERLI)
        self.assertTrue(iban_gecerli_tr(GECERLI))
        self.assertTrue(iban_gecerli_tr("TR33 0006 1005 1978 6457 8413 26"))
        self.assertFalse(iban_gecerli_tr("TR330006100519786457841327"))                         # kontrol hanesi bozuk
        self.assertFalse(iban_gecerli_tr("TR33000610051978645784132"))                          # 25 karakter
        self.assertFalse(iban_gecerli_tr("DE89370400440532013000"))                             # TR değil
        self.assertFalse(iban_gecerli_tr("TR33000610051978645784132A"))
        self.assertEqual(iban_bicimli(GECERLI), "TR33 0006 1005 1978 6457 8413 26")
        self.assertEqual(iban_bicimli(""), "")
        ziraat = _tr_iban("00010")
        self.assertTrue(iban_gecerli_tr(ziraat))
        self.assertEqual((iban_banka_adi(ziraat), iban_banka_adi(_tr_iban("00062")), iban_banka_adi(GECERLI)), ("ZİRAAT BANKASI", "GARANTİ BBVA", ""))
        with self.assertRaises(ValidationError):
            iban_dogrula("TR12")
        iban_dogrula("")                                                                        # boş serbest


class ServisTest(TestCase):
    def test_kaydet_oneri_ve_varsayilanlar(self):
        p = personel_kur(ad="ayşe", soyad="kaya", iban=" tr" + _tr_iban("00010")[2:].lower() + " ")
        self.assertEqual((p.iban, p.banka_adi, p.hesap_sahibi), (_tr_iban("00010"), "ZİRAAT BANKASI", "AYŞE KAYA"))
        personel_servis.personel_guncelle(p, ad="ayşe", soyad="kaya", ise_giris_tarihi=p.ise_giris_tarihi, iban=_tr_iban("00062"),
                                          banka_adi="garanti şube x", hesap_sahibi="mehmet kaya")
        p.refresh_from_db()
        self.assertEqual((p.banka_adi, p.hesap_sahibi), ("GARANTİ ŞUBE X", "MEHMET KAYA"))        # elle değer korunur
        with self.assertRaisesMessage(personel_servis.PersonelHatasi, "Geçersiz IBAN"):
            personel_servis.personel_guncelle(p, ad="ayşe", soyad="kaya", ise_giris_tarihi=p.ise_giris_tarihi, iban="TR33 0006 1005 1978 6457 8413 27")
        p.refresh_from_db()
        self.assertEqual(p.iban, _tr_iban("00062"))                                              # hata: kayıt değişmedi
        personel_servis.personel_guncelle(p, ad="ayşe", soyad="kaya", ise_giris_tarihi=p.ise_giris_tarihi, iban="")
        p.refresh_from_db()
        self.assertEqual((p.iban, p.banka_adi, p.hesap_sahibi), ("", "", ""))                    # IBAN silinince öneriler de boş
        q = personel_kur(ad="veli", soyad="can")
        self.assertEqual((q.iban, q.banka_adi, q.hesap_sahibi), ("", "", ""))


class EkranTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.yon = User.objects.create_superuser("ibyon", password="x")
        cls.p = personel_kur(ad="hasan", soyad="demir", ise_giris_tarihi=date(2024, 3, 1))
        cls.u = hesap_olustur(cls.p, kullanici_adi="hasand", sifre="cokGizli!2026", ilk_giris_sifre=False)               # mesai hesabı (personel yetkisi yok)
        cls.ik = User.objects.create_user("ikci", password="x")
        EkranYetki.objects.create(kullanici=cls.ik, ekran_kod="personel")

    def govde(self, **ek):
        v = {"ad": "hasan", "soyad": "demir", "ise_giris_tarihi": "2024-03-01"}
        v.update(ek)
        return v

    def test_form_kaydet_hata_ve_detay(self):
        self.client.force_login(self.ik)
        url = reverse("core:personel_duzenle", args=[self.p.pk])
        r = self.client.get(url)
        self.assertContains(r, "Banka Bilgileri")
        self.assertContains(r, 'id="banka-kodlari"')
        r = self.client.post(url, self.govde(iban="TR33 0006 1005 1978 6457 8413 27"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Geçersiz IBAN")
        self.p.refresh_from_db()
        self.assertEqual(self.p.iban, "")
        r = self.client.post(url, self.govde(iban=" tr33 0006 1005 1978 6457 8413 26 "), follow=True)
        self.p.refresh_from_db()
        self.assertEqual((self.p.iban, self.p.hesap_sahibi, self.p.banka_adi), (GECERLI, "HASAN DEMİR", ""))
        self.assertContains(r, "TR33 0006 1005 1978 6457 8413 26")                                # detay: 4'lü gruplar
        self.assertContains(r, "HASAN DEMİR")
        r = self.client.post(url, self.govde(iban=_tr_iban("00010")), follow=True)
        self.assertContains(r, "ZİRAAT BANKASI")                                                   # öneri kaydedildi
        r = self.client.post(url, self.govde(iban=_tr_iban("00010"), banka_adi="Ziraat Çankaya"), follow=True)
        self.assertContains(r, "ZİRAAT ÇANKAYA")
        r = self.client.get(reverse("core:personeller"))                                           # listede IBAN yok
        self.assertContains(r, "HASAN DEMİR")
        self.assertNotContains(r, "TR33")
        self.assertNotContains(r, "IBAN")

    def test_mesai_hesabi_goremez_duzenleyemez(self):
        personel_servis.personel_guncelle(self.p, ad="hasan", soyad="demir", ise_giris_tarihi=date(2024, 3, 1), iban=GECERLI)
        self.client.force_login(self.u)
        self.assertEqual(self.client.get(reverse("core:personel_detay", args=[self.p.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse("core:personel_duzenle", args=[self.p.pk])).status_code, 403)
        self.assertEqual(self.client.post(reverse("core:personel_duzenle", args=[self.p.pk]), self.govde(iban=_tr_iban("00010"))).status_code, 403)
        r = self.client.get(reverse("core:mesaim"))
        self.assertEqual(r.status_code, 200)
        self.assertNotContains(r, "TR33")
        self.assertNotContains(r, "IBAN")
        self.p.refresh_from_db()
        self.assertEqual(self.p.iban, GECERLI)

    def test_excel_dokumunde_iban_sutunu(self):
        personel_servis.personel_guncelle(self.p, ad="hasan", soyad="demir", ise_giris_tarihi=date(2024, 3, 1), iban=GECERLI)
        ws = load_workbook(BytesIO(dokum_xlsx(2026, 6, bugun=date(2026, 7, 1)))).active
        basliklar = [c.value for c in ws[5]]
        self.assertIn("IBAN", basliklar)
        satir = next(r for r in ws.iter_rows(min_row=6, values_only=True) if r[1] == "HASAN DEMİR")
        self.assertEqual(satir[basliklar.index("IBAN")], "TR33 0006 1005 1978 6457 8413 26")
