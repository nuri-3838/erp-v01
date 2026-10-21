"""ATÖLYE > Personel İş Takibi (2026-10-14): Şu an (açık oturumlar + boşta, dünden kalan rozeti), Geçmiş (oturumlar + oturumsuz ofis/emirsiz kayıtlar; süre —,
"Süresiz kayıt: N"), kişi özeti (ekip işinde katıldığı tam + adam-saniye oranlı pay, fire %, verimlilik = oturum süresi ÷ mesai, açık oturum/mesai
şimdiye kadar), istasyon özeti, Excel 3 sayfa, filtreler, sabit sorgu sayısı, yetki (yalnız yönetici; menüde yalnız yöneticiye)."""
import datetime
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.models import EkranYetki, IsOturumu, IsOturumuUyesi, MesaiKaydi, OperasyonKaydi
from core.services import is_takibi
from core.services.uretim import emirsiz_uret, istasyon_emri_uret
from core.tarih import tr_bugun
from core.tests.test_atolye import IZINLI
from core.tests.test_is_oturumu import OturumTaban
from core.yetki import kullanici_menusu

D = Decimal


class TakipTaban(OturumTaban):
    def kaydir(self, oturum, baslangic_dk, bitis_dk=None, uye_dk=None):
        """Test oturumunu kurgulanmış zamana taşır (başlangıç/bitiş 'şimdi'den N dk önce); üyelerin katılım/ayrılma zamanı da aynı kaydırılır."""
        simdi = timezone.now()
        bas = simdi - datetime.timedelta(minutes=baslangic_dk)
        bit = simdi - datetime.timedelta(minutes=bitis_dk) if bitis_dk is not None else None
        IsOturumu.objects.filter(pk=oturum.pk).update(baslangic_zamani=bas, bitis_zamani=bit)
        for u in IsOturumuUyesi.objects.filter(oturum=oturum):
            dk = (uye_dk or {}).get(u.personel_id)
            IsOturumuUyesi.objects.filter(pk=u.pk).update(katilim_zamani=bas, ayrilma_zamani=(simdi - datetime.timedelta(minutes=dk)) if dk is not None else bit)

    def mesai_kaydir(self, personel, giris_dk, cikis_dk=None):
        simdi = timezone.now()
        MesaiKaydi.objects.filter(personel=personel, silindi=False).update(
            giris_zamani=simdi - datetime.timedelta(minutes=giris_dk), cikis_zamani=(simdi - datetime.timedelta(minutes=cikis_dk)) if cikis_dk is not None else None)

    def takip(self, **q):
        self.client.force_login(self.yon)
        return self.client.get(reverse("core:atolye_takip"), q)


class SuAnTest(TakipTaban):
    def test_acik_oturum_bosta_ve_dunden_kalan(self):
        self.baslat(self.u2)
        self.baslat(self.u3)
        self.baslat()                                                                                 # errahman + ayşe ekip, mehmet boşta
        self.basla(self.ie_kesim, ekip=[self.p2.pk])
        o = self.acik_oturum(self.ie_kesim)
        IsOturumu.objects.filter(pk=o.pk).update(is_tarihi=tr_bugun() - datetime.timedelta(days=1))
        MesaiKaydi.objects.filter(personel=self.p3).update(is_tarihi=tr_bugun() - datetime.timedelta(days=1))
        d = is_takibi.su_an()
        self.assertEqual(len(d["oturumlar"]), 1)
        s = d["oturumlar"][0]
        self.assertEqual((s["oturum"].pk, s["cikti"].pk, s["ozet"]["kisi"], s["dunku"]), (o.pk, self.kesilmis.pk, 2, True))
        self.assertEqual([b["personel"].pk for b in d["bosta"]], [self.p3.pk])
        self.assertTrue(d["bosta"][0]["dunku"])
        self.assertEqual(d["mesaide"], 3)
        r = self.takip()
        h = r.content.decode()
        self.assertContains(r, "ERRAHMAN OPERATOR")
        self.assertContains(r, "ekip: AYSE MESAİ")                                                   # kişi kartı: ekip arkadaşları (kendisi hariç)
        self.assertContains(r, "ekip: ERRAHMAN OPERATOR")
        self.assertContains(r, "MEHMET USTA")
        self.assertContains(r, "dünden kalan")
        self.assertContains(r, "60 sn")
        self.assertContains(r, self.ie_kesim.no)
        self.assertEqual((h.count('class="pk-kart calisiyor"'), h.count('class="pk-kart bosta"'), h.count('class="pk-kart yok"')), (2, 1, 0))
        self.assertContains(r, 'data-baslangic="')                                                    # canlı sayaç
        k = is_takibi.su_an_kartlar()
        self.assertEqual(k["sayac"], {"calisiyor": 2, "bosta": 1, "yok": 0})
        self.assertEqual([(x["personel"].pk, x["durum"]) for x in k["kartlar"]], [(self.p2.pk, "calisiyor"), (self.p.pk, "calisiyor"), (self.p3.pk, "bosta")])
        self.assertEqual([x.pk for x in k["kartlar"][1]["ekip"]], [self.p2.pk])
        # ayşe mesaisini bitirince oturumdan ayrılır: boşta değil (mesaide değil), ekip metninde kalır
        self.mesai_bitir(self.u2)
        d = is_takibi.su_an()
        self.assertEqual([b["personel"].pk for b in d["bosta"]], [self.p3.pk])
        self.assertEqual(d["mesaide"], 2)
        k = is_takibi.su_an_kartlar()                                                                 # ayşe atölye hesaplı → gri (mesaide değil) kart
        self.assertEqual([(x["personel"].pk, x["durum"]) for x in k["kartlar"]], [(self.p.pk, "calisiyor"), (self.p3.pk, "bosta"), (self.p2.pk, "yok")])
        self.assertContains(self.takip(), 'class="pk-kart yok"')

    def test_gri_kartlar_yalniz_atolye_veya_uretim_personeli(self):
        from core.tests.ik_yardimci import personel_kur
        ofis = personel_kur(ad="ofis", soyad="kisi", ise_giris_tarihi=datetime.date(2020, 1, 1))                  # hesapsız, departmansız
        uretim = personel_kur(ad="uretim", soyad="isci", ise_giris_tarihi=datetime.date(2020, 1, 1), departman="üretim")
        k = is_takibi.su_an_kartlar()
        gri = {x["personel"].pk for x in k["kartlar"] if x["durum"] == "yok"}
        self.assertIn(uretim.pk, gri)                                                                 # ÜRETİM departmanı
        self.assertIn(self.p.pk, gri)                                                                 # atölye hesabı
        self.assertNotIn(ofis.pk, gri)                                                                # ofis: yalnız mesaideyse
        from core.services import mesai as mesai_servis
        mesai_servis.manuel_ekle(ofis, giris_zamani=timezone.now() - datetime.timedelta(hours=1), duzeltme_notu="test")
        k = is_takibi.su_an_kartlar()
        self.assertEqual([x["durum"] for x in k["kartlar"] if x["personel"].pk == ofis.pk], ["bosta"])

    def test_sorgu_sayisi_sabit(self):
        self.baslat(self.u2)
        self.baslat(self.u3)
        self.baslat()
        self.basla(self.ie_kesim, ekip=[self.p2.pk])
        with self.assertNumQueries(4):
            is_takibi.su_an()
        with self.assertNumQueries(5):
            is_takibi.su_an_kartlar()
        self.client.force_login(self.u3)
        self.basla(self.ie_bukum)                                                                     # mehmet ikinci oturum
        with self.assertNumQueries(4):
            is_takibi.su_an()
        with self.assertNumQueries(5):
            is_takibi.su_an_kartlar()


class GecmisTest(TakipTaban):
    def kur(self):
        """errahman+ayşe ekip 60 dk (sağlam 4, fire 1); errahman tek 30 dk (sağlam 2); mehmet iptal; ofis kaydı (personelsiz); emirsiz (mehmet)."""
        self.baslat(self.u2)
        self.baslat(self.u3)
        self.baslat()
        self.basla(self.ie_kesim, ekip=[self.p2.pk])
        o1 = self.acik_oturum(self.ie_kesim)
        self.bitir(o1, "4", fire="1")
        o1.refresh_from_db()
        self.kaydir(o1, 120, 60, uye_dk={self.p.pk: 60, self.p2.pk: 90})                              # errahman 60 dk, ayşe 30 dk → pay 2/3, 1/3
        self.basla(self.ie_kesim)
        o2 = self.acik_oturum(self.ie_kesim)
        self.bitir(o2, "2")
        o2.refresh_from_db()
        self.kaydir(o2, 50, 20)                                                                       # 30 dk
        self.client.force_login(self.u3)
        self.basla(self.ie_bukum)
        o3 = self.acik_oturum(self.ie_bukum)
        self.client.post(reverse("core:atolye_oturum_iptal", args=[o3.pk]), HTTP_X_REAL_IP=IZINLI)
        self.kaydir(o3, 15, 10)
        ofis, _ = istasyon_emri_uret(self.ie_bukum, miktar=D("3"), tarih=tr_bugun())                  # ofis "Üretildi": personelsiz, süresiz
        emirsiz, _ = emirsiz_uret(operasyon_id=self.kesim_op.pk, depo_id=self.depo.pk, tarih=tr_bugun(), saglam="6", fire="2", personel=self.p3)
        self.mesai_kaydir(self.p, 240)                                                                # errahman 240 dk mesai (açık → şimdiye kadar)
        self.mesai_kaydir(self.p2, 200, 20)                                                           # ayşe 180 dk
        self.mesai_kaydir(self.p3, 100)                                                               # mehmet 100 dk
        return o1, o2, o3, ofis, emirsiz

    def test_satirlar_ozetler_ve_sayac(self):
        o1, o2, o3, ofis, emirsiz = self.kur()
        bugun = tr_bugun()
        sat = is_takibi.gecmis(baslangic=bugun, bitis=bugun)
        self.assertEqual(len(sat), 5)
        satir = {s["kayit"].pk if s["kayit"] else ("iptal", s["oturum"].pk): s for s in sat}
        s1 = satir[o1.kayit_id]
        self.assertEqual((s1["personel_metni"], s1["saglam"], s1["fire"], s1["fire_yuzde"], s1["sure_saniye"], s1["kaynak"]), ("ERRAHMAN OPERATOR, AYSE MESAİ", D("4"), D("1"), D("20.0"), 3600, "oturum"))
        self.assertEqual([(p["personel"].pk, p["saniye"]) for p in s1["paylar"]], [(self.p.pk, 3600), (self.p2.pk, 1800)])
        self.assertEqual(satir[("iptal", o3.pk)]["sure_metni"], "—")
        so = satir[ofis.pk]
        self.assertEqual((so["personel_metni"], so["sure_metni"], so["kaynak"], so["saglam"], so["fire"]), ("—", "—", "ofis", D("3"), D("0")))
        se = satir[emirsiz.pk]
        self.assertEqual((se["personel_metni"], se["kaynak"], se["kaynak_metni"], se["saglam"], se["fire"], se["fire_yuzde"]), ("MEHMET USTA", "emirsiz", "Emirsiz", D("6"), D("2"), D("25.0")))
        self.assertEqual(is_takibi.toplamlar(sat), {"satir": 5, "saglam": D("15"), "fire": D("3"), "suresiz": 2, "oturum": 3})                      # oturum sayacı iptali de sayar (satır sayısı)
        # kişi özeti
        kisi = {k["personel"].pk: k for k in is_takibi.kisi_ozeti(sat, baslangic=bugun, bitis=bugun)}
        e, a, m = kisi[self.p.pk], kisi[self.p2.pk], kisi[self.p3.pk]
        self.assertEqual((e["saglam"], e["fire"], e["pay"], e["is_saniye"], e["oturum"], e["suresiz"]), (D("6"), D("1"), D("4.667"), 5400, 2, 0))   # 4×2/3 + 2
        self.assertEqual((a["saglam"], a["pay"], a["is_saniye"]), (D("4"), D("1.333"), 1800))
        self.assertEqual((m["saglam"], m["fire"], m["oturum"], m["suresiz"], m["is_saniye"]), (D("6"), D("2"), 0, 1, 0))                             # iptal: süre/oturum sayılmaz
        self.assertEqual((e["mesai_saniye"] // 60, a["mesai_saniye"] // 60, m["mesai_saniye"] // 60), (240, 180, 100))
        self.assertEqual((e["verimlilik"], a["verimlilik"], m["verimlilik"]), (D("37.5"), D("16.7"), D("0.0")))
        # istasyon özeti
        ist = {k["istasyon"].pk: k for k in is_takibi.istasyon_ozeti(sat)}
        lz, bk = ist[self.lazer.pk], ist[self.bukum.pk]
        self.assertEqual((lz["saglam"], lz["fire"], lz["oturum"], lz["suresiz"], lz["sure_saniye"], lz["adam_saat"], lz["adet_saat"]), (D("12"), D("3"), 2, 1, 5400, D("2.00"), D("4.0")))
        self.assertEqual((bk["saglam"], bk["oturum"], bk["suresiz"], bk["sure_saniye"], bk["adet_saat"]), (D("3"), 0, 1, 0, None))                                     # iptal oturum sayılmaz

    def test_filtreler(self):
        o1, o2, o3, ofis, emirsiz = self.kur()
        bugun = tr_bugun()
        self.assertEqual(len(is_takibi.gecmis(baslangic=bugun, bitis=bugun, personel_id=self.p2.pk)), 1)
        self.assertEqual(len(is_takibi.gecmis(baslangic=bugun, bitis=bugun, personel_id=self.p3.pk)), 2)               # iptal oturumu + emirsiz
        self.assertEqual(len(is_takibi.gecmis(baslangic=bugun, bitis=bugun, istasyon_id=self.bukum.pk)), 2)            # iptal + ofis
        self.assertEqual([s["kaynak"] for s in is_takibi.gecmis(baslangic=bugun, bitis=bugun, kaynak="emirsiz")], ["emirsiz"])
        self.assertEqual(len(is_takibi.gecmis(baslangic=bugun, bitis=bugun, kaynak="ofis")), 1)
        self.assertEqual(len(is_takibi.gecmis(baslangic=bugun, bitis=bugun, kaynak="oturum")), 3)
        self.assertEqual(is_takibi.gecmis(baslangic=bugun - datetime.timedelta(days=3), bitis=bugun - datetime.timedelta(days=1)), [])

    def test_geri_alinan_kayit_sure_sayilir_adet_sayilmaz(self):
        o1, o2, *_ = self.kur()
        from core.services.uretim import operasyon_kaydi_sil
        operasyon_kaydi_sil(OperasyonKaydi.objects.get(pk=o2.kayit_id), onayli_geri_al=True)
        bugun = tr_bugun()
        s = next(x for x in is_takibi.gecmis(baslangic=bugun, bitis=bugun) if x["oturum"] is not None and x["oturum"].pk == o2.pk)
        self.assertEqual((s["saglam"], s["geri_alindi"], s["sure_saniye"]), (D("0"), True, 1800))

    def test_ekran_excel_ve_sorgu_sayisi(self):
        self.kur()
        r = self.takip(sekme="gecmis")
        self.assertEqual(r.status_code, 200)
        h = r.content.decode()
        for metin in ("Kişi özeti", "İstasyon özeti", "süresiz kayıt", "verimlilik", "ERRAHMAN OPERATOR, AYSE MESAİ", "Emirsiz", "İptal (iş yapılmadı)", "oturum süresi ÷ mesai",
                      "Ayrıntılı tablo", "Daha fazla filtre", "Bu hafta", "Özel tarih", "zaman çizelgesi"):
            self.assertContains(r, metin)
        self.assertContains(r, 'data-href="/uretim/kayitlar/')
        self.assertContains(r, "%37,5")
        self.assertIn('<span class="aktif">Bu hafta</span>', h)                                        # varsayılan aralık
        self.assertNotIn('<details class="pt-daha" style="flex-basis:100%" open', h)                   # istasyon/kaynak seçilmeden kapalı
        self.assertIn('<details class="pt-tablo" style="margin-top:1.2rem">', h)                      # ayrıntılı tablo varsayılan kapalı
        self.assertIn('class="kk-cubuk orta"', h)                                                     # errahman %37,5 → turuncu
        self.assertIn('class="kk-cubuk dusuk"', h)                                                    # ayşe %16,7 / mehmet %0 → kırmızı
        self.assertIn(f'id="zc-{self.p.pk}" hidden', h)                                               # çizelge kapalı gelir
        self.assertRegex(h, r'class="zc-blok" href="/uretim/kayitlar/\d+/" style="left:[\d.]+%;width:[\d.]+%;background:#')
        self.assertIn('class="zc-blok iptal"', h)
        self.assertIn('class="zc-mesai acik"', h)                                                     # errahman'ın açık mesaisi çizgili
        self.assertContains(r, 'data-et="Sağlam"')                                                    # dar ekranda etiketli sütun
        r = self.takip(sekme="gecmis", istasyon=self.lazer.pk)
        self.assertIn('<details class="pt-daha" style="flex-basis:100%" open', r.content.decode())    # filtre etkinse açık
        r = self.takip(sekme="gecmis", aralik="dun")
        self.assertIn('<span class="aktif">Dün</span>', r.content.decode())
        self.assertContains(r, "Bu aralıkta kayıt yok")
        r = self.takip(sekme="gecmis", kaynak="emirsiz")
        self.assertContains(r, "MEHMET USTA")
        self.assertNotContains(r, "ERRAHMAN OPERATOR, AYSE MESAİ")
        r = self.takip(sekme="gecmis", xlsx="1")
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Geçmiş", "Kişi özeti", "İstasyon özeti"])
        g = list(wb["Geçmiş"].iter_rows(min_row=5, values_only=True))
        self.assertEqual(len(g), 5)
        k = {row[0]: row for row in wb["Kişi özeti"].iter_rows(min_row=5, values_only=True)}
        self.assertEqual((k["ERRAHMAN OPERATOR"][1], k["ERRAHMAN OPERATOR"][2], k["ERRAHMAN OPERATOR"][7], k["ERRAHMAN OPERATOR"][9]), (6, 4.667, 90, 37.5))
        self.assertEqual(len(list(wb["İstasyon özeti"].iter_rows(min_row=5, values_only=True))), 2)
        # sorgu sayısı satır sayısından bağımsız
        bugun = tr_bugun()
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        with CaptureQueriesContext(connection) as q1:
            sat = is_takibi.gecmis(baslangic=bugun, bitis=bugun)
            is_takibi.kisi_ozeti(sat, baslangic=bugun, bitis=bugun)
            is_takibi.istasyon_ozeti(sat)
        for _ in range(3):
            emirsiz_uret(operasyon_id=self.kesim_op.pk, depo_id=self.depo.pk, tarih=bugun, saglam="1", personel=self.p3)
        self.client.force_login(self.u3)
        self.basla(self.ie_bukum)
        self.bitir(self.acik_oturum(self.ie_bukum), "1")
        with CaptureQueriesContext(connection) as q2:
            sat = is_takibi.gecmis(baslangic=bugun, bitis=bugun)
            is_takibi.kisi_ozeti(sat, baslangic=bugun, bitis=bugun)
            is_takibi.istasyon_ozeti(sat)
        self.assertEqual(len(q1), len(q2))
        self.assertLessEqual(len(q2), 6)
        self.takip(sekme="gecmis")                                                                    # ısınma (oturum/giriş sorguları)
        with CaptureQueriesContext(connection) as q3:
            self.takip(sekme="gecmis")
        for _ in range(3):
            emirsiz_uret(operasyon_id=self.kesim_op.pk, depo_id=self.depo.pk, tarih=bugun, saglam="1", personel=self.p2)
        with CaptureQueriesContext(connection) as q4:
            self.takip(sekme="gecmis")
        self.assertEqual(len(q3), len(q4))                                                            # sayfa sorgu sayısı satır sayısından bağımsız

    def test_zaman_cizelgesi(self):
        o1, o2, o3, ofis, emirsiz = self.kur()
        bugun = tr_bugun()
        sat = is_takibi.gecmis(baslangic=bugun, bitis=bugun)
        mesai = is_takibi.mesai_araliklari(bugun, bugun)
        self.assertEqual(sorted(p.pk for p in mesai), sorted([self.p.pk, self.p2.pk, self.p3.pk]))
        c = is_takibi.zaman_cizelgesi(sat, mesai, baslangic=bugun, bitis=bugun)
        e = c[self.p.pk]
        self.assertFalse(e["cok_uzun"])
        self.assertEqual(len(e["gunler"]), 1)
        g = e["gunler"][0]
        self.assertEqual(g["tarih"], bugun)
        self.assertEqual(len(g["mesai"]), 1)
        self.assertTrue(g["mesai"][0]["acik"])
        self.assertEqual([b["kod"] for b in g["bloklar"]], [self.lazer.kod, self.lazer.kod])            # o1 (120→60 dk önce), o2 (50→20 dk önce)
        b1, b2 = g["bloklar"]
        self.assertLess(float(b1["sol"]), float(b2["sol"]))
        self.assertTrue(0 <= float(b1["sol"]) < float(b1["sol"]) + float(b1["en"]) <= 100.01)
        self.assertNotIn(",", b1["sol"] + b1["en"])                                                # CSS için noktalı
        self.assertIn("sağlam 4 / fire 1", b1["ipucu"])
        self.assertTrue(len(g["saat_etiketleri"]) >= 2)
        m = c[self.p3.pk]
        self.assertTrue(m["gunler"][0]["bloklar"][0]["iptal"])
        self.assertEqual([s["kayit"].pk for s in m["suresiz"]], [emirsiz.pk])
        # 31 günden uzun aralık: çizelge yok
        c = is_takibi.zaman_cizelgesi(sat, mesai, baslangic=bugun - datetime.timedelta(days=40), bitis=bugun)
        self.assertTrue(c[self.p.pk]["cok_uzun"])
        self.assertEqual(c[self.p.pk]["gunler"], [])
        # kişi özeti aynı mesai verisiyle aynı sonucu verir (ek sorgu yok)
        with self.assertNumQueries(0):
            k = is_takibi.kisi_ozeti(sat, baslangic=bugun, bitis=bugun, mesai=mesai)
        self.assertEqual({x["personel"].pk: x["mesai_saniye"] // 60 for x in k}, {self.p.pk: 240, self.p2.pk: 180, self.p3.pk: 100})


class YetkiTest(TakipTaban):
    def test_yalniz_yonetici_ve_menu(self):
        self.client.force_login(self.u)                                                               # operatör
        self.assertEqual(self.client.get(reverse("core:atolye_takip")).status_code, 403)
        yetkili = User.objects.create_user("takipyetkili", password="x")
        EkranYetki.objects.create(kullanici=yetkili, ekran_kod="personel_is_takibi")
        EkranYetki.objects.create(kullanici=yetkili, ekran_kod="atolye")
        self.client.force_login(yetkili)
        self.assertEqual(self.client.get(reverse("core:atolye_takip")).status_code, 403)             # ekran yetkisi verilse de yönetici değil
        menu = {e.kod for m in kullanici_menusu(yetkili) for e in m.ekranlar}
        self.assertIn("atolye", menu)
        self.assertNotIn("personel_is_takibi", menu)
        menu = {e.kod for m in kullanici_menusu(self.yon) for e in m.ekranlar}
        self.assertIn("personel_is_takibi", menu)
        self.assertEqual(self.takip().status_code, 200)
        self.assertEqual(self.client.get(reverse("core:atolye")).status_code, 200)                     # operatör ekranı değişmedi
