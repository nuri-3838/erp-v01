"""CARİ RAPORLARI (Cariler > Raporlar): Yaşlandırma, Vadesi geçenler, Tarih itibarıyla bakiye, Cari ciro özeti — hepsi salt-okunur, filtreli, Excel çıktılı.
Kilitler: ``cari_ozet.acik_faturalar`` ile eşdeğerlik (FIFO), tarih itibarıyla kesim, dilim sınırları, sekme/PB ayrımı, ciro iade/yalnız-KDV kuralları ve
``FaturaSatir.tutar_tl`` ile aynı yuvarlama, SORGU SAYISININ cari/fatura sayısından bağımsızlığı, yazmama, yetki ve menü."""
import datetime
import re
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Cari, CariKategori, EkranYetki, Fatura, FaturaSatir, FaturaTipi, Kur, YevmiyeFisi, YevmiyeSatir
from core.services import cari_ozet, cari_rapor
from core.services.fatura import fatura_olustur
from core.services.finans import kasa_olustur
from core.services.kasa_hareket import hareket_olustur as kasa_hareket
from core.tarih import tr_bugun
from core.tests.test_fatura import FaturaTestTemel, _hesap

D = datetime.date
Dc = Decimal
T = D(2026, 6, 1)


def _yazma(sorgular):
    return [q["sql"][:80] for q in sorgular if q["sql"].lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE"))]


class RaporTemel(FaturaTestTemel):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        for gun in range(1, 32):
            Kur.objects.get_or_create(tarih=D(2026, 3, gun), defaults={"usd_alis": Dc("30"), "eur_alis": Dc("35")})
        _hesap("100.01", "KASA")
        cls.kasa = kasa_olustur(ad="K1", muhasebe_kodu="100.01")
        cls.su = User.objects.create_superuser("rpsu", password="x")
        cls.kat_ted = CariKategori.objects.create(kod="320", ad="TEDARİKÇİLER")
        cls.kat_ted_alt = CariKategori.objects.create(kod="10", ad="HAMMADDE", ust=cls.kat_ted)
        cls.kat_mus = CariKategori.objects.create(kod="120", ad="MÜŞTERİLER")
        cls.kat_mus_alt = CariKategori.objects.create(kod="10", ad="YURTİÇİ", ust=cls.kat_mus)
        Cari.objects.filter(pk=cls.tedarikci.pk).update(kategori=cls.kat_ted_alt)
        Cari.objects.filter(pk=cls.musteri.pk).update(kategori=cls.kat_mus_alt)
        _hesap("320.10.0002", "TEDARİKÇİ USD", kalem="KVYK")
        cls.usd = Cari.objects.create(kod="320-10-0002", unvan="TEDARİKÇİ USD", para_birimi="USD", muhasebe_kodu="320.10.0002", kategori=cls.kat_ted_alt)

    def _alis(self, cari, tarih, miktar="10", vade=None, no="", **kw):
        return fatura_olustur(tip_id=self.alis.pk, cari_id=cari.pk, tarih=tarih, satirlar=self._satir(miktar, "100"), vade_tarihi=vade, fatura_no=no, **kw)

    def _satis(self, cari, tarih, miktar="10", **kw):
        return fatura_olustur(tip_id=self.satis.pk, cari_id=cari.pk, tarih=tarih, satirlar=self._satir(miktar, "100"), **kw)


class DilimTest(RaporTemel):
    def test_dilim_anahtari_sinirlari(self):
        a = cari_rapor.dilim_anahtari
        self.assertEqual(a(None, 0), "vadesiz")
        self.assertEqual([a(T, g) for g in (-5, 0)], ["gelmemis", "gelmemis"])                    # vade günü dahil: vadesi gelmemiş
        self.assertEqual([a(T, g) for g in (1, 30)], ["d30", "d30"])
        self.assertEqual([a(T, g) for g in (31, 60)], ["d60", "d60"])
        self.assertEqual([a(T, g) for g in (61, 90)], ["d90", "d90"])
        self.assertEqual([a(T, g) for g in (91, 400)], ["d90p", "d90p"])

    def test_yaslandirma_dilimleri_tutarlari_ve_toplam(self):
        self._alis(self.tedarikci, D(2026, 3, 10), "10", vade=D(2026, 2, 20))      # 1.200 — 101 gün → 90+
        self._alis(self.tedarikci, D(2026, 3, 10), "5", vade=D(2026, 4, 17))       # 600 — 45 gün → 31-60
        self._alis(self.tedarikci, D(2026, 3, 10), "2", vade=D(2026, 5, 22))       # 240 — 10 gün → 1-30
        self._alis(self.tedarikci, D(2026, 3, 10), "1", vade=T)                    # 120 — vade günü → gelmemiş
        self._alis(self.tedarikci, D(2026, 3, 10), "1", vade=D(2026, 6, 10))       # 120 — gelmemiş
        self._alis(self.tedarikci, D(2026, 3, 11), "3", vade=None)                 # 360 — vadesi hiç yok (cari koşulu da yok)
        r = cari_rapor.yaslandirma(T, "borc")
        self.assertEqual([t["pb"] for t in r["tablolar"]], ["TRY"])
        satir = r["tablolar"][0]["satirlar"][0]
        self.assertEqual(satir["cari"].pk, self.tedarikci.pk)
        self.assertEqual(satir["dilimler"], {"gelmemis": Dc("240.00"), "d30": Dc("240.00"), "d60": Dc("600.00"), "d90": Dc("0"), "d90p": Dc("1200.00"), "vadesiz": Dc("360.00")})
        self.assertEqual((satir["toplam"], satir["adet"]), (Dc("2640.00"), 6))
        top = r["tablolar"][0]["toplam"]
        self.assertEqual((top["toplam"], top["adet"], top["hucreler"]), (Dc("2640.00"), 6, [Dc("240.00"), Dc("240.00"), Dc("600.00"), Dc("0"), Dc("1200.00"), Dc("360.00")]))
        self.assertEqual(r["sekme_adetleri"], {"alacak": 0, "borc": 6})

    def test_odeme_en_eski_faturadan_dusulur_ve_tarih_itibariyla_kesim(self):
        f1 = self._alis(self.tedarikci, D(2026, 3, 10), "10", vade=D(2026, 4, 10), no="A1")        # 1.200
        f2 = self._alis(self.tedarikci, D(2026, 3, 11), "5", vade=D(2026, 4, 11), no="A2")         # 600
        kasa_hareket(kasa=self.kasa, tip="cari_odeme", karsi=self.tedarikci, tutar="1300", tarih=D(2026, 3, 20))
        self._alis(self.tedarikci, D(2026, 3, 25), "1", vade=D(2026, 4, 25), no="A3")              # 120, ödemeden sonra kesilen
        # 15.03 itibarıyla: ödeme ve A3 henüz yok → A1 + A2 tam açık
        r = cari_rapor.acik_faturalar_toplu(D(2026, 3, 15))
        self.assertEqual({s["fatura"].pk: s["kalan"] for s in r}, {f1.pk: Dc("1200.00"), f2.pk: Dc("600.00")})
        # 31.03 itibarıyla: 1.300 ödeme A1'i kapatır, A2'den 100 düşer; A3 açık
        r = cari_rapor.acik_faturalar_toplu(D(2026, 3, 31))
        self.assertEqual(sorted((s["fatura"].fatura_no, s["kalan"]) for s in r), [("A2", Dc("500.00")), ("A3", Dc("120.00"))])
        self.assertEqual(cari_rapor.acik_faturalar_toplu(D(2026, 3, 9)), [])                         # hiçbiri henüz kesilmedi

    def test_sekme_pb_ve_kategori_ayrimi(self):
        self._alis(self.tedarikci, D(2026, 3, 10), "10", vade=D(2026, 4, 1))
        self._satis(self.musteri, D(2026, 3, 10), "10", vade_tarihi=D(2026, 4, 1))
        fatura_olustur(tip_id=self.alis.pk, cari_id=self.usd.pk, tarih=D(2026, 3, 10), satirlar=self._satir("10", "100"), para_birimi="USD", vade_tarihi=D(2026, 4, 1))
        fatura_olustur(tip_id=self.alis.pk, cari_id=self.usd.pk, tarih=D(2026, 3, 10), satirlar=self._satir("1", "100"), para_birimi="USD", vade_tarihi=D(2026, 5, 1))
        borc = cari_rapor.yaslandirma(T, "borc")
        self.assertEqual([t["pb"] for t in borc["tablolar"]], ["TRY", "USD"])                         # PB ayrı tablolar, TRY önce
        self.assertEqual([t["toplam"]["toplam"] for t in borc["tablolar"]], [Dc("1200.00"), Dc("1320.00")])   # USD: 10×100 + 1×100, KDV dahil; TL ile TOPLANMAZ
        self.assertEqual(borc["pb_secenekleri"], ["TRY", "USD"])
        self.assertEqual([t["pb"] for t in cari_rapor.yaslandirma(T, "borc", pb="USD")["tablolar"]], ["USD"])
        alacak = cari_rapor.yaslandirma(T, "alacak")
        self.assertEqual([s["cari"].pk for t in alacak["tablolar"] for s in t["satirlar"]], [self.musteri.pk])
        self.assertEqual(alacak["sekme_adetleri"], {"alacak": 1, "borc": 3})
        kat = cari_rapor.yaslandirma(T, "borc", kategori_id=self.kat_mus_alt.pk)
        self.assertEqual(kat["tablolar"], [])                                                          # müşteri kategorisinde borç sekmesi boş

    def test_acik_faturalar_toplu_acik_faturalar_ile_esdeger(self):
        self._alis(self.tedarikci, D(2026, 3, 10), "10", vade=D(2026, 4, 10))
        self._alis(self.tedarikci, D(2026, 3, 11), "5")
        self._satis(self.musteri, D(2026, 3, 12), "4")
        self._satis(self.musteri, D(2026, 3, 13), "2")
        fatura_olustur(tip_id=self.alis.pk, cari_id=self.usd.pk, tarih=D(2026, 3, 10), satirlar=self._satir("3", "100"), para_birimi="USD")
        kasa_hareket(kasa=self.kasa, tip="cari_odeme", karsi=self.tedarikci, tutar="1500", tarih=D(2026, 3, 20))
        kasa_hareket(kasa=self.kasa, tip="cari_tahsilat", karsi=self.musteri, tutar="700", tarih=D(2026, 3, 21))
        bugun = D(2026, 12, 31)
        toplu = {(s["cari"].pk, s["fatura"].pk): (s["pb"], s["tutar"], s["kalan"], s["vade"], s["gecikme"], s["tarih"]) for s in cari_rapor.acik_faturalar_toplu(bugun)}
        tek = {}
        for c in (self.tedarikci, self.musteri, self.usd):
            for k in cari_ozet.acik_faturalar(c, bugun)["satirlar"]:
                tek[(c.pk, k["fatura"].pk)] = (k["pb"], k["tutar"], k["kalan"], k["vade"], k["gecikme"], k["tarih"])
        self.assertTrue(tek)
        self.assertEqual(toplu, tek)


class VadesiGecenlerTest(RaporTemel):
    def test_siralama_filtre_ve_toplam(self):
        a = self._alis(self.tedarikci, D(2026, 3, 10), "10", vade=D(2026, 4, 1), no="A")          # 1.200 — 61 gün
        b = self._alis(self.tedarikci, D(2026, 3, 10), "5", vade=D(2026, 5, 1), no="B")           # 600 — 31 gün
        self._alis(self.tedarikci, D(2026, 3, 10), "1", vade=T, no="C")                           # vade günü: GEÇMEZ
        self._alis(self.tedarikci, D(2026, 3, 10), "1", vade=None, no="D")                        # vadesiz: geçmez
        m = self._satis(self.musteri, D(2026, 3, 10), "2", vade_tarihi=D(2026, 3, 1))              # 240 — 92 gün (müşteri)
        u = fatura_olustur(tip_id=self.alis.pk, cari_id=self.usd.pk, tarih=D(2026, 3, 10), satirlar=self._satir("1", "100"), para_birimi="USD", vade_tarihi=D(2026, 5, 1))
        r = cari_rapor.vadesi_gecenler(T)
        self.assertEqual([(s["fatura"].pk, s["gecikme"]) for s in r["satirlar"]], [(m.pk, 92), (a.pk, 61), (b.pk, 31), (u.pk, 31)])    # gecikme azalan, eşitte kalan azalan
        self.assertEqual(r["toplam"], [("TRY", Dc("2040.00")), ("USD", Dc("120.00"))])
        self.assertEqual([s["fatura"].pk for s in cari_rapor.vadesi_gecenler(T, sekme="alacak")["satirlar"]], [m.pk])
        self.assertEqual([s["fatura"].pk for s in cari_rapor.vadesi_gecenler(T, sekme="borc", pb="USD")["satirlar"]], [u.pk])
        self.assertEqual(cari_rapor.vadesi_gecenler(D(2026, 3, 9))["satirlar"], [])                # o tarihte hiçbir fatura henüz kesilmemiş
        self.assertEqual({s["fatura"].pk for s in cari_rapor.vadesi_gecenler(D(2026, 4, 2))["satirlar"]}, {m.pk, a.pk})


class BakiyeTarihTest(RaporTemel):
    def test_tarih_itibariyla_bakiye_ve_filtreler(self):
        self._alis(self.tedarikci, D(2026, 3, 10), "10")                                           # −1.200 (alacak)
        self._satis(self.musteri, D(2026, 3, 12), "5")                                              # +600 (borç)
        kasa_hareket(kasa=self.kasa, tip="cari_odeme", karsi=self.tedarikci, tutar="200", tarih=D(2026, 3, 20))
        fatura_olustur(tip_id=self.alis.pk, cari_id=self.usd.pk, tarih=D(2026, 3, 10), satirlar=self._satir("10", "100"), para_birimi="USD")
        r15 = {s["cari"].pk: s for s in cari_rapor.bakiye_listesi(D(2026, 3, 15))["satirlar"]}
        self.assertEqual((r15[self.tedarikci.pk]["tl"], r15[self.musteri.pk]["tl"]), (Dc("-1200.00"), Dc("600.00")))
        r31 = cari_rapor.bakiye_listesi(D(2026, 3, 31))
        s = {x["cari"].pk: x for x in r31["satirlar"]}
        self.assertEqual(s[self.tedarikci.pk]["tl"], Dc("-1000.00"))                                # 20.03 ödemesi dahil
        self.assertEqual(s[self.usd.pk]["doviz"], [("USD", Dc("-1200.00"))])                          # döviz havuzu
        self.assertEqual(s[self.usd.pk]["son"], D(2026, 3, 10))
        self.assertEqual(cari_rapor.bakiye_listesi(D(2026, 3, 9))["satirlar"], [])                    # hiçbir hareket yok → bakiyesizler gizli
        self.assertEqual(len(cari_rapor.bakiye_listesi(D(2026, 3, 9), bakiyesizleri_gizle=False)["satirlar"]), Cari.objects.exclude(muhasebe_kodu="").count())
        self.assertEqual((r31["tl_borc"], r31["tl_alacak"], r31["tl_net"]), (Dc("600.00"), Dc("1000.00") + Dc("0") + s[self.usd.pk]["tl"].copy_abs(), r31["tl_borc"] - r31["tl_alacak"]))
        self.assertEqual([d[0] for d in r31["doviz_toplam"]], ["USD"])
        kat = cari_rapor.bakiye_listesi(D(2026, 3, 31), kategori_id=self.kat_mus_alt.pk)
        self.assertEqual([x["cari"].pk for x in kat["satirlar"]], [self.musteri.pk])
        self.assertEqual([x["cari"].pk for x in cari_rapor.bakiye_listesi(D(2026, 3, 31), tur="ticari")["satirlar"]].count(self.musteri.pk), 1)

    def test_tarihsiz_bakiyeler_ile_esdeger(self):
        self._alis(self.tedarikci, D(2026, 3, 10), "10")
        self._satis(self.musteri, D(2026, 3, 12), "5")
        cariler = [self.tedarikci, self.musteri, self.usd]
        self.assertEqual(cari_ozet.bakiyeler(cariler, tarih=D(2099, 1, 1)), cari_ozet.bakiyeler(cariler))


class CiroTest(RaporTemel):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.satis_iade = FaturaTipi.objects.create(ad="ALIŞ FATURASI-SATIŞ İADE", yon=FaturaTipi.Yon.ALIS, maliyet_fisi=FaturaTipi.MaliyetFisi.SATIS_IADE)
        cls.alis_iade = FaturaTipi.objects.create(ad="SATIŞ FATURASI-ALIŞ İADE", yon=FaturaTipi.Yon.SATIS, maliyet_fisi=FaturaTipi.MaliyetFisi.ALIS_IADE)
        cls.kdv_alis = FaturaTipi.objects.create(ad="ALIŞ FATURASI-KDV", yon=FaturaTipi.Yon.ALIS, yalniz_kdv=True)

    def _f(self, tip, cari, tarih, miktar, fiyat, pb="TRY", kur="1", durum=Fatura.Durum.ONAYLI):
        f = Fatura.objects.create(tip=tip, yon=tip.yon, durum=durum, cari=cari, tarih=tarih, para_birimi=pb, kur=Dc(kur))
        FaturaSatir.objects.create(fatura=f, stok=self.stok, miktar=Dc(miktar), birim_fiyat=Dc(fiyat), kdv=self.kdv)
        return f

    def test_ciro_brut_iade_net_belge_ve_kurallar(self):
        self._f(self.satis, self.musteri, D(2026, 3, 5), "10", "100")                                   # 1.000
        self._f(self.satis, self.musteri, D(2026, 3, 6), "1", "100", "USD", "30.5")                      # 100 USD × 30,5 = 3.050 TL
        self._f(self.satis_iade, self.musteri, D(2026, 3, 7), "2", "100")                               # satış iadesi 200 (ALIŞ yönlü tip) → satıştan düşer
        self._f(self.alis, self.tedarikci, D(2026, 3, 8), "5", "100")                                    # alış 500
        self._f(self.alis_iade, self.tedarikci, D(2026, 3, 9), "1", "100")                              # alış iadesi 100 (SATIŞ yönlü tip) → alıştan düşer
        self._f(self.kdv_alis, self.tedarikci, D(2026, 3, 9), "50", "100")                              # yalnız KDV tipi: DIŞARIDA
        self._f(self.satis, self.musteri, D(2026, 3, 10), "99", "100", durum=Fatura.Durum.TASLAK)        # taslak: dışarıda
        self._f(self.satis, self.musteri, D(2026, 4, 10), "99", "100")                                   # dönem dışı
        r = cari_rapor.ciro(D(2026, 3, 1), D(2026, 3, 31))
        s = {x["cari"].pk: x for x in r["satirlar"]}
        m, t = s[self.musteri.pk], s[self.tedarikci.pk]
        self.assertEqual((m["satis"], m["satis_iade"], m["satis_net"], m["satis_belge"]), (Dc("4050.00"), Dc("200.00"), Dc("3850.00"), 3))
        self.assertEqual((m["alis"], m["alis_net"], m["alis_belge"]), (Dc("0"), Dc("0"), 0))
        self.assertEqual((t["alis"], t["alis_iade"], t["alis_net"], t["alis_belge"], t["satis_net"]), (Dc("500.00"), Dc("100.00"), Dc("400.00"), 2, Dc("0")))
        self.assertEqual((m["toplam"], t["toplam"]), (Dc("3850.00"), Dc("400.00")))
        self.assertEqual([x["cari"].pk for x in r["satirlar"]], [self.musteri.pk, self.tedarikci.pk])    # en yüksekten
        self.assertEqual((r["toplam"]["satis_net"], r["toplam"]["alis_net"], r["toplam"]["toplam"], r["toplam"]["satis_belge"], r["toplam"]["alis_belge"]),
                         (Dc("3850.00"), Dc("400.00"), Dc("4250.00"), 3, 2))
        self.assertEqual([x["cari"].pk for x in cari_rapor.ciro(D(2026, 3, 1), D(2026, 3, 31), sirala="alis")["satirlar"]], [self.tedarikci.pk, self.musteri.pk])
        self.assertEqual([x["cari"].pk for x in cari_rapor.ciro(D(2026, 3, 1), D(2026, 3, 31), kategori_id=self.kat_ted_alt.pk)["satirlar"]], [self.tedarikci.pk])
        self.assertEqual([x["cari"].pk for x in cari_rapor.ciro(D(2026, 3, 1), D(2026, 3, 31), tur="ticari")["satirlar"]].__len__(), 2)

    def test_sql_tutari_fatura_satir_tutar_tl_ile_ayni_yuvarlama(self):
        f = self._f(self.satis, self.musteri, D(2026, 3, 5), "3", "33.333333", "USD", "30.555555")
        self._f(self.satis, self.musteri, D(2026, 3, 6), "7", "0.145", "EUR", "35.123456")
        beklenen = sum((s.tutar_tl for s in FaturaSatir.objects.filter(fatura__cari=self.musteri).select_related("fatura")), Dc("0"))
        r = cari_rapor.ciro(D(2026, 3, 1), D(2026, 3, 31))
        self.assertEqual(r["satirlar"][0]["satis"], beklenen)
        self.assertEqual(f.satirlar.first().tutar_tl, Dc("3055.56"))                                  # önce tutar 2 hane (99,999999 → 100,00), sonra × kur → 2 hane

    def test_onaylanmis_faturali_hizmet_ciro_calismasi_gercek_servis_faturasi(self):
        self._satis(self.musteri, D(2026, 3, 10), "10")                                              # 1.000 (KDV hariç) — gerçek servis akışı
        self._alis(self.tedarikci, D(2026, 3, 10), "4")
        r = cari_rapor.ciro(D(2026, 3, 1), D(2026, 3, 31))
        self.assertEqual((r["toplam"]["satis_net"], r["toplam"]["alis_net"]), (Dc("1000.00"), Dc("400.00")))


class SorguSayisiTest(RaporTemel):
    """Sorgu sayısı cari/fatura/ödeme sayısından BAĞIMSIZ (sabit)."""

    def _doldur(self, adet, baslangic):
        for i in range(baslangic, baslangic + adet):
            kod = f"320.20.{i:04d}"
            _hesap(kod, f"TOPLU {i}", kalem="KVYK")
            c = Cari.objects.create(kod=f"320-20-{i:04d}", unvan=f"TOPLU CARİ {i}", para_birimi="TRY", muhasebe_kodu=kod, kategori=self.kat_ted_alt)
            self._alis(c, D(2026, 3, 10), "10", vade=D(2026, 4, 1) + datetime.timedelta(days=i))
            self._alis(c, D(2026, 3, 11), "3")
            kasa_hareket(kasa=self.kasa, tip="cari_odeme", karsi=c, tutar="500", tarih=D(2026, 3, 20))

    def _sayilar(self):
        sonuc = {}
        for ad, fn in (("yaslandirma", lambda: cari_rapor.yaslandirma(T, "borc")), ("vadesi", lambda: cari_rapor.vadesi_gecenler(T)),
                       ("bakiye", lambda: cari_rapor.bakiye_listesi(T)), ("ciro", lambda: cari_rapor.ciro(D(2026, 1, 1), D(2026, 12, 31)))):
            fn()
            with CaptureQueriesContext(connection) as q:
                fn()
            sonuc[ad] = len(q)
            self.assertEqual(_yazma(q.captured_queries), [], ad)
        return sonuc

    def test_servis_sorgu_sayisi_sabit_ve_az(self):
        self._doldur(3, 100)
        az = self._sayilar()
        self._doldur(40, 200)
        cok = self._sayilar()
        self.assertEqual(az, cok)
        for ad, n in cok.items():
            self.assertLessEqual(n, 6, ad)

    def test_ekran_sorgu_sayisi_sabit_ve_yazmaz(self):
        self.client.force_login(self.su)
        urller = {"yaslandirma": reverse("core:cari_yaslandirma") + "?sekme=borc&tarih=2026-06-01", "vadesi": reverse("core:cari_vadesi_gecenler") + "?tarih=2026-06-01",
                  "bakiye": reverse("core:cari_bakiye_tarih") + "?tarih=2026-06-01", "ciro": reverse("core:cari_ciro") + "?bas=2026-01-01&bit=2026-12-31"}

        def say():
            sonuc = {}
            for ad, u in urller.items():
                self.client.get(u)
                with CaptureQueriesContext(connection) as q:
                    self.assertEqual(self.client.get(u).status_code, 200)
                sonuc[ad] = len(q)
                self.assertEqual(_yazma(q.captured_queries), [], ad)
            return sonuc
        self._doldur(3, 100)
        az = say()
        self._doldur(40, 200)
        cok = say()
        self.assertEqual(az, cok)
        for ad, n in cok.items():
            self.assertLessEqual(n, 20, ad)

    def test_rapor_hicbir_kayit_yazmaz(self):
        self._doldur(2, 100)
        once = (Fatura.objects.count(), YevmiyeFisi.objects.count(), YevmiyeSatir.objects.count(), Cari.objects.count())
        self.client.force_login(self.su)
        for u in (reverse("core:cari_yaslandirma"), reverse("core:cari_vadesi_gecenler"), reverse("core:cari_bakiye_tarih"), reverse("core:cari_ciro")):
            self.client.get(u)
            self.client.get(u + "?xlsx=1")
        self.assertEqual(once, (Fatura.objects.count(), YevmiyeFisi.objects.count(), YevmiyeSatir.objects.count(), Cari.objects.count()))


class EkranExcelYetkiTest(RaporTemel):
    def _veri(self):
        self._alis(self.tedarikci, D(2026, 3, 10), "10", vade=D(2026, 4, 1), no="A1")
        self._satis(self.musteri, D(2026, 3, 10), "5", vade_tarihi=D(2026, 4, 5))
        fatura_olustur(tip_id=self.alis.pk, cari_id=self.usd.pk, tarih=D(2026, 3, 10), satirlar=self._satir("1", "100"), para_birimi="USD", vade_tarihi=D(2026, 4, 1))

    def test_ekranlar_icerik_ve_cari_linkleri(self):
        self._veri()
        self.client.force_login(self.su)
        h = self.client.get(reverse("core:cari_yaslandirma"), {"tarih": "2026-06-01", "sekme": "borc"}).content.decode()
        for beklenen in ("Yaşlandırma", "90+ gün", "Vadesi yok", "Vadesi gelmemiş", "TEDARİKÇİ A", "TEDARİKÇİ USD", reverse("core:cari_detay", args=[self.tedarikci.pk]),
                         "Alacak (müşteri)", "Borç (tedarikçi)", "1.200,00", "120,00"):
            self.assertIn(beklenen, h)
        self.assertEqual(len(re.findall(r"<h2 class=\"rp-baslik\">", h)), 2)                         # TRY + USD ayrı tablolar
        h = self.client.get(reverse("core:cari_vadesi_gecenler"), {"tarih": "2026-06-01"}).content.decode()
        self.assertIn("A1", h)
        self.assertIn(reverse("core:cari_detay", args=[self.tedarikci.pk]), h)
        self.assertIn("gün", h)
        h = self.client.get(reverse("core:cari_bakiye_tarih"), {"tarih": "2026-06-01"}).content.decode()
        self.assertIn("1.200,00 A", h)
        self.assertIn("120,00 USD A", h)
        h = self.client.get(reverse("core:cari_ciro"), {"bas": "2026-01-01", "bit": "2026-12-31"}).content.decode()
        self.assertIn("MÜŞTERİ A", h)
        self.assertIn("500,00", h)

    def test_varsayilan_tarih_bugun(self):
        self.client.force_login(self.su)
        for ad in ("core:cari_yaslandirma", "core:cari_vadesi_gecenler", "core:cari_bakiye_tarih"):
            self.assertEqual(self.client.get(reverse(ad)).context["tarih"], tr_bugun(), ad)
        self.assertEqual(self.client.get(reverse("core:cari_yaslandirma"), {"tarih": "bozuk"}).context["tarih"], tr_bugun())

    def test_excel_ciktilari(self):
        self._veri()
        self.client.force_login(self.su)
        r = self.client.get(reverse("core:cari_yaslandirma"), {"tarih": "2026-06-01", "sekme": "borc", "xlsx": "1"})
        self.assertIn("spreadsheetml", r["Content-Type"])
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Borç TRY", "Borç USD"])                                    # PB başına sayfa
        ws = wb["Borç TRY"]
        satirlar = list(ws.iter_rows(min_row=4, values_only=True))
        self.assertEqual(satirlar[0][:2], ("Cari Kodu", "Ünvan"))
        self.assertEqual((satirlar[1][1], satirlar[1][-2]), ("TEDARİKÇİ A", 1200.0))
        self.assertEqual(satirlar[-1][1], "TOPLAM")
        ws = load_workbook(BytesIO(self.client.get(reverse("core:cari_vadesi_gecenler"), {"tarih": "2026-06-01", "xlsx": "1"}).content)).active
        satirlar = list(ws.iter_rows(min_row=5, values_only=True))
        self.assertEqual((satirlar[0][3], satirlar[0][6]), ("A1", 61))
        ws = load_workbook(BytesIO(self.client.get(reverse("core:cari_bakiye_tarih"), {"tarih": "2026-06-01", "xlsx": "1"}).content)).active
        self.assertIn("TEDARİKÇİ A", [row[1] for row in ws.iter_rows(min_row=5, values_only=True)])
        ws = load_workbook(BytesIO(self.client.get(reverse("core:cari_ciro"), {"bas": "2026-01-01", "bit": "2026-12-31", "xlsx": "1"}).content)).active
        satirlar = list(ws.iter_rows(min_row=4, values_only=True))
        self.assertEqual(satirlar[0][5], "Net Satış")
        self.assertEqual(satirlar[-1][1], "TOPLAM")

    def test_yetki_ayri_kodlar_otomatik_atanmaz(self):
        kullanici = User.objects.create_user("sadece_cariler", password="x")
        EkranYetki.objects.create(kullanici=kullanici, ekran_kod="cariler")
        self.client.force_login(kullanici)
        self.assertEqual(self.client.get(reverse("core:cariler")).status_code, 200)
        for ad in ("core:cari_yaslandirma", "core:cari_vadesi_gecenler", "core:cari_bakiye_tarih", "core:cari_ciro"):
            self.assertEqual(self.client.get(reverse(ad)).status_code, 403, ad)                       # 'cariler' yetkisi rapor yetkisi VERMEZ
        for kod, ad in (("cari_yaslandirma", "core:cari_yaslandirma"), ("cari_vadesi_gecenler", "core:cari_vadesi_gecenler"),
                        ("cari_bakiye_tarih", "core:cari_bakiye_tarih"), ("cari_ciro", "core:cari_ciro")):
            k = User.objects.create_user(f"u_{kod}", password="x")
            EkranYetki.objects.create(kullanici=k, ekran_kod=kod)
            self.client.force_login(k)
            self.assertEqual(self.client.get(reverse(ad)).status_code, 200, ad)
            self.assertEqual(self.client.get(reverse(ad), {"xlsx": "1"}).status_code, 200, ad)

    def test_menu_raporlar_alt_menusu_yetkiye_gore(self):
        def menu(kullanici, yol):
            self.client.force_login(kullanici)
            h = self.client.get(yol).content.decode()
            return h[h.index('<nav class="yan-menu">'):h.index("</nav>", h.index('<nav class="yan-menu">'))]
        m = menu(self.su, reverse("core:cari_ciro"))
        gruplar = re.findall(r'<details class="alt-menu[^"]*"[^>]*>.*?</details>', m, re.S)
        self.assertEqual(len(gruplar), 4)                                                              # Muhasebe Raporlar + Finans Raporlar + Cariler (Raporlar + Tanımlar)
        raporlar = next(g for g in gruplar if "Yaşlandırma" in g)                                      # Cariler > Raporlar
        for ad in ("Yaşlandırma", "Vadesi Geçenler", "Tarih İtibarıyla Bakiye", "Cari Ciro Özeti"):
            self.assertIn(ad, raporlar)
        self.assertRegex(raporlar.split(">")[0] + ">", r"\bopen\b")                                    # içindeki sayfada AÇIK
        self.assertRegex(raporlar, rf'<a href="{re.escape(reverse("core:cari_ciro"))}" class="aktif">')
        self.assertNotRegex(next(g for g in gruplar if "Tanımlar" in g).split(">")[0] + ">", r"\bopen\b")
        m = menu(self.su, reverse("core:cariler"))
        self.assertNotRegex(next(g for g in re.findall(r'<details class="alt-menu[^"]*"[^>]*>.*?</details>', m, re.S) if "Yaşlandırma" in g).split(">")[0] + ">", r"\bopen\b")
        # yalnız 'cariler' yetkisi → Raporlar HİÇ görünmez; yalnız ciro yetkisi → Raporlar içinde yalnız Ciro
        k1 = User.objects.create_user("m1", password="x")
        EkranYetki.objects.create(kullanici=k1, ekran_kod="cariler")
        self.assertNotIn("Raporlar", menu(k1, reverse("core:cariler")))
        k2 = User.objects.create_user("m2", password="x")
        EkranYetki.objects.create(kullanici=k2, ekran_kod="cari_ciro")
        g = re.findall(r'<details class="alt-menu[^"]*"[^>]*>.*?</details>', menu(k2, reverse("core:cari_ciro")), re.S)
        self.assertEqual(len(g), 1)
        self.assertIn("Cari Ciro Özeti", g[0])
        self.assertNotIn("Yaşlandırma", g[0])

    def test_sablonlarda_onclick_yok(self):
        import pathlib
        kok = pathlib.Path(__file__).resolve().parent.parent / "templates" / "core"
        for ad in ("cari_rapor_yaslandirma", "cari_rapor_vadesi_gecenler", "cari_rapor_bakiye", "cari_rapor_ciro"):
            self.assertNotIn("window.location", (kok / f"{ad}.html").read_text(encoding="utf-8"))
