"""Stok Sayımı — maliyeti belirsiz kartta FAZLA için "Maliyet öner" / elle maliyet: öneri sırası (son hareket → reçete → boş), elle/önerilen maliyetle Düzelt'e açılış (loglu, yalnız yönetici),
onayda KESİN giriş (tahmini/geçici işareti yok; hareket ve fiş açıklamasında "SAYIM – tahmini maliyet (kaynak: …)"), sonraki sayımda Düzelt engellenmez, EKSİK belirsiz kartta engelli ("önce faturayı bağlayın"),
maliyet kancası çalışmaya devam eder."""
import datetime
import io
from decimal import Decimal

from django.contrib.auth.models import User
from django.urls import reverse
from openpyxl import load_workbook

from core.models import EkranYetki, IsIstasyonu, StokHareket, StokSayimMaliyetLogu
from core.services import stok_sayim as ss
from core.services.hareket import eldeki_miktar, hareket_ekle
from core.services.uretim import operasyon_olustur
from core.tests.test_stok_sayim_onay import BUGUN, GECEN_AY, OnayTaban

D = Decimal
G, C = StokHareket.Tur.GIRIS, StokHareket.Tur.CIKIS


class MaliyetTaban(OnayTaban):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.personel = User.objects.create_user("sayici", password="x")
        EkranYetki.objects.create(kullanici=cls.personel, ekran_kod="stok_sayim")
        cls.istasyon = IsIstasyonu.objects.create(kod="MNT", ad="MONTAJ")

    def fazla_sayim(self, stok, sayilan="2", kapat=False):
        s = ss.sayim_ac(depo_id=self.depo.pk, tarih=BUGUN, kullanici=self.yon)
        sat = s.satirlar.filter(stok=stok).first() or ss.satir_ekle(s, stok_id=stok.pk, kullanici=self.yon)
        ss.girdi_ekle(sat, kullanici=self.personel, miktar=sayilan)
        if kapat:
            ss.giris_kapat(s, self.yon)
        return s, sat

    def satir_sonuc(self, s, sat):
        (d,) = ss.fark_satirlari(s, satir_idleri=[sat.pk])["satirlar"]
        return d


class OneriTest(MaliyetTaban):
    def test_a_son_hareket_birim_maliyeti_ve_kaynak(self):
        k = self.kart("151")
        self.giris(k, "10", "1000")                                                           # 100 / adet
        hareket_ekle(stok_id=k.pk, depo_id=self.depo.pk, tarih=GECEN_AY + datetime.timedelta(days=1), tur=C, miktar=D("10"))    # stok 0 → ortalama yok
        s, sat = self.fazla_sayim(k)
        d = self.satir_sonuc(s, sat)
        self.assertEqual((d["fark"], d["belirsiz"], d["tutar"]), (D("2"), True, None))
        self.assertIn("elle maliyet", d["neden"])
        o = ss.maliyet_onerisi(k)
        self.assertEqual(o["deger"], D("100.000000"))
        self.assertTrue(o["kaynak"].startswith("Son hareket "))
        self.assertIn((GECEN_AY + datetime.timedelta(days=1)).strftime("%d.%m.%Y"), o["kaynak"])
        self.assertFalse(o["kismi"])

    def kur_recete(self, kismi=False):
        ham, civata = self.kart("150", ad="ham"), self.kart("150", ad="civata")
        self.giris(ham, "10", "100")                                                         # 10 / adet
        self.giris(civata, "10", "50")                                                       # 5 / adet
        cikti = self.kart("151", ad="üretilen")
        satirlar = [(ham, D("3")), (civata, D("2"))]
        if kismi:
            satirlar.append((self.kart("150", ad="maliyetsiz girdi"), D("1")))
        operasyon_olustur(istasyon_id=self.istasyon.pk, cikti_id=cikti.pk, cikti_miktar=D("1"), satirlar=satirlar)
        return cikti

    def test_b_recete_maliyeti_tam_ve_kismi(self):
        cikti = self.kur_recete()
        o = ss.maliyet_onerisi(cikti)
        self.assertEqual((o["deger"], o["kaynak"], o["kismi"]), (D("40.000000"), "Reçete (malzeme)", False))      # 3×10 + 2×5
        kismi = self.kur_recete(kismi=True)
        o = ss.maliyet_onerisi(kismi)
        self.assertEqual(o["deger"], D("40.000000"))
        self.assertTrue(o["kismi"])
        self.assertIn("KISMİ", o["kaynak"])
        self.assertIn("1 girdinin maliyeti yok", o["kaynak"])

    def test_c_oneri_yoksa_bos(self):
        k = self.kart("151")
        o = ss.maliyet_onerisi(k)
        self.assertEqual((o["deger"], o["kaynak"]), (None, ""))

    def test_son_hareket_recetenin_onune_gecer(self):
        cikti = self.kur_recete()
        self.giris(cikti, "4", "800")                                                        # 200 / adet (reçete 40 olsa da)
        hareket_ekle(stok_id=cikti.pk, depo_id=self.depo.pk, tarih=GECEN_AY + datetime.timedelta(days=1), tur=C, miktar=D("4"))
        self.assertEqual(ss.maliyet_onerisi(cikti)["deger"], D("200.000000"))


class ElleMaliyetTest(MaliyetTaban):
    def belirsiz_kart(self):
        k = self.kart("151")
        self.giris(k, "10", "1000")
        hareket_ekle(stok_id=k.pk, depo_id=self.depo.pk, tarih=GECEN_AY + datetime.timedelta(days=1), tur=C, miktar=D("10"))
        return k

    def test_kaydet_duzelt_acilir_ve_loglanir(self):
        k = self.belirsiz_kart()
        s, sat = self.fazla_sayim(k)
        with self.assertRaisesMessage(ss.SayimHatasi, "elle maliyet girin"):
            ss.karar_ver(sat, "DUZELT", kullanici=self.yon)
        oneri = ss.maliyet_onerisi(k)
        sat = ss.maliyet_kaydet(sat, kullanici=self.yon, birim="12,5", kaynak="Elle giriş", oneri=oneri)
        self.assertEqual((sat.manuel_birim_maliyet_try, sat.maliyet_kaynagi, sat.maliyet_giren_id), (D("12.5"), "Elle giriş", self.yon.pk))
        d = self.satir_sonuc(s, sat)
        self.assertEqual((d["belirsiz"], d["elle"], d["birim_maliyet"], d["tutar"], d["manuel_kaynak"]), (False, True, D("12.5"), D("25.00"), "Elle giriş"))
        ss.karar_ver(sat, "DUZELT", kullanici=self.yon)                                       # artık Düzelt verilebilir
        ss.maliyet_kaydet(sat, kullanici=self.yon, birim="13", kaynak="Son hareket test")
        l1, l2 = StokSayimMaliyetLogu.objects.filter(satir=sat).order_by("id")
        self.assertEqual((l1.kullanici_id, l1.eski, l1.yeni, l1.kaynak, l1.oneri, l1.oneri_kaynagi.startswith("Son hareket")), (self.yon.pk, None, D("12.5"), "Elle giriş", D("100.000000"), True))
        self.assertEqual((l2.eski, l2.yeni, l2.kaynak), (D("12.5"), D("13"), "Son hareket test"))

    def test_kosullar(self):
        k = self.belirsiz_kart()
        s = ss.sayim_ac(depo_id=self.depo.pk, tarih=BUGUN, kullanici=self.yon)
        sat = ss.satir_ekle(s, stok_id=k.pk, kullanici=self.yon)
        with self.assertRaisesMessage(ss.SayimHatasi, "sayılmadı"):
            ss.maliyet_kaydet(sat, kullanici=self.yon, birim="5")
        ss.girdi_ekle(sat, kullanici=self.personel, miktar="0")
        with self.assertRaisesMessage(ss.SayimHatasi, "fark yok"):
            ss.maliyet_kaydet(sat, kullanici=self.yon, birim="5")
        ss.girdi_ekle(sat, kullanici=self.personel, miktar="2")
        for kotu, mesaj in (("0", "sıfırdan büyük"), ("-3", "geçerli bir sayı"), ("abc", "geçerli bir sayı")):
            with self.assertRaisesMessage(ss.SayimHatasi, mesaj):
                ss.maliyet_kaydet(sat, kullanici=self.yon, birim=kotu)
        self.assertEqual(StokSayimMaliyetLogu.objects.count(), 0)
        belli = self.kart("151")                                                              # ortalaması olan kartta elle maliyet girilmez
        self.giris(belli, "5", "500")
        s2, sat2 = self.fazla_sayim(belli, sayilan="7")
        with self.assertRaisesMessage(ss.SayimHatasi, "ortalama maliyeti var"):
            ss.maliyet_kaydet(sat2, kullanici=self.yon, birim="5")
        ss.giris_kapat(s, self.yon)
        ss.giris_ac(s, self.yon)
        ss.sayim_sil(s, self.yon)
        with self.assertRaisesMessage(ss.SayimHatasi, "düzenlenemez"):
            ss.maliyet_kaydet(sat, kullanici=self.yon, birim="5")

    def test_eksik_belirsiz_kartta_engelli(self):
        f = self.kart("151")
        self.giris(f, "5", tutar=None) if False else hareket_ekle(stok_id=f.pk, depo_id=self.depo.pk, tarih=GECEN_AY, tur=G, miktar=D("5"))       # fiyatsız giriş: ortalama yok
        s, sat = self.fazla_sayim(f, sayilan="3")
        d = self.satir_sonuc(s, sat)
        self.assertEqual((d["fark"], d["belirsiz"]), (D("-2"), True))
        self.assertIn("önce faturayı bağlayın", d["neden"])
        with self.assertRaisesMessage(ss.SayimHatasi, "eksik için elle maliyet girilemez"):
            ss.maliyet_kaydet(sat, kullanici=self.yon, birim="5")
        with self.assertRaisesMessage(ss.SayimHatasi, "(önce faturayı bağlayın)"):
            ss.karar_ver(sat, "DUZELT", kullanici=self.yon)
        self.assertEqual(ss.karar_ver(sat, "ARASTIR", kullanici=self.yon).karar, "ARASTIR")

    def test_kaldir_karari_sifirlar(self):
        k = self.belirsiz_kart()
        s, sat = self.fazla_sayim(k)
        ss.maliyet_kaydet(sat, kullanici=self.yon, birim="12")
        ss.karar_ver(sat, "DUZELT", kullanici=self.yon, not_="x")
        sat = ss.maliyet_kaldir(sat, kullanici=self.yon)
        self.assertEqual((sat.manuel_birim_maliyet_try, sat.maliyet_kaynagi, sat.karar), (None, "", "BEKLIYOR"))
        self.assertEqual(list(StokSayimMaliyetLogu.objects.filter(satir=sat).order_by("id").values_list("yeni", flat=True)), [D("12.000000"), None])

    def test_karisik_kart_ortalamasi_varsa_belirli_sayilir(self):
        k = self.kart("151")
        self.giris(k, "5", "500")                                                              # fiyatlı 100/adet
        hareket_ekle(stok_id=k.pk, depo_id=self.depo.pk, tarih=GECEN_AY, tur=G, miktar=D("3"))  # + fiyatsız giriş (irsaliye benzeri)
        s, sat = self.fazla_sayim(k, sayilan="9")
        d = self.satir_sonuc(s, sat)
        self.assertEqual((d["belirsiz"], d["elle"]), (False, False))                          # kullanılabilir ortalama var → Düzelt engellenmez
        ss.karar_ver(sat, "DUZELT", kullanici=self.yon)


class OnayVeKancaTest(MaliyetTaban):
    def kur(self):
        k = self.kart("151", ad="bükülmüş test")
        self.giris(k, "10", "1000")
        hareket_ekle(stok_id=k.pk, depo_id=self.depo.pk, tarih=GECEN_AY + datetime.timedelta(days=1), tur=C, miktar=D("10"))
        s, sat = self.fazla_sayim(k)
        ss.maliyet_kaydet(sat, kullanici=self.yon, birim="12,5", kaynak="Elle giriş")
        ss.karar_ver(sat, "DUZELT", kullanici=self.yon)
        ss.giris_kapat(s, self.yon)
        return k, s, sat

    def test_onayda_kesin_giris_ve_aciklamalar(self):
        k, s, sat = self.kur()
        s = ss.onayla(s, self.yon)
        h = StokHareket.objects.get(kaynak="SAYIM")
        self.assertEqual((h.tur, h.miktar, h.giris_tutar_try, h.giris_tutar_usd, h.maliyet_durumu, h.giris_tahmini), (G, D("2"), D("25.00"), D("0.83"), "KESIN", False))
        self.assertIn("SAYIM – TAHMİNİ MALİYET (KAYNAK: ELLE GİRİŞ)", h.aciklama)
        self.assertEqual(eldeki_miktar(k), D("2"))
        k.refresh_from_db()
        self.assertEqual(k.ort_maliyet_try, D("12.5"))
        satirlar = {x.hesap_id: x for x in s.fis.satirlar.filter(silindi=False)}
        self.assertEqual((satirlar["151.10"].borc, satirlar["623"].alacak), (D("25.00"), D("25.00")))
        self.assertIn("SAYIM – TAHMİNİ MALİYET 12.5 (KAYNAK: ELLE GİRİŞ)", satirlar["151.10"].aciklama)
        self.assertEqual(s.satirlar.get(stok=k).onay_tutar_try, D("25.00"))
        # sonraki sayım: kartın artık ortalaması var → Düzelt engellenmez (elle maliyet gerekmez)
        s2, sat2 = self.fazla_sayim(k, sayilan="3")
        d = self.satir_sonuc(s2, sat2)
        self.assertEqual((d["belirsiz"], d["elle"], d["birim_maliyet"]), (False, False, D("12.5")))
        ss.karar_ver(sat2, "DUZELT", kullanici=self.yon)

    def test_donmus_ekran_ve_stok_karti_aciklamasi(self):
        k, s, sat = self.kur()
        s = ss.onayla(s, self.yon)
        self.client.force_login(self.yon)
        r = self.client.get(reverse("core:stok_sayim_belge", args=[s.pk]))
        d = next(x for x in r.context["satirlar"] if x["stok"].pk == k.pk)
        self.assertEqual((d["elle"], d["birim_maliyet"], d["manuel_kaynak"], d["tutar"]), (True, D("12.5"), "Elle giriş", D("25.00")))
        self.assertContains(r, "elle/önerilen")
        self.assertContains(self.client.get(reverse("core:stok_detay", args=[k.pk])), "SAYIM – TAHMİNİ MALİYET")

    def test_kanca_diger_satirlar_degisince_fis_yerinde_yenilenir_elle_giris_sabit(self):
        k, s, sat = self.kur()
        normal = self.kart("151", ad="normal kart")
        self.giris(normal, "10", "1000")                                                       # ort 100
        s2 = ss.sayim_ac(depo_id=self.depo.pk, tarih=BUGUN, kullanici=self.yon)
        for st, m in ((k, "2"), (normal, "8")):
            sat_x = s2.satirlar.filter(stok=st).first() or ss.satir_ekle(s2, stok_id=st.pk, kullanici=self.yon)
            ss.girdi_ekle(sat_x, kullanici=self.personel, miktar=m)
            if st == k:
                ss.maliyet_kaydet(sat_x, kullanici=self.yon, birim="12,5")
            ss.karar_ver(sat_x, "DUZELT", kullanici=self.yon)
        ss.giris_kapat(s2, self.yon)
        s2 = ss.onayla(s2, self.yon)
        fis_pk = s2.fis_id
        self.assertEqual(StokHareket.objects.get(sayim_satir__sayim=s2, stok=normal).tutar_try, D("200.00"))
        self.giris(normal, "10", "2000", tarih=GECEN_AY + datetime.timedelta(days=5))         # geç gelen pahalı giriş → ort 150
        s2.refresh_from_db()
        self.assertEqual(s2.fis_id, fis_pk)                                                    # aynı fiş yerinde güncellendi
        self.assertEqual(StokHareket.objects.get(sayim_satir__sayim=s2, stok=normal).tutar_try, D("300.00"))
        satirlar = list(s2.fis.satirlar.filter(silindi=False))
        self.assertEqual(sum(x.borc for x in satirlar if x.hesap_id == "623"), D("300.00"))        # normal kartın eksiği güncellendi
        self.assertEqual(sum(x.alacak for x in satirlar if x.hesap_id == "623"), D("25.00"))       # elle maliyetli fazla giriş SABİT
        self.assertEqual(sum(x.borc for x in s2.fis.satirlar.filter(silindi=False)), sum(x.alacak for x in s2.fis.satirlar.filter(silindi=False)))


class EkranTest(MaliyetTaban):
    def kur(self):
        k = self.kart("151", ad="ekran kart")
        self.giris(k, "10", "1000")
        hareket_ekle(stok_id=k.pk, depo_id=self.depo.pk, tarih=GECEN_AY + datetime.timedelta(days=1), tur=C, miktar=D("10"))
        s, sat = self.fazla_sayim(k, kapat=True)
        return k, s, sat

    def test_oneri_gorunur_kabul_elle_kaldir_ve_yetki(self):
        k, s, sat = self.kur()
        self.client.force_login(self.yon)
        belge = reverse("core:stok_sayim_belge", args=[s.pk])
        r = self.client.get(belge)
        self.assertContains(r, "Öneri:")
        self.assertContains(r, "Son hareket")
        self.assertContains(r, "Kabul et")
        self.assertContains(r, "elle birim maliyet")
        url = reverse("core:stok_sayim_maliyet", args=[s.pk])
        r = self.client.post(url, {"satir": sat.pk, "islem": "kabul", "goster": "hepsi"}, follow=True)
        self.assertContains(r, "birim maliyet kaydedildi")
        sat.refresh_from_db()
        self.assertEqual((sat.manuel_birim_maliyet_try, sat.maliyet_kaynagi.startswith("Son hareket"), sat.maliyet_giren_id), (D("100"), True, self.yon.pk))
        self.assertContains(r, "elle/önerilen")
        self.assertContains(r, "maliyet geçmişi (1)")
        self.client.post(url, {"satir": sat.pk, "islem": "elle", "birim": "15,75"})
        sat.refresh_from_db()
        self.assertEqual((sat.manuel_birim_maliyet_try, sat.maliyet_kaynagi), (D("15.75"), "Elle giriş"))
        r = self.client.post(reverse("core:stok_sayim_karar", args=[s.pk]), {f"karar_{sat.pk}": "DUZELT", "goster": "hepsi"}, follow=True)
        self.assertContains(r, "kararı kaydedildi")
        r = self.client.get(belge)
        self.assertContains(r, "kesin giriş")                                                  # onay panelinde elle maliyetli satır notu
        self.client.post(url, {"satir": sat.pk, "islem": "kaldir"})
        sat.refresh_from_db()
        self.assertEqual((sat.manuel_birim_maliyet_try, sat.karar), (None, "BEKLIYOR"))
        self.assertEqual(StokSayimMaliyetLogu.objects.filter(satir=sat).count(), 3)
        r = self.client.post(url, {"satir": sat.pk, "islem": "elle", "birim": "-1"}, follow=True)
        self.assertContains(r, "geçerli bir sayı")
        self.assertEqual(self.client.post(url, {"satir": "x", "islem": "kabul"}).status_code, 404)
        self.client.force_login(self.personel)
        self.assertEqual(self.client.post(url, {"satir": sat.pk, "islem": "kabul"}).status_code, 403)
        self.assertEqual(StokSayimMaliyetLogu.objects.filter(satir=sat).count(), 3)

    def test_oneri_yoksa_elle_ve_excel_notu(self):
        k = self.kart("151", ad="öneri yok")
        s, sat = self.fazla_sayim(k, kapat=True)
        self.client.force_login(self.yon)
        r = self.client.get(reverse("core:stok_sayim_belge", args=[s.pk]))
        self.assertContains(r, "Öneri yok")
        self.client.post(reverse("core:stok_sayim_maliyet", args=[s.pk]), {"satir": sat.pk, "islem": "kabul"})          # öneri yok → kabul edilemez
        sat.refresh_from_db()
        self.assertIsNone(sat.manuel_birim_maliyet_try)
        self.client.post(reverse("core:stok_sayim_maliyet", args=[s.pk]), {"satir": sat.pk, "islem": "elle", "birim": "9"})
        x = self.client.get(reverse("core:stok_sayim_excel_indir", args=[s.pk]), {"tur": "fark"})
        ws = load_workbook(io.BytesIO(x.content)).active
        self.assertEqual([c.value for c in ws[4]][-1], "Maliyet notu")
        satir = next(row for row in ws.iter_rows(min_row=5, values_only=True) if row[0] == k.kod)
        self.assertIn("elle/önerilen maliyet: Elle giriş", satir[-1])
