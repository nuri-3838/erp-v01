"""ÜRÜN AĞACI — NEREDE KULLANILIYOR görünümü: bir stoğu kullanan bitmiş ürünler (1 adet ürün başına KESİRLİ tüketim), yol, yan çıktı notu,
seriye göre gruplama. Zincir (canlı veriyi örnekler): 150-10-0007 profil → C serisi 6 model (ayak kesimi) ; ortak ara parça → A serisi 5 model."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Kategori, Stok
from core.services import urun_agaci as ua
from core.services.uretim import ana_cikti_payi, operasyon_olustur
from core.tests.test_uretim import _istasyon
from core.tests.test_urun_agaci_gorunum import GorunumBase

D = Decimal


class KullanimBase(GorunumBase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.p7 = cls._stok("150-10-0007", cls.kat_girdi, cls.boy, satinalma=True, uretim=False)
        cls.p9 = cls._stok("150-10-0009", cls.kat_girdi, cls.boy, satinalma=True, uretim=False)
        cls.paket = _istasyon("PAKET")
        # C serisi: 6 model, i. modelde i+1 adet ayak (3 ayak/boy)
        cls.c_modeller, cls.ayaklar = [], []
        for i in range(6):
            ayak = cls._stok(f"151-22-{i:04d}", cls.kat_ana)
            operasyon_olustur(istasyon_id=cls.kesim.pk, cikti_id=ayak.pk, cikti_miktar=D("3"), satirlar=[(cls.p7, D("1"))])
            model = cls._stok(f"152-22-{i + 10:04d}", cls.kat_ana)
            operasyon_olustur(istasyon_id=cls.montaj.pk, cikti_id=model.pk, cikti_miktar=D("1"), satirlar=[(ayak, D(str(i + 1)))])
            cls.ayaklar.append(ayak)
            cls.c_modeller.append(model)
        # A serisi: 5 model, hepsi ortak ara parçayı (1. basamak, 2/boy, p9'dan) kullanır
        cls.ortak = cls._stok("151-10-0026", cls.kat_ana)
        operasyon_olustur(istasyon_id=cls.kesim.pk, cikti_id=cls.ortak.pk, cikti_miktar=D("2"), satirlar=[(cls.p9, D("1"))])
        cls.a_modeller = []
        for i in range(5):
            m = cls._stok(f"152-10-{i + 10:04d}", cls.kat_ana)
            operasyon_olustur(istasyon_id=cls.montaj.pk, cikti_id=m.pk, cikti_miktar=D("1"), satirlar=[(cls.ortak, D("1"))])
            cls.a_modeller.append(m)
        cls.bos = cls._stok("150-99-9999", cls.kat_girdi, satinalma=True, uretim=False)     # hiçbir zincirde yok


class KullanimServisTest(KullanimBase):
    def sonuc(self, stok):
        return ua.nerede_kullaniliyor(ua.graf_yukle(), stok)

    def test_hammadde_c_serisinin_alti_modeli(self):
        s = self.sonuc(self.p7)
        self.assertEqual(s["sayi"], 6)
        self.assertEqual([g["seri"] for g in s["gruplar"]], ["C"])
        satir = {r["urun"].kod: r for r in s["gruplar"][0]["satirlar"]}
        self.assertEqual(set(satir), {m.kod for m in self.c_modeller})
        for i, m in enumerate(self.c_modeller):
            self.assertEqual(satir[m.kod]["tuketim"], D(i + 1) / D(3))                 # kesirli, yuvarlama yok
            self.assertEqual([x.kod for x in satir[m.kod]["yol"]], ["150-10-0007", self.ayaklar[i].kod, m.kod])

    def test_ortak_ara_parca_a_serisinin_bes_modeli(self):
        s = self.sonuc(self.ortak)
        self.assertEqual((s["sayi"], [g["seri"] for g in s["gruplar"]]), (5, ["A"]))
        self.assertTrue(all(r["tuketim"] == D("1") for r in s["gruplar"][0]["satirlar"]))
        s2 = self.sonuc(self.p9)                                                       # ham profil: ara parça üzerinden
        self.assertEqual(s2["sayi"], 5)
        r = s2["gruplar"][0]["satirlar"][0]
        self.assertEqual(r["tuketim"], D("0.5"))
        self.assertEqual([x.kod for x in r["yol"]], ["150-10-0009", "151-10-0026", r["urun"].kod])

    def test_seri_gruplari_siralidir_a_once_c(self):
        for kod in ("152-77-0000", "152-22-9000", "152-10-9000"):                      # Diğer, Çift çıkış, A tipi
            m = self._stok(kod, self.kat_ana)
            operasyon_olustur(istasyon_id=self.montaj.pk, cikti_id=m.pk, cikti_miktar=D("1"), satirlar=[(self.civata, D("1"))])
        s = self.sonuc(self.civata)
        seriler = [g["seri"] for g in s["gruplar"]]
        self.assertEqual(seriler, sorted(seriler, key=lambda x: ("A", "C", "DIGER").index(x)))
        self.assertIn("DIGER", seriler)

    def test_kullanilmayan_stok(self):
        s = self.sonuc(self.bos)
        self.assertEqual((s["sayi"], s["gruplar"], s["bitmis_urun"]), (0, [], False))

    def test_bitmis_urun_kendisi(self):
        s = self.sonuc(self.c_modeller[0])
        self.assertEqual(s["sayi"], 0)
        self.assertTrue(s["bitmis_urun"])

    def test_yan_cikti_notu(self):
        """AYAK55, AYAK66 kesiminin yan çıktısı: hem blok notu hem (ürün AYAK66'yı da kullanıyorsa) satır notu."""
        s = self.sonuc(self.yan)
        self.assertEqual([x.kod for x in s["yan_cikti_ureten"]], ["AYAK66"])
        m1 = self._stok("152-22-8001", self.kat_ana)
        operasyon_olustur(istasyon_id=self.montaj.pk, cikti_id=m1.pk, cikti_miktar=D("1"), satirlar=[(self.yan, D("2"))])
        m2 = self._stok("152-22-8002", self.kat_ana)
        operasyon_olustur(istasyon_id=self.montaj.pk, cikti_id=m2.pk, cikti_miktar=D("1"), satirlar=[(self.yan, D("1")), (self.ana, D("3"))])
        s = self.sonuc(self.yan)
        satir = {r["urun"].kod: r for g in s["gruplar"] for r in g["satirlar"]}
        self.assertEqual(satir["152-22-8001"]["yan_cikti_ureten"], [])                  # zincirinde AYAK66 yok
        self.assertEqual([x.kod for x in satir["152-22-8002"]["yan_cikti_ureten"]], ["AYAK66"])

    def test_ana_cikti_yan_payi_tuketime_yansir(self):
        m = self._stok("152-22-8003", self.kat_ana)
        operasyon_olustur(istasyon_id=self.montaj.pk, cikti_id=m.pk, cikti_miktar=D("1"), satirlar=[(self.ana, D("3"))])
        pay = ana_cikti_payi(self.op.cikti_miktar, self.op.boy_mm, [self.op.yan_ciktilar.get()])
        satirlar = [x for g in self.sonuc(self.profil)["gruplar"] for x in g["satirlar"]]
        r = next(x for x in satirlar if x["urun"].pk == m.pk)
        self.assertEqual(r["tuketim"], pay)                                              # 3 ana ayak = 1 boy × pay

    def test_urun_gorunumuyle_ayni_rakam(self):
        """Nerede kullanılıyor'daki tüketim, Ürün/Maliyet görünümündeki birim tüketimle aynı."""
        g = ua.graf_yukle()
        for m in self.c_modeller:
            nerede = {r["urun"].pk: r["tuketim"] for gr in ua.nerede_kullaniliyor(g, self.p7)["gruplar"] for r in gr["satirlar"]}[m.pk]
            maliyet = {s["stok"].pk: s for gr in ua.urun_maliyet(g, m, D("1"))["gruplar"] for s in gr["satirlar"]}[self.p7.pk]["tuketim"]
            self.assertEqual(nerede, maliyet)


class KullanimEkranTest(KullanimBase):
    def setUp(self):
        self.client.force_login(User.objects.create_superuser("ua_k", password="x"))
        self.url = reverse("core:urun_agaci")

    def al(self, **p):
        return self.client.get(self.url, {"gorunum": "kullanim", **p})

    def test_secici_gruplu_ve_zincir_disi_stoklar(self):
        h = self.al().content.decode()
        for g in ("Hammadde / hazır stok", "Ara parçalar", "Hiçbir zincirde olmayan stoklar"):
            self.assertIn(f'<optgroup label="{g}">', h)
        self.assertIn("150-99-9999", h)
        self.assertContains(self.al(), 'class="ua-seg aktif"')

    def test_sonuc_tablosu(self):
        r = self.al(stok=self.p7.pk)
        self.assertContains(r, "6</strong> bitmiş ürün zincirinde")
        self.assertContains(r, "Çift çıkış · 6 ürün")
        self.assertContains(r, "152-22-0010")
        self.assertContains(r, "0,3333 BOY")
        self.assertContains(r, "xlsx=1")

    def test_kullanilmayan_mesaji(self):
        self.assertContains(self.al(stok=self.bos.pk), "Bu stok hiçbir ürün zincirinde yok")

    def test_gecersiz_stok(self):
        self.assertEqual(self.al(stok="abc").status_code, 200)
        self.assertEqual(self.al(stok="999999").context["sonuc"], None)

    def test_excel(self):
        r = self.al(stok=self.p7.pk, xlsx="1")
        self.assertIn("spreadsheetml", r["Content-Type"])
        ws = load_workbook(BytesIO(r.content)).active
        self.assertTrue(ws["A1"].value.startswith("Nerede kullanılıyor"))
        self.assertEqual([c.value for c in ws[4]][:4], ["Seri", "Ürün kodu", "Ürün adı", "Tüketim"])
        self.assertEqual(ws.max_row, 4 + 6)
        self.assertIn("→", ws["E5"].value)

    def test_sorgu_sayisi_sabit(self):
        with CaptureQueriesContext(connection) as az:
            self.al(stok=self.p7.pk)
        for i in range(6):
            ayak = self._stok(f"151-22-9{i}", self.kat_ana)
            operasyon_olustur(istasyon_id=self.kesim.pk, cikti_id=ayak.pk, cikti_miktar=D("3"), satirlar=[(self.p7, D("1"))])
            m = self._stok(f"152-22-9{i}", self.kat_ana)
            operasyon_olustur(istasyon_id=self.montaj.pk, cikti_id=m.pk, cikti_miktar=D("1"), satirlar=[(ayak, D("2"))])
        with CaptureQueriesContext(connection) as cok:
            r = self.al(stok=self.p7.pk)
        self.assertEqual(r.context["sonuc"]["sayi"], 12)
        self.assertEqual(len(az), len(cok))
