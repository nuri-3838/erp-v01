"""ATÖLYE operatör ekranı iyileştirmeleri (2026-10-12): malzeme durumu (bugünkü kullanılabilir stoğa göre yapılabilir adet; varyant bazında),
"Malzeme yok" soluk kart + Başla uyarısı (engel yok), sıralama (hazır önce, sonra termin/tarih), istasyon kartlarında açık iş sayısı (yoksa soluk),
başlıkta ad soyad (kullanıcı adı/TC yok) + Çıkış (açık mesai/iş uyarısı), kart IE no küçük / "Stok" rozeti, Bitir: sağlam + fire alanları BOŞ gelir
(placeholder "Kalan: X"; boş → "Adet girin"), fire girdisi tüketilir ama stoğa girmez, tamamlanan sağlam kadar; İş Oturumları'nda fire sütunu."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.urls import reverse
from openpyxl import load_workbook

from core.models import IsOturumu, OperasyonKaydi, OperasyonKaydiCikti, StokHareket
from core.services import atolye, is_oturumu
from core.services.hareket import hareket_ekle
from core.services.uretim import uretim_emri_olustur
from core.tests.test_atolye import AtolyeTaban
from core.tests.test_uretim import _istasyon

D = Decimal


class MalzemeVeSiralamaTest(AtolyeTaban):
    def test_malzeme_durumu_siralama_ve_basla_uyarisi(self):
        emir2 = uretim_emri_olustur(kalemler=[{"hedef_urun_id": self.bukulmus.pk, "hedef_miktar": D("2")}], depo_id=self.depo.pk, tarih=date(2026, 1, 5))
        satirlar = atolye.acik_emirler(self.lazer)
        self.assertEqual([s["ie"].uretim_emri_id for s in satirlar], [emir2.pk, self.emir.pk])             # ikisi de hazır → tarihi yakın önce
        self.assertEqual([(s["malzeme"]["hazir"], s["malzeme"]["durum"]) for s in satirlar], [(D("2"), "HAZIR"), (D("10"), "HAZIR")])
        self.assertEqual(satirlar[1]["malzeme"]["metin"], "Malzeme hazır: 10 adet yapılabilir")
        # profil 5'e düşer: emir1 için kullanılabilir = 5 − emir2'nin ayırması 1 = 4 profil → 8 adet (KISMI); emir2 için 5 − 5 = 0 → YOK
        hareket_ekle(stok_id=self.profil.pk, depo_id=self.depo.pk, tarih=date(2026, 1, 2), tur=StokHareket.Tur.CIKIS, miktar=D("95"))
        satirlar = atolye.acik_emirler(self.lazer)
        self.assertEqual([s["ie"].uretim_emri_id for s in satirlar], [self.emir.pk, emir2.pk])             # KISMI, YOK'tan önce
        self.assertEqual([(s["malzeme"]["hazir"], s["malzeme"]["durum"], s["malzeme"]["metin"]) for s in satirlar],
                         [(D("8"), "KISMI", "Kısmen hazır: 8 / 10 — eksik: ham profil (8 adetlik) — satınalma bekleniyor"),
                          (D("0"), "YOK", "Malzeme yok: ham profil yok — satınalma bekleniyor")])
        bukum = atolye.acik_emirler(self.bukum)                                                               # kesilmiş: ara mamul, bu ÜS'de kesim emri
        ie2 = emir2.istasyon_emirleri.get(operasyon=self.kesim_op)
        self.assertEqual([s["malzeme"]["metin"] for s in bukum],                                                # ikisi de YOK → termin: emir2 önce
                         [f"Malzeme yok: kesilmiş yok — önce LAZER LAZER İSTASYONU ({ie2.no})",
                          f"Malzeme yok: kesilmiş yok — önce LAZER LAZER İSTASYONU ({self.ie_kesim.no})"])
        self.baslat()
        r = self.client.get(reverse("core:atolye_istasyon", args=[self.lazer.pk]))
        self.assertContains(r, "🟠 Kısmen hazır: 8 / 10 — eksik: ham profil (8 adetlik) — satınalma bekleniyor")
        self.assertContains(r, "⛔ Malzeme yok: ham profil yok — satınalma bekleniyor")
        self.assertContains(r, "malzeme-kismi")
        self.assertContains(r, "malzeme-yok")
        self.assertContains(r, "Yine de başlansın", count=1)                                                  # yalnız YOK kartında
        self.assertNotContains(r, "kalandan fazlası yazılamaz")                                                # kart notu kalktı (varyantlı emir yok)
        r = self.basla(emir2.istasyon_emirleri.get(operasyon=self.kesim_op))                                  # engel yok, yalnız uyarı
        self.assertTrue(any("iş başladı" in m for m in self.mesajlar(r)), self.mesajlar(r))

    def test_fason_siparisinden_bekleniyor_ve_fazla_yasak_notu(self):
        from core.models import Cari, FasonSiparisi, FasonSiparisRezerv, FasonSiparisSatiri
        cari = Cari.objects.create(kod="FSN", unvan="FASON LAZER", para_birimi="TRY")
        fs = FasonSiparisi.objects.create(yil=2026, sira=1, no="FS-2026-0001", cari=cari, tarih=date(2026, 1, 5), termin=date(2026, 1, 20),
                                          durum=FasonSiparisi.Durum.GONDERILDI)
        satir = FasonSiparisSatiri.objects.create(siparis=fs, stok=self.kesilmis, adet=D("20"), sira=10)
        FasonSiparisRezerv.objects.create(satir=satir, uretim_emri=self.emir, miktar=D("6"))
        self.baslat()
        r = self.client.get(reverse("core:atolye_istasyon", args=[self.bukum.pk]))
        self.assertContains(r, "📤 6 adet fason siparişinden bekleniyor (FS-2026-0001, termin 20.01.2026)")
        self.assertEqual(atolye.acik_emirler(self.bukum)[0]["fason_bekleyen"], ["6 adet fason siparişinden bekleniyor (FS-2026-0001, termin 20.01.2026)"])
        self.assertEqual(atolye.acik_emirler(self.lazer)[0]["fason_bekleyen"], [])                           # kesim emrinin girdisi profil: eşleşmez
        self.assertFalse(r.context["fazla_yasak_not"])
        self.assertEqual(atolye._varyant_on_eki(None), "")


class EkranTest(AtolyeTaban):
    def test_istasyon_kartlari_baslik_cikis_ve_rozetler(self):
        _istasyon("PAKET")                                                                                    # açık işi olmayan istasyon
        self.baslat()
        r = self.client.get(reverse("core:atolye") + "?sec=1")
        self.assertContains(r, "1 açık iş", count=2)                                                          # LAZER, BUKUM
        self.assertContains(r, "açık iş yok", count=1)
        self.assertContains(r, 'class="bos"', count=1)
        self.assertContains(r, self.p.ad_soyad)
        self.assertNotContains(r, self.u.username)                                                            # kullanıcı adı (TC) hiçbir yerde yok
        self.assertContains(r, reverse("logout"))
        self.assertContains(r, "Mesainiz açık; çıkış yapsanız da mesai açık kalır")
        r = self.client.get(reverse("core:atolye_istasyon", args=[self.lazer.pk]))
        self.assertContains(r, 'class="ie-no"')
        self.assertContains(r, '<span class="rozet-stok">Stok</span>')
        self.assertNotContains(r, "Stok için üretim")
        self.assertNotContains(r, self.u.username)
        self.basla(self.ie_kesim)
        r = self.client.get(reverse("core:atolye_istasyon", args=[self.lazer.pk]))
        self.assertContains(r, "Açık işiniz var; çıkış yapsanız da iş oturumu açık kalır")
        self.assertContains(r, 'placeholder="Kalan: 10"')
        self.assertNotContains(r, 'name="miktar" inputmode="numeric" pattern="[0-9]*" value=')              # sağlam adet ön dolu değil
        self.assertContains(r, 'name="fire"')
        self.assertContains(r, "atBitirOnay(this)")
        self.client.force_login(self.yon)                                                                     # yönetici izleme: kullanıcı adı yok
        r = self.client.get(reverse("core:atolye"))
        self.assertContains(r, "Yönetici")
        self.assertNotContains(r, "atyon")
        self.assertContains(r, "açık iş yok")


class BitirFireTest(AtolyeTaban):
    def test_bos_adet_reddi_ve_fire(self):
        self.baslat()
        self.basla(self.ie_kesim)
        o = self.acik_oturum(self.ie_kesim)
        r = self.bitir(o, "")
        self.assertTrue(any("Adet girin" in m for m in self.mesajlar(r)), self.mesajlar(r))
        r = self.bitir(o, "4", fire="x")
        self.assertTrue(any("Fire tam sayı olmalı" in m for m in self.mesajlar(r)), self.mesajlar(r))
        o.refresh_from_db()
        self.assertEqual(o.durum, "ACIK")
        self.assertFalse(OperasyonKaydi.objects.filter(istasyon_emri=self.ie_kesim).exists())
        r = self.bitir(o, "4", fire="2")
        self.assertTrue(any("4 sağlam + 2 fire yazıldı" in m for m in self.mesajlar(r)), self.mesajlar(r))
        k = OperasyonKaydi.objects.get(istasyon_emri=self.ie_kesim, silindi=False)
        self.assertEqual((k.hedef_cikti_miktari, k.gelen, k.durum), (D("6"), {str(self.kesilmis.pk): "4"}, "ONAYLI"))
        self.assertEqual(StokHareket.objects.get(operasyon_kaydi=k, tur="GIRIS").miktar, D("4"))            # stoğa yalnız sağlam
        self.assertEqual(StokHareket.objects.get(operasyon_kaydi=k, tur="CIKIS").miktar, D("3"))            # girdi: 6 adet = 3 çalıştırma × 1 profil
        c = OperasyonKaydiCikti.objects.get(kayit=k, ana_mi=True)
        self.assertEqual((c.beklenen_miktar, c.miktar), (D("6"), D("4")))                                  # fire = beklenen − gelen = 2
        self.ie_kesim.refresh_from_db()
        self.assertEqual(self.ie_kesim.tamamlanan, D("2"))                                                  # sağlam 4 adet = 2 çalıştırma
        o.refresh_from_db()
        self.assertEqual((o.durum, o.miktar, o.fire), ("BITTI", D("4"), D("2")))
        self.assertEqual(is_oturumu.ozet(IsOturumu.objects.get(pk=o.pk))["fire"], D("2"))
        self.assertContains(r, "Son kaydım")
        # firesiz bitir: eski mesaj, fire boş/0
        self.uret(self.ie_kesim, "2")
        o2 = IsOturumu.objects.filter(istasyon_emri=self.ie_kesim).latest("pk")
        self.assertEqual((o2.miktar, o2.fire), (D("2"), D("0")))
        # geri alma fire'lı kaydı da aynen sarar
        r = self.client.post(reverse("core:atolye_geri_al", args=[o2.kayit_id]), HTTP_X_REAL_IP="10.0.0.5", follow=True)
        r = self.client.post(reverse("core:atolye_geri_al", args=[k.pk]), HTTP_X_REAL_IP="10.0.0.5", follow=True)
        self.assertTrue(any("geri alındı" in m for m in self.mesajlar(r)), self.mesajlar(r))
        self.ie_kesim.refresh_from_db()
        self.assertEqual(self.ie_kesim.tamamlanan, D("0"))
        self.assertFalse(StokHareket.objects.filter(operasyon_kaydi=k, silindi=False).exists())
        # ofis raporu: Fire sütunu + Excel
        self.client.force_login(self.yon)
        r = self.client.get(reverse("core:is_oturumlari"))
        self.assertContains(r, '<th class="sag">Fire</th>')
        ws = load_workbook(BytesIO(self.client.get(reverse("core:is_oturumlari") + "?xlsx=1").content)).active
        self.assertEqual([c.value for c in ws[4]][13:16], ["Adet", "Fire", "Adet/saat"])
        satir = next(row for row in ws.iter_rows(min_row=5, values_only=True) if row[17] == k.no)                 # 17 = Kayıt No
        self.assertEqual((float(satir[13]), float(satir[14])), (4.0, 2.0))
