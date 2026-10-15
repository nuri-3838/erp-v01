"""Faturalar liste ekranı Excel çıktısı: FİLTRELENMİŞ listenin TAMAMI (sayfalama yok), ekrandaki kolonlar + TL karşılıkları + toplam satırları. Yalnız dizilim/biçim — hesaplar fatura_liste'de."""
from __future__ import annotations

import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.services import fatura_odeme

BASLIK_DOLGU = PatternFill("solid", fgColor="15294D")
KOLONLAR = ("Tarih", "Fatura No", "Cari", "Tip", "Matrah (TL)", "KDV (TL)", "Toplam (fatura PB)", "PB", "Kur", "Toplam (TL)", "Vade", "Ödeme Durumu", "Kalan", "Kalan PB", "PDF")
GENISLIK = (12, 22, 40, 14, 16, 14, 18, 6, 12, 16, 12, 16, 14, 9, 8)
PARA = "#,##0.00;[Red]-#,##0.00"


def _num(x):
    return None if x is None else float(x)


def liste_xlsx(satirlar, toplam, baslik="Faturalar", aciklama="") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = baslik[:31]
    ws["A1"] = baslik
    ws["A1"].font = Font(bold=True, size=13, color="15294D")
    ws["A2"] = f"{datetime.date.today():%d.%m.%Y}" + (f" · {aciklama}" if aciklama else "") + f" · {len(satirlar)} fatura · Tutarlar: Matrah ve KDV TL karşılığıdır (fatura kuruyla)."
    ws["A2"].font = Font(color="5B6678")
    for i, k in enumerate(KOLONLAR, start=1):
        c = ws.cell(row=4, column=i, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BASLIK_DOLGU
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(GENISLIK, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    r = 5
    for s in satirlar:
        od = s["odeme"]
        etiket = s["tip_etiket"] + (" (yalnız KDV)" if s["yalniz_kdv"] else "") + (" · Stok+Hesap" if s["karma"] else "")
        ws.append([s["tarih"], s["fatura_no"], s["cari_unvan"], etiket, _num(s["matrah_tl"]), _num(s["kdv_tl"]), _num(s["toplam"]), s["pb"], _num(s["kur"]), _num(s["toplam_tl"]),
                   s["vade"], "—" if s["taslak"] else (fatura_odeme.DURUM_ADLARI[od["durum"]] + (f" (vadesi geçti, {od['gecikme']} gün)" if od["gecikme"] > 0 else "") if od else "—"),
                   _num(od["kalan"]) if od and od["kalan"] else None, od["pb"] if od and od["kalan"] else "", "Var" if s["pdf_ek_pk"] else ("—" if s["taslak"] else "Yok")])
        for col in (5, 6, 7, 9, 10, 13):
            ws.cell(row=r, column=col).number_format = PARA
        for col in (1, 11):
            ws.cell(row=r, column=col).number_format = "DD.MM.YYYY"
        r += 1
    r += 1

    def satir(etiket, **hucreler):
        nonlocal r
        ws.cell(row=r, column=3, value=etiket).font = Font(bold=True)
        for col, deger in hucreler.items():
            c = ws.cell(row=r, column=int(col[1:]), value=deger)
            c.font = Font(bold=True)
            c.number_format = PARA
        r += 1

    t = toplam["tl"]
    satir("TOPLAM (TL karşılığı)", c5=float(t["matrah"]), c6=float(t["kdv"]), c10=float(t["toplam"]))
    for pb, d in toplam["doviz"]:
        satir(f"{pb} alt toplam (fatura para birimi)", c5=float(d["matrah"]), c6=float(d["kdv"]), c7=float(d["toplam"]))
    for pb, tutar in toplam["acik"]:
        satir(f"Açık bakiye ({pb})", c13=float(tutar))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
