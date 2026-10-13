"""CARİLER Faz 2: Excel çıktıları (cari listesi, cari ekstresi) ve mutabakat mektubu verisi. Yalnız dizilim/biçim — hesaplar cari_ozet / raporlar'da."""
from __future__ import annotations

import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.services import raporlar

SIFIR = Decimal("0")
ESIK = Decimal("0.005")
BASLIK_DOLGU = PatternFill("solid", fgColor="15294D")
LISTE_KOLONLAR = ("Cari Kodu", "Ünvan", "Kategori", "Şehir", "Ülke", "Cari PB", "Bakiye TL", "Yön", "Döviz Bakiyeleri", "Son Hareket", "Durum")
LISTE_GENISLIK = (14, 44, 24, 16, 14, 8, 16, 8, 30, 12, 8)


def yon_metni(net) -> str:
    """Bakiye yönü (ekstreyle aynı): pozitif = Borç (cari bize borçlu), negatif = Alacak (biz cariye borçluyuz)."""
    return "Borç" if net > 0 else ("Alacak" if net < 0 else "")


def doviz_metni(doviz) -> str:
    """[(pb, net)] → '1.234,56 USD Borç; 500,00 EUR Alacak'."""
    from core.templatetags.core_extras import tr_para
    return "; ".join(f"{tr_para(abs(n))} {pb} {yon_metni(n)}" for pb, n in doviz)


def _baslik(ws, satir, kolonlar, genislik):
    for i, k in enumerate(kolonlar, start=1):
        c = ws.cell(row=satir, column=i, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BASLIK_DOLGU
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(genislik, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{satir + 1}"          # (ws.cell() satırı oluşturup append'i kaydırırdı)


def liste_xlsx(satirlar, aciklama="") -> bytes:
    """Filtrelenmiş cari listesi (``cari_ozet.liste_getir`` satırları)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Cariler"
    ws["A1"] = "Cari Listesi"
    ws["A1"].font = Font(bold=True, size=13, color="15294D")
    ws["A2"] = f"{datetime.date.today():%d.%m.%Y}" + (f" · {aciklama}" if aciklama else "") + " · Bakiye: pozitif = Borç (cari bize borçlu), negatif = Alacak"
    ws["A2"].font = Font(color="5B6678")
    _baslik(ws, 4, LISTE_KOLONLAR, LISTE_GENISLIK)
    r = 5
    for s in satirlar:
        c = s["cari"]
        ws.append([c.kod, c.unvan, c.kategori.ad if c.kategori_id else "", c.sehir.ad if c.sehir_id else "", c.ulke.ad if c.ulke_id else "", c.para_birimi,
                   float(s["tl"]), yon_metni(s["tl"]), doviz_metni(s["doviz"]), s["son"], "Aktif" if c.aktif else "Pasif"])
        ws.cell(row=r, column=7).number_format = "#,##0.00;[Red]-#,##0.00"
        ws.cell(row=r, column=10).number_format = "DD.MM.YYYY"
        r += 1
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


EKSTRE_KOLONLAR = ("Tarih", "Fiş No", "Açıklama", "Borç (TL)", "Alacak (TL)", "Bakiye (TL)", "Yön", "PB", "Döviz Borç", "Döviz Alacak", "Döviz Bakiye")
EKSTRE_GENISLIK = (11, 11, 52, 15, 15, 16, 7, 6, 14, 14, 14)


def ekstre_xlsx(cari, ekstre) -> bytes:
    """Cari ekstresi (``raporlar.ekstre_devirli``): devir + dönem hareketleri + toplam. Sayılar sayı hücresidir (Excel'de toplanabilir)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ekstre"
    ws["A1"] = f"Cari Ekstresi — {cari.kod} {cari.unvan}"
    ws["A1"].font = Font(bold=True, size=13, color="15294D")
    ws["A2"] = f"Hesap: {ekstre.hesap_kodu} · Dönem: {ekstre.baslangic:%d.%m.%Y} – {ekstre.bitis:%d.%m.%Y} · Bakiye: pozitif = Borç (cari bize borçlu), negatif = Alacak"
    ws["A2"].font = Font(color="5B6678")
    _baslik(ws, 4, EKSTRE_KOLONLAR, EKSTRE_GENISLIK)
    ws.append([None, None, "Devir (açılış bakiyesi)", None, None, float(ekstre.acilis), yon_metni(ekstre.acilis)])
    ws.cell(row=5, column=3).font = Font(italic=True)
    r = 6
    for s in ekstre.satirlar:
        ws.append([s.tarih, f"{s.fis_yil}/{s.fis_no}", (s.fis_aciklama or "") + (f" · {s.satir_aciklama}" if s.satir_aciklama else ""),
                   float(s.borc) or None, float(s.alacak) or None, float(s.yur_bakiye), yon_metni(s.yur_bakiye), s.pb,
                   float(s.dvz_borc) or None, float(s.dvz_alacak) or None, float(s.dvz_yur_bakiye)])
        ws.cell(row=r, column=1).number_format = "DD.MM.YYYY"
        for col in (4, 5, 6, 9, 10, 11):
            ws.cell(row=r, column=col).number_format = "#,##0.00;[Red]-#,##0.00"
        r += 1
    ws.append([None, None, "KAPANIŞ", float(ekstre.toplam_borc), float(ekstre.toplam_alacak), float(ekstre.kapanis_bakiye), yon_metni(ekstre.kapanis_bakiye)])
    for col in range(1, 8):
        ws.cell(row=r, column=col).font = Font(bold=True)
    for col in (4, 5, 6):
        ws.cell(row=r, column=col).number_format = "#,##0.00;[Red]-#,##0.00"
    ws.cell(row=5, column=6).number_format = "#,##0.00;[Red]-#,##0.00"
    for pb, t in ekstre.dvz_toplamlar.items():
        if pb != "TRY":
            r += 1
            ws.append([None, None, f"Döviz kapanış ({pb})", None, None, None, None, pb, float(t["borc"]), float(t["alacak"]), float(t["bakiye"])])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def mutabakat_baglami(cari, tarih: datetime.date) -> dict:
    """Belirli tarih (dahil) itibarıyla bakiye: TL net + döviz havuzları. Yön: net > 0 → bizim ALACAĞIMIZ var (cari borçlu); net < 0 → BORCUMUZ var."""
    if not cari.muhasebe_kodu:
        return {"tl": SIFIR, "doviz": [], "hesap": ""}
    tl, dvz = raporlar._devir(cari.muhasebe_kodu, tarih + datetime.timedelta(days=1))
    doviz = sorted((pb, n) for pb, n in dvz.items() if pb != "TRY" and abs(n) >= ESIK)
    return {"tl": tl, "doviz": doviz, "hesap": cari.muhasebe_kodu}
