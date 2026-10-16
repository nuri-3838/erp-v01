"""STOK SAYIMI — Excel indir / yükle / rapor çıktıları (openpyxl).

Giriş şablonu (sayım girenler): A "Satır Kimliği" (gizli), B stok kodu, C ürün, D varyant, E birim, [F sistem (kesme) — KÖR belgede YOK], Sayılan. Yükleme: dosyanın 1. satırındaki sayım no belgeyle
aynı olmalı; "Sayılan" dolu satırlar TOPLAMSAL giriş olur (kaynak EXCEL); aynı dosya aynı belgeye iki kez yüklenemez (içerik özeti). Metin hücreleri formül enjeksiyonuna karşı güvenli yazılır.
Sınırlar: yalnız .xlsx, en çok 2 MB, en çok 20.000 satır."""
from __future__ import annotations

import hashlib
import io
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.sayi import SayiHatasi, parse_tr

MAKS_BOYUT = 2 * 1024 * 1024
MAKS_SATIR = 20000
BASLIK_KIMLIK, BASLIK_SAYILAN = "Satır Kimliği", "Sayılan"
KALIN = Font(bold=True)
BASLIK_DOLGU = PatternFill("solid", fgColor="DDE5F3")


class ExcelHatasi(ValueError):
    pass


def guvenli(deger):
    """Metni hücreye güvenle yazar: formül gibi başlayanlar (=, +, -, @) başına kesme alır."""
    if isinstance(deger, str) and deger[:1] in ("=", "+", "-", "@"):
        return "'" + deger
    return deger


def _sayi(x):
    return None if x is None else float(x)


def _bayt(wb) -> bytes:
    cikti = io.BytesIO()
    wb.save(cikti)
    return cikti.getvalue()


def _basliklar(ws, satir, basliklar):
    for i, b in enumerate(basliklar, 1):
        h = ws.cell(row=satir, column=i, value=b)
        h.font, h.fill = KALIN, BASLIK_DOLGU
        h.alignment = Alignment(wrap_text=True, vertical="center")


def _genislik(ws, genislikler):
    for i, g in enumerate(genislikler, 1):
        ws.column_dimensions[get_column_letter(i)].width = g


# ----------------------------------------------------------------------------------------------------------------------- giriş şablonu
def giris_sablonu(sayim, satirlar, *, sistem_goster) -> bytes:
    """``satirlar``: [{"satir": StokSayimSatiri, "etiket": varyant etiketi, "birim": ad}]. ``sistem_goster`` False ise (KÖR belge) sistem sütunu dosyaya HİÇ yazılmaz."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sayım"
    ws["A1"], ws["B1"] = "Sayım No", sayim.no
    ws["A2"], ws["B2"] = "Depo", f"{sayim.depo.kod} · {sayim.depo.ad}"
    ws["A3"], ws["B3"] = "Sayım tarihi", sayim.tarih.strftime("%d.%m.%Y")
    ws["A4"], ws["B4"] = "Not", "Yalnızca \"Sayılan\" sütununu doldurun. Girişler mevcut sayımlara EKLENİR (toplamsal). Boş bırakılan satırlar atlanır."
    for r in range(1, 5):
        ws.cell(row=r, column=1).font = KALIN
    basliklar = [BASLIK_KIMLIK, "Stok Kodu", "Ürün", "Varyant", "Birim"] + (["Sistem (kesme anı)"] if sistem_goster else []) + [BASLIK_SAYILAN]
    _basliklar(ws, 6, basliklar)
    sayilan_kolon = len(basliklar)
    for i, r in enumerate(satirlar, start=7):
        s = r["satir"]
        ws.cell(row=i, column=1, value=s.pk)
        ws.cell(row=i, column=2, value=guvenli(s.stok.kod))
        ws.cell(row=i, column=3, value=guvenli(s.stok.ad))
        ws.cell(row=i, column=4, value=guvenli(r.get("etiket") or ""))
        ws.cell(row=i, column=5, value=guvenli(r.get("birim") or ""))
        if sistem_goster:
            ws.cell(row=i, column=6, value=_sayi(s.sistem_snapshot))
    son = 6 + len(satirlar)
    if satirlar:
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True, errorTitle="Geçersiz miktar", error="Sayılan miktar 0 ya da daha büyük bir sayı olmalı.")
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(sayilan_kolon)}7:{get_column_letter(sayilan_kolon)}{son}")
    ws.column_dimensions["A"].hidden = True
    _genislik(ws, [12, 16, 48, 18, 8] + ([16] if sistem_goster else []) + [14])
    ws.freeze_panes = "A7"
    return _bayt(wb)


# ----------------------------------------------------------------------------------------------------------------------- yükleme
def _ondalik(deger, satir_no):
    if isinstance(deger, bool):
        raise ExcelHatasi("geçerli bir sayı değil")
    if isinstance(deger, (int, float, Decimal)):
        try:
            d = Decimal(str(deger))
        except InvalidOperation:
            raise ExcelHatasi("geçerli bir sayı değil")
    else:
        try:
            d = parse_tr(str(deger).strip())
        except SayiHatasi:
            raise ExcelHatasi("geçerli bir sayı değil")
    if not d.is_finite():
        raise ExcelHatasi("geçerli bir sayı değil")
    if d < 0:
        raise ExcelHatasi("negatif olamaz")
    return d


def oku(icerik: bytes, sayim, satir_haritasi) -> dict:
    """Yüklenen dosyayı çözer. ``satir_haritasi``: {satir id: StokSayimSatiri} (belgenin satırları). Döner: {"satirlar": [(satir_id, kod, miktar)], "hatalar": [(excel satırı, mesaj)], "ozet": içerik özeti}."""
    if len(icerik) > MAKS_BOYUT:
        raise ExcelHatasi("Dosya 2 MB'den büyük.")
    if icerik[:2] != b"PK":
        raise ExcelHatasi("Yalnızca .xlsx dosyası yüklenebilir.")
    try:
        wb = load_workbook(io.BytesIO(icerik), read_only=True, data_only=True)
    except Exception:                                                                    # noqa: BLE001 — bozuk/şifreli dosya
        raise ExcelHatasi("Dosya okunamadı (geçerli bir .xlsx değil).")
    ws = wb.worksheets[0]
    satirlar = list(ws.iter_rows(min_row=1, max_row=MAKS_SATIR + 10, values_only=True))
    if len(satirlar) > MAKS_SATIR + 6:
        raise ExcelHatasi("Dosyada çok fazla satır var (en çok 20.000).")
    if not satirlar or str((satirlar[0][1] if len(satirlar[0]) > 1 else "") or "").strip() != sayim.no:
        raise ExcelHatasi(f"Bu dosya {sayim.no} belgesine ait değil (1. satırdaki sayım no uyuşmuyor).")
    baslik_no = next((i for i, r in enumerate(satirlar[:15]) if r and str(r[0] or "").strip() == BASLIK_KIMLIK), None)
    if baslik_no is None:
        raise ExcelHatasi(f"\"{BASLIK_KIMLIK}\" başlığı bulunamadı; indirdiğiniz şablonu kullanın.")
    basliklar = [str(h or "").strip() for h in satirlar[baslik_no]]
    if BASLIK_SAYILAN not in basliklar:
        raise ExcelHatasi(f"\"{BASLIK_SAYILAN}\" sütunu bulunamadı.")
    sayilan_kolon = basliklar.index(BASLIK_SAYILAN)
    sonuc, hatalar, gorulen = [], [], set()
    for no, r in enumerate(satirlar[baslik_no + 1:], start=baslik_no + 2):
        if not r or all(c in (None, "") for c in r):
            continue
        ham = r[sayilan_kolon] if len(r) > sayilan_kolon else None
        if ham in (None, "") or (isinstance(ham, str) and not ham.strip()):
            continue
        try:
            satir_id = int(r[0])
        except (TypeError, ValueError):
            hatalar.append((no, "satır kimliği okunamadı (şablonun A sütununu değiştirmeyin)"))
            continue
        satir = satir_haritasi.get(satir_id)
        if satir is None:
            hatalar.append((no, "bu satır sayım belgesinde yok"))
            continue
        if satir_id in gorulen:
            hatalar.append((no, f"{satir.stok.kod}: aynı satır dosyada birden fazla"))
            continue
        gorulen.add(satir_id)
        try:
            sonuc.append((satir_id, satir.stok.kod, _ondalik(ham, no)))
        except ExcelHatasi as e:
            hatalar.append((no, f"{satir.stok.kod}: {e}"))
    ozet = hashlib.sha256(",".join(f"{a}:{m}" for a, _k, m in sorted(sonuc)).encode()).hexdigest()[:10]
    return {"satirlar": sonuc, "hatalar": hatalar, "ozet": ozet}


# ----------------------------------------------------------------------------------------------------------------------- yönetici çıktıları
def fark_raporu(sayim, satirlar, ozet) -> bytes:
    """Fark / karar ekranının Excel'i (yönetici)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Fark"
    ws["A1"] = f"{sayim.no} — {sayim.depo.kod} {sayim.depo.ad} — sayım tarihi {sayim.tarih:%d.%m.%Y} — durum: {sayim.get_durum_display()}"
    ws["A1"].font = KALIN
    ws["A2"] = (f"Sayılan satır {ozet.get('sayilan_adet', 0)} · fazla tutarı {ozet.get('fazla_tutar', 0)} · eksik tutarı {ozet.get('eksik_tutar', 0)} · "
                f"Düzelt: {ozet.get('duzelt_adet', 0)} satır")
    _basliklar(ws, 4, ["Stok Kodu", "Ürün", "Varyant", "Sistem (kesme)", "Sistem (referans)", "Sayılan", "Fark", "Birim maliyet (TL)", "Fark tutarı (TL)", "Karar", "Not", "Maliyet uyarısı"])
    for i, d in enumerate(satirlar, start=5):
        ws.append([guvenli(d["stok"].kod), guvenli(d["stok"].ad), guvenli(d.get("etiket") or ""), _sayi(d["snapshot"]), _sayi(d["referans"]), _sayi(d["sayilan"]), _sayi(d["fark"]),
                   _sayi(d["birim_maliyet"]), _sayi(d["tutar"]), d["satir"].get_karar_display(), guvenli(d["satir"].karar_notu), guvenli(d["neden"])])
    _genislik(ws, [16, 44, 18, 14, 14, 12, 10, 16, 16, 12, 28, 30])
    ws.freeze_panes = "A5"
    return _bayt(wb)


def gecmis_raporu(satirlar) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sayım geçmişi"
    _basliklar(ws, 1, ["Sayım No", "Tarih", "Depo", "Kapsam", "Durum", "Kör", "Satır", "Sayılan satır", "Düzelt", "Araştır", "Fazla tutarı (TL)", "Eksik tutarı (TL)", "Fiş"])
    for r in satirlar:
        s = r["sayim"]
        ws.append([s.no, s.tarih, f"{s.depo.kod} {s.depo.ad}", s.kategori.ad if s.kategori_id else "Tümü", s.get_durum_display(), "Evet" if s.kor else "", r["satir"], r["sayilan"], r["duzelt"],
                   r["arastir"], _sayi(r["fazla_tutar"]), _sayi(r["eksik_tutar"]), f"{r['fis'].yil}/{r['fis'].fis_no}" if r["fis"] is not None else ""])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "DD.MM.YYYY"
    _genislik(ws, [16, 12, 26, 22, 14, 6, 8, 12, 8, 8, 16, 16, 12])
    ws.freeze_panes = "A2"
    return _bayt(wb)


def stok_gecmisi_raporu(stok, satirlar) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Stok fark geçmişi"
    ws["A1"] = f"{stok.kod} — {stok.ad}"
    ws["A1"].font = KALIN
    _basliklar(ws, 3, ["Sayım No", "Sayım tarihi", "Depo", "Varyant", "Belge durumu", "Sistem (referans)", "Sayılan", "Fark", "Karar", "Düzeltme tutarı (TL)"])
    from core.services.varyant import varyant_kisa
    for r in satirlar:
        s = r["sayim"]
        ws.append([s.no, s.tarih, s.depo.kod, varyant_kisa(r["varyant"]) if r["varyant"] is not None else "", s.get_durum_display(), _sayi(r["referans"]), _sayi(r["sayilan"]), _sayi(r["fark"]),
                   r["satir"].get_karar_display(), _sayi(r["tutar"])])
    for row in ws.iter_rows(min_row=4, min_col=2, max_col=2):
        row[0].number_format = "DD.MM.YYYY"
    _genislik(ws, [16, 12, 10, 18, 14, 16, 12, 10, 12, 18])
    return _bayt(wb)
