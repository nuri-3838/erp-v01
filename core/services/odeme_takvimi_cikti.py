"""Ödeme Takvimi Excel çıktısı (yalnız dizilim/biçim — hesaplar ``odeme_takvimi``'nde). 1. sayfa: ekrandaki gruplu tablo (düz liste); 2. sayfa: özet kartları + grup toplamları + mutabakat."""
from __future__ import annotations

import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASLIK_DOLGU = PatternFill("solid", fgColor="15294D")
KIRMIZI = PatternFill("solid", fgColor="FDECEF")
GRUP_DOLGU = PatternFill("solid", fgColor="EEF3FB")
PARA = "#,##0.00;[Red]-#,##0.00"
KOLONLAR = ("Grup", "Tarih", "Yön", "Kaynak", "Cari / Açıklama", "Belge No", "Döviz Tutarı", "PB", "TL", "Kural ile", "Not")
GENISLIK = (22, 12, 8, 22, 54, 24, 16, 6, 16, 10, 46)


def _num(x):
    return None if x is None else float(x)


def takvim_xlsx(veri, grup_sonucu, mutabakat, *, bas, bit, aciklama="") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ödeme Takvimi"
    ws["A1"] = "Ödeme Takvimi"
    ws["A1"].font = Font(bold=True, size=13, color="15294D")
    kur = " · ".join(f"{pb} {k:,.4f}" for pb, k in veri["kurlar"].items())
    ws["A2"] = f"{veri['bugun']:%d.%m.%Y} · {bas:%d.%m.%Y} – {bit:%d.%m.%Y}" + (f" · {aciklama}" if aciklama else "") + f" · TL karşılıkları bugünkü kurla ({kur}, {veri['kur_tarihi']:%d.%m.%Y})"
    ws["A2"].font = Font(color="5B6678")
    for i, k in enumerate(KOLONLAR, start=1):
        c = ws.cell(row=4, column=i, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BASLIK_DOLGU
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(GENISLIK, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    r = 5

    def blok(ad, rows, kum=None, negatif=False):
        nonlocal r
        ws.cell(row=r, column=1, value=ad).font = Font(bold=True)
        cikis = sum((x["tl"] for x in rows if x["yon"] == "CIKIS"), 0)
        giris = sum((x["tl"] for x in rows if x["yon"] == "GIRIS"), 0)
        ws.cell(row=r, column=5, value=f"Çıkış {float(cikis):,.2f} · Giriş {float(giris):,.2f}" + (f" · Kümülatif {float(kum):,.2f}" if kum is not None else "")).font = Font(bold=True)
        for col in range(1, 12):
            ws.cell(row=r, column=col).fill = KIRMIZI if negatif else GRUP_DOLGU
        r += 1
        for x in rows:
            ws.append([ad, x["tarih"], "Çıkış" if x["yon"] == "CIKIS" else "Giriş", x["kaynak_ad"], f"{x['baslik']}" + (f" — {x['cari_ad']}" if x["cari_ad"] and x["cari_ad"] not in x["baslik"] else "") + (f" · {x['aciklama']}" if x["aciklama"] else ""),
                       x["belge_no"], _num(x["tutar"]), x["pb"], _num(x["tl"]), "tahmini" if x.get("tahmini") else ("kural ile" if x["kural"] else ""), "TL karşılığı için kur yok" if x["tl_yok"] else ""])
            ws.cell(row=r, column=2).number_format = "DD.MM.YYYY"
            for col in (7, 9):
                ws.cell(row=r, column=col).number_format = PARA
            r += 1

    for g in grup_sonucu["gruplar"]:
        blok(g["ad"], g["satirlar"], g["kumulatif"], g["negatif"])
    if grup_sonucu["vadesiz"]:
        blok("Vadesiz", grup_sonucu["vadesiz"]["satirlar"])

    ws2 = wb.create_sheet("Özet ve Mutabakat")
    ws2.column_dimensions["A"].width = 52
    for col in "BCDE":
        ws2.column_dimensions[col].width = 20
    ws2.column_dimensions["F"].width = 52
    ws2["A1"] = "Özet"
    ws2["A1"].font = Font(bold=True, size=13, color="15294D")
    satirlar = [("Bugünkü banka + kasa (TL)", veri["nakit_tl"]), ("Dönem çıkışı (gecikmiş dahil, vadesiz hariç)", grup_sonucu["cikis"]), ("  ‣ gecikmiş çıkış", grup_sonucu["gecikmis_cikis"]),
                ("Dönem girişi", grup_sonucu["giris"]), ("Dönem sonu tahmini bakiye", grup_sonucu["donem_sonu"]), ("Net ihtiyaç (en düşük kümülatif bakiye)", grup_sonucu["net_ihtiyac"]),
                ("Vadesiz çıkış (kümülatife dahil değil)", grup_sonucu["vadesiz"]["cikis"] if grup_sonucu["vadesiz"] else 0)]
    for i, (ad, tutar) in enumerate(satirlar, start=2):
        ws2.cell(row=i, column=1, value=ad)
        c = ws2.cell(row=i, column=2, value=_num(tutar))
        c.number_format = PARA
    rr = len(satirlar) + 4
    ws2.cell(row=rr - 1, column=1, value="Grup toplamları").font = Font(bold=True)
    for j, k in enumerate(("Grup", "Çıkış", "Giriş", "Net", "Kümülatif"), start=1):
        c = ws2.cell(row=rr, column=j, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BASLIK_DOLGU
    for g in grup_sonucu["gruplar"]:
        rr += 1
        ws2.cell(row=rr, column=1, value=g["ad"])
        for j, k in enumerate(("cikis", "giris", "net", "kumulatif"), start=2):
            c = ws2.cell(row=rr, column=j, value=_num(g[k]))
            c.number_format = PARA
    rr += 3
    ws2.cell(row=rr - 1, column=1, value="Mutabakat (takvim ↔ defter bakiyeleri, TL)").font = Font(bold=True)
    for j, k in enumerate(("Hesap / kalem", "Defter", "Takvim", "Fark", "", "Not"), start=1):
        c = ws2.cell(row=rr, column=j, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BASLIK_DOLGU
    for b in mutabakat:
        rr += 1
        ws2.cell(row=rr, column=1, value=b["ad"]).font = Font(bold=True)
        for j, k in enumerate(("defter", "takvim", "fark"), start=2):
            c = ws2.cell(row=rr, column=j, value=_num(b[k]))
            c.number_format = PARA
            c.font = Font(bold=True)
        for k in b["farkli"]:
            rr += 1
            ws2.cell(row=rr, column=1, value="   " + k["ad"])
            for j, key in enumerate(("defter", "takvim", "fark"), start=2):
                c = ws2.cell(row=rr, column=j, value=_num(k[key]))
                c.number_format = PARA
            ws2.cell(row=rr, column=6, value=k["not"])
    cikti = BytesIO()
    wb.save(cikti)
    return cikti.getvalue()
