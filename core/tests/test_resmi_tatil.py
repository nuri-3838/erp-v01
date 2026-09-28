"""İNSAN KAYNAKLARI > Resmî Tatiller — seed (2026/2027 sabit tatiller), model kısıtı, servis
CRUD, yoklama/aylık özet entegrasyonu (resmi_tatil/tatil_calisma sayaçları, izin/Pazar ile
etkileşim), izin gün önerisinin (pazarsiz_gun) tatili düşmesi, puantaj dökümü yeni sütunları
ve yetki (ayrı "resmi_tatil" ekranı). Kapsam dışı: bordro/ücret hesabı, arife/yarım gün."""
import importlib
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.apps import apps as canli_apps
from django.contrib.auth.models import User
from django.db import IntegrityError, transaction
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.metin import buyuk_harf_tr
from core.models import EkranYetki, PersonelDevam, ResmiTatil
from core.services.personel_devam import aylik_ozet
from core.services.personel_dokum import aylik_dokum, dokum_xlsx
from core.services.personel_izin import izin_ekle, pazarsiz_gun
from core.services.resmi_tatil import (
    ResmiTatilHatasi, dini_bayram_var_mi, gunun_tatili, tatil_ekle, tatil_gunleri_seti,
    tatil_guncelle, tatil_sil,
)
from core.tests.ik_yardimci import personel_kur

AY_SONU = date(2026, 9, 30)          # Eylül 2026'da hiç sabit resmî tatil YOK (regresyon ayı)


def _kayit(p, tarih, durum="GELDI"):
    return PersonelDevam.objects.create(personel=p, tarih=tarih, durum=durum)


class SeedMigrasyonuTest(TestCase):
    def _migrasyon(self):
        return importlib.import_module("core.migrations.0141_resmi_tatil_seed")

    def test_2026_2027_sabit_tatiller_seed_edilmis(self):
        self.assertEqual(ResmiTatil.objects.filter(silindi=False).count(), 14)
        self.assertTrue(ResmiTatil.objects.filter(
            tarih=date(2026, 10, 29), ad=buyuk_harf_tr("29 Ekim Cumhuriyet Bayramı")).exists())
        self.assertTrue(ResmiTatil.objects.filter(tarih=date(2027, 1, 1)).exists())
        self.assertFalse(ResmiTatil.objects.filter(tarih=date(2025, 10, 29)).exists())

    def test_ileri_idempotent(self):
        migrasyon = self._migrasyon()
        onceki = ResmiTatil.objects.filter(silindi=False).count()
        migrasyon._ileri(canli_apps, None)
        self.assertEqual(ResmiTatil.objects.filter(silindi=False).count(), onceki)

    def test_geri_alma_sabit_tatilleri_siler(self):
        migrasyon = self._migrasyon()
        migrasyon._geri_al(canli_apps, None)
        self.assertEqual(ResmiTatil.objects.filter(tarih__year__in=(2026, 2027)).count(), 0)


class ResmiTatilDBKisitTest(TestCase):
    def test_ayni_tarih_iki_aktif_kayit_yasak(self):
        ResmiTatil.objects.create(tarih=date(2028, 3, 3), ad="TEST")
        with self.assertRaises(IntegrityError), transaction.atomic():
            ResmiTatil.objects.create(tarih=date(2028, 3, 3), ad="TEST2")

    def test_silinmis_kayit_benzersizligi_engellemez(self):
        t = ResmiTatil.objects.create(tarih=date(2028, 3, 3), ad="TEST")
        t.silindi = True
        t.save(update_fields=["silindi"])
        ResmiTatil.objects.create(tarih=date(2028, 3, 3), ad="TEST2")


class TatilServisTest(TestCase):
    def test_ekle_ad_buyuk_harf_ve_benzersizlik(self):
        t = tatil_ekle(tarih=date(2028, 3, 3), ad="test bayramı")
        self.assertEqual(t.ad, buyuk_harf_tr("test bayramı"))
        with self.assertRaises(ResmiTatilHatasi):
            tatil_ekle(tarih=date(2028, 3, 3), ad="başka")

    def test_bos_ad_hata(self):
        with self.assertRaises(ResmiTatilHatasi):
            tatil_ekle(tarih=date(2028, 3, 3), ad="   ")

    def test_guncelle_ve_sil(self):
        t = tatil_ekle(tarih=date(2028, 3, 3), ad="test")
        tatil_guncelle(t, tarih=date(2028, 3, 4), ad="test2")
        t.refresh_from_db()
        self.assertEqual((t.tarih, t.ad), (date(2028, 3, 4), "TEST2"))
        tatil_sil(t)
        t.refresh_from_db()
        self.assertTrue(t.silindi)
        tatil_sil(t)                                       # ikinci çağrı hata vermemeli

    def test_silinmis_kayit_duzenlenemez(self):
        t = tatil_ekle(tarih=date(2028, 3, 3), ad="test")
        tatil_sil(t)
        with self.assertRaises(ResmiTatilHatasi):
            tatil_guncelle(t, tarih=date(2028, 3, 3), ad="test")

    def test_dini_bayram_var_mi_sabit_tatiller_sayilmaz(self):
        self.assertFalse(dini_bayram_var_mi(2026))          # yalnız sabit tatiller seed edildi
        tatil_ekle(tarih=date(2026, 3, 20), ad="Ramazan Bayramı")
        self.assertTrue(dini_bayram_var_mi(2026))
        self.assertFalse(dini_bayram_var_mi(2027))           # 2027 hâlâ eksik

    def test_gunun_tatili(self):
        self.assertIsNotNone(gunun_tatili(date(2026, 10, 29)))
        self.assertIsNone(gunun_tatili(date(2026, 10, 28)))

    def test_tatil_gunleri_seti(self):
        s = tatil_gunleri_seti(date(2026, 1, 1), date(2026, 12, 31))
        self.assertEqual(len(s), 7)
        self.assertIn(date(2026, 5, 1), s)


class TatilYoklamaEntegrasyonuTest(TestCase):
    def test_tatil_gunu_kayitsiz_girilmemis_sayilmaz(self):
        personel_kur(ise_giris_tarihi=date(2020, 1, 1))
        tatil_ekle(tarih=date(2026, 9, 15), ad="test tatil")           # Salı
        (s,) = aylik_ozet(2026, 9, bugun=AY_SONU)
        self.assertEqual(s.resmi_tatil, 1)
        self.assertEqual(s.girilmemis, 26 - 1)                          # 30 gün - 4 Pazar = 26

    def test_pazar_ve_tatil_ayni_gunde_yalniz_tatil_calisma(self):
        p = personel_kur(ise_giris_tarihi=date(2020, 1, 1))
        tatil_ekle(tarih=date(2026, 9, 6), ad="test tatil")             # Pazar
        _kayit(p, date(2026, 9, 6), "GELDI")
        (s,) = aylik_ozet(2026, 9, bugun=AY_SONU)
        self.assertEqual((s.tatil_calisma, s.pazar_calisma, s.geldi), (1, 0, 1))

    def test_tatil_yarim_gun_kaydi_da_tatil_calisma(self):
        p = personel_kur(ise_giris_tarihi=date(2020, 1, 1))
        tatil_ekle(tarih=date(2026, 9, 15), ad="test tatil")
        _kayit(p, date(2026, 9, 15), "YARIM_GUN")
        (s,) = aylik_ozet(2026, 9, bugun=AY_SONU)
        self.assertEqual((s.tatil_calisma, s.yarim, s.resmi_tatil), (1, 1, 0))

    def test_tatil_izin_araligindaysa_izin_sayilmaz_tatil_sayilir(self):
        p = personel_kur(ise_giris_tarihi=date(2020, 1, 1))
        tatil_ekle(tarih=date(2026, 9, 8), ad="test tatil")             # Salı, Pzt(7)-Çar(9) aralığında
        izin_ekle(p, tur="YILLIK", baslangic=date(2026, 9, 7), bitis=date(2026, 9, 9))
        (s,) = aylik_ozet(2026, 9, bugun=AY_SONU)
        self.assertEqual(s.resmi_tatil, 1)
        self.assertEqual((s.izinli, s.yillik), (2, 2))                  # 3 gün izin - 1 tatil günü

    def test_tatil_gunu_gelmedi_kaydi_hala_gelmedi_sayilir(self):
        p = personel_kur(ise_giris_tarihi=date(2020, 1, 1))
        tatil_ekle(tarih=date(2026, 9, 15), ad="test tatil")
        _kayit(p, date(2026, 9, 15), "GELMEDI")
        (s,) = aylik_ozet(2026, 9, bugun=AY_SONU)
        self.assertEqual((s.gelmedi, s.resmi_tatil, s.tatil_calisma), (1, 0, 0))

    def test_tatilsiz_ay_regresyon_degismedi(self):
        p = personel_kur(ise_giris_tarihi=date(2020, 1, 1))
        _kayit(p, date(2026, 9, 1), "GELDI")
        (s,) = aylik_ozet(2026, 9, bugun=AY_SONU)
        self.assertEqual((s.geldi, s.resmi_tatil, s.tatil_calisma, s.girilmemis), (1, 0, 0, 25))


class PazarsizGunTatilTest(TestCase):
    def test_tatil_gunu_dusulur(self):
        tatil_ekle(tarih=date(2026, 9, 15), ad="test tatil")            # Salı
        self.assertEqual(pazarsiz_gun(date(2026, 9, 14), date(2026, 9, 16)), Decimal(2))

    def test_pazar_ve_tatil_birlikte_dusulur(self):
        tatil_ekle(tarih=date(2026, 9, 8), ad="test tatil")             # Salı
        self.assertEqual(pazarsiz_gun(date(2026, 9, 6), date(2026, 9, 9)), Decimal(2))  # 4 gün - Pazar - tatil

    def test_tatilsiz_aralik_degismedi(self):
        self.assertEqual(pazarsiz_gun(date(2026, 9, 14), date(2026, 9, 16)), Decimal(3))


class DokumTatilSutunTest(TestCase):
    def test_yeni_sutunlar_dogru_sirada_ve_degerde(self):
        p = personel_kur(ad="test", soyad="tatil", ise_giris_tarihi=date(2020, 1, 1))
        tatil_ekle(tarih=date(2026, 9, 15), ad="test tatil")
        _kayit(p, date(2026, 9, 15), "GELDI")

        xbytes = dokum_xlsx(2026, 9, bugun=AY_SONU)
        wb = load_workbook(BytesIO(xbytes))
        ws = wb.active
        baslik = [ws.cell(row=5, column=i).value for i in range(1, 20)]
        self.assertEqual(baslik[8], "Pazar Çalışması")
        self.assertEqual(baslik[9], "Resmî Tatil")
        self.assertEqual(baslik[10], "Tatil Çalışması")
        self.assertEqual(baslik[11], "Gelmedi (Devamsız)")
        self.assertEqual(baslik[-1], "Açıklama")

        self.assertEqual(ws.cell(row=6, column=10).value, 0)            # Resmî Tatil
        self.assertEqual(ws.cell(row=6, column=11).value, 1)            # Tatil Çalışması

        toplam_satiri = 7
        self.assertEqual(ws.cell(row=toplam_satiri, column=1).value, "TOPLAM")
        self.assertEqual(ws.cell(row=toplam_satiri, column=11).value, 1)


class ResmiTatilViewTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.yon = User.objects.create_superuser("rtyon", password="x")
        cls.yetkili = User.objects.create_user("rtyetkili", password="x")
        EkranYetki.objects.create(kullanici=cls.yetkili, ekran_kod="resmi_tatil")
        cls.bos = User.objects.create_user("rtbos", password="x")

    def test_anonim_302_yetkisiz_403(self):
        urller = [reverse("core:resmi_tatiller"), reverse("core:resmi_tatil_ekle")]
        for u in urller:
            self.assertEqual(self.client.get(u).status_code, 302, u)
        self.client.force_login(self.bos)
        for u in urller:
            self.assertEqual(self.client.get(u).status_code, 403, u)

    def test_yetkili_liste_gorur_dini_bayram_uyarisi_ve_ekler(self):
        self.client.force_login(self.yetkili)
        r = self.client.get(reverse("core:resmi_tatiller"), {"yil": 2026})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Ramazan/Kurban Bayramı tanımlı değil")
        r2 = self.client.post(
            reverse("core:resmi_tatil_ekle"), {"tarih": "2028-03-03", "ad": "test bayram"})
        self.assertRedirects(r2, reverse("core:resmi_tatiller") + "?yil=2028")
        self.assertTrue(ResmiTatil.objects.filter(tarih=date(2028, 3, 3)).exists())

    def test_dini_bayram_eklenince_uyari_kaybolur(self):
        tatil_ekle(tarih=date(2026, 3, 20), ad="Ramazan Bayramı")
        self.client.force_login(self.yetkili)
        r = self.client.get(reverse("core:resmi_tatiller"), {"yil": 2026})
        self.assertNotContains(r, "tanımlı değil")

    def test_yonetici_de_erisir(self):
        self.client.force_login(self.yon)
        self.assertEqual(self.client.get(reverse("core:resmi_tatiller")).status_code, 200)
