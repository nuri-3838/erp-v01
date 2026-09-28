"""İNSAN KAYNAKLARI > Aylık Muhasebe Dökümü (Excel) — YALNIZ SAYIM + ÜCRET BİLGİSİ.

Bordro/brüt/SGK günü/vergi/prim hesabı, mesai giriş/çıkış YOKTUR. Ayın personel bazlı devam
sayımını (core.services.personel_devam.aylik_ozet — resmî tatil/tatil çalışması dahil) o
ayın sonu — ya da ay içinde işten çıkıldıysa çıkış günü — itibarıyla geçerli ücret bilgisiyle
(core.services.personel_ucret) bir araya getirir. Hiçbir değer hesaplanıp saklanmaz; her
indirmede canlı üretilir.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from io import BytesIO
from typing import Optional
from zoneinfo import ZoneInfo

from django.utils import dateformat, timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from core.models import Personel, PersonelUcret
from core.sayi import format_tr
from core.services import personel_devam as devam_servis
from core.services import personel_ucret as ucret_servis
from core.services.firma import firma_bilgisi_getir
from core.tarih import ay_araligi, tr_bugun

_TR = ZoneInfo("Europe/Istanbul")
_FIRMA_VARSAYILAN = "SEMTA EV GEREÇLERİ SAN. VE TİC. LTD. ŞTİ."

_SUTUNLAR = (
    ("Sıra", 6), ("Ad Soyad", 24), ("TC Kimlik No", 14), ("İşe Giriş", 12),
    ("İşten Çıkış", 12), ("Ücret (Net TL)", 16), ("Çalıştığı Gün", 12), ("Yarım Gün", 10),
    ("Pazar Çalışması", 14), ("Resmî Tatil", 12), ("Tatil Çalışması", 14),
    ("Gelmedi (Devamsız)", 16), ("Yıllık İzin", 11), ("Mazeret İzni", 12),
    ("Ücretsiz İzin", 12), ("Rapor", 8), ("Diğer İzin", 10), ("Girilmemiş Gün", 13),
    ("Açıklama", 40),
)
_SARI = PatternFill("solid", fgColor="FFF9C4")
_KIRMIZI = PatternFill("solid", fgColor="FFCDD2")
_BASLIK_DOLGU = PatternFill("solid", fgColor="21385F")
_BASLIK_YAZI = Font(bold=True, color="FFFFFF")


def _ucret_metni(ucret: Optional[PersonelUcret]) -> str:
    if ucret is None:
        return "TANIMSIZ"
    if ucret.tip == PersonelUcret.Tip.ASGARI:
        return "Asgari Ücret"
    return format_tr(ucret.net_tutar, 2)


def _ucret_degisiklikleri(personel, ilk, son) -> list:
    """Ay içinde başlayan ücret kayıtlarından, bir ÖNCESİ olanlar için "önce/sonra" notu üretir.
    İlk kez tanımlanan ücret (önceki kayıt yok) bir değişiklik sayılmaz, not eklenmez."""
    notlar = []
    kayitlar = (PersonelUcret.objects.filter(
                    personel=personel, silindi=False,
                    gecerlilik_baslangic__gte=ilk, gecerlilik_baslangic__lte=son)
                .order_by("gecerlilik_baslangic", "id"))
    for k in kayitlar:
        onceki = (PersonelUcret.objects.filter(
                      personel=personel, silindi=False,
                      gecerlilik_baslangic__lt=k.gecerlilik_baslangic)
                  .order_by("-gecerlilik_baslangic", "-id").first())
        if onceki is None:
            continue
        notlar.append(
            f"{k.gecerlilik_baslangic:%d.%m}'dan itibaren {_ucret_metni(k)} "
            f"(önce {_ucret_metni(onceki)}).")
    return notlar


@dataclass(frozen=True)
class DokumSatiri:
    sira: int
    personel: Personel
    ise_giris: date
    isten_cikis: Optional[date]
    ucret_metni: str
    ucret_tanimsiz: bool
    aciklama: str
    ozet: "devam_servis.AylikSatir"


def aylik_dokum(yil: int, ay: int, *, bugun=None) -> list:
    """[DokumSatiri] — sıra numarası core.services.personel_devam.aylik_ozet sırasıyla (ad/soyad)."""
    bugun = bugun or tr_bugun()
    ilk, son = ay_araligi(yil, ay)
    sonuc = []
    for i, s in enumerate(devam_servis.aylik_ozet(yil, ay, bugun=bugun), start=1):
        p = s.personel
        cikis_ay_icinde = p.isten_cikis_tarihi is not None and ilk <= p.isten_cikis_tarihi <= son
        esas_tarih = p.isten_cikis_tarihi if cikis_ay_icinde else son
        ucret = ucret_servis.gecerli_ucret(p, esas_tarih)

        notlar = []
        if ilk <= p.ise_giris_tarihi <= son:
            notlar.append(f"İşe başlama: {p.ise_giris_tarihi:%d.%m.%Y}.")
        if cikis_ay_icinde:
            notlar.append(f"İşten çıkış: {p.isten_cikis_tarihi:%d.%m.%Y}.")
        notlar.extend(_ucret_degisiklikleri(p, ilk, son))

        sonuc.append(DokumSatiri(
            sira=i, personel=p, ise_giris=p.ise_giris_tarihi,
            isten_cikis=(p.isten_cikis_tarihi if cikis_ay_icinde else None),
            ucret_metni=_ucret_metni(ucret), ucret_tanimsiz=(ucret is None),
            aciklama=" ".join(notlar), ozet=s))
    return sonuc


def _baslik_yaz(ws, satir, metin, *, kalin, boyut, renk="000000", italik=False):
    son_harf = get_column_letter(len(_SUTUNLAR))
    ws.merge_cells(f"A{satir}:{son_harf}{satir}")
    hucre = ws.cell(row=satir, column=1, value=metin)
    hucre.font = Font(bold=kalin, size=boyut, italic=italik, color=renk)
    return hucre


def dokum_xlsx(yil: int, ay: int, *, bugun=None, kullanici=None) -> bytes:
    """Ayın puantaj dökümünü XLSX bayt dizisi olarak üretir (hiçbir şey diske/DB'ye yazılmaz)."""
    bugun = bugun or tr_bugun()
    ilk, son = ay_araligi(yil, ay)
    satirlar = aylik_dokum(yil, ay, bugun=bugun)
    toplam = devam_servis.aylik_toplam([s.ozet for s in satirlar])
    ay_tamamlanmadi = bugun < son

    firma = (firma_bilgisi_getir().unvan or "").strip() or _FIRMA_VARSAYILAN
    ay_adi = dateformat.format(ilk, "F Y")
    uretim_zamani = timezone.now().astimezone(_TR)
    uretici = ""
    if kullanici is not None:
        uretici = kullanici.get_full_name() or kullanici.username

    wb = Workbook()
    ws = wb.active
    ws.title = "Puantaj"
    toplam_sutun = len(_SUTUNLAR)

    _baslik_yaz(ws, 1, firma, kalin=True, boyut=13)
    baslik2 = f"Personel Puantaj Dökümü — {ay_adi}"
    if ay_tamamlanmadi:
        baslik2 += f" ({bugun:%d.%m.%Y} tarihine kadar — ay tamamlanmadı)"
    _baslik_yaz(ws, 2, baslik2, kalin=True, boyut=11)
    uretim_metni = f"Üretim: {uretim_zamani:%d.%m.%Y %H:%M}"
    if uretici:
        uretim_metni += f" — {uretici}"
    _baslik_yaz(ws, 3, uretim_metni, kalin=False, boyut=9, renk="666666", italik=True)

    baslik_satiri = 5
    for idx, (baslik, genislik) in enumerate(_SUTUNLAR, start=1):
        hucre = ws.cell(row=baslik_satiri, column=idx, value=baslik)
        hucre.font = _BASLIK_YAZI
        hucre.fill = _BASLIK_DOLGU
        hucre.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = genislik
    ws.freeze_panes = f"A{baslik_satiri + 1}"

    satir_no = baslik_satiri
    for s in satirlar:
        satir_no += 1
        o = s.ozet
        degerler = [
            s.sira, s.personel.ad_soyad, s.personel.tc_kimlik_no, s.ise_giris, s.isten_cikis,
            s.ucret_metni, o.geldi, o.yarim, o.pazar_calisma, o.resmi_tatil, o.tatil_calisma,
            o.gelmedi, o.yillik, o.mazeret, o.ucretsiz, o.rapor, o.diger, o.girilmemis,
            s.aciklama,
        ]
        for col, deger in enumerate(degerler, start=1):
            hucre = ws.cell(row=satir_no, column=col, value=deger)
            if col in (4, 5):
                hucre.number_format = "DD.MM.YYYY"
            elif col == 3:
                hucre.number_format = "@"          # TC: metin — baştaki sıfır/bilimsel gösterim yok
        if s.ucret_tanimsiz:
            dolgu = _KIRMIZI
        elif o.girilmemis > 0:
            dolgu = _SARI
        else:
            dolgu = None
        if dolgu:
            for col in range(1, toplam_sutun + 1):
                ws.cell(row=satir_no, column=col).fill = dolgu

    satir_no += 1
    ws.cell(row=satir_no, column=1, value="TOPLAM")
    ws.cell(row=satir_no, column=2, value=f"{len(satirlar)} kişi")
    for col, deger in (
        (7, toplam.geldi), (8, toplam.yarim), (9, toplam.pazar_calisma),
        (10, toplam.resmi_tatil), (11, toplam.tatil_calisma), (12, toplam.gelmedi),
        (13, toplam.yillik), (14, toplam.mazeret), (15, toplam.ucretsiz), (16, toplam.rapor),
        (17, toplam.diger), (18, toplam.girilmemis),
    ):
        ws.cell(row=satir_no, column=col, value=deger)
    for col in range(1, toplam_sutun + 1):
        ws.cell(row=satir_no, column=col).font = Font(bold=True)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.5)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
