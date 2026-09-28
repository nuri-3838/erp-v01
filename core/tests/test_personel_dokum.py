"""İNSAN KAYNAKLARI > Aylık Muhasebe Dökümü (Excel) — yalnız SAYIM + ÜCRET bilgisi, bordro YOK.

`core.services.personel_devam.aylik_ozet`'e eklenen izin-türü/pazar-çalışma sayaçları,
`core.services.personel_dokum.aylik_dokum` (ücret metni + açıklama notları) ve
`dokum_xlsx` (gerçek XLSX üretimi, openpyxl ile geri okunarak doğrulanır) + view/yetki
(personel_devam VE personel_ucret ekranlarının İKİSİ de gerekli)."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import EkranYetki, PersonelDevam
from core.services.personel_devam import aylik_ozet
from core.services.personel_dokum import aylik_dokum, dokum_xlsx
from core.services.personel_izin import izin_ekle
from core.services.personel_ucret import ucret_ekle
from core.tests.ik_yardimci import personel_kur

AY_SONU = date(2026, 9, 30)          # Eylül 2026'da 4 Pazar: 6, 13, 20, 27


def _kayit(p, tarih, durum="GELDI"):
    return PersonelDevam.objects.create(personel=p, tarih=tarih, durum=durum)


class AylikOzetIzinTuruVePazarTest(TestCase):
    def test_izin_turu_bazli_sayaclar_ve_toplamla_tutarlilik(self):
        p = personel_kur(ise_giris_tarihi=date(2020, 1, 1))
        izin_ekle(p, tur="YILLIK", baslangic=date(2026, 9, 7), bitis=date(2026, 9, 9))     # Pzt-Çar, 3 gün
        izin_ekle(p, tur="RAPOR", baslangic=date(2026, 9, 10), bitis=date(2026, 9, 11))    # 2 gün
        izin_ekle(p, tur="MAZERET", baslangic=date(2026, 9, 12), bitis=date(2026, 9, 14))  # Cmt,Paz,Pzt -> 2 gün (Pazar hariç)
        (s,) = aylik_ozet(2026, 9, bugun=AY_SONU)
        self.assertEqual((s.yillik, s.rapor, s.mazeret, s.ucretsiz, s.diger), (3, 2, 2, 0, 0))
        self.assertEqual(s.izinli, s.yillik + s.ucretsiz + s.mazeret + s.diger)
        self.assertEqual(s.raporlu, s.rapor)

    def test_ucretsiz_ve_diger_izin_sayaclara_gider(self):
        p = personel_kur(ise_giris_tarihi=date(2020, 1, 1))
        izin_ekle(p, tur="UCRETSIZ", baslangic=date(2026, 9, 2), bitis=date(2026, 9, 2))
        izin_ekle(p, tur="DIGER", baslangic=date(2026, 9, 3), bitis=date(2026, 9, 3))
        (s,) = aylik_ozet(2026, 9, bugun=AY_SONU)
        self.assertEqual((s.ucretsiz, s.diger, s.izinli), (1, 1, 2))

    def test_pazar_calisma_sayaci_ayri_ve_geldi_yarima_dahil(self):
        p = personel_kur(ise_giris_tarihi=date(2020, 1, 1))
        _kayit(p, date(2026, 9, 6), "GELDI")           # Pazar
        _kayit(p, date(2026, 9, 13), "YARIM_GUN")      # Pazar
        _kayit(p, date(2026, 9, 7), "GELDI")           # Pazartesi — pazar_calisma'ya girmez
        (s,) = aylik_ozet(2026, 9, bugun=AY_SONU)
        self.assertEqual(s.pazar_calisma, 2)
        self.assertEqual(s.geldi, 2)
        self.assertEqual(s.yarim, 1)

    def test_pazar_gelmedi_kaydi_pazar_calismaya_girmez(self):
        p = personel_kur(ise_giris_tarihi=date(2020, 1, 1))
        _kayit(p, date(2026, 9, 6), "GELMEDI")
        (s,) = aylik_ozet(2026, 9, bugun=AY_SONU)
        self.assertEqual(s.pazar_calisma, 0)
        self.assertEqual(s.gelmedi, 1)

    def test_eski_alanlar_degismedi_regresyon(self):
        p = personel_kur(ise_giris_tarihi=date(2020, 1, 1))
        _kayit(p, date(2026, 9, 1), "GELDI")
        _kayit(p, date(2026, 9, 2), "GELDI")
        _kayit(p, date(2026, 9, 3), "YARIM_GUN")
        _kayit(p, date(2026, 9, 4), "GELMEDI")
        (s,) = aylik_ozet(2026, 9, bugun=AY_SONU)
        self.assertEqual((s.geldi, s.yarim, s.gelmedi), (2, 1, 1))


class AylikDokumTest(TestCase):
    def test_ucret_asgari_net_tanimsiz(self):
        a = personel_kur(ad="asgari", ise_giris_tarihi=date(2020, 1, 1))
        ucret_ekle(a, gecerlilik_baslangic=date(2026, 1, 1), tip="ASGARI")
        b = personel_kur(ad="net", ise_giris_tarihi=date(2020, 1, 1))
        ucret_ekle(b, gecerlilik_baslangic=date(2026, 1, 1), tip="NET", net_tutar=Decimal("45000"))
        c = personel_kur(ad="tanimsiz", ise_giris_tarihi=date(2020, 1, 1))
        satirlar = {s.personel.pk: s for s in aylik_dokum(2026, 9, bugun=AY_SONU)}
        self.assertEqual(satirlar[a.pk].ucret_metni, "Asgari Ücret")
        self.assertFalse(satirlar[a.pk].ucret_tanimsiz)
        self.assertEqual(satirlar[b.pk].ucret_metni, "45.000,00")
        self.assertFalse(satirlar[b.pk].ucret_tanimsiz)
        self.assertEqual(satirlar[c.pk].ucret_metni, "TANIMSIZ")
        self.assertTrue(satirlar[c.pk].ucret_tanimsiz)

    def test_ay_ici_zam_aciklamasi_ve_guncel_ucret_ay_sonu(self):
        p = personel_kur(ad="zam", ise_giris_tarihi=date(2020, 1, 1))
        ucret_ekle(p, gecerlilik_baslangic=date(2026, 1, 1), tip="NET", net_tutar=Decimal("45000"))
        ucret_ekle(p, gecerlilik_baslangic=date(2026, 9, 15), tip="NET", net_tutar=Decimal("48000"))
        (s,) = [x for x in aylik_dokum(2026, 9, bugun=AY_SONU) if x.personel.pk == p.pk]
        self.assertIn("15.09'dan itibaren 48.000,00 (önce 45.000,00).", s.aciklama)
        self.assertEqual(s.ucret_metni, "48.000,00")

    def test_ilk_kez_tanimlanan_ucret_degisiklik_sayilmaz(self):
        p = personel_kur(ad="yenihire", ise_giris_tarihi=date(2026, 9, 10))
        ucret_ekle(p, gecerlilik_baslangic=date(2026, 9, 10), tip="ASGARI")
        (s,) = [x for x in aylik_dokum(2026, 9, bugun=AY_SONU) if x.personel.pk == p.pk]
        self.assertNotIn("önce", s.aciklama)

    def test_ay_ici_ise_giris_aciklamada(self):
        p = personel_kur(ad="giren", ise_giris_tarihi=date(2026, 9, 10))
        (s,) = [x for x in aylik_dokum(2026, 9, bugun=AY_SONU) if x.personel.pk == p.pk]
        self.assertIn("İşe başlama: 10.09.2026.", s.aciklama)
        self.assertIsNone(s.isten_cikis)

    def test_ay_ici_isten_cikis_aciklamada_ve_alanda(self):
        p = personel_kur(ad="cikan", ise_giris_tarihi=date(2020, 1, 1), isten_cikis_tarihi=date(2026, 9, 20))
        (s,) = [x for x in aylik_dokum(2026, 9, bugun=AY_SONU) if x.personel.pk == p.pk]
        self.assertIn("İşten çıkış: 20.09.2026.", s.aciklama)
        self.assertEqual(s.isten_cikis, date(2026, 9, 20))

    def test_ay_disinda_giris_cikis_aciklama_uretmez(self):
        p = personel_kur(ad="eskiden", ise_giris_tarihi=date(2020, 1, 1))
        (s,) = [x for x in aylik_dokum(2026, 9, bugun=AY_SONU) if x.personel.pk == p.pk]
        self.assertNotIn("İşe başlama", s.aciklama)
        self.assertIsNone(s.isten_cikis)

    def test_cikis_ay_icindeyse_ucret_cikis_gunu_itibariyla(self):
        p = personel_kur(ad="cikankisi", ise_giris_tarihi=date(2020, 1, 1),
                         isten_cikis_tarihi=date(2026, 9, 15))
        ucret_ekle(p, gecerlilik_baslangic=date(2026, 1, 1), tip="NET", net_tutar=Decimal("40000"))
        ucret_ekle(p, gecerlilik_baslangic=date(2026, 9, 20), tip="NET", net_tutar=Decimal("50000"))
        (s,) = [x for x in aylik_dokum(2026, 9, bugun=AY_SONU) if x.personel.pk == p.pk]
        self.assertEqual(s.ucret_metni, "40.000,00")

    def test_sira_ad_soyad_sirasiyla(self):
        personel_kur(ad="berk", ise_giris_tarihi=date(2020, 1, 1))
        personel_kur(ad="ali", ise_giris_tarihi=date(2020, 1, 1))
        satirlar = aylik_dokum(2026, 9, bugun=AY_SONU)
        self.assertEqual([s.sira for s in satirlar], [1, 2])
        self.assertEqual([s.personel.ad for s in satirlar], ["ALİ", "BERK"])   # TR büyük harf


class DokumXlsxTest(TestCase):
    def test_basliklar_sutun_sirasi_tc_metin_ve_toplam(self):
        p = personel_kur(ad="ahmet", soyad="veli", tc_kimlik_no="10000000146",
                         ise_giris_tarihi=date(2020, 1, 1))
        ucret_ekle(p, gecerlilik_baslangic=date(2026, 1, 1), tip="NET", net_tutar=Decimal("45000"))
        _kayit(p, date(2026, 9, 1), "GELDI")
        _kayit(p, date(2026, 9, 2), "GELMEDI")

        xbytes = dokum_xlsx(2026, 9, bugun=AY_SONU)
        wb = load_workbook(BytesIO(xbytes))
        ws = wb.active

        self.assertIn("Personel Puantaj Dökümü", ws["A2"].value)
        self.assertIn("Eylül 2026", ws["A2"].value)
        self.assertNotIn("tamamlanmadı", ws["A2"].value)     # bugun=ay sonu -> ay tamamlandı

        beklenen_baslik = ["Sıra", "Ad Soyad", "TC Kimlik No", "İşe Giriş", "İşten Çıkış",
                           "Ücret (Net TL)", "Çalıştığı Gün", "Yarım Gün", "Pazar Çalışması",
                           "Resmî Tatil", "Tatil Çalışması", "Gelmedi (Devamsız)", "Yıllık İzin",
                           "Mazeret İzni", "Ücretsiz İzin", "Rapor", "Diğer İzin",
                           "Girilmemiş Gün", "Açıklama"]
        baslik_satiri = 5
        gercek = [ws.cell(row=baslik_satiri, column=i + 1).value for i in range(len(beklenen_baslik))]
        self.assertEqual(gercek, beklenen_baslik)

        veri_satiri = baslik_satiri + 1
        self.assertEqual(ws.cell(row=veri_satiri, column=2).value, "AHMET VELİ")
        tc_hucre = ws.cell(row=veri_satiri, column=3)
        self.assertEqual(tc_hucre.value, "10000000146")
        self.assertEqual(tc_hucre.number_format, "@")
        self.assertEqual(ws.cell(row=veri_satiri, column=6).value, "45.000,00")
        self.assertEqual(ws.cell(row=veri_satiri, column=7).value, 1)      # Çalıştığı Gün
        self.assertEqual(ws.cell(row=veri_satiri, column=12).value, 1)     # Gelmedi

        toplam_satiri = veri_satiri + 1
        self.assertEqual(ws.cell(row=toplam_satiri, column=1).value, "TOPLAM")
        self.assertEqual(ws.cell(row=toplam_satiri, column=7).value, 1)
        self.assertEqual(ws.cell(row=toplam_satiri, column=12).value, 1)

    def test_ay_tamamlanmadiysa_baslikta_belirtilir(self):
        personel_kur(ise_giris_tarihi=date(2020, 1, 1))
        xbytes = dokum_xlsx(2026, 9, bugun=date(2026, 9, 19))
        wb = load_workbook(BytesIO(xbytes))
        ws = wb.active
        self.assertIn("19.09.2026 tarihine kadar", ws["A2"].value)
        self.assertIn("ay tamamlanmadı", ws["A2"].value)

    def _satir_bul(self, ws, tc):
        """TC (metin sütunu) ile satır arar — ad/soyad TR büyük harfe (İ/I) çevrildiği için
        ASCII string eşleşmesi güvenilir değil; bkz. core.metin.buyuk_harf_tr."""
        row = 6
        while ws.cell(row=row, column=2).value:
            if ws.cell(row=row, column=3).value == tc:
                return row
            row += 1
        raise AssertionError(f"TC {tc} satırı bulunamadı")

    def test_girilmemis_satir_sari_tanimsiz_ucret_kirmizi(self):
        a = personel_kur(ad="sarikisi", soyad="test", tc_kimlik_no="10000000146",
                         ise_giris_tarihi=date(2020, 1, 1))
        ucret_ekle(a, gecerlilik_baslangic=date(2026, 1, 1), tip="ASGARI")
        b = personel_kur(ad="kirmizikisi", soyad="test", tc_kimlik_no="10000000214",
                         ise_giris_tarihi=date(2020, 1, 1))

        xbytes = dokum_xlsx(2026, 9, bugun=AY_SONU)
        wb = load_workbook(BytesIO(xbytes))
        ws = wb.active
        a_row = self._satir_bul(ws, "10000000146")
        b_row = self._satir_bul(ws, "10000000214")
        self.assertEqual(ws.cell(row=a_row, column=1).fill.fgColor.rgb, "00FFF9C4")   # girilmemiş var, ücret tanımlı -> sarı
        self.assertEqual(ws.cell(row=b_row, column=1).fill.fgColor.rgb, "00FFCDD2")   # ücret tanımsız -> kırmızı


class DokumViewTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.yon = User.objects.create_superuser("dkyon", password="x")
        cls.devamci = User.objects.create_user("dkdevamci", password="x")
        EkranYetki.objects.create(kullanici=cls.devamci, ekran_kod="personel_devam")
        cls.ucretci = User.objects.create_user("dkucretci", password="x")
        EkranYetki.objects.create(kullanici=cls.ucretci, ekran_kod="personel_ucret")
        cls.ikisi = User.objects.create_user("dkikisi", password="x")
        EkranYetki.objects.create(kullanici=cls.ikisi, ekran_kod="personel_devam")
        EkranYetki.objects.create(kullanici=cls.ikisi, ekran_kod="personel_ucret")
        cls.bos = User.objects.create_user("dkbos", password="x")

    def test_anonim_302_tek_yetkiler_403(self):
        url = reverse("core:yoklama_aylik_dokum")
        self.assertEqual(self.client.get(url).status_code, 302)
        for kullanici in (self.bos, self.devamci, self.ucretci):
            self.client.force_login(kullanici)
            self.assertEqual(self.client.get(url).status_code, 403, kullanici.username)

    def test_iki_yetki_de_varsa_xlsx_indirir(self):
        self.client.force_login(self.ikisi)
        r = self.client.get(reverse("core:yoklama_aylik_dokum"), {"ay": "2026-09"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(
            r["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("puantaj_2026-09.xlsx", r["Content-Disposition"])
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.active.title, "Puantaj")

    def test_yonetici_yetkisiz_de_indirir(self):
        self.client.force_login(self.yon)
        self.assertEqual(self.client.get(reverse("core:yoklama_aylik_dokum")).status_code, 200)

    def test_buton_yalniz_iki_yetkiyle_gorunur(self):
        url = reverse("core:yoklama_aylik")
        self.client.force_login(self.devamci)
        r = self.client.get(url)
        self.assertEqual(r.status_code, 200)
        self.assertNotContains(r, "Muhasebe Dökümü")
        self.client.force_login(self.ikisi)
        r2 = self.client.get(url)
        self.assertContains(r2, "Muhasebe Dökümü")

    def test_girilmemis_gun_uyarisi_indirmeyi_engellemez(self):
        personel_kur(ise_giris_tarihi=date(2020, 1, 1))     # hiç yoklama girilmemiş -> girilmemis > 0
        self.client.force_login(self.ikisi)
        r = self.client.get(reverse("core:yoklama_aylik"), {"ay": "2026-09"})
        self.assertContains(r, "gün yoklama girilmemiş")
        r2 = self.client.get(reverse("core:yoklama_aylik_dokum"), {"ay": "2026-09"})
        self.assertEqual(r2.status_code, 200)
