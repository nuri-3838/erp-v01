"""ÜRÜN AĞACI maliyeti: birim maliyet önceliği (ortalama > stok kartı alış fiyatı > yok), alış fiyatının TCMB kuruyla çevrimi, kaynak
rozet/özet sayıları, TL ve USD'nin her yerde ayrı gösterimi (Ürün/Maliyet, Karşılaştır, Excel)."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.models import Kur, Stok
from core.services import urun_agaci as ua
from core.tests.test_urun_agaci_kullanim import KullanimBase

D = Decimal
BUGUN = date(2026, 10, 10)                 # YanCiktiBase kur kaydı: USD 40


def kart(kod, fiyat, pb="TRY", **kw):
    return Stok(kod=kod, alis_fiyati=D(str(fiyat)) if fiyat is not None else None, alis_fiyati_pb=pb, **kw)


class BirimMaliyetTest(KullanimBase):
    KURLAR = {"USD": (D("40"), date(2026, 10, 9)), "EUR": (D("44"), date(2026, 10, 9))}

    def test_ortalama_kartin_onune_gecer(self):
        s = kart("X", 999, ort_maliyet_try=D("5"), ort_maliyet_usd=D("0.125"))
        b = ua.birim_maliyet(s, self.KURLAR)
        self.assertEqual((b["kaynak"], b["try"], b["usd"]), ("ORTALAMA", D("5"), D("0.125")))

    def test_kart_try(self):
        b = ua.birim_maliyet(kart("X", 80, "TRY"), self.KURLAR)
        self.assertEqual((b["kaynak"], b["try"], b["usd"]), ("KART", D("80"), D("2")))
        self.assertEqual(b["kurlar"], ("USD",))

    def test_kart_usd(self):
        b = ua.birim_maliyet(kart("X", 2, "USD"), self.KURLAR)
        self.assertEqual((b["kaynak"], b["try"], b["usd"]), ("KART", D("80"), D("2")))

    def test_kart_eur(self):
        b = ua.birim_maliyet(kart("X", 10, "EUR"), self.KURLAR)
        self.assertEqual((b["try"], b["usd"]), (D("440"), D("11")))                 # TL = fiyat × EUR; USD = TL / USD
        self.assertEqual(b["kurlar"], ("EUR", "USD"))

    def test_yuvarlama_round_half_up_6_hane(self):
        b = ua.birim_maliyet(kart("X", 1, "TRY"), {"USD": (D("3"), date(2026, 10, 9))})
        self.assertEqual(b["usd"], D("0.333333"))
        b = ua.birim_maliyet(kart("X", 2, "TRY"), {"USD": (D("3"), date(2026, 10, 9))})
        self.assertEqual(b["usd"], D("0.666667"))

    def test_ikisi_de_yoksa_yok(self):
        for s in (kart("X", None), kart("X", 0)):
            b = ua.birim_maliyet(s, self.KURLAR)
            self.assertEqual((b["kaynak"], b["try"], b["usd"], b["uyari"]), ("YOK", None, None, ""))

    def test_kur_yoksa_uyari_ve_yok(self):
        b = ua.birim_maliyet(kart("X", 10, "TRY"), {})
        self.assertEqual((b["kaynak"], b["try"]), ("YOK", None))
        self.assertIn("USD için TCMB kuru bulunamadı", b["uyari"])
        b = ua.birim_maliyet(kart("X", 10, "GBP"), self.KURLAR)                      # GBP kuru yok
        self.assertEqual(b["kaynak"], "YOK")
        self.assertIn("GBP", b["uyari"])

    def test_kur_haritasi_kaydirmali_tarih(self):
        Kur.objects.create(tarih=date(2026, 10, 12), usd_alis=D("41"), eur_alis=D("45"))
        h = ua.kur_haritasi(date(2026, 10, 14))                                      # 14'ünde kur yok → son yayımlanan (12)
        self.assertEqual(h["USD"], (D("41"), date(2026, 10, 12)))
        self.assertEqual(h["EUR"], (D("45"), date(2026, 10, 12)))
        self.assertNotIn("GBP", h)
        self.assertEqual(ua.kur_haritasi(date(2026, 10, 30)), {})                    # 7 günden eski → yok

    def test_kur_notu(self):
        self.assertEqual(ua.kur_notu(self.KURLAR, {"USD"}), "Alış fiyatı çevrimi: TCMB 09.10.2026, USD 40,0000")
        self.assertEqual(ua.kur_notu(self.KURLAR, {"USD", "EUR"}), "Alış fiyatı çevrimi: TCMB 09.10.2026, USD 40,0000, EUR 44,0000")
        self.assertEqual(ua.kur_notu(self.KURLAR, set()), "")


class UrunMaliyetKaynakTest(KullanimBase):
    """mamul_a: PROFIL 0,5 boy (ortalama) + CIVATA 4 (ortalama) + BOYA 1 (kartta alış fiyatı)."""

    def setUp(self):
        self.giris(self.profil, 2, 200, 5)
        self.giris(self.civata, 25, 50, 1.25)
        Stok.objects.filter(pk=self.civata.pk).update(alis_fiyati=D("999"))          # ortalama olduğu için YOK SAYILMALI
        Stok.objects.filter(pk=self.boya.pk).update(alis_fiyati=D("10"), alis_fiyati_pb="TRY")

    def maliyet(self, bugun=BUGUN, miktar=D("1")):
        return ua.urun_maliyet(ua.graf_yukle(bugun), self.mamul_a, miktar)

    def satirlar(self, m):
        return {s["stok"].kod: s for g in m["gruplar"] for s in g["satirlar"]}

    def test_oncelik_ve_kaynaklar(self):
        m = self.maliyet()
        s = self.satirlar(m)
        self.assertEqual((s["PROFIL"]["kaynak"], s["CIVATA"]["kaynak"], s["BOYA"]["kaynak"]), ("ORTALAMA", "ORTALAMA", "KART"))
        self.assertEqual(s["CIVATA"]["tutar_try"], D("8"))                           # kart fiyatı (999) değil ortalama
        self.assertEqual((s["BOYA"]["ort_try"], s["BOYA"]["ort_usd"]), (D("10"), D("0.25")))
        self.assertEqual((m["ortalama_sayi"], m["kart_sayi"], m["maliyetsiz"]), (2, 1, 0))
        self.assertEqual(m["toplam_try"], D("68"))                                   # 50 + 8 + 10
        self.assertAlmostEqual(m["kart_pay"], D("10") / D("68") * 100, places=6)
        self.assertEqual(m["kur_notu"], "Alış fiyatı çevrimi: TCMB 10.10.2026, USD 40,0000")
        self.assertEqual(m["usd_eksik"], 0)

    def test_tl_usd_ayri_toplamlar_ve_miktar(self):
        m = self.maliyet(miktar=D("10"))
        self.assertEqual(m["miktar_toplam_try"], m["toplam_try"] * 10)
        self.assertEqual(m["miktar_toplam_usd"], m["toplam_usd"] * 10)
        for g in m["gruplar"]:
            self.assertEqual((g["toplam_try_n"], g["toplam_usd_n"]), (g["toplam_try"] * 10, g["toplam_usd"] * 10))
        self.assertEqual(sum(g["toplam_usd"] for g in m["gruplar"]), m["toplam_usd"])
        s = self.satirlar(m)["BOYA"]
        self.assertEqual((s["tutar_try_n"], s["tutar_usd_n"]), (D("100"), D("2.5")))

    def test_kur_yokken_maliyet_yok_ve_uyari(self):
        m = self.maliyet(bugun=date(2030, 1, 1))
        s = self.satirlar(m)["BOYA"]
        self.assertEqual((s["kaynak"], s["maliyet_yok"], s["tutar_try"]), ("YOK", True, None))
        self.assertEqual((m["kart_sayi"], m["maliyetsiz"]), (0, 1))
        self.assertEqual(m["toplam_try"], D("58"))
        self.assertEqual(len(m["kur_uyarilari"]), 1)
        self.assertIn("BOYA", m["kur_uyarilari"][0])
        self.assertEqual(m["kur_notu"], "")

    def test_usd_eksik_sayaci(self):
        Stok.objects.filter(pk=self.civata.pk).update(ort_maliyet_usd=None)         # TL var, USD yok
        m = self.maliyet()
        self.assertEqual(m["usd_eksik"], 1)
        self.assertEqual(m["maliyetsiz"], 0)

    def test_karsilastir_tl_usd_urun_maliyetiyle_ayni(self):
        g = ua.graf_yukle(BUGUN)
        urunler = [self.mamul_a, self.mamul_c, self.a_modeller[0], self.c_modeller[1]]
        for pb, alan in (("TL", "toplam_try"), ("USD", "toplam_usd")):
            m = ua.karsilastir(g, urunler, "maliyet", pb)
            self.assertEqual(m["pb"], pb)
            for i, u in enumerate(urunler):
                um = ua.urun_maliyet(g, u, D("1"))
                self.assertEqual(m["toplamlar"][i]["toplam"], um[alan], (pb, u.kod))
                if pb == "TL":
                    self.assertEqual(m["toplamlar"][i]["maliyetsiz"], um["maliyetsiz"])
        tl = ua.karsilastir(g, urunler, "maliyet", "TL")
        satir = {s["stok"].kod: s for gr in tl["gruplar"] for s in gr["satirlar"]}
        self.assertEqual(satir["BOYA"]["kaynak"], "KART")
        self.assertEqual(satir["PROFIL"]["kaynak"], "ORTALAMA")
        self.assertEqual((tl["kart_kalem"], tl["kur_notu"]), (1, "Alış fiyatı çevrimi: TCMB 10.10.2026, USD 40,0000"))
        self.assertEqual(ua.karsilastir(g, urunler, "maliyet", "ABC")["pb"], "TL")    # geçersiz para birimi TL'ye düşer


class AlisFiyatiEkranTest(KullanimBase):
    def setUp(self):
        Kur.objects.get_or_create(tarih=timezone.localdate(), defaults={"usd_alis": D("40")})
        self.giris(self.profil, 2, 200, 5)
        self.giris(self.civata, 25, 50, 1.25)
        Stok.objects.filter(pk=self.boya.pk).update(alis_fiyati=D("10"), alis_fiyati_pb="TRY")
        self.client.force_login(User.objects.create_superuser("ua_af", password="x"))
        self.url = reverse("core:urun_agaci")

    def al(self, **p):
        return self.client.get(self.url, p)

    def test_maliyet_ekrani_rozet_ve_ozet(self):
        h = self.al(urun=self.mamul_a.pk, miktar="10", sekme="maliyet").content.decode()
        self.assertIn("kaynak-ortalama", h)
        self.assertIn("kaynak-kart", h)
        self.assertIn("Alış fiyatı (kart)", h)
        self.assertIn("<strong>2</strong> kalem ortalama", h)
        self.assertIn("<strong>1</strong> kalem alış fiyatı", h)
        self.assertIn("toplamın %14,7'i", h)
        self.assertIn("<strong>0</strong> kalem maliyetsiz", h)
        self.assertIn("Alış fiyatı çevrimi: TCMB", h)
        self.assertIn("USD 40,0000", h)
        self.assertIn("680,00", h)                                                   # 10 adet TL
        self.assertIn("Tutar USD ×10", h)

    def test_kur_yoksa_ekranda_uyari(self):
        Kur.objects.all().delete()
        h = self.al(urun=self.mamul_a.pk, sekme="maliyet").content.decode()
        self.assertIn("için TCMB kuru bulunamadı", h)
        self.assertIn("1 kalemin maliyeti yok — toplam eksik", h)

    def test_karsilastir_tl_usd_anahtari(self):
        h = self.al(gorunum="karsilastir", seri="A", mod="maliyet", pb="TL").content.decode()
        self.assertIn("ÜRÜN BAŞINA TOPLAM (TL)", h)
        self.assertIn("Maliyet matrisi (TL)", h)
        self.assertIn('class="aktif">TL<', h)
        self.assertIn("alış fiyatından hesaplandı", h)
        self.assertIn("*", h)                                                        # kart kaynaklı hücre işareti
        h = self.al(gorunum="karsilastir", seri="A", mod="maliyet", pb="USD").content.decode()
        self.assertIn("ÜRÜN BAŞINA TOPLAM (USD)", h)
        self.assertIn('class="aktif">USD<', h)

    def test_miktar_modunda_para_anahtari_yok(self):
        h = self.al(gorunum="karsilastir", seri="A", mod="miktar").content.decode()
        self.assertNotIn("Para birimi", h)

    def test_excel_urun_maliyeti_tl_usd_ayri_sutun(self):
        r = self.al(urun=self.mamul_a.pk, miktar="10", sekme="maliyet", xlsx="1")
        ws = load_workbook(BytesIO(r.content)).active
        basliklar = [c.value for c in ws[4]]
        for b in ("Kaynak", "Birim maliyet TL", "Birim maliyet USD", "Tutar TL", "Tutar USD", "Tutar TL (10 adet)", "Tutar USD (10 adet)"):
            self.assertIn(b, basliklar)
        satir = {row[1]: row for row in ws.iter_rows(min_row=5, values_only=True) if row[1]}
        self.assertEqual(satir["BOYA"][4], "Alış fiyatı (kart)")
        self.assertEqual(satir["PROFIL"][4], "Ortalama")
        self.assertEqual(satir["BOYA"][8], D("10"))                                  # Tutar TL: sayı hücresi
        self.assertEqual(satir["BOYA"][9], D("0.25"))
        self.assertIn("2 kalem ortalama · 1 kalem alış fiyatı · 0 kalem maliyetsiz", ws["A2"].value)
        self.assertIn("Alış fiyatı çevrimi", ws["A2"].value)

    def test_excel_karsilastir_iki_sayfa(self):
        r = self.al(gorunum="karsilastir", seri="A", mod="maliyet", xlsx="1")
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Maliyet TL", "Maliyet USD"])
        for ad in wb.sheetnames:
            ws = wb[ad]
            self.assertIn("Kaynak", [c.value for c in ws[4]])
            son = [row for row in ws.iter_rows(min_row=5, values_only=True) if row[0] and str(row[0]).startswith("ÜRÜN BAŞINA TOPLAM")][0]
            self.assertIn(ad.split()[-1], son[0])
        tl = [row for row in wb["Maliyet TL"].iter_rows(min_row=5, values_only=True) if row[0] and str(row[0]).startswith("ÜRÜN BAŞINA")][0]
        usd = [row for row in wb["Maliyet USD"].iter_rows(min_row=5, values_only=True) if row[0] and str(row[0]).startswith("ÜRÜN BAŞINA")][0]
        self.assertTrue(all(isinstance(v, (int, float, D)) for v in tl[5:]))
        self.assertTrue(all(u <= t for t, u in zip(tl[5:], usd[5:])))
        self.assertEqual((tl[5], usd[5]), (68, 1.7))                                # 152-10-0001: 68 TL; 50/100 + 8 + 10 → USD 1,25 + 0,2 + 0,25
        r = self.al(gorunum="karsilastir", seri="A", mod="miktar", xlsx="1")
        self.assertEqual(load_workbook(BytesIO(r.content)).sheetnames, ["Miktar"])
