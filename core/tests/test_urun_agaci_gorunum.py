"""ÜRÜN AĞACI — Ürün görünümü: Ağaç · Malzeme ve Stok · Maliyet (+ Excel). Hepsi Operasyon zincirinden hesaplanır (salt-okunur).
Zincir: PROFIL (BOY) → ON21 (kesim, 2 adet/boy, tam boy) → 152-10-0001 (+4 CIVATA, +1 BOYA[stoksuz/maliyetsiz]);
        PROFIL → AYAK66 (3 adet/boy, yan çıktı AYAK55, boy payı) → 152-22-0001 (+2 CIVATA)."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from core.models import EkranYetki, Kategori, Stok, StokHareket
from core.services import urun_agaci as ua
from core.services.hareket import hareket_ekle
from core.services.uretim import (
    ana_cikti_payi, ihtiyac_hesapla, operasyon_kaydi_onayla, operasyon_olustur,
)
from core.tests.test_uretim_yan_cikti import YanCiktiBase

D = Decimal


class GorunumBase(YanCiktiBase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.kat_hird = Kategori.objects.create(kod="HI", ad="HIRDAVAT")
        cls.civata = cls._stok("CIVATA", cls.kat_hird, satinalma=True, uretim=False)
        cls.boya = cls._stok("BOYA", cls.kat_hird, satinalma=True, uretim=False)
        cls.on21 = cls._stok("ON21", cls.kat_ana)
        cls.mamul_a = cls._stok("152-10-0001", cls.kat_ana)
        cls.mamul_c = cls._stok("152-22-0001", cls.kat_ana)
        cls.montaj = cls._istasyon_al("MONTAJ")
        operasyon_olustur(istasyon_id=cls.kesim.pk, cikti_id=cls.on21.pk, cikti_miktar=D("2"), satirlar=[(cls.profil, D("1"))])
        cls.op_a = operasyon_olustur(istasyon_id=cls.montaj.pk, cikti_id=cls.mamul_a.pk, cikti_miktar=D("1"),
                                     satirlar=[(cls.on21, D("1")), (cls.civata, D("4")), (cls.boya, D("1"))])
        cls.op_c = operasyon_olustur(istasyon_id=cls.montaj.pk, cikti_id=cls.mamul_c.pk, cikti_miktar=D("1"),
                                     satirlar=[(cls.ana, D("3")), (cls.civata, D("2"))])

    @classmethod
    def _istasyon_al(cls, kod):
        from core.tests.test_uretim import _istasyon
        return _istasyon(kod)

    def giris(self, stok, miktar, tl, usd=None):
        return hareket_ekle(stok_id=stok.pk, depo_id=self.depo.pk, tarih=date(2026, 1, 1), tur=StokHareket.Tur.GIRIS, miktar=D(str(miktar)),
                            giris_tutar_try=D(str(tl)), giris_tutar_usd=(D(str(usd)) if usd is not None else None))

    def maliyetli(self):
        self.giris(self.profil, 2, 200, 5)            # ort 100 TL / 2,5 USD
        self.giris(self.civata, 25, 50, 1.25)         # ort 2 TL / 0,05 USD


class IhtiyacSecenekTest(GorunumBase):
    def test_varsayilan_davranis_eskisiyle_ayni(self):
        e = ihtiyac_hesapla([(self.mamul_a, D("5"))])
        y = ihtiyac_hesapla([(self.mamul_a, D("5"))], boy_yuvarla=True, pay_dus=False)

        def ozet(s):
            return [(o["stok"].pk, o["ihtiyac"], o["calistirma_sayisi"], o["uretilecek_miktar"], o["fazla_miktar"], o["tam"]) for o in s["ozet"]]
        self.assertEqual(ozet(e), ozet(y))
        self.assertEqual(ozet(e), ozet(ihtiyac_hesapla([(self.mamul_a, D("5"))], graf=ua.graf_yukle())))   # bellek grafı aynı sonucu verir
        o = {x["stok"].kod: x for x in e["ozet"]}
        self.assertEqual((o["ON21"]["calistirma_sayisi"], o["PROFIL"]["ihtiyac"]), (D("3"), D("3")))      # 5 adet → 2,5 → 3 boy

    def test_kesirli_tuketim_yuvarlamasiz(self):
        """1 adet 2+1 ön ayak (2 adet/boy) → 0,5 boy; varsayılan davranış 1 boy."""
        k = ihtiyac_hesapla([(self.mamul_a, D("1"))], boy_yuvarla=False)
        self.assertEqual({o["stok"].kod: o["ihtiyac"] for o in k["ozet"]}["PROFIL"], D("0.5"))
        v = ihtiyac_hesapla([(self.mamul_a, D("1"))])
        self.assertEqual({o["stok"].kod: o["ihtiyac"] for o in v["ozet"]}["PROFIL"], D("1"))

    def test_yan_cikti_pay_dusumu(self):
        """6+6 SAĞ: 1 boy → 3 ana adet (+1 yan); ana başına girdi = 1/3 × boy payı."""
        yan = [self.op.yan_ciktilar.get()]
        pay = ana_cikti_payi(self.op.cikti_miktar, self.op.boy_mm, yan)
        self.assertEqual(pay, D("3") * D("1292.60") / (D("3") * D("1292.60") + D("1063.53")))
        tam = ihtiyac_hesapla([(self.ana, D("1"))], boy_yuvarla=False)["ozet"][0]["ihtiyac"]
        net = ihtiyac_hesapla([(self.ana, D("1"))], boy_yuvarla=False, pay_dus=True)["ozet"][0]["ihtiyac"]
        self.assertEqual(tam, D("1") / D("3"))
        self.assertEqual(net, D("1") / D("3") * pay)
        self.assertTrue(D("0.26") < net < D("0.27"))

    def test_pay_onay_ani_snapshot_ile_ayni_kural(self):
        self.profil_gir(1, 300)
        kayit = self.kayit("3")
        operasyon_kaydi_onayla(kayit)
        ana = kayit.ciktilar.get(ana_mi=True)
        pay = ana_cikti_payi(self.op.cikti_miktar, self.op.boy_mm, [self.op.yan_ciktilar.get()])
        self.assertAlmostEqual(ana.pay_orani, pay, places=9)

    def test_pay_eksik_boyda_ana_yuzde_100(self):
        class Y:
            miktar, boy_mm = D("1"), None
        self.assertEqual(ana_cikti_payi(D("3"), D("1292.6"), [Y()]), D("1"))
        self.assertEqual(ana_cikti_payi(D("3"), D("1292.6"), []), D("1"))

    def test_agac_yapraklari_birim_tuketimle_tutarli(self):
        g = ua.graf_yukle()
        agac = ihtiyac_hesapla([(self.mamul_c, D("1"))], boy_yuvarla=False, pay_dus=True, graf=g)["agac"][0]
        toplam = {}

        def gez(d):
            if d["yaprak"]:
                toplam[d["stok"].pk] = toplam.get(d["stok"].pk, D("0")) + d["miktar"]
            for c in d["cocuklar"]:
                gez(c)
        gez(agac)
        self.assertEqual(toplam, {pk: m for pk, (_, m) in g.yapraklar(self.mamul_c).items()})


class MalzemeMaliyetTest(GorunumBase):
    def test_malzeme_eksik_ve_yeterli(self):
        self.giris(self.profil, 2, 200)
        self.giris(self.civata, 25, 50)
        self.giris(self.boya, 10, 10)
        g = ua.graf_yukle()
        m = ua.urun_malzeme(g, self.mamul_a, D("5"))
        satir = {s["stok"].kod: s for gr in m["gruplar"] for s in gr["satirlar"]}
        self.assertEqual(set(satir), {"PROFIL", "CIVATA", "BOYA"})
        self.assertEqual((satir["PROFIL"]["gerekli"], satir["PROFIL"]["eldeki"], satir["PROFIL"]["eksik"]), (D("3"), D("2"), D("1")))   # tam boy
        self.assertEqual((satir["CIVATA"]["gerekli"], satir["CIVATA"]["eksik"]), (D("20"), D("0")))
        self.assertEqual((m["yeterli"], m["eksik"]), (2, 1))
        self.assertEqual({gr["kategori"] for gr in m["gruplar"]}, {"ÜT GİRDİ", "HIRDAVAT"})            # alt kategoriye göre gruplu
        # toplam gerekli, ağaç hesabıyla aynı
        ozet = {o["stok"].pk: o["ihtiyac"] for o in ihtiyac_hesapla([(self.mamul_a, D("5"))])["ozet"] if o["yaprak"]}
        self.assertEqual({s["stok"].pk: s["gerekli"] for s in satir.values()}, ozet)

    def test_depo_filtresi(self):
        from core.models import Depo
        d2 = Depo.objects.create(kod="D2", ad="İKİNCİ DEPO")
        self.giris(self.profil, 2, 200)
        hareket_ekle(stok_id=self.profil.pk, depo_id=d2.pk, tarih=date(2026, 1, 1), tur=StokHareket.Tur.GIRIS, miktar=D("5"),
                     giris_tutar_try=D("500"))
        g = ua.graf_yukle()
        tum = {s["stok"].kod: s for gr in ua.urun_malzeme(g, self.mamul_a, D("1"))["gruplar"] for s in gr["satirlar"]}
        sadece = {s["stok"].kod: s for gr in ua.urun_malzeme(g, self.mamul_a, D("1"), d2)["gruplar"] for s in gr["satirlar"]}
        self.assertEqual((tum["PROFIL"]["eldeki"], sadece["PROFIL"]["eldeki"]), (D("7"), D("5")))

    def test_maliyet_bir_adet_kesirli(self):
        self.maliyetli()
        g = ua.graf_yukle()
        m = ua.urun_maliyet(g, self.mamul_a, D("10"))
        satir = {s["stok"].kod: s for gr in m["gruplar"] for s in gr["satirlar"]}
        self.assertEqual(satir["PROFIL"]["tuketim"], D("0.5"))                       # yuvarlama yok
        self.assertEqual(satir["PROFIL"]["tutar_try"], D("0.5") * satir["PROFIL"]["ort_try"])
        self.assertEqual(satir["PROFIL"]["tutar_try"], D("50"))
        self.assertEqual(satir["CIVATA"]["tutar_try"], D("8"))
        self.assertEqual(m["toplam_try"], D("58"))
        self.assertEqual(m["miktar_toplam_try"], D("580"))
        self.assertAlmostEqual(satir["PROFIL"]["pay"], D("50") / D("58") * 100, places=6)
        self.assertEqual(m["toplam_usd"], satir["PROFIL"]["tutar_usd"] + satir["CIVATA"]["tutar_usd"])
        self.assertEqual(sum(gr["toplam_try"] for gr in m["gruplar"]), m["toplam_try"])  # ara toplamlar genel toplamı verir

    def test_maliyetsiz_kalem_uyarisi(self):
        self.maliyetli()
        m = ua.urun_maliyet(ua.graf_yukle(), self.mamul_a, D("1"))
        satir = {s["stok"].kod: s for gr in m["gruplar"] for s in gr["satirlar"]}
        self.assertTrue(satir["BOYA"]["maliyet_yok"])
        self.assertIsNone(satir["BOYA"]["tutar_try"])
        self.assertEqual(m["maliyetsiz"], 1)
        self.assertEqual(m["toplam_try"], D("58"))                                   # maliyetsiz kalem toplama girmez

    def test_yan_cikti_pay_maliyette(self):
        self.maliyetli()
        m = ua.urun_maliyet(ua.graf_yukle(), self.mamul_c, D("1"))
        satir = {s["stok"].kod: s for gr in m["gruplar"] for s in gr["satirlar"]}
        pay = ana_cikti_payi(self.op.cikti_miktar, self.op.boy_mm, [self.op.yan_ciktilar.get()])
        self.assertEqual(satir["PROFIL"]["tuketim"], D("1") * pay)                  # 3 AYAK66 = 1 boy × pay


class EkranTest(GorunumBase):
    def setUp(self):
        self.su = User.objects.create_superuser("ua_g", password="x")
        self.client.force_login(self.su)
        self.url = reverse("core:urun_agaci")

    def al(self, **p):
        return self.client.get(self.url, p)

    def test_gruplu_secici(self):
        h = self.al().content.decode()
        self.assertIn('<optgroup label="Bitmiş ürünler – A tipi">', h)
        self.assertIn('<optgroup label="Bitmiş ürünler – Çift çıkış">', h)
        self.assertIn('<optgroup label="Ara parçalar">', h)
        self.assertLess(h.index("152-10-0001"), h.index('label="Ara parçalar"'))

    def test_agac_rozetleri(self):
        r = self.al(urun=self.mamul_c.pk)
        self.assertContains(r, "yan çıktı")
        r = self.al(urun=self.mamul_a.pk)
        self.assertContains(r, "tam boy")
        self.assertContains(r, 'class="ua-sekme aktif"')

    def test_malzeme_sekmesi(self):
        self.giris(self.profil, 2, 200)
        r = self.al(urun=self.mamul_a.pk, miktar="5", sekme="malzeme")
        self.assertEqual(r.status_code, 200)
        h = r.content.decode()
        self.assertIn("kalem yeterli", h)
        self.assertIn("kalem eksik", h)
        self.assertIn("Açık kayıtlara ayrılan stok düşülmez", h)
        self.assertIn('class="ua-s eksik"', h)
        self.assertIsNone(r.context["kokler"])

    def test_maliyet_sekmesi_ve_uyari(self):
        self.maliyetli()
        h = self.al(urun=self.mamul_a.pk, miktar="10", sekme="maliyet").content.decode()
        self.assertIn("1 kalemin maliyeti yok — toplam eksik", h)
        self.assertIn("maliyet yok", h)
        self.assertIn("58,00 TL", h)
        self.assertIn("580,00 TL", h)

    def test_gecersiz_sekme_agaca_duser(self):
        r = self.al(urun=self.mamul_a.pk, sekme="x")
        self.assertEqual(r.context["sekme"], "agac")

    def test_excel_ciktilari(self):
        self.maliyetli()
        for sekme, baslik, ilk in (("agac", "Ürün Ağacı", "Seviye"), ("malzeme", "Malzeme ve Stok", "Kategori"), ("maliyet", "Maliyet", "Kategori")):
            r = self.al(urun=self.mamul_a.pk, miktar="10", sekme=sekme, xlsx="1")
            self.assertEqual(r.status_code, 200)
            self.assertIn("spreadsheetml", r["Content-Type"])
            self.assertIn(".xlsx", r["Content-Disposition"])
            ws = load_workbook(BytesIO(r.content)).active
            self.assertTrue(ws["A1"].value.startswith(baslik), ws["A1"].value)
            self.assertEqual(ws["A4"].value, ilk)
            self.assertGreater(ws.max_row, 5)
        ws = load_workbook(BytesIO(self.al(urun=self.mamul_a.pk, miktar="10", sekme="maliyet", xlsx="1").content)).active
        degerler = [c.value for row in ws.iter_rows(min_row=5) for c in row]
        self.assertIn("maliyet yok", degerler)
        self.assertTrue(any(isinstance(v, (int, float, Decimal)) and v == 580 for v in degerler))     # N adet genel toplam

    def test_excel_sonuc_yoksa_sayfa(self):
        r = self.al(xlsx="1")
        self.assertEqual(r.status_code, 200)
        self.assertIn("text/html", r["Content-Type"])

    def test_yazdir_ve_excel_dugmeleri(self):
        r = self.al(urun=self.mamul_a.pk, sekme="maliyet")
        self.assertContains(r, "xlsx=1")
        self.assertContains(r, "window.print()")

    def test_yetkisiz_kullanici(self):
        u = User.objects.create_user("ua_y", password="x")
        self.client.force_login(u)
        self.assertEqual(self.al().status_code, 403)
        EkranYetki.objects.create(kullanici=u, ekran_kod="urun_agaci")
        self.assertEqual(self.al().status_code, 200)

    def test_sorgu_sayisi_sabit(self):
        self.maliyetli()
        sayilar = {}
        for sekme in ("agac", "malzeme", "maliyet"):
            with CaptureQueriesContext(connection) as az:
                self.al(urun=self.mamul_a.pk, miktar="3", sekme=sekme)
            sayilar[sekme] = len(az)
        for i in range(8):                                                      # zinciri büyüt: yeni girdi, kesim, kategori
            kat = Kategori.objects.create(kod=f"K{i}", ad=f"KAT {i}")
            girdi = self._stok(f"EKS-G{i}", kat, satinalma=True, uretim=False)
            ara = self._stok(f"EKS-A{i}", kat)
            operasyon_olustur(istasyon_id=self.kesim.pk, cikti_id=ara.pk, cikti_miktar=D("2"), satirlar=[(girdi, D("1"))])
            self.giris(girdi, 5, 10)
        op = self.op_a
        from core.services.uretim import operasyon_guncelle
        operasyon_guncelle(op, istasyon_id=self.montaj.pk, cikti_miktar=D("1"),
                           satirlar=[(self.on21, D("1")), (self.civata, D("4")), (self.boya, D("1"))] +
                                    [(Stok.objects.get(kod=f"EKS-A{i}"), D("1")) for i in range(8)])
        for sekme in ("agac", "malzeme", "maliyet"):
            with CaptureQueriesContext(connection) as cok:
                r = self.al(urun=self.mamul_a.pk, miktar="3", sekme=sekme)
            self.assertEqual(r.status_code, 200)
            self.assertEqual(len(cok), sayilar[sekme], sekme)
            self.assertLessEqual(len(cok), 25)
