"""CARİ RAPORLARI Excel çıktıları (yalnız dizilim/biçim — hesaplar ``cari_rapor``da). Ekrandaki süzgeçlerle AYNI veri."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from core.services.cari_cikti import _baslik, doviz_metni, yon_metni
from core.services.cari_rapor import DILIMLER, SEKMELER

PARA = "#,##0.00;[Red]-#,##0.00"
TARIH = "DD.MM.YYYY"


def _kitap(baslik, aciklama):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = baslik
    ws["A1"].font = Font(bold=True, size=13, color="15294D")
    ws["A2"] = aciklama
    ws["A2"].font = Font(color="5B6678")
    return wb, ws


def _bayt(wb) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _bicim(ws, satir, sutunlar, bicim):
    for s in sutunlar:
        ws.cell(row=satir, column=s).number_format = bicim


def yaslandirma_xlsx(rapor) -> bytes:
    """Seçili sekmenin PB başına bir sayfası (ekrandaki tablolar)."""
    sekme_ad = dict(SEKMELER)[rapor["sekme"]]
    kolonlar = ("Cari Kodu", "Ünvan") + tuple(ad for _k, ad in DILIMLER) + ("Toplam", "Açık Fatura")
    genislik = (14, 44) + (16,) * len(DILIMLER) + (16, 11)
    wb = None
    for t in rapor["tablolar"] or [None]:
        if wb is None:
            wb, ws = _kitap(f"Yaşlandırma — {sekme_ad}", f"{rapor['tarih']:%d.%m.%Y} itibarıyla · açık fatura kalanları (ödemeler en eski faturadan düşülür) · PB'ler ayrı")
        else:
            ws = wb.create_sheet()
        ws.title = f"{sekme_ad.split()[0]} {t['pb']}" if t else "Yaşlandırma"
        if t is None:
            ws["A4"] = "Açık fatura yok."
            continue
        ws["A3"] = f"Para birimi: {t['pb']}"
        ws["A3"].font = Font(bold=True)
        _baslik(ws, 4, kolonlar, genislik)
        r = 5
        for s in t["satirlar"]:
            ws.append([s["cari"].kod, s["cari"].unvan] + [float(s["dilimler"][k]) for k, _ in DILIMLER] + [float(s["toplam"]), s["adet"]])
            _bicim(ws, r, range(3, 3 + len(DILIMLER) + 1), PARA)
            r += 1
        top = t["toplam"]
        ws.append(["", "TOPLAM"] + [float(top[k]) for k, _ in DILIMLER] + [float(top["toplam"]), top["adet"]])
        _bicim(ws, r, range(3, 3 + len(DILIMLER) + 1), PARA)
        for c in ws[r]:
            c.font = Font(bold=True)
    return _bayt(wb)


def vadesi_gecenler_xlsx(rapor) -> bytes:
    wb, ws = _kitap("Vadesi Geçen Açık Faturalar", f"{rapor['tarih']:%d.%m.%Y} itibarıyla · gecikme gününe göre")
    ws.title = "Vadesi Geçenler"
    _baslik(ws, 4, ("Cari Kodu", "Ünvan", "Sekme", "Fatura No", "Fatura Tarihi", "Vade", "Gecikme (gün)", "PB", "Kalan"), (14, 44, 18, 16, 13, 12, 13, 6, 16))
    r = 5
    sekme_ad = dict(SEKMELER)
    for s in rapor["satirlar"]:
        ws.append([s["cari"].kod, s["cari"].unvan, sekme_ad[s["sekme"]], s["fatura"].fatura_no, s["tarih"], s["vade"], s["gecikme"], s["pb"], float(s["kalan"])])
        _bicim(ws, r, (5, 6), TARIH)
        _bicim(ws, r, (9,), PARA)
        r += 1
    for pb, t in rapor["toplam"]:
        ws.append(["", "TOPLAM", "", "", "", "", "", pb, float(t)])
        _bicim(ws, r, (9,), PARA)
        for c in ws[r]:
            c.font = Font(bold=True)
        r += 1
    return _bayt(wb)


def bakiye_xlsx(rapor) -> bytes:
    wb, ws = _kitap("Tarih İtibarıyla Cari Bakiyeleri", f"{rapor['tarih']:%d.%m.%Y} itibarıyla · pozitif = Borç (cari bize borçlu), negatif = Alacak")
    ws.title = "Bakiyeler"
    _baslik(ws, 4, ("Cari Kodu", "Ünvan", "Kategori", "Bakiye TL", "Yön", "Döviz Bakiyeleri", "Son Hareket"), (14, 44, 24, 17, 8, 34, 13))
    r = 5
    for s in rapor["satirlar"]:
        c = s["cari"]
        ws.append([c.kod, c.unvan, c.kategori.ad if c.kategori_id else "", float(s["tl"]), yon_metni(s["tl"]), doviz_metni(s["doviz"]), s["son"]])
        _bicim(ws, r, (4,), PARA)
        _bicim(ws, r, (7,), TARIH)
        r += 1
    ws.append(["", "TOPLAM (TL borç bakiyeler)", "", float(rapor["tl_borc"])])
    ws.append(["", "TOPLAM (TL alacak bakiyeler)", "", float(rapor["tl_alacak"])])
    ws.append(["", "NET (TL)", "", float(rapor["tl_net"]), yon_metni(rapor["tl_net"])])
    for pb, borc, alacak, net in rapor["doviz_toplam"]:
        ws.append(["", f"{pb} — borç / alacak / net", "", float(borc), "", f"{alacak} / {net}"])
    for rr in range(r, ws.max_row + 1):
        _bicim(ws, rr, (4,), PARA)
        for c in ws[rr]:
            c.font = Font(bold=True)
    return _bayt(wb)


def ciro_xlsx(rapor) -> bytes:
    wb, ws = _kitap("Cari Ciro Özeti", f"{rapor['baslangic']:%d.%m.%Y} – {rapor['bitis']:%d.%m.%Y} · KDV hariç, TL · onaylı faturalar · iadeler ilgili taraftan düşülür")
    ws.title = "Ciro"
    _baslik(ws, 4, ("Cari Kodu", "Ünvan", "Kategori", "Satış (brüt)", "Satış İade", "Net Satış", "Satış Belge", "Alış (brüt)", "Alış İade", "Net Alış", "Alış Belge", "Toplam Ciro"),
            (14, 44, 24, 17, 15, 17, 11, 17, 15, 17, 11, 17))
    r = 5

    def satir(s, kod="", unvan="", kategori=""):
        return [kod, unvan, kategori, float(s["satis"]), float(s["satis_iade"]), float(s["satis_net"]), s["satis_belge"],
                float(s["alis"]), float(s["alis_iade"]), float(s["alis_net"]), s["alis_belge"], float(s["toplam"])]
    for s in rapor["satirlar"]:
        c = s["cari"]
        ws.append(satir(s, c.kod, c.unvan, c.kategori.ad if c.kategori_id else ""))
        _bicim(ws, r, (4, 5, 6, 8, 9, 10, 12), PARA)
        r += 1
    ws.append(satir(rapor["toplam"], "", "TOPLAM"))
    _bicim(ws, r, (4, 5, 6, 8, 9, 10, 12), PARA)
    for c in ws[r]:
        c.font = Font(bold=True)
    return _bayt(wb)
