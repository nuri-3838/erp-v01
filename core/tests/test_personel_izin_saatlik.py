"""SAATLİK izin (2026-10-12): form "Tam gün / Saatlik"; saatlikte tek gün + saat aralığı, süre dakika, bitiş > başlangıç; yıllık izinden gün
karşılığı (dakika / günlük çalışma saati — IkAyari, varsayılan 7,5) düşer, kalan "12 gün 3 sa" biçiminde; diğer türlerde de saatlik; ücretsiz
saatlik izin bordro ekranında saat bazlı kesinti önerisi ve puantaj Excel'inde sütun; devam/yoklama: saatlik izinli gün Geldi sayılır + rozet,
aylık özette saatlik izin satırı; mesai kaydıyla çakışma uyarısı (engel yok); liste/kart geçmişinde saat aralığı ve süre; mevcut kayıtlar tam gün."""
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from unittest import mock

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone as dj_tz
from openpyxl import load_workbook

from core.models import EkranYetki, IkAyari, MesaiKaydi, PersonelDevam, PersonelIzin
from core.services import personel_devam as devam_servis
from core.services import personel_izin as izin_servis
from core.services.personel_dokum import dokum_xlsx
from core.services.personel_izin import PersonelIzinHatasi, bakiye, gun_saat_metni, izin_ekle, izin_guncelle
from core.tarih import _TR
from core.tests.ik_yardimci import personel_kur

D = Decimal
GUN = date(2026, 6, 10)                                                                  # Çarşamba


class SaatlikTaban(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.p = personel_kur(ad="ali", soyad="veli", ise_giris_tarihi=date(2020, 3, 15))
        cls.yon = User.objects.create_superuser("szyon", password="x")

    def saatlik(self, bas="09:00", bit="11:30", tarih=GUN, tur="YILLIK", **kw):
        return izin_ekle(self.p, tur=tur, baslangic=tarih, saatlik=True, baslangic_saati=bas, bitis_saati=bit, kullanici=self.yon, **kw)


class ServisTest(SaatlikTaban):
    def test_kayit_sure_gun_karsiligi_ve_dogrulama(self):
        self.assertEqual(IkAyari.get().gunluk_calisma_saati, D("7.5"))
        i = self.saatlik()                                                                 # 2 sa 30 dk = 150 dk
        self.assertEqual((i.saatlik, i.baslangic, i.bitis, i.baslangic_saati, i.bitis_saati, i.dakika, i.gun),
                         (True, GUN, GUN, time(9, 0), time(11, 30), 150, D("0.333")))      # 150 / 450
        self.assertEqual(izin_servis.sure_metni(150), "2 sa 30 dk")
        self.assertEqual(izin_servis.sure_metni(45), "45 dk")
        self.assertEqual(izin_servis.sure_metni(180), "3 sa")
        with self.assertRaisesMessage(PersonelIzinHatasi, "Bitiş saati başlangıç saatinden sonra olmalı"):
            self.saatlik("13:00", "13:00", tarih=GUN + timedelta(days=1))
        with self.assertRaisesMessage(PersonelIzinHatasi, "bitiş saati zorunlu"):
            izin_ekle(self.p, tur="YILLIK", baslangic=GUN + timedelta(days=1), saatlik=True, baslangic_saati="09:00")
        with self.assertRaisesMessage(PersonelIzinHatasi, "geçerli değil"):
            self.saatlik("9h", "10:00", tarih=GUN + timedelta(days=1))
        # tam gün kayıt eski davranış: saatlik alanlar boş, gün yarım adımlı
        t = izin_ekle(self.p, tur="YILLIK", baslangic=GUN + timedelta(days=5), bitis=GUN + timedelta(days=6), kullanici=self.yon)
        self.assertEqual((t.saatlik, t.dakika, t.baslangic_saati, t.gun), (False, 0, None, D("2")))

    def test_gunluk_saat_ayari_ve_cevrim(self):
        izin_servis.gunluk_saat_ayarla("8", kullanici=self.yon)
        self.assertEqual(izin_servis.gunluk_saat(), D("8"))
        i = self.saatlik("08:00", "12:00")                                                 # 240 / 480
        self.assertEqual(i.gun, D("0.5"))
        izin_servis.gunluk_saat_ayarla("7,5")                                              # ayar değişince eski kayıt sabit kalır
        i.refresh_from_db()
        self.assertEqual(i.gun, D("0.5"))
        with self.assertRaisesMessage(PersonelIzinHatasi, "1 ile 24"):
            izin_servis.gunluk_saat_ayarla("0")
        with self.assertRaisesMessage(PersonelIzinHatasi, "geçerli bir sayı"):
            izin_servis.gunluk_saat_ayarla("abc")

    def test_bakiye_gun_ve_saat_metni(self):
        self.assertEqual(gun_saat_metni(D("12.4"), D("7.5")), "12 gün 3 sa")
        self.assertEqual(gun_saat_metni(D("12"), D("7.5")), "12 gün")
        self.assertEqual(gun_saat_metni(D("0.467"), D("7.5")), "3,5 sa")
        self.assertEqual(gun_saat_metni(D("-2.533"), D("7.5")), "-2 gün 4 sa")
        self.assertEqual(gun_saat_metni(D("0"), D("7.5")), "0 gün")
        once = bakiye(self.p, bugun=GUN).kalan
        self.saatlik("09:00", "12:00")                                                     # 3 sa = 0,4 gün
        b = bakiye(self.p, bugun=GUN)
        self.assertEqual((b.kullanilan, once - b.kalan, b.gunluk_saat), (D("0.4"), D("0.4"), D("7.5")))
        self.assertTrue(b.kalan_metni.endswith(" gün 4,5 sa") or b.kalan_metni.endswith("gün 4,5 sa"), b.kalan_metni)   # 0,6 gün × 7,5 = 4,5 sa
        self.saatlik("13:00", "15:00", tur="MAZERET")                                      # mazeret: bakiyeden düşmez
        self.assertEqual(bakiye(self.p, bugun=GUN).kullanilan, D("0.4"))

    def test_cakisma_kurallari(self):
        self.saatlik("09:00", "11:00")
        self.saatlik("13:00", "14:00")                                                     # aynı gün, kesişmeyen saat: serbest
        with self.assertRaisesMessage(PersonelIzinHatasi, "09:00–11:00 saatlik"):
            self.saatlik("10:30", "12:00")                                                 # saat kesişmesi
        with self.assertRaisesMessage(PersonelIzinHatasi, "saatlik"):
            izin_ekle(self.p, tur="YILLIK", baslangic=GUN, bitis=GUN + timedelta(days=1))  # tam gün, saatlik izinli güne değiyor
        izin_ekle(self.p, tur="YILLIK", baslangic=GUN + timedelta(days=1), bitis=GUN + timedelta(days=2))
        with self.assertRaisesMessage(PersonelIzinHatasi, "tarihli"):
            self.saatlik("09:00", "10:00", tarih=GUN + timedelta(days=2))                  # tam gün izinli güne saatlik
        i = PersonelIzin.objects.get(saatlik=True, baslangic_saati=time(9, 0))
        izin_guncelle(i, tur="YILLIK", baslangic=GUN, saatlik=True, baslangic_saati="08:30", bitis_saati="11:00", kullanici=self.yon)   # kendisiyle çakışmaz
        i.refresh_from_db()
        self.assertEqual((i.dakika, i.gun), (150, D("0.333")))
        izin_guncelle(i, tur="YILLIK", baslangic=GUN + timedelta(days=3), bitis=GUN + timedelta(days=3), kullanici=self.yon)   # saatlik → tam gün (Cumartesi)
        i.refresh_from_db()
        self.assertEqual((i.saatlik, i.dakika, i.baslangic_saati, i.gun), (False, 0, None, D("1")))

    def test_mesai_cakismasi_uyari_listesi(self):
        giris = datetime.combine(GUN, time(8, 0), tzinfo=_TR)
        m = MesaiKaydi.objects.create(personel=self.p, is_tarihi=GUN, giris_zamani=giris, cikis_zamani=giris + timedelta(hours=9))
        self.assertEqual([x.pk for x in izin_servis.mesai_cakismasi(self.p, GUN, time(9, 0), time(11, 0))], [m.pk])
        self.assertEqual(izin_servis.mesai_cakismasi(self.p, GUN, time(18, 0), time(19, 0)), [])
        self.assertEqual(izin_servis.mesai_cakismasi(self.p, GUN + timedelta(days=1), time(9, 0), time(11, 0)), [])

    def test_ucretsiz_saatlik_ozet(self):
        self.saatlik("09:00", "11:00", tur="UCRETSIZ")
        self.saatlik("14:00", "15:30", tur="UCRETSIZ", tarih=GUN + timedelta(days=1))
        self.saatlik("09:00", "10:00", tur="YILLIK", tarih=GUN + timedelta(days=2))
        o = izin_servis.ucretsiz_saatlik_ozet(2026, 6)
        self.assertEqual([(r["personel"].pk, r["dakika"], r["saat"], r["kayit_sayisi"], r["sure_metni"]) for r in o], [(self.p.pk, 210, D("3.5"), 2, "3 sa 30 dk")])
        self.assertEqual(izin_servis.ucretsiz_saatlik_ozet(2026, 7), [])


class DevamTest(SaatlikTaban):
    def setUp(self):
        self._m = mock.patch("core.services.personel_devam.tr_bugun", return_value=GUN + timedelta(days=20))
        self._m.start()
        self.addCleanup(self._m.stop)

    def test_gunluk_satir_geldi_ve_rozet_aylik_ozet(self):
        self.saatlik("09:00", "11:00")                                                     # 2 sa
        self.saatlik("14:00", "15:00", tur="UCRETSIZ")                                     # 1 sa ücretsiz
        s = next(x for x in devam_servis.gun_satirlari(GUN) if x.personel.pk == self.p.pk)
        self.assertEqual((s.kilitli, s.izin_turu, s.saatlik_dakika, s.saatlik_metni, s.durum), (False, None, 180, "3 sa", ""))
        kaydedilen, _, _ = devam_servis.gunluk_kaydet(GUN, {self.p.pk: ("GELDI", "")}, kullanici=self.yon)
        self.assertEqual(kaydedilen, 1)                                                    # saatlik izin günü yoklama girilir
        s = next(x for x in devam_servis.gun_satirlari(GUN) if x.personel.pk == self.p.pk)
        self.assertEqual(s.durum, "GELDI")
        izin_ekle(self.p, tur="YILLIK", baslangic=GUN + timedelta(days=1), bitis=GUN + timedelta(days=1), kullanici=self.yon)   # tam gün: kilitli
        s2 = next(x for x in devam_servis.gun_satirlari(GUN + timedelta(days=1)) if x.personel.pk == self.p.pk)
        self.assertEqual((s2.kilitli, s2.durum), (True, "IZINLI"))
        o = next(x for x in devam_servis.aylik_ozet(2026, 6) if x.personel.pk == self.p.pk)
        self.assertEqual((o.geldi, o.izinli, o.yillik, o.saatlik_dakika, o.ucretsiz_dakika, o.saatlik_saat, o.ucretsiz_saat),
                         (1, 1, 1, 180, 60, D("3"), D("1")))
        t = devam_servis.aylik_toplam(devam_servis.aylik_ozet(2026, 6))
        self.assertEqual((t.saatlik_dakika, t.ucretsiz_dakika), (180, 60))
        # ekranlar
        self.client.force_login(self.yon)
        r = self.client.get(reverse("core:yoklama") + f"?tarih={GUN.isoformat()}")
        self.assertContains(r, "⏱ 3 sa izin")
        self.assertContains(r, f'name="durum_{self.p.pk}"')                                 # kilitli değil: durum seçilir
        r = self.client.get(reverse("core:yoklama_aylik") + "?ay=2026-06")
        self.assertContains(r, "Saatlik İzin")
        self.assertContains(r, "Saatlik izin: 3,00 saat (ücretsiz: 1,00 saat)")
        ws = load_workbook(BytesIO(dokum_xlsx(2026, 6, bugun=GUN + timedelta(days=20)))).active
        basliklar = [c.value for c in ws[5]]
        self.assertEqual(basliklar[-2:], ["Saatlik İzin (sa)", "Ücretsiz Saatlik (sa)"])
        satir = next(r for r in ws.iter_rows(min_row=6, values_only=True) if r[1] == self.p.ad_soyad)
        self.assertEqual((satir[-2], satir[-1]), (D("3"), D("1")))


class EkranTest(SaatlikTaban):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.ik = User.objects.create_user("szik", password="x")
        for kod in ("personel", "personel_izinleri", "bordro"):
            EkranYetki.objects.create(kullanici=cls.ik, ekran_kod=kod)

    def test_form_saatlik_kayit_hata_liste_ve_kart(self):
        self.client.force_login(self.ik)
        r = self.client.get(reverse("core:izin_ekle"))
        self.assertContains(r, "Tam gün")
        self.assertContains(r, "Saatlik")
        self.assertContains(r, 'name="baslangic_saati"')
        govde = {"personel": self.p.pk, "tur": "YILLIK", "sure_turu": "SAAT", "baslangic": GUN.isoformat(), "baslangic_saati": "11:00", "bitis_saati": "09:00"}
        r = self.client.post(reverse("core:izin_ekle"), govde)
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Bitiş saati başlangıç saatinden sonra olmalı")
        self.assertFalse(PersonelIzin.objects.exists())
        giris = datetime.combine(GUN, time(8, 0), tzinfo=_TR)
        MesaiKaydi.objects.create(personel=self.p, is_tarihi=GUN, giris_zamani=giris, cikis_zamani=giris + timedelta(hours=9))
        govde.update(baslangic_saati="09:00", bitis_saati="11:30")
        r = self.client.post(reverse("core:izin_ekle"), govde, follow=True)
        i = PersonelIzin.objects.get()
        self.assertEqual((i.saatlik, i.dakika, i.gun, i.bitis), (True, 150, D("0.333"), GUN))
        mesajlar = [str(m) for m in r.context["messages"]]
        self.assertTrue(any("2 sa 30 dk" in m for m in mesajlar), mesajlar)
        self.assertTrue(any("mesai kaydıyla çakışıyor" in m and "08:00–17:00" in m for m in mesajlar), mesajlar)   # uyarı, engel yok
        r = self.client.get(reverse("core:izinler"))
        self.assertContains(r, "09:00–11:30")
        self.assertContains(r, "2 sa 30 dk")
        r = self.client.get(reverse("core:personel_detay", args=[self.p.pk]), {"sekme": "izin"})
        self.assertContains(r, "09:00–11:30")
        self.assertContains(r, "2 sa 30 dk")
        self.assertContains(r, "gün 5 sa")                                                 # kalan: … gün 5 sa (7,5 − 2,5)
        r = self.client.get(reverse("core:izin_duzenle", args=[i.pk]))
        self.assertRegex(r.content.decode(), r'value="SAAT"[^>]*checked')
        self.assertContains(r, 'value="09:00"')
        # tam gün kayıt: saat sütunları boş
        self.client.post(reverse("core:izin_ekle"), {"personel": self.p.pk, "tur": "YILLIK", "sure_turu": "GUN", "baslangic": (GUN + timedelta(days=5)).isoformat(),
                                                     "bitis": (GUN + timedelta(days=5)).isoformat()})
        self.assertEqual(PersonelIzin.objects.filter(saatlik=False).count(), 1)
        r = self.client.get(reverse("core:izin_bakiyeleri"))
        self.assertContains(r, "gün 5 sa")

    def test_ayar_ekrani_ve_bordro_paneli(self):
        self.client.force_login(self.yon)
        r = self.client.get(reverse("core:mesai_ayarlari"))
        self.assertContains(r, "Günlük Çalışma Süresi")
        self.assertContains(r, "7,50 saat = 1 gün")
        r = self.client.post(reverse("core:ik_ayari_kaydet"), {"gunluk_calisma_saati": "8"}, follow=True)
        self.assertEqual(IkAyari.get().gunluk_calisma_saati, D("8"))
        self.assertContains(r, "8,00 saat = 1 gün")
        self.client.force_login(self.ik)
        self.assertEqual(self.client.post(reverse("core:ik_ayari_kaydet"), {"gunluk_calisma_saati": "9"}).status_code, 403)   # yalnız yönetici
        self.saatlik("09:00", "12:00", tur="UCRETSIZ")
        r = self.client.get(reverse("core:bordro_ekle") + "?yil=2026&ay=6")
        self.assertEqual(r.status_code, 200)
        r = self.client.post(reverse("core:bordro_ekle"), {"yil": "2026", "ay": "6", "tahakkuk_tarihi": "2026-06-30", "form-TOTAL_FORMS": "0",
                                                           "form-INITIAL_FORMS": "0", "form-MIN_NUM_FORMS": "0", "form-MAX_NUM_FORMS": "1000"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Ücretsiz Saatlik İzinler — 6/2026")
        self.assertContains(r, self.p.ad_soyad)
        self.assertContains(r, "3 sa")
        self.assertContains(r, "8,00 sa")
