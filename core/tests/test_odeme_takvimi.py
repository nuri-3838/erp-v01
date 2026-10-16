"""Finans › Raporlar › Ödeme Takvimi (2026-10-16): kaynaklar (alış/satış faturası açık kalanı — listedeki FIFO, faturasız cari borcu, KK nakit akışı, verilen/alınan çek, vergi/SGK yasal vade,
emanet vadesi, POS bloke), "5G kuralı" (cariden gelen para faturayı kapatmaz), gruplama (gecikmiş / gün-hafta-ay / vadesiz), kümülatif bakiye, TL (bugünkü kur), mutabakat, Excel, eski adres yönlendirmesi,
yetki, SALT-OKUNUR (sayfa açılırken yazma yok). Tarih bağımlılığı yok: tüm hesaplar ``bugun`` parametresiyle sabitlenir."""
import datetime
from decimal import Decimal
from io import BytesIO

from django.db import connection
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Cari, CekSenet, Kur
from core.services import fatura_odeme, odeme_takvimi as ot
from core.services.yevmiye import SatirGirdi, fis_olustur
from core.tests import test_fatura_liste as tfl
from core.tests.test_fatura import _hesap

D = datetime.date
Dc = Decimal
BUGUN = D(2026, 3, 25)


def _satirlar(bugun=BUGUN, **kw):
    return ot.tum_satirlar(bugun)["satirlar"]


def _bul(satirlar, **kw):
    return [r for r in satirlar if all(r.get(k) == v for k, v in kw.items())]


class TakvimTaban(tfl.ListeTaban):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        _hesap("360.20", "ÖDENECEK GELİR VERGİSİ STOPAJI", kalem="KVYK")
        _hesap("360.10.0710", "7/10 TEVKİFAT ÖDENECEK VERGİ", kalem="KVYK")
        _hesap("361", "ÖDENECEK SOSYAL GÜVENLİK KESİNTİLERİ", kalem="KVYK")
        _hesap("336.10.0001", "EMANET KİŞİ", kalem="KVYK")
        _hesap("108.10", "HALKÖDE POS")
        for g in (D(2026, 4, 1), D(2026, 4, 5), D(2026, 4, 30)):
            Kur.objects.get_or_create(tarih=g, defaults={"usd_alis": Dc("30")})

    def vergi_fisi(self, hesap, tutar, tarih):
        fis_olustur(tarih=tarih, satirlar=[SatirGirdi(hesap_kodu=hesap, taraf="A", islem_tutari=str(tutar)), SatirGirdi(hesap_kodu="100.01", taraf="B", islem_tutari=str(tutar))])


class FaturaKaynaklariTest(TakvimTaban):
    def test_alis_faturasi_acik_kalan_vadede_ve_link(self):
        f = self.a_fatura(no="TK-1", vade_tarihi=D(2026, 4, 10))
        (r,) = _bul(_satirlar(), kaynak="FATURA_ALIS")
        self.assertEqual((r["tarih"], r["yon"], r["tutar"], r["tl"], r["belge_no"], r["pb"]), (D(2026, 4, 10), "CIKIS", Dc("1200.00"), Dc("1200.00"), "TK-1", "TRY"))
        self.assertEqual(r["url"], ("core:fatura_detay", f.pk))
        self.assertEqual(r["cari_id"], self.tedarikci.pk)

    def test_kismi_odeme_kalan_kadar_kapanan_satir_yok(self):
        f = self.a_fatura(no="TK-2", vade_tarihi=D(2026, 4, 10))
        self.odeme("500")
        (r,) = _bul(_satirlar(), kaynak="FATURA_ALIS")
        self.assertEqual(r["tutar"], Dc("700.00"))                                         # listedeki FIFO kalanı
        self.assertEqual(r["tutar"], fatura_odeme.odeme_durumlari([self.tedarikci], BUGUN)[f.pk]["kalan"])
        self.odeme("700")
        self.assertEqual(_bul(_satirlar(), kaynak="FATURA_ALIS"), [])

    def test_satis_faturasi_giris(self):
        from core.services.fatura import fatura_olustur
        self.a_fatura(no="ST-0", miktar="100", fiyat="10")                                  # stok girişi
        f = fatura_olustur(tip_id=self.satis.pk, cari_id=self.musteri.pk, tarih=D(2026, 3, 10), fatura_no="TS-1", vade_tarihi=D(2026, 4, 5), satirlar=self._satir("5", "100"), depo_id=None)
        (r,) = _bul(_satirlar(), kaynak="FATURA_SATIS")
        self.assertEqual((r["yon"], r["tarih"], r["tutar"], r["belge_no"]), ("GIRIS", D(2026, 4, 5), Dc("600.00"), "TS-1"))
        self.assertEqual(r["url"], ("core:fatura_detay", f.pk))

    def test_doviz_fatura_bugunku_kurla_tl(self):
        from core.services.fatura import fatura_olustur
        fatura_olustur(tip_id=self.alis.pk, cari_id=self.usd.pk, tarih=D(2026, 3, 10), fatura_no="TK-USD", para_birimi="USD", vade_tarihi=D(2026, 4, 20), satirlar=self._satir("10", "100"))
        v = ot.tum_satirlar(BUGUN)
        (r,) = _bul(v["satirlar"], kaynak="FATURA_ALIS")
        self.assertEqual((r["pb"], r["tutar"]), ("USD", Dc("1200.00")))
        self.assertEqual(r["tl"], (Dc("1200.00") * v["kurlar"]["USD"]).quantize(Dc("0.01")))      # bugünkü (son TCMB alış) kur

    def test_vade_yoksa_vadesiz(self):
        self.a_fatura(no="TK-3")
        (r,) = _bul(_satirlar(), kaynak="FATURA_ALIS")
        self.assertIsNone(r["tarih"])
        self.assertIsNone(r["tarih"])


class BesGKuraliTest(TakvimTaban):
    def test_cariden_gelen_para_faturayi_kapatmis_saymaz(self):
        """5G vakası: cariden gelen (alacak) ve geri giden (borç) banka hareketi birbirini götürür; fatura açık kalır."""
        f = self.a_fatura(no="5G-1", vade_tarihi=D(2026, 11, 30), miktar="10", fiyat="100")
        from core.services.kasa_hareket import hareket_olustur
        hareket_olustur(kasa=self.kasa, tip="cari_tahsilat", karsi=self.tedarikci, tutar="1200", tarih=D(2026, 3, 11))          # cariden gelen
        self.odeme("1200", D(2026, 3, 12))                                                                                       # geri giden
        liste = fatura_odeme.odeme_durumlari([self.tedarikci], BUGUN)[f.pk]
        self.assertEqual(liste["durum"], "ACIK")                                                                                 # liste FIFO'su ile takvim AYNI (kural fatura_odeme.fifo_kalemler içinde)
        # kasa hareketi KASA kaynaklı → nakit iade kuralı (BANKA/KASA/KREDI_KARTI/CEK_SENET): takvimde de fatura açık
        (r,) = _bul(_satirlar(), kaynak="FATURA_ALIS")
        self.assertEqual((r["tutar"], r["tarih"]), (Dc("1200.00"), D(2026, 11, 30)))
        self.assertEqual(_bul(_satirlar(), kaynak="CARI_BORC"), [])


class FaturasizCariBorcuTest(TakvimTaban):
    def test_faturasiz_alacak_vadesiz_ve_odeme_once_faturaya(self):
        _hesap("120.99", "DİĞER")
        self.a_fatura(no="FB-1", miktar="10", fiyat="10", vade_tarihi=D(2026, 4, 1))                                           # 120 TL
        fis_olustur(tarih=D(2026, 3, 12), satirlar=[SatirGirdi(hesap_kodu="120.99", taraf="B", islem_tutari="1000"), SatirGirdi(hesap_kodu="320.10.0001", taraf="A", islem_tutari="1000")], kaynak="FASON_TAHAKKUK")
        self.odeme("120", D(2026, 3, 15))                                                                                       # faturayı kapatır
        s = _satirlar()
        self.assertEqual(_bul(s, kaynak="FATURA_ALIS"), [])
        (r,) = _bul(s, kaynak="CARI_BORC")
        self.assertEqual((r["tarih"], r["tutar"], r["cari_id"]), (None, Dc("1000.00"), self.tedarikci.pk))
        self.odeme("300", D(2026, 3, 16))                                                                                       # artan ödeme tahakkuktan düşer
        (r,) = _bul(_satirlar(), kaynak="CARI_BORC")
        self.assertEqual(r["tutar"], Dc("700.00"))


class VergiSgkTest(TakvimTaban):
    def test_yasal_vade_kurallari_tatil_kaydirmali(self):
        self.assertEqual(ot.yasal_vade("kdv", 2026, 3), D(2026, 4, 28))
        self.assertEqual(ot.yasal_vade("muhtasar", 2026, 3), D(2026, 4, 27))                 # 26.04.2026 pazar → pazartesi
        self.assertEqual(ot.yasal_vade("sgk", 2026, 3), D(2026, 4, 30))
        self.assertEqual(ot.yasal_vade("sgk", 2026, 9), D(2026, 11, 2))                      # 31.10.2026 cumartesi → pazartesi

    def test_tahakkuk_odeme_dusumu_ve_kural_etiketi(self):
        self.vergi_fisi("360.20", "3068.77", D(2026, 3, 31))
        self.vergi_fisi("361", "27823.84", D(2026, 3, 31))
        self.vergi_fisi("360.10.0710", "556288.66", D(2026, 3, 31))
        s = _bul(_satirlar(), kaynak="VERGI")
        d = {r["baslik"]: r for r in s}
        self.assertEqual(len(s), 3)
        self.assertTrue(all(r["kural"] for r in s))
        self.assertEqual((d["Muhtasar — gelir vergisi stopajı (360.20)"]["tarih"], d["Muhtasar — gelir vergisi stopajı (360.20)"]["tutar"]), (D(2026, 4, 27), Dc("3068.77")))
        self.assertEqual(d["KDV tevkifatı / KDV-2 (360.10)"]["tarih"], D(2026, 4, 28))
        self.assertEqual(d["SGK prim borcu (361)"]["tarih"], D(2026, 4, 30))
        # ödeme tek alt hesaba yazılsa da grup olarak en eski aydan düşer
        fis_olustur(tarih=D(2026, 4, 5), satirlar=[SatirGirdi(hesap_kodu="361", taraf="B", islem_tutari="823.84"), SatirGirdi(hesap_kodu="100.01", taraf="A", islem_tutari="823.84")])
        sgk = _bul(_satirlar(), kaynak="VERGI", baslik="SGK prim borcu (361)")
        self.assertEqual(sgk[0]["tutar"], Dc("27000.00"))

    def test_mutabakat_360_361_fark_sifir(self):
        self.vergi_fisi("360.20", "1000", D(2026, 3, 31))
        self.vergi_fisi("361", "2000", D(2026, 3, 31))
        v = ot.tum_satirlar(BUGUN)
        blok = next(b for b in ot.mutabakat(v) if b["ad"].startswith("360 / 361"))
        self.assertEqual((blok["defter"], blok["takvim"], blok["fark"]), (Dc("3000.00"), Dc("3000.00"), Dc("0.00")))


class DigerKaynaklarTest(TakvimTaban):
    def test_emanet_vadesi_ve_usd_bugunku_kur(self):
        c = Cari.objects.create(kod="336-10-0001", unvan="EMANET KİŞİ", para_birimi="USD", muhasebe_kodu="336.10.0001", notlar="EMANET (336.10) - 20.000 USD - GERİ ÖDEME VADESİ: 31.10.2026")
        self.assertEqual(ot.emanet_vadesi(c.notlar), D(2026, 10, 31))
        self.assertIsNone(ot.emanet_vadesi("vade yazılmamış"))
        fis_olustur(tarih=D(2026, 3, 10), satirlar=[SatirGirdi(hesap_kodu="100.01", taraf="B", islem_tutari="20000", islem_pb="USD", islem_kuru="30"),
                                                    SatirGirdi(hesap_kodu="336.10.0001", taraf="A", islem_tutari="20000", islem_pb="USD", islem_kuru="30")])
        v = ot.tum_satirlar(BUGUN)
        (r,) = _bul(v["satirlar"], kaynak="EMANET")
        self.assertEqual((r["tarih"], r["tutar"], r["pb"]), (D(2026, 10, 31), Dc("20000.00"), "USD"))
        self.assertEqual(r["tl"], (Dc("20000") * v["kurlar"]["USD"]).quantize(Dc("0.01")))

    def test_verilen_ve_alinan_cek(self):
        CekSenet.objects.create(tip="CEK", yon="VERILEN", tutar=Dc("790000"), para_birimi="TRY", vade=D(2026, 12, 31), durum="VERILDI", belge_no="CK-1", cari=self.tedarikci)
        CekSenet.objects.create(tip="CEK", yon="ALINAN", tutar=Dc("5000"), para_birimi="TRY", vade=D(2026, 4, 10), durum="PORTFOYDE", belge_no="CK-2", cari=self.musteri)
        CekSenet.objects.create(tip="CEK", yon="VERILEN", tutar=Dc("100"), para_birimi="TRY", vade=D(2026, 2, 1), durum="ODENDI", belge_no="CK-3", cari=self.tedarikci)
        s = _satirlar()
        (v,) = _bul(s, kaynak="CEK_VERILEN")
        (a,) = _bul(s, kaynak="CEK_ALINAN")
        self.assertEqual((v["yon"], v["tarih"], v["tutar"], v["belge_no"]), ("CIKIS", D(2026, 12, 31), Dc("790000.00"), "CK-1"))
        self.assertEqual((a["yon"], a["tarih"], a["tutar"]), ("GIRIS", D(2026, 4, 10), Dc("5000.00")))                              # ödenmiş çek yok

    def test_pos_bloke_giris(self):
        fis_olustur(tarih=D(2026, 3, 20), satirlar=[SatirGirdi(hesap_kodu="108.10", taraf="B", islem_tutari="43000"), SatirGirdi(hesap_kodu="100.01", taraf="A", islem_tutari="43000")])
        (r,) = _bul(_satirlar(), kaynak="POS")
        self.assertEqual((r["yon"], r["tarih"], r["tutar"]), ("GIRIS", BUGUN, Dc("43000.00")))


class GruplamaTest(TakvimTaban):
    def test_gecikmis_en_ustte_kumulatif_ve_vadesiz_dahil_degil(self):
        self.a_fatura(no="G-1", vade_tarihi=D(2026, 3, 10))                         # gecikmiş
        self.a_fatura(no="G-2", tarih=D(2026, 3, 11), vade_tarihi=D(2026, 4, 5))
        self.a_fatura(no="G-3", tarih=D(2026, 3, 12))                               # vadesiz
        v = ot.tum_satirlar(BUGUN)
        g = ot.grupla(v["satirlar"], bugun=BUGUN, bas=BUGUN, bit=D(2026, 9, 25), grup="ay", acilis=Dc("1000"))
        self.assertEqual([x["ad"] for x in g["gruplar"]], ["Gecikmiş", "04.2026"])
        self.assertEqual((g["gruplar"][0]["cikis"], g["gruplar"][0]["kumulatif"]), (Dc("1200.00"), Dc("-200.00")))
        self.assertTrue(g["gruplar"][0]["negatif"])
        self.assertEqual(g["gruplar"][1]["kumulatif"], Dc("-1400.00"))
        self.assertEqual(g["vadesiz"]["cikis"], Dc("1200.00"))                         # kümülatife dahil değil
        self.assertEqual((g["cikis"], g["net_ihtiyac"], g["donem_sonu"]), (Dc("2400.00"), Dc("1400.00"), Dc("-1400.00")))

    def test_gun_hafta_ay_anahtarlari(self):
        self.assertEqual(ot._grup_anahtari(D(2026, 3, 25), "gun")[1], "25.03.2026")
        self.assertEqual(ot._grup_anahtari(D(2026, 3, 25), "hafta")[1], "23.03 – 29.03.2026")
        self.assertEqual(ot._grup_anahtari(D(2026, 3, 25), "ay")[1], "03.2026")

    def test_filtreler(self):
        self.a_fatura(no="F-1", vade_tarihi=D(2026, 4, 5))
        self.vergi_fisi("361", "100", D(2026, 3, 31))
        s = _satirlar()
        self.assertEqual({r["kaynak"] for r in ot.filtrele(s, kaynaklar=["VERGI"])}, {"VERGI"})
        self.assertEqual({r["yon"] for r in ot.filtrele(s, yon="GIRIS")}, set())
        self.assertEqual({r["cari_id"] for r in ot.filtrele(s, cari_id=self.tedarikci.pk)}, {self.tedarikci.pk})


class EkranTest(TakvimTaban):
    def get(self, **q):
        self.client.force_login(self.su)
        yazma = []

        def engel(execute, sql, params, many, context):
            if sql.lstrip()[:6].upper() in ("INSERT", "UPDATE", "DELETE"):
                yazma.append(sql[:80])
            return execute(sql, params, many, context)
        with connection.execute_wrapper(engel):
            r = self.client.get(reverse("core:odeme_takvimi"), q)
        self.assertEqual(yazma, [], "Sayfa açılırken veritabanına YAZMA denendi")
        return r

    def test_sayfa_acilir_kartlar_ve_yazma_yok(self):
        self.a_fatura(no="E-1", vade_tarihi=D(2026, 4, 5))
        r = self.get(bas="2026-01-01", bit="2026-12-31", grup="ay")
        self.assertEqual(r.status_code, 200)
        for metin in ("Ödeme Takvimi", "Bugünkü banka + kasa", "Dönem çıkışı", "Dönem girişi", "Net ihtiyaç", "Mutabakat", "bugünkü kurla"):
            self.assertContains(r, metin)
        self.assertContains(r, "E-1")
        self.assertContains(r, "excel=1")                                                  # Excel bağlantısı

    def test_excel_ciktisi(self):
        self.a_fatura(no="E-2", vade_tarihi=D(2026, 4, 5))
        r = self.get(excel="1", bas="2026-01-01", bit="2026-12-31")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Ödeme Takvimi", "Özet ve Mutabakat"])
        basliklar = [c.value for c in wb["Ödeme Takvimi"][4]]
        self.assertEqual(basliklar[:6], ["Grup", "Tarih", "Yön", "Kaynak", "Cari / Açıklama", "Belge No"])

    def test_eski_kk_adresi_yonlenir_ve_menu(self):
        self.client.force_login(self.su)
        r = self.client.get("/finans/kredi-karti/nakit-akisi/")
        self.assertEqual(r.status_code, 301)
        self.assertEqual(r["Location"], reverse("core:kredi_karti_nakit_akisi"))
        self.assertEqual(reverse("core:kredi_karti_nakit_akisi"), "/finans/raporlar/kart-nakit-akisi/")
        self.assertEqual(self.client.get(reverse("core:kredi_karti_nakit_akisi")).status_code, 200)
        from core.moduller import MODULLER
        finans = next(m for m in MODULLER if m.kod == "FINANS")
        grup = next(o for o in finans.ogeler if getattr(o, "alt_menu", False))
        self.assertEqual((grup.ad, [e.kod for e in grup.ekranlar]), ("Raporlar", ["odeme_takvimi", "kk_nakit_akisi"]))

    def test_yetkisiz_403(self):
        from django.contrib.auth.models import User
        u = User.objects.create_user("yetkisiz", password="x")
        self.client.force_login(u)
        self.assertEqual(self.client.get(reverse("core:odeme_takvimi")).status_code, 403)


class DamgaTahminiTest(TakvimTaban):
    """Damga vergisi: tahakkuk hesabı YOK, ödeme günü 770.11'e yazılır → son ödenen dönemin tutarı sonraki aylar için "tahmini" satır olur; gerçek ödeme girilince kendiliğinden düşer."""

    BUGUN = D(2026, 10, 16)

    def setUp(self):
        _hesap("770.11", "DAMGA VERGİSİ")
        for g in (D(2026, 2, 20), D(2026, 8, 18), D(2026, 9, 22), D(2026, 9, 23), D(2026, 10, 15), D(2026, 10, 16)):
            Kur.objects.get_or_create(tarih=g, defaults={"usd_alis": Dc("30")})

    def damga(self, tur, donem, tutar, tarih=D(2026, 9, 22)):
        fis_olustur(tarih=tarih, satirlar=[SatirGirdi(hesap_kodu="770.11", taraf="B", islem_tutari=str(tutar), aciklama=f"BEYANNAME DAMGA VERGİSİ {tur} {donem}"),
                                           SatirGirdi(hesap_kodu="100.01", taraf="A", islem_tutari=str(tutar))])

    def damgalar(self):
        return [r for r in ot.tum_satirlar(self.BUGUN)["satirlar"] if r["ek"] == "Damga vergisi (tahmini)"]

    def test_muhtasar_26sinda_kdv_28inde_son_odenen_tutarla(self):
        self.damga("MUHTASAR", "07/2026", "900.00", D(2026, 8, 18))
        self.damga("MUHTASAR", "08/2026", "939.70")
        self.damga("KDV1", "08/2026", "791.00", D(2026, 9, 23))
        s = self.damgalar()
        muh = [r for r in s if r["baslik"].startswith("Damga vergisi — muhtasar")]
        kdv = [r for r in s if r["baslik"].startswith("Damga vergisi — KDV")]
        self.assertEqual(muh[0]["tarih"], D(2026, 10, 26))                                   # 09/2026 dönemi, ayın 26'sı (pazartesi)
        self.assertEqual(muh[0]["tutar"], Dc("939.70"))                                      # EN SON ödenen dönem (08/2026), 07/2026 değil
        self.assertEqual(kdv[0]["tarih"], D(2026, 10, 28))
        self.assertEqual(kdv[0]["tutar"], Dc("791.00"))
        self.assertTrue(all(r["tahmini"] and r["kural"] and r["kaynak"] == "VERGI" and r["yon"] == "CIKIS" for r in s))
        self.assertEqual((muh[1]["tarih"], kdv[1]["tarih"]), (D(2026, 11, 26), D(2026, 11, 30)))   # 28.11 cumartesi → pazartesi; her ay bir satır
        self.assertEqual(len(muh), 12)

    def test_gercek_odeme_girilince_tahmini_satir_kendiliginden_duser(self):
        self.damga("MUHTASAR", "08/2026", "939.70")
        self.assertEqual(self.damgalar()[0]["aciklama"].split(" · ")[0], "09/2026 dönemi")
        self.damga("MUHTASAR", "09/2026", "960.00", D(2026, 10, 15))                         # Eylül muhtasar damgası ödendi
        s = self.damgalar()
        self.assertEqual((s[0]["aciklama"].split(" · ")[0], s[0]["tarih"], s[0]["tutar"]), ("10/2026 dönemi", D(2026, 11, 26), Dc("960.00")))   # sonraki dönem, yeni tutar
        self.assertFalse([r for r in s if r["tarih"] == D(2026, 10, 26)])

    def test_odeme_yoksa_veya_kdv2_ise_tahmin_yok(self):
        self.assertEqual(self.damgalar(), [])                                                # hiç ödeme yok
        self.damga("KDV2", "08/2026", "790.22")                                              # KDV2 yalnız tevkifatlı aylarda — tahmin edilmez
        self.assertEqual(self.damgalar(), [])

    def test_cok_eski_donem_tahmin_edilmez(self):
        self.damga("MUHTASAR", "01/2026", "939.70", D(2026, 2, 20))                          # 02-07/2026 kayıt girilmemiş: 60 günden eski vadeler atlanır
        s = self.damgalar()
        self.assertGreaterEqual(min(r["tarih"] for r in s), self.BUGUN - datetime.timedelta(days=60))

    def test_mutabakat_360_361_etkilenmez_ve_ekranda_etiket(self):
        self.damga("MUHTASAR", "08/2026", "939.70")
        veri = ot.tum_satirlar(self.BUGUN)
        for b in ot.mutabakat(veri, veri["atlanan_kartlar"]):
            if b["ad"].startswith("360"):
                self.assertEqual(b["fark"], Dc("0.00"))
        self.client.force_login(self.su)
        r = self.client.get(reverse("core:odeme_takvimi"), {"bas": "2026-10-16", "bit": "2027-03-31"})
        self.assertContains(r, "Damga vergisi — muhtasar (tahmini)")
        self.assertContains(r, ">tahmini<")
        x = self.client.get(reverse("core:odeme_takvimi"), {"bas": "2026-10-16", "bit": "2027-03-31", "excel": "1"})
        ws = load_workbook(BytesIO(x.content))["Ödeme Takvimi"]
        etiketler = [c.value for row in ws.iter_rows(min_row=5) for c in row if c.value == "tahmini"]
        self.assertTrue(etiketler)
