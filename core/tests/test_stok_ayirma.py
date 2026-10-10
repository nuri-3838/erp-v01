"""Üretim siparişi STOK AYIRMA altyapısı (plan adım 1 — docs/uretim-siparisi-plan.md): ``StokAyirma`` modeli (sipariş×stok tek aktif satır,
miktar ≥ 0), ``stok_ayirma.ayirma_ayarla/ayirma_ekle`` (tek yazma noktası; 0 → satır kapanır), kullanılabilir = eldeki − açık siparişlerin ayrılanı
(EKSİ olabilir — yumuşak kural, hiçbir çıkış engellenmez; ``haric_emir`` kendi ayırmasını sayılmaz), ürün ağacı Malzeme ve Stok sekmesinde Ayrılan /
Kullanılabilir / kullanılabilirden Eksik sütunları (servis, ekran, Excel). Eşdeğerlik referansı değişmez (ayırma yokken rakamlar aynı)."""
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.db import IntegrityError, transaction
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Depo, StokAyirma, StokHareket, UretimEmri
from core.services import stok_ayirma as sa
from core.services import urun_agaci as ua
from core.services.hareket import hareket_ekle
from core.services.uretim import uretim_emri_olustur
from core.services.urun_agaci_xlsx import malzeme_xlsx
from core.tests.test_urun_agaci_gorunum import GorunumBase

D = Decimal


class AyirmaBase(GorunumBase):
    def setUp(self):
        self.emir1 = uretim_emri_olustur(kalemler=[{"hedef_urun_id": self.mamul_a.pk, "hedef_miktar": "1"}], depo_id=self.depo.pk, tarih=date(2026, 10, 1))
        self.emir2 = uretim_emri_olustur(kalemler=[{"hedef_urun_id": self.mamul_c.pk, "hedef_miktar": "1"}], depo_id=self.depo.pk, tarih=date(2026, 10, 1))


class ServisTest(AyirmaBase):
    def test_ayarla_ekle_sifirla_negatif(self):
        a = sa.ayirma_ayarla(self.emir1, self.profil, "4")
        self.assertEqual((a.uretim_emri_id, a.stok_id, a.miktar, a.silindi), (self.emir1.pk, self.profil.pk, D("4"), False))
        b = sa.ayirma_ayarla(self.emir1.pk, self.profil.pk, D("2.5"))                      # pk ile de çağrılır; aynı satır güncellenir
        self.assertEqual((b.pk, b.miktar), (a.pk, D("2.5")))
        self.assertEqual(sa.ayrilan_miktar(self.emir1, self.profil), D("2.5"))
        c = sa.ayirma_ekle(self.emir1, self.profil, D("1.5"))
        self.assertEqual((c.pk, c.miktar), (a.pk, D("4")))
        self.assertIsNone(sa.ayirma_ekle(self.emir1, self.profil, D("-9")))                 # sıfırın altına inmez → satır kapanır
        self.assertEqual(sa.ayrilan_miktar(self.emir1, self.profil), D("0"))
        self.assertTrue(StokAyirma.objects.get(pk=a.pk).silindi)
        d = sa.ayirma_ayarla(self.emir1, self.profil, "3")                                  # kapanmış satır yerine yeni aktif satır
        self.assertNotEqual(d.pk, a.pk)
        self.assertEqual(StokAyirma.objects.filter(uretim_emri=self.emir1, stok=self.profil, silindi=False).count(), 1)
        self.assertIsNone(sa.ayirma_ayarla(self.emir2, self.profil, 0))                     # hiç satır yokken 0 → hiçbir şey yazılmaz
        with self.assertRaises(sa.AyirmaHatasi):
            sa.ayirma_ayarla(self.emir1, self.profil, "-1")
        self.assertEqual([x.stok.kod for x in sa.emir_ayirmalari(self.emir1)], ["PROFIL"])

    def test_tek_aktif_satir_kisiti(self):
        sa.ayirma_ayarla(self.emir1, self.profil, "1")
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                StokAyirma.objects.create(uretim_emri=self.emir1, stok=self.profil, miktar=D("2"))
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                StokAyirma.objects.create(uretim_emri=self.emir2, stok=self.profil, miktar=D("-1"))
        StokAyirma.objects.create(uretim_emri=self.emir2, stok=self.profil, miktar=D("2"))   # başka sipariş: serbest

    def test_kullanilabilir_toplam_haric_ve_eksi(self):
        self.giris(self.profil, 10, 1000)
        sa.ayirma_ayarla(self.emir1, self.profil, "4")
        sa.ayirma_ayarla(self.emir2, self.profil, "3")
        k = sa.toplu_kullanilabilir([self.profil.pk, self.civata.pk])
        self.assertEqual(k[self.profil.pk], {"eldeki": D("10"), "ayrilan": D("7"), "kullanilabilir": D("3")})
        self.assertEqual(k[self.civata.pk], {"eldeki": D("0"), "ayrilan": D("0"), "kullanilabilir": D("0")})
        self.assertEqual(sa.toplu_kullanilabilir([self.profil.pk], haric_emir=self.emir1)[self.profil.pk]["kullanilabilir"], D("7"))
        self.assertEqual(sa.kullanilabilir(self.profil, haric_emir=self.emir2.pk), D("6"))
        self.assertEqual(sa.ayirma_haritasi([]), {})
        # silinmiş ayırma ve silinmiş sipariş sayılmaz
        sa.ayirma_ayarla(self.emir2, self.profil, 0)
        self.assertEqual(sa.kullanilabilir(self.profil), D("6"))
        UretimEmri.objects.filter(pk=self.emir1.pk).update(silindi=True)
        self.assertEqual(sa.kullanilabilir(self.profil), D("10"))
        UretimEmri.objects.filter(pk=self.emir1.pk).update(silindi=False)
        for durum in (UretimEmri.Durum.KAPALI, UretimEmri.Durum.IPTAL):              # kapalı / iptal sipariş de sayılmaz
            UretimEmri.objects.filter(pk=self.emir1.pk).update(durum=durum)
            self.assertEqual((sa.kullanilabilir(self.profil), sa.ayrilan_miktar(self.emir1, self.profil)), (D("10"), D("0")), durum)
        UretimEmri.objects.filter(pk=self.emir1.pk).update(durum=UretimEmri.Durum.ACIK)
        self.assertEqual(sa.kullanilabilir(self.profil), D("6"))
        # yumuşak kural: ayrılmış stok başka çıkışla tüketilebilir → kullanılabilir EKSİ
        hareket_ekle(stok_id=self.profil.pk, depo_id=self.depo.pk, tarih=date(2026, 10, 2), tur=StokHareket.Tur.CIKIS, miktar=D("8"))
        self.assertEqual(sa.kullanilabilir(self.profil), D("-2"))
        # verilen eldeki haritası yeniden sorgulanmaz (depo filtreli kullanım)
        self.assertEqual(sa.toplu_kullanilabilir([self.profil.pk], eldeki={self.profil.pk: D("100")})[self.profil.pk]["kullanilabilir"], D("96"))

    def test_urun_malzeme_sutunlari(self):
        self.giris(self.profil, 2, 200)
        self.giris(self.civata, 25, 50)
        sa.ayirma_ayarla(self.emir2, self.profil, "1")                                     # başka siparişe ayrılmış 1 boy
        sa.ayirma_ayarla(self.emir2, self.civata, "30")                                    # eldekinden fazla ayrılmış (fiilen tüketilmiş)
        g = ua.graf_yukle()
        m = ua.urun_malzeme(g, self.mamul_a, D("5"))                                        # gerekli: PROFIL 3 boy, CIVATA 20, BOYA 5
        satir = {s["stok"].kod: s for gr in m["gruplar"] for s in gr["satirlar"]}
        p, c, b = satir["PROFIL"], satir["CIVATA"], satir["BOYA"]
        self.assertEqual((p["gerekli"], p["eldeki"], p["ayrilan"], p["kullanilabilir"], p["eksik"], p["asim"]), (D("3"), D("2"), D("1"), D("1"), D("2"), False))
        self.assertEqual((c["eldeki"], c["ayrilan"], c["kullanilabilir"], c["eksik"], c["asim"]), (D("25"), D("30"), D("-5"), D("25"), True))
        self.assertEqual((b["ayrilan"], b["kullanilabilir"], b["eksik"]), (D("0"), D("0"), D("5")))
        self.assertEqual((m["yeterli"], m["eksik"], m["ayrilan_sayi"], m["asim_sayi"]), (0, 3, 2, 1))
        # ayırma yokken rakamlar eskisiyle aynı (eşdeğerlik)
        sa.ayirma_ayarla(self.emir2, self.profil, 0)
        sa.ayirma_ayarla(self.emir2, self.civata, 0)
        m = ua.urun_malzeme(g, self.mamul_a, D("5"))
        satir = {s["stok"].kod: s for gr in m["gruplar"] for s in gr["satirlar"]}
        self.assertEqual((satir["PROFIL"]["eksik"], satir["CIVATA"]["eksik"], m["yeterli"], m["eksik"], m["ayrilan_sayi"]), (D("1"), D("0"), 1, 2, 0))   # BOYA stoksuz
        # depo filtresi: eldeki o depodan, ayrılan tüm depolardan
        d2 = Depo.objects.create(kod="D2", ad="İKİNCİ DEPO")
        hareket_ekle(stok_id=self.profil.pk, depo_id=d2.pk, tarih=date(2026, 1, 1), tur=StokHareket.Tur.GIRIS, miktar=D("5"), giris_tutar_try=D("500"))
        sa.ayirma_ayarla(self.emir2, self.profil, "6")
        s = {x["stok"].kod: x for gr in ua.urun_malzeme(g, self.mamul_a, D("1"), d2)["gruplar"] for x in gr["satirlar"]}["PROFIL"]
        self.assertEqual((s["eldeki"], s["ayrilan"], s["kullanilabilir"], s["eksik"]), (D("5"), D("6"), D("-1"), D("2")))   # gerekli 1 boy − (−1)

    def test_excel_sutunlari(self):
        self.giris(self.profil, 2, 200)
        sa.ayirma_ayarla(self.emir2, self.profil, "1")
        g = ua.graf_yukle()
        ws = load_workbook(BytesIO(malzeme_xlsx(self.mamul_a, D("5"), ua.urun_malzeme(g, self.mamul_a, D("5"))))).active
        self.assertEqual([c.value for c in ws[4]], ["Kategori", "Kod", "Ad", "Birim", "Gerekli", "Eldeki stok", "Ayrılan", "Kullanılabilir", "Eksik"])
        satir = {row[1]: row for row in ws.iter_rows(min_row=5, values_only=True) if row[1]}
        self.assertEqual([D(str(v)) for v in satir["PROFIL"][4:9]], [D("3"), D("2"), D("1"), D("1"), D("2")])
        self.assertIn("kullanılabilir = eldeki − açık üretim siparişlerine ayrılan", ws["A2"].value)


class EkranTest(AyirmaBase):
    def setUp(self):
        super().setUp()
        self.client.force_login(User.objects.create_superuser("ua_sa", password="x"))
        self.url = reverse("core:urun_agaci")

    def test_malzeme_sekmesi_sutunlari_ve_uyari(self):
        self.giris(self.profil, 2, 200)
        self.giris(self.civata, 25, 50)
        sa.ayirma_ayarla(self.emir2, self.profil, "1")
        sa.ayirma_ayarla(self.emir2, self.civata, "30")
        r = self.client.get(self.url, {"urun": self.mamul_a.pk, "miktar": "5", "sekme": "malzeme"})
        h = r.content.decode()
        for parca in ("<th class=\"sag\">Ayrılan</th>", "<th class=\"sag\">Kullanılabilir</th>", "<strong>2</strong> kalemde üretim siparişine ayrılan stok var",
                      "⚠ 1 kalemde ayrılan stok eldekini aşıyor", 'class="sag ua-asim"', "Kullanılabilir = eldeki − açık üretim siparişlerine ayrılan",
                      "<strong>3</strong> kalem eksik"):
            self.assertIn(parca, h, parca)
        self.assertEqual(r.context["malzeme"]["asim_sayi"], 1)
        r = self.client.get(self.url, {"urun": self.mamul_a.pk, "miktar": "5", "sekme": "malzeme", "xlsx": "1"})
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn("Kullanılabilir", [c.value for c in load_workbook(BytesIO(r.content)).active[4]])
