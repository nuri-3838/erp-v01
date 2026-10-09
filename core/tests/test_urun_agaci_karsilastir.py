"""ÜRÜN AĞACI — KARŞILAŞTIR görünümü: yaprak malzemeler × ürünler matrisi (Miktar / Maliyet), seri ya da elle en fazla 8 ürün; hücre değerleri Ürün
görünümüyle aynı hesaptan gelir."""
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook

from core.services import urun_agaci as ua
from core.services.uretim import operasyon_olustur
from core.tests.test_urun_agaci_kullanim import KullanimBase

D = Decimal


class KarsilastirBase(KullanimBase):
    def setUp(self):
        self.giris(self.profil, 2, 200, 5)
        self.giris(self.civata, 25, 50, 1.25)
        self.giris(self.p9, 4, 40, 1)            # ort 10 TL
        self.giris(self.p7, 3, 90, 2)            # ort 30 TL
        # BOYA'nın (mamul_a) ve diğerlerinin maliyeti yok


class MatrisServisTest(KarsilastirBase):
    def test_a_serisi_matris_boyutu(self):
        g = ua.graf_yukle()
        urunler, uyari = ua.karsilastir_urunleri(g, "A")
        self.assertEqual(uyari, "")
        self.assertEqual(len(urunler), 6)                                              # 152-10-0001 + 5 model
        m = ua.karsilastir(g, urunler, "miktar")
        satir = {s["stok"].kod: s for gr in m["gruplar"] for s in gr["satirlar"]}
        self.assertEqual(set(satir), {"PROFIL", "CIVATA", "BOYA", "150-10-0009"})
        self.assertEqual(m["satir_sayisi"], 4)
        self.assertTrue(all(len(s["hucreler"]) == 6 for s in satir.values()))
        self.assertEqual([u.kod for u in urunler], sorted(u.kod for u in urunler))

    def test_c_serisi_matris_boyutu(self):
        g = ua.graf_yukle()
        urunler, _ = ua.karsilastir_urunleri(g, "C")
        self.assertEqual(len(urunler), 7)
        m = ua.karsilastir(g, urunler)
        self.assertEqual(m["satir_sayisi"], 3)                                         # PROFIL, CIVATA, 150-10-0007

    def test_kullanilmayan_hucre_bos(self):
        g = ua.graf_yukle()
        urunler, _ = ua.karsilastir_urunleri(g, "A")
        m = ua.karsilastir(g, urunler)
        satir = {s["stok"].kod: s for gr in m["gruplar"] for s in gr["satirlar"]}
        idx = {u.kod: i for i, u in enumerate(urunler)}
        self.assertIsNotNone(satir["PROFIL"]["hucreler"][idx["152-10-0001"]])
        self.assertIsNone(satir["PROFIL"]["hucreler"][idx["152-10-0010"]])
        self.assertIsNone(satir["150-10-0009"]["hucreler"][idx["152-10-0001"]])
        self.assertEqual(satir["150-10-0009"]["hucreler"][idx["152-10-0010"]]["miktar"], D("0.5"))

    def test_hucreler_urun_gorunumuyle_ayni(self):
        g = ua.graf_yukle()
        urunler, _ = ua.karsilastir_urunleri(g, "", [self.mamul_a.pk, self.mamul_c.pk, self.c_modeller[2].pk, self.a_modeller[0].pk])
        mik = ua.karsilastir(g, urunler, "miktar")
        mal = ua.karsilastir(g, urunler, "maliyet")
        for i, u in enumerate(urunler):
            urun_m = ua.urun_maliyet(g, u, D("1"))
            urun_satir = {s["stok"].pk: s for gr in urun_m["gruplar"] for s in gr["satirlar"]}
            for gr in mik["gruplar"]:
                for s in gr["satirlar"]:
                    h = s["hucreler"][i]
                    if h is None:
                        self.assertNotIn(s["stok"].pk, urun_satir)
                    else:
                        self.assertEqual(h["deger"], urun_satir[s["stok"].pk]["tuketim"])
            for gr in mal["gruplar"]:
                for s in gr["satirlar"]:
                    h = s["hucreler"][i]
                    if h is not None and not h["maliyet_yok"]:
                        self.assertEqual(h["deger"], urun_satir[s["stok"].pk]["tutar_try"])
            self.assertEqual(mal["toplamlar"][i]["toplam"], urun_m["toplam_try"])
            self.assertEqual(mal["toplamlar"][i]["maliyetsiz"], urun_m["maliyetsiz"])

    def test_maliyet_toplam_satiri(self):
        g = ua.graf_yukle()
        urunler, _ = ua.karsilastir_urunleri(g, "A")
        m = ua.karsilastir(g, urunler, "maliyet")
        self.assertEqual(len(m["toplamlar"]), 6)
        idx = {u.kod: i for i, u in enumerate(urunler)}
        self.assertEqual(m["toplamlar"][idx["152-10-0010"]]["toplam"], D("5"))         # 0,5 boy × 10 TL
        self.assertEqual(m["toplamlar"][idx["152-10-0001"]]["maliyetsiz"], 1)           # BOYA
        self.assertIsNone(ua.karsilastir(g, urunler, "miktar")["toplamlar"])

    def test_sekiz_urun_siniri(self):
        g = ua.graf_yukle()
        tum = [m.pk for m in self.c_modeller + self.a_modeller]                        # 11 ürün
        urunler, uyari = ua.karsilastir_urunleri(g, "", tum)
        self.assertEqual(len(urunler), 8)
        self.assertIn("En fazla 8", uyari)
        self.assertEqual([u.pk for u in urunler], tum[:8])
        urunler, uyari = ua.karsilastir_urunleri(g, "", tum[:8])
        self.assertEqual((len(urunler), uyari), (8, ""))

    def test_gecersiz_ve_tekrar_secim(self):
        g = ua.graf_yukle()
        urunler, _ = ua.karsilastir_urunleri(g, "", [999999, self.bos.pk, self.mamul_a.pk, self.mamul_a.pk])
        self.assertEqual([u.pk for u in urunler], [self.mamul_a.pk])


class MatrisEkranTest(KarsilastirBase):
    def setUp(self):
        super().setUp()
        self.client.force_login(User.objects.create_superuser("ua_m", password="x"))
        self.url = reverse("core:urun_agaci")

    def al(self, **p):
        return self.client.get(self.url, {"gorunum": "karsilastir", **p})

    def test_acilis_bos(self):
        r = self.al()
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Bir seri seç ya da en fazla 8 ürün işaretle")
        self.assertContains(r, 'class="ua-seg aktif"')
        self.assertContains(r, "Bitmiş ürünler – A tipi")

    def test_a_serisi_miktar_ve_maliyet(self):
        r = self.al(seri="A")
        h = r.content.decode()
        self.assertIn('class="ua-matris"', h)
        self.assertEqual(r.context["sonuc"]["satir_sayisi"], 4)
        self.assertIn("0,5", h)
        self.assertNotIn("ÜRÜN BAŞINA TOPLAM", h)
        h = self.al(seri="A", mod="maliyet").content.decode()
        self.assertIn("ÜRÜN BAŞINA TOPLAM (TL)", h)
        self.assertIn("5,00", h)
        self.assertIn("kalem maliyetsiz", h)

    def test_elle_secim_ve_sinir_uyarisi(self):
        r = self.al(urunler=[self.mamul_a.pk, self.mamul_c.pk])
        self.assertEqual(len(r.context["sonuc"]["urunler"]), 2)
        tum = [m.pk for m in self.c_modeller + self.a_modeller]
        r = self.al(urunler=tum)
        self.assertContains(r, "En fazla 8 ürün karşılaştırılabilir")
        self.assertEqual(len(r.context["sonuc"]["urunler"]), 8)

    def test_gecersiz_parametreler(self):
        self.assertEqual(self.al(seri="Z", mod="x", urunler=["a", ""]).status_code, 200)

    def test_excel_miktar_ve_maliyet(self):
        for mod, satir in (("miktar", 8), ("maliyet", 10)):
            r = self.al(seri="A", mod=mod, xlsx="1")
            self.assertIn("spreadsheetml", r["Content-Type"])
            ws = load_workbook(BytesIO(r.content)).active
            self.assertEqual(ws["A1"].value, "Karşılaştırma")
            self.assertEqual([c.value for c in ws[4]][:4], ["Kategori", "Kod", "Ad", "Birim"])
            self.assertEqual(len([c for c in ws[4] if c.value]), (4 if mod == "miktar" else 5) + 6)   # 6 ürün sütunu (+ Kaynak)
            self.assertEqual(ws.max_row, satir)                                      # 4 malzeme (+ maliyette toplam + maliyetsiz notu)

    def test_sorgu_sayisi_sabit(self):
        sayi = {}
        for mod in ("miktar", "maliyet"):
            with CaptureQueriesContext(connection) as az:
                self.al(seri="C", mod=mod)
            sayi[mod] = len(az)
        for i in range(6):
            ayak = self._stok(f"151-22-8{i}", self.kat_ana)
            operasyon_olustur(istasyon_id=self.kesim.pk, cikti_id=ayak.pk, cikti_miktar=D("3"), satirlar=[(self.p7, D("1"))])
            m = self._stok(f"152-22-8{i}", self.kat_ana)
            operasyon_olustur(istasyon_id=self.montaj.pk, cikti_id=m.pk, cikti_miktar=D("1"), satirlar=[(ayak, D("2")), (self.boya, D("1"))])
        for mod in ("miktar", "maliyet"):
            with CaptureQueriesContext(connection) as cok:
                r = self.al(seri="C", mod=mod)
            self.assertEqual(len(r.context["sonuc"]["urunler"]), 13)
            self.assertEqual(len(cok), sayi[mod], mod)
