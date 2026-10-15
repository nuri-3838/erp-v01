"""Muhasebe › Raporlar › KDV Listeleri (2026-10-16): İndirilecek KDV Listesi (alışlar) ve Satışlara Ait Fatura Listesi — dönem (yıl/ay), GİB kolonları (seri/sıra ayrımı, VKN, cins, miktar, KDV hariç, KDV, dönem),
döviz (fiş kuru), gerçek tarih (açıklamada "GERÇEK TARİH"), tevkifat (indirilecek = KDV'nin tamamı, tevkifat ayrı kolon), yalnız KDV (faturadaki matrah), faturasız 191/391 hareketleri, mutabakat (liste +
faturasız = dönem hareketi, fark 0), uyarılar (eksik VKN/no, mükerrer), satışta iade/ihraç ayrı + Trendyol alıcı, Excel şablon kolon sırası, yetki, SALT-OKUNUR."""
import datetime
from decimal import Decimal
from io import BytesIO

from django.db import connection
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Cari, Fatura, FaturaTipi, HesapPlani, KategoriHesap, KdvOrani, Stok
from core.services import kdv_listeleri as kl
from core.services.fatura import fatura_olustur
from core.services.yevmiye import SatirGirdi, fis_olustur
from core.tests import test_fatura_form_fis_karsilastirma as kt
from core.sayi import yuvarla

D = datetime.date
Dc = Decimal


class KdvTaban(kt.KarsilastirmaTaban):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        from core.models import Kur
        for g in (D(2026, 1, 1), D(2026, 3, 11), D(2026, 3, 12), D(2026, 3, 15), D(2026, 4, 2)):
            Kur.objects.get_or_create(tarih=g, defaults={"usd_alis": Dc("30"), "eur_alis": Dc("35")})
        from core.tests.test_fatura import _hesap
        if not HesapPlani.objects.filter(hesap_kodu="100.01").exists():
            _hesap("100.01", "KASA")
        Cari.objects.filter(pk=cls.tedarikci.pk).update(vkn_tckn="1234567890")
        Cari.objects.filter(pk=cls.eur.pk).update(vkn_tckn="9876543210")
        Cari.objects.filter(pk=cls.musteri.pk).update(vkn_tckn="5555555555")

    def _alis(self, no, satirlar, tip=None, cari=None, tarih=D(2026, 3, 10), **kw):
        return fatura_olustur(tip_id=(tip or self.alis).pk, cari_id=(cari or self.tedarikci).pk, tarih=tarih, fatura_no=no, satirlar=satirlar,
                              depo_id=(self.depo.pk if kw.pop("depo", True) else None), kullanici=self.su, **kw)


class IndirilecekTest(KdvTaban):
    def test_karma_alis_kolonlar_ve_mutabakat(self):
        f = self._alis("OTT2026000011765", [self._stok("100", "50"), self._hesap_satir("730.09", "200")], aciklama="ilk")
        v = kl.indirilecek(2026, 3)
        (r,) = v["satirlar"]
        self.assertEqual((r["seri"], r["sira_no"], r["belge_tarihi"], r["donem"], r["vkn"], r["satici"]), ("OTT", "2026000011765", D(2026, 3, 10), "03/2026", "1234567890", "TEDARİKÇİ A"))
        self.assertEqual((r["matrah_tl"], r["kdv_tl"], r["indirilecek_tl"]), (Dc("5200.00"), Dc("1040.00"), Dc("1040.00")))
        self.assertIn("ALÜMİNYUM", r["cins"])
        self.assertEqual(r["miktar"], "100 AD")
        m = v["mutabakat"]
        self.assertEqual((m["liste"], m["faturasiz"], m["hareket"], m["fark"], m["kalemler"], m["tamam"]), (Dc("1040.00"), Dc("0.00"), Dc("1040.00"), Dc("0.00"), [], True))
        self.assertEqual(f.fis.satirlar.filter(hesap_id="191").first().borc, Dc("1040.00"))

    def test_donem_disi_ve_taslak_dahil_degil(self):
        self._alis("D-1", [self._stok("10", "10")], tarih=D(2026, 4, 2))
        self.assertEqual(kl.indirilecek(2026, 3)["satirlar"], [])
        self.assertEqual(len(kl.indirilecek(2026, 4)["satirlar"]), 1)

    def test_yalniz_kdv_tipinde_faturadaki_matrah_gosterilir(self):
        self._alis("MDA2026000000011", [self._stok("10", "1000")], tip=self.yalniz_kdv, cari=self.tedarikci, depo=False)
        (r,) = kl.indirilecek(2026, 3)["satirlar"]
        self.assertEqual((r["matrah_tl"], r["kdv_tl"]), (Dc("10000.00"), Dc("2000.00")))               # matrah fişe yazılmaz ama belge bedeli gösterilir
        self.assertIn(("kdv", "Yalnız KDV tipi: matrah fişe yazılmadı (faturadaki matrah gösterilir)", "yalnız KDV"), r["rozetler"])
        self.assertTrue(kl.indirilecek(2026, 3)["mutabakat"]["tamam"])

    def test_doviz_fatura_fis_kuruyla_tl(self):
        f = self._alis("EUR2026000000007", [self._stok("12", "13,5"), self._stok("3", "10,333333")], cari=self.eur, para_birimi="EUR")
        (r,) = kl.indirilecek(2026, 3)["satirlar"]
        kur = f.kur
        beklenen_matrah = yuvarla(Dc("12") * Dc("13.5") * kur, 2) + yuvarla(Dc("3") * Dc("10.333333") * kur, 2)
        self.assertEqual(r["matrah_tl"], beklenen_matrah)
        fis_191 = f.fis.satirlar.get(hesap_id="191").borc
        self.assertEqual(r["kdv_tl"], fis_191)                                                         # fişle kuruşu kuruşuna aynı
        self.assertEqual((r["pb"], r["kur"]), ("EUR", kur))
        self.assertTrue(kl.indirilecek(2026, 3)["mutabakat"]["tamam"])

    def test_tevkifatli_indirilecek_kdv_tamami_tevkifat_ayri(self):
        Stok.objects.filter(pk=self.stok.pk).update(tevkifat=self.tev)                                 # 5/10
        self._alis("TV-1", [self._stok("10", "100")])
        (r,) = kl.indirilecek(2026, 3)["satirlar"]
        self.assertEqual((r["kdv_tl"], r["tevkifat_tl"], r["indirilecek_tl"]), (Dc("200.00"), Dc("100.00"), Dc("200.00")))
        self.assertTrue(any(k == "tevkifat" for k, _a, _b in r["rozetler"]))
        t = kl.indirilecek(2026, 3)["toplam"]
        self.assertEqual((t["kdv"], t["tevkifat"]), (Dc("200.00"), Dc("100.00")))
        self.assertTrue(kl.indirilecek(2026, 3)["mutabakat"]["tamam"])                                   # 191'e KDV'nin tamamı yazılır

    def test_gercek_tarih_ve_isaretli_01_01(self):
        self._alis("GT-1", [self._stok("1", "100")], tarih=D(2026, 1, 1), aciklama="GERÇEK TARİH: 04.12.2025. test gideri")
        self._alis("GT-2", [self._stok("1", "100")], tarih=D(2026, 1, 1), aciklama="tarih yok")
        v = kl.indirilecek(2026, 1)
        a, b = sorted(v["satirlar"], key=lambda r: r["belge_no"])
        self.assertEqual((a["belge_tarihi"], a["donem"], a["tarih_isaretli"]), (D(2025, 12, 4), "01/2026", False))
        self.assertEqual((b["belge_tarihi"], b["tarih_isaretli"]), (D(2026, 1, 1), True))
        self.assertEqual([r["belge_no"] for r in v["uyarilar"]["gercek_tarih"]], ["GT-1"])
        self.assertEqual([r["belge_no"] for r in v["uyarilar"]["tarih_isaretli"]], ["GT-2"])
        self.assertEqual(kl.gercek_tarih("x GERÇEK TARİH: 5.1.2025 y"), D(2025, 1, 5))
        self.assertIsNone(kl.gercek_tarih("yok"))

    def test_faturasiz_191_hareketi_ayri_tabloda_ve_mutabakat(self):
        self._alis("FS-1", [self._stok("10", "100")])
        fis_olustur(tarih=D(2026, 3, 12), aciklama="GÜMRÜK KDV", satirlar=[SatirGirdi(hesap_kodu="191", taraf="B", islem_tutari="500"), SatirGirdi(hesap_kodu="100.01", taraf="A", islem_tutari="500")])
        v = kl.indirilecek(2026, 3)
        (x,) = v["faturasiz"]
        self.assertEqual((x["borc"], x["alacak"], x["aciklama"]), (Dc("500.00"), Dc("0.00"), "GÜMRÜK KDV"))
        m = v["mutabakat"]
        self.assertEqual((m["liste"], m["faturasiz"], m["toplam"], m["hareket"], m["fark"], m["tamam"]), (Dc("200.00"), Dc("500.00"), Dc("700.00"), Dc("700.00"), Dc("0.00"), True))

    def test_uyarilar_eksik_vkn_no_ve_mukerrer(self):
        Cari.objects.filter(pk=self.tedarikci.pk).update(vkn_tckn="")
        self._alis("MK-1", [self._stok("1", "10")])
        self._alis("MK-1", [self._stok("1", "20")], mukerrer_onayli=True)
        f = fatura_olustur(tip_id=self.alis.pk, cari_id=self.tedarikci.pk, tarih=D(2026, 3, 11), fatura_no="", faturasiz=True, satirlar=[self._stok("1", "30")], depo_id=self.depo.pk, kullanici=self.su)
        u = kl.indirilecek(2026, 3)["uyarilar"]
        self.assertEqual(len(u["eksik_vkn"]), 3)
        self.assertEqual([r["fatura"].pk for r in u["eksik_no"]], [f.pk])
        self.assertEqual(len(u["mukerrer"]), 1)
        self.assertEqual(len(u["mukerrer"][0]), 2)

    def test_seri_sira_ayrimi(self):
        self.assertEqual(kl.seri_sira("OTT2026000011765"), ("OTT", "2026000011765"))
        self.assertEqual(kl.seri_sira("ABC-1"), ("", "ABC-1"))
        self.assertEqual(kl.seri_sira(""), ("", ""))


class SatisTest(KdvTaban):
    def _hazirla(self):
        self.stok_girisi("500", "50")
        KdvOrani.objects.get_or_create(oran=Dc("0"), defaults={"aciklama": "KDV YOK"})
        self.satis_ihrac = FaturaTipi.objects.create(ad="SATIŞ FATURASI-İHRAÇ KAYITLI", yon=FaturaTipi.Yon.SATIS)
        KategoriHesap.objects.create(kategori=self.alt, fatura_tipi=self.satis_ihrac, hesap=HesapPlani.objects.get(hesap_kodu="600"))

    def _satis(self, no, tip, cari, fiyat="100", miktar="10", **kw):
        return fatura_olustur(tip_id=tip.pk, cari_id=cari.pk, tarih=D(2026, 3, 10), fatura_no=no, satirlar=[self._stok(miktar, fiyat)], depo_id=self.depo.pk, kullanici=self.su, **kw)

    def test_normal_iade_ihrac_ayri_ve_mutabakat(self):
        self._hazirla()
        self._satis("AVR2026000000001", self.satis, self.musteri)
        self._satis("SMT2026000000001", self.alis_iade, self.tedarikci, fiyat="50", miktar="6")
        self._satis("IHR2026000000001", self.satis_ihrac, self.musteri, fiyat="10", miktar="5")
        v = kl.satislar(2026, 3)
        self.assertEqual((len(v["normal"]), len(v["iade"]), len(v["ihrac"])), (1, 1, 1))
        n, i, h = v["normal"][0], v["iade"][0], v["ihrac"][0]
        self.assertEqual((n["matrah_tl"], n["hesaplanan_tl"], n["seri"], n["sira_no"], n["vkn"]), (Dc("1000.00"), Dc("200.00"), "AVR", "2026000000001", "5555555555"))
        self.assertEqual((i["matrah_tl"], i["hesaplanan_tl"], i["tur"]), (Dc("300.00"), Dc("60.00"), "iade"))
        self.assertEqual((h["hesaplanan_tl"], h["tur"]), (Dc("0.00"), "ihrac"))
        m = v["mutabakat"]
        self.assertEqual((m["liste"], m["hareket"], m["fark"], m["tamam"]), (Dc("260.00"), Dc("260.00"), Dc("0.00"), True))

    def test_trendyol_alici_aciklamadan_tckn_bos_isaretli(self):
        self._hazirla()
        self._satis("AVR2026000000002", self.satis, self.trendyol, aciklama="Trendyol e-arşiv AVR2026000000002 · alıcı ÖZCAN ÖZDEMİR · Sip. 11658030736 · e-arşiv tarihi 30.09.2026")
        (r,) = kl.satislar(2026, 3)["normal"]
        self.assertEqual((r["alici"], r["vkn"], r["vkn_eksik"]), ("ÖZCAN ÖZDEMİR", "", True))
        self.assertTrue(any(k == "alici" for k, _a, _b in r["rozetler"]))
        self.assertEqual(len(kl.satislar(2026, 3)["uyarilar"]["eksik_vkn"]), 1)
        self.assertEqual(kl.alici_adi("Trendyol e-arşiv X · alıcı ATACAN ATALAN (TCKN 13976454616) · Sip. 1"), ("ATACAN ATALAN", "13976454616"))
        self.assertEqual(kl.alici_adi("alıcı AHMET HAGY"), ("AHMET HAGY", ""))
        self.assertEqual(kl.alici_adi("alıcı yok burada"[:0]), ("", ""))

    def test_faturasiz_391_hareketi(self):
        self._hazirla()
        self._satis("AVR2026000000003", self.satis, self.musteri)
        fis_olustur(tarih=D(2026, 3, 15), aciklama="KDV MAHSUP", satirlar=[SatirGirdi(hesap_kodu="391", taraf="B", islem_tutari="50"), SatirGirdi(hesap_kodu="191", taraf="A", islem_tutari="50")])
        v = kl.satislar(2026, 3)
        (x,) = v["faturasiz"]
        self.assertEqual((x["borc"], x["alacak"]), (Dc("50.00"), Dc("0.00")))
        m = v["mutabakat"]
        self.assertEqual((m["liste"], m["faturasiz"], m["hareket"], m["fark"]), (Dc("200.00"), Dc("0.00"), Dc("200.00"), Dc("0.00")))      # 391 alacak: mahsup borç satırı alacağı etkilemez
        self.assertEqual(m["karsi_taraf"], Dc("50.00"))


class EkranVeExcelTest(KdvTaban):
    def get(self, **q):
        self.client.force_login(self.su)
        yazma = []

        def engel(execute, sql, params, many, context):
            if sql.lstrip()[:6].upper() in ("INSERT", "UPDATE", "DELETE"):
                yazma.append(sql[:80])
            return execute(sql, params, many, context)
        with connection.execute_wrapper(engel):
            r = self.client.get(reverse("core:kdv_listeleri"), q)
        self.assertEqual(yazma, [], "Sayfa açılırken veritabanına YAZMA denendi")
        return r

    def test_sayfa_iki_sekme_ve_yazma_yok(self):
        self._alis("EK-1", [self._stok("10", "100")])
        r = self.get(yil="2026", ay="3")
        self.assertEqual(r.status_code, 200)
        for m in ("İndirilecek KDV Listesi", "Satışlara Ait Fatura Listesi", "Mutabakat", "EK-1"):
            self.assertContains(r, m)
        r2 = self.get(yil="2026", ay="3", sekme="satis")
        self.assertEqual(r2.status_code, 200)
        self.assertContains(r2, "Hesaplanan KDV")
        self.assertEqual(self.get(yil="bozuk", ay="99").status_code, 200)                              # geçersiz dönem → bu ay

    def test_excel_gib_kolon_sirasi(self):
        self._alis("OTT2026000011765", [self._stok("10", "100")])
        r = self.get(yil="2026", ay="3", excel="1")
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["İndirilecek KDV Listesi", "Özet ve Mutabakat"])
        ws = wb["İndirilecek KDV Listesi"]
        self.assertEqual([c.value for c in ws[1]], ["Sıra No", "Belge Tarihi", "Belge Seri No", "Belge Sıra No", "Satıcının Adı Soyadı / Unvanı", "Satıcının VKN / TCKN",
                                                    "Mal / Hizmet Cinsi", "Miktar", "KDV Hariç Tutar", "KDV Tutarı", "KDV'nin İndirildiği Dönem"])
        satir = [c.value for c in ws[2]]
        self.assertEqual((satir[0], satir[2], satir[3], satir[5], satir[8], satir[9], satir[10]), (1, "OTT", "2026000011765", "1234567890", 1000.0, 200.0, "03/2026"))
        s = load_workbook(BytesIO(self.get(yil="2026", ay="3", sekme="satis", excel="1").content))
        self.assertEqual(s.sheetnames, ["Satışlara Ait Fatura Listesi", "İade ve İhraç Kayıtlı", "Özet ve Mutabakat"])
        self.assertEqual([c.value for c in s["Satışlara Ait Fatura Listesi"][1]], ["Sıra No", "Belge Tarihi", "Belge Seri No", "Belge Sıra No", "Alıcının Adı Soyadı / Unvanı", "Alıcının VKN / TCKN", "KDV Hariç Tutar", "KDV Tutarı"])

    def test_menu_ve_yetki(self):
        from core.moduller import MODULLER
        muh = next(m for m in MODULLER if m.kod == "MUHASEBE")
        grup = next(o for o in muh.ogeler if getattr(o, "alt_menu", False))
        self.assertEqual((grup.ad, [e.kod for e in grup.ekranlar]), ("Raporlar", ["kdv_listeleri"]))
        from django.contrib.auth.models import User
        self.client.force_login(User.objects.create_user("yetkisiz2", password="x"))
        self.assertEqual(self.client.get(reverse("core:kdv_listeleri")).status_code, 403)
