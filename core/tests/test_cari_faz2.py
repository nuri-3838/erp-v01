"""Cariler FAZ 2 (2026-10-13): liste (bakiye TL/döviz, son hareket, filtreler, arama, sıralama, sayfalama, Excel, sabit sorgu sayısı), detay (özet şeridi, belgeler,
açık faturalar FIFO), ekstre PDF/Excel, mutabakat mektubu PDF (A4 siyah-beyaz), ölü bağlantı taraması."""
import datetime
import re
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.template.loader import render_to_string
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Cari, CariKategori, CariYetkili, Fatura, HesapPlani, Kur, Sehir, TeklifSiparis, Ulke
from core.services import cari_cikti, cari_ozet
from core.services.cari import cari_olustur, cari_pasife_al
from core.services.fatura import fatura_olustur
from core.services.finans import kasa_olustur
from core.services.kasa_hareket import hareket_olustur as kasa_hareket
from core.services.teklif_siparis import teklif_siparis_olustur
from core.services.yevmiye import SatirGirdi, fis_olustur
from core.tests.test_fatura import FaturaTestTemel, _hesap

D = datetime.date
Dc = Decimal


def _hesaplar():
    for kod, ad, kalem in (("120", "ALICILAR", "DV"), ("320", "SATICILAR", "KVYK"), ("335", "PERSONELE BORÇLAR", "KVYK"), ("100", "KASA", "DV")):
        HesapPlani.objects.get_or_create(hesap_kodu=kod, defaults=dict(hesap_adi=ad, rapor_grubu="BILANCO", rapor_kalemi=kalem, parasal=True))
    for gun in range(1, 32):                                          # fiş tarihleri için günlük kur (TCMB benzeri)
        Kur.objects.get_or_create(tarih=D(2026, 3, gun), defaults={"usd_alis": Dc("30")})


def _kat(ust_kod, ust_ad, alt_kod="10", alt_ad="GENEL"):
    ust = CariKategori.objects.get_or_create(kod=ust_kod, ust=None, defaults={"ad": ust_ad})[0]
    return CariKategori.objects.get_or_create(kod=alt_kod, ust=ust, defaults={"ad": alt_ad})[0]


def _fis(c, taraf, tutar, tarih=D(2026, 3, 1), pb="TRY", kur="1"):
    """Cari hesabına tek yönlü hareket; karşı taraf kasa (100) TL karşılığıyla."""
    tl = (Dc(tutar) * Dc(kur)).quantize(Dc("0.01"))
    karsi = "B" if taraf == "A" else "A"
    return fis_olustur(tarih=tarih, kullanici=None, satirlar=[SatirGirdi(c.muhasebe_kodu, taraf, Dc(tutar), pb, Dc(kur)), SatirGirdi("100", karsi, tl)])


class ListeTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        _hesaplar()
        cls.su = User.objects.create_superuser("f2su", password="x")
        ted, mus, per = _kat("320", "TEDARİKÇİLER", "10", "HAMMADDE"), _kat("120", "MÜŞTERİLER", "10", "YURTİÇİ"), _kat("335", "PERSONELLER", "10", "PERSONEL")
        tr = Ulke.objects.create(kod="TR", ad="TÜRKİYE")
        cls.kayseri = Sehir.objects.create(ulke=tr, ad="KAYSERİ", kod="38")
        cls.ankara = Sehir.objects.create(ulke=tr, ad="ANKARA", kod="06")
        cls.a = cari_olustur(unvan="alfa tedarik", kategori_id=ted.pk, para_birimi="TRY", sehir_id=cls.kayseri.pk, ulke_id=tr.pk, telefon="0 (352) 224 11 01",
                             eposta="satis@alfa.com", vkn_tckn="1234567890")
        cls.b = cari_olustur(unvan="beta müşteri", kategori_id=mus.pk, para_birimi="TRY", sehir_id=cls.ankara.pk, ulke_id=tr.pk)
        cls.c = cari_olustur(unvan="ceyda personel", kategori_id=per.pk, para_birimi="TRY")
        cls.d = cari_olustur(unvan="delta dolar", kategori_id=ted.pk, para_birimi="USD")
        cls.e = cari_olustur(unvan="epsilon pasif", kategori_id=ted.pk, para_birimi="TRY")
        _fis(cls.a, "A", "1000", D(2026, 3, 5))                      # tedarikçi: alacak 1.000 → bakiye −1.000 (A)
        _fis(cls.a, "A", "200", D(2026, 3, 9))
        _fis(cls.b, "B", "500", D(2026, 3, 7))                       # müşteri: borç 500 → +500 (B)
        _fis(cls.d, "A", "100", D(2026, 3, 8), "USD", "30")          # dolar cari: 100 USD alacak
        _fis(cls.e, "B", "50", D(2026, 3, 2))
        cari_pasife_al(cls.e)
        CariYetkili.objects.create(cari=cls.b, ad_soyad="ZEYNEP YETKİLİ", telefon="+90 532 111 22 33", eposta="zeynep@beta.com")

    def _liste(self, **q):
        self.client.force_login(self.su)
        return self.client.get(reverse("core:cariler"), q)

    def _unvanlar(self, r):
        return [s["cari"].unvan for s in r.context["satirlar"]]

    def test_sutunlar_bakiye_doviz_son_hareket_ve_pasif_gizli(self):
        r = self._liste()
        self.assertEqual(self._unvanlar(r), ["ALFA TEDARİK", "BETA MÜŞTERİ", "CEYDA PERSONEL", "DELTA DOLAR"])        # pasif gizli
        s = {x["cari"].unvan: x for x in r.context["satirlar"]}
        self.assertEqual((s["ALFA TEDARİK"]["tl"], s["BETA MÜŞTERİ"]["tl"], s["CEYDA PERSONEL"]["tl"]), (Dc("-1200.00"), Dc("500.00"), Dc("0")))
        self.assertEqual(s["DELTA DOLAR"]["doviz_metin"], ["100,00 USD A"])
        self.assertEqual((s["ALFA TEDARİK"]["son"], s["CEYDA PERSONEL"]["son"]), (D(2026, 3, 9), None))
        self.assertContains(r, "1.200,00 A")
        self.assertContains(r, "100,00 USD A")
        self.assertContains(r, "09.03.2026")
        for baslik in ("Cari", "Kategori", "Şehir / Ülke", "Bakiye TL", "Bakiye Döviz", "Son Hareket"):
            self.assertContains(r, baslik)
        self.assertNotContains(r, "VKN</th>")                                                                        # VKN/telefon sütunu YOK
        self.assertNotContains(r, "Telefon</th>")
        self.assertNotContains(r, "1234567890")
        self.assertContains(r, 'data-href="%s"' % reverse("core:cari_detay", args=[self.a.pk]))
        self.assertContains(r, 'class="satir-bag"')
        self.assertNotContains(r, "onclick=")

    def test_filtreler(self):
        self.assertEqual(self._unvanlar(self._liste(bakiye="bakiyeli")), ["ALFA TEDARİK", "BETA MÜŞTERİ", "DELTA DOLAR"])
        self.assertEqual(self._unvanlar(self._liste(bakiye="borclu")), ["BETA MÜŞTERİ"])
        self.assertEqual(self._unvanlar(self._liste(bakiye="alacakli")), ["ALFA TEDARİK", "DELTA DOLAR"])
        self.assertEqual(self._unvanlar(self._liste(doviz="1")), ["DELTA DOLAR"])
        self.assertEqual(self._unvanlar(self._liste(tur="ticari")), ["ALFA TEDARİK", "BETA MÜŞTERİ", "DELTA DOLAR"])
        self.assertEqual(self._unvanlar(self._liste(tur="diger")), ["CEYDA PERSONEL"])
        self.assertEqual(self._unvanlar(self._liste(pasif="1")), ["ALFA TEDARİK", "BETA MÜŞTERİ", "CEYDA PERSONEL", "DELTA DOLAR", "EPSİLON PASİF"])
        self.assertEqual(self._unvanlar(self._liste(pasif="1", bakiye="borclu")), ["BETA MÜŞTERİ", "EPSİLON PASİF"])
        self.assertEqual(self._unvanlar(self._liste(sehir=self.kayseri.pk)), ["ALFA TEDARİK"])
        r = self._liste()
        self.assertContains(r, "Pasif dahil (1)")
        self.assertContains(r, "Dövizli")
        self.assertNotContains(r, "VKN eksik")                                                                       # istenmeyen filtre eklenmedi

    def test_arama_gorunmeyen_alanlarda(self):
        self.assertEqual(self._unvanlar(self._liste(ara="224 11 01")), ["ALFA TEDARİK"])                              # telefon (boşluklu)
        self.assertEqual(self._unvanlar(self._liste(ara="3522241101")), ["ALFA TEDARİK"])                             # telefon (bitişik)
        self.assertEqual(self._unvanlar(self._liste(ara="0352 224")), ["ALFA TEDARİK"])                               # yerel biçim
        self.assertEqual(self._unvanlar(self._liste(ara="alfa.com")), ["ALFA TEDARİK"])                               # e-posta
        self.assertEqual(self._unvanlar(self._liste(ara="1234567890")), ["ALFA TEDARİK"])                             # VKN
        self.assertEqual(self._unvanlar(self._liste(ara="zeynep")), ["BETA MÜŞTERİ"])                                 # yetkili adı
        self.assertEqual(self._unvanlar(self._liste(ara="532 111")), ["BETA MÜŞTERİ"])                                # yetkili telefonu
        self.assertEqual(self._unvanlar(self._liste(ara="ankara")), ["BETA MÜŞTERİ"])                                 # şehir
        self.assertEqual(self._unvanlar(self._liste(ara="320-10-0001")), ["ALFA TEDARİK"])                            # kod
        self.assertEqual(self._unvanlar(self._liste(ara="yokboyle")), [])

    def test_siralama_ve_basliklar(self):
        self.assertEqual(self._unvanlar(self._liste(sirala="bakiye", yon="asc")), ["DELTA DOLAR", "ALFA TEDARİK", "CEYDA PERSONEL", "BETA MÜŞTERİ"])   # −3.000, −1.200, 0, +500
        self.assertEqual(self._unvanlar(self._liste(sirala="bakiye", yon="desc")), ["BETA MÜŞTERİ", "CEYDA PERSONEL", "ALFA TEDARİK", "DELTA DOLAR"])
        self.assertEqual(self._unvanlar(self._liste(sirala="cari", yon="desc"))[0], "DELTA DOLAR")
        hareket = self._unvanlar(self._liste(sirala="hareket", yon="desc"))
        self.assertEqual(hareket[:3], ["ALFA TEDARİK", "DELTA DOLAR", "BETA MÜŞTERİ"])
        self.assertEqual(hareket[-1], "CEYDA PERSONEL")                                                               # hareketsiz hep sonda
        self.assertEqual(self._unvanlar(self._liste(sirala="hareket", yon="asc"))[-1], "CEYDA PERSONEL")
        r = self._liste(sirala="bakiye", yon="asc")
        self.assertContains(r, "Bakiye TL ▲")
        self.assertContains(r, "sirala=bakiye&amp;yon=desc")                                                          # aynı başlığa tıklayınca yön döner
        # Türkçe alfabe: Ç, İ, Ş doğru yerde
        from core.services.cari_ozet import tr_anahtar
        self.assertEqual(sorted(["ŞİRKET", "SİRKE", "ÇAM", "CAM", "DAL", "IŞIK", "İZMİR"], key=tr_anahtar), ["CAM", "ÇAM", "DAL", "IŞIK", "İZMİR", "SİRKE", "ŞİRKET"])

    def test_excel_filtrelenmis_listenin_tamami(self):
        self.client.force_login(self.su)
        r = self.client.get(reverse("core:cariler"), {"bakiye": "alacakli", "xlsx": "1"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        ws = load_workbook(BytesIO(r.content)).active
        basliklar = [c.value for c in ws[4]]
        self.assertEqual(basliklar[:2], ["Cari Kodu", "Ünvan"])
        self.assertNotIn("VKN", " ".join(basliklar))
        satirlar = [row for row in ws.iter_rows(min_row=5, values_only=True) if row[0]]
        self.assertEqual([x[1] for x in satirlar], ["ALFA TEDARİK", "DELTA DOLAR"])
        self.assertEqual(satirlar[0][6], -1200.0)
        self.assertEqual(satirlar[0][7], "Alacak")
        self.assertEqual(satirlar[1][8], "100,00 USD Alacak")

    def test_sayfalama_ve_sabit_sorgu_sayisi(self):
        ted = _kat("320", "TEDARİKÇİLER", "10", "HAMMADDE")

        def sorgu():
            with CaptureQueriesContext(connection) as q:
                cari_ozet.liste_getir(pasif=True)
            return len(q)
        az = sorgu()
        for i in range(60):
            c = cari_olustur(unvan=f"toplu cari {i:02d}", kategori_id=ted.pk, para_birimi="TRY")
            _fis(c, "A", "10", D(2026, 3, 1))
        cok = sorgu()
        self.assertEqual(az, cok)                                                                                    # bakiye/son hareket N+1 değil
        self.assertLessEqual(cok, 6)
        r = self._liste()
        self.assertEqual(len(r.context["satirlar"]), 50)
        self.assertContains(r, "sayfa 1 / 2")
        r = self._liste(sayfa="2")
        self.assertEqual(len(r.context["satirlar"]), 14)
        self.assertContains(r, "sayfa=1")
        with CaptureQueriesContext(connection) as q:
            self._liste()
        self.assertLessEqual(len(q), 25)


class DetayEkstreTest(FaturaTestTemel):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.su = User.objects.create_superuser("f2detay", password="x")
        _hesap("100.01", "KASA")
        cls.kasa = kasa_olustur(ad="K1", muhasebe_kodu="100.01")
        for gun in range(1, 32):
            Kur.objects.get_or_create(tarih=D(2026, 3, gun), defaults={"usd_alis": Dc("30")})

    def _faturalar_ve_odeme(self, odeme="1300"):
        f1 = fatura_olustur(tip_id=self.alis.pk, cari_id=self.tedarikci.pk, tarih=D(2026, 3, 10), satirlar=self._satir("10", "100"))      # 1.000 + %20 = 1.200
        f2 = fatura_olustur(tip_id=self.alis.pk, cari_id=self.tedarikci.pk, tarih=D(2026, 3, 11), satirlar=self._satir("5", "100"))       # 600
        kasa_hareket(kasa=self.kasa, tip="cari_odeme", karsi=self.tedarikci, tutar=odeme, tarih=D(2026, 3, 12))
        return f1, f2

    def test_acik_faturalar_fifo(self):
        f1, f2 = self._faturalar_ve_odeme("1300")
        af = cari_ozet.acik_faturalar(self.tedarikci, bugun=D(2026, 4, 30))
        self.assertEqual([(k["fatura"].pk, k["kalan"], k["tutar"]) for k in af["satirlar"]], [(f2.pk, Dc("500.00"), Dc("600.00"))])   # en eski tam kapandı
        self.assertEqual(af["toplam"], {"TRY": Dc("500.00")})
        self.assertEqual(af["esitlenmemis"], {})
        Fatura.objects.filter(pk=f2.pk).update(vade_tarihi=D(2026, 4, 10))
        af = cari_ozet.acik_faturalar(self.tedarikci, bugun=D(2026, 4, 30))
        self.assertEqual((af["satirlar"][0]["gecikme"], af["vadesi_gecen"]), (20, {"TRY": Dc("500.00")}))
        # fazla ödeme → eşleşmeyen avans
        kasa_hareket(kasa=self.kasa, tip="cari_odeme", karsi=self.tedarikci, tutar="700", tarih=D(2026, 3, 13))
        af = cari_ozet.acik_faturalar(self.tedarikci, bugun=D(2026, 4, 30))
        self.assertEqual((af["satirlar"], af["esitlenmemis"]), ([], {"TRY": Dc("200.00")}))

    def test_ozet_serit_belgeler_ve_acik_fatura_bolumu(self):
        f1, f2 = self._faturalar_ve_odeme("1300")
        ts = teklif_siparis_olustur(belge_tur=TeklifSiparis.BelgeTur.SIPARIS, yon=TeklifSiparis.Yon.ALIS, cari_id=self.tedarikci.pk, tarih=D(2026, 3, 1),
                                    satirlar=[{"stok_id": self.stok.pk, "miktar": "2", "birim_fiyat": "100"}])
        o = cari_ozet.ozet(self.tedarikci)
        self.assertEqual((o["tl"], o["acik_fatura_adedi"], o["son"]["tarih"]), (Dc("-500.00"), 1, D(2026, 3, 12)))
        b = cari_ozet.belgeler(self.tedarikci)
        self.assertEqual([x["tur"] for x in b], ["Fatura (ALIŞ FATURASI)", "Fatura (ALIŞ FATURASI)", "Sipariş (Alış)"])      # yeni üstte, tek liste
        self.assertEqual(b[2]["tutar"], Dc("240.00"))                                                                       # 2 × 100 + KDV %20
        self.client.force_login(self.su)
        r = self.client.get(reverse("core:cari_detay", args=[self.tedarikci.pk]))
        self.assertContains(r, "Açık Faturalar")
        self.assertContains(r, "500,00 A")
        self.assertContains(r, "Belgeler (3)")
        self.assertContains(r, reverse("core:fatura_detay", args=[f1.pk]))
        self.assertContains(r, reverse("core:teklif_siparis_detay", args=[ts.pk]))
        self.assertContains(r, "Son hareket")
        self.assertNotContains(r, "Faz 4")

    def test_detay_baglantilari_olu_degil(self):
        self._faturalar_ve_odeme()
        self.client.force_login(self.su)
        r = self.client.get(reverse("core:cari_detay", args=[self.tedarikci.pk]))
        yollar = sorted(set(re.findall(r'href="(/[^"#]*)"', r.content.decode())))
        self.assertGreaterEqual(len(yollar), 8)
        for y in yollar:
            if y in ("/logout/",):
                continue
            kod = self.client.get(y).status_code
            self.assertIn(kod, (200, 302), f"{y} → {kod}")

    def test_ekstre_pdf_excel_ve_mutabakat(self):
        self._faturalar_ve_odeme()
        self.client.force_login(self.su)
        e = reverse("core:cari_ekstresi", args=[self.tedarikci.pk])
        r = self.client.get(e, {"baslangic": "2026-03-01", "bitis": "2026-03-31"})
        self.assertContains(r, "PDF")
        self.assertContains(r, "Mutabakat PDF")
        self.assertContains(r, reverse("core:cari_ekstre_xlsx", args=[self.tedarikci.pk]) + "?baslangic=2026-03-01&amp;bitis=2026-03-31")
        r = self.client.get(reverse("core:cari_ekstre_pdf", args=[self.tedarikci.pk]), {"baslangic": "2026-03-01", "bitis": "2026-03-31"})
        self.assertEqual((r.status_code, r["Content-Type"]), (200, "application/pdf"))
        self.assertTrue(r.content.startswith(b"%PDF"))
        r = self.client.get(reverse("core:cari_ekstre_xlsx", args=[self.tedarikci.pk]), {"baslangic": "2026-03-01", "bitis": "2026-03-31"})
        ws = load_workbook(BytesIO(r.content)).active
        satirlar = [row for row in ws.iter_rows(min_row=5, values_only=True)]
        self.assertEqual(satirlar[0][2], "Devir (açılış bakiyesi)")
        hareketler = [x for x in satirlar if x[0] and isinstance(x[0], (datetime.date, datetime.datetime))]
        self.assertEqual(len(hareketler), 3)
        self.assertEqual(satirlar[-1][2], "KAPANIŞ")
        self.assertEqual((satirlar[-1][3], satirlar[-1][4], satirlar[-1][5]), (1300.0, 1800.0, -500.0))
        # tarih aralığı dışı: yalnız devir
        r = self.client.get(reverse("core:cari_ekstre_xlsx", args=[self.tedarikci.pk]), {"baslangic": "2026-04-01", "bitis": "2026-04-30"})
        ws = load_workbook(BytesIO(r.content)).active
        self.assertEqual([row for row in ws.iter_rows(min_row=5, values_only=True)][0][5], -500.0)
        # mutabakat
        r = self.client.get(reverse("core:cari_mutabakat_pdf", args=[self.tedarikci.pk]), {"tarih": "2026-03-31"})
        self.assertEqual((r.status_code, r["Content-Type"]), (200, "application/pdf"))
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertIn("mutabakat", r["Content-Disposition"].lower())

    def test_mutabakat_baglami_ve_sablon(self):
        self._faturalar_ve_odeme()
        b = cari_cikti.mutabakat_baglami(self.tedarikci, D(2026, 3, 10))
        self.assertEqual(b["tl"], Dc("-1200.00"))                                                                    # 10.03 dahil: yalnız 1. fatura
        self.assertEqual(cari_cikti.mutabakat_baglami(self.tedarikci, D(2026, 3, 9))["tl"], Dc("0"))
        self.assertEqual(cari_cikti.mutabakat_baglami(self.tedarikci, D(2026, 3, 31))["tl"], Dc("-500.00"))
        html = render_to_string("core/cari_mutabakat_pdf.html", {
            "cari": self.tedarikci, "tarih": D(2026, 3, 31), "belge_tarihi": D(2026, 4, 1), "firma": None, "logo_b64": None,
            **cari_cikti.mutabakat_baglami(self.tedarikci, D(2026, 3, 31))})
        self.assertIn("size: A4 portrait", html)
        self.assertIn("grayscale(100%)", html)
        self.assertIn("500,00", html)
        self.assertNotIn("-500,00", html)
        self.assertIn("Bizim borcumuz", html)
        for metin in ("MUTABIKIZ", "MUTABIK DEĞİLİZ", "İmza / Kaşe", "31.03.2026", self.tedarikci.unvan):
            self.assertIn(metin, html)
        self.assertEqual(re.findall(r"#[0-9a-fA-F]{3,6}\b", html.split("</style>")[0]), ["#000"] * len(re.findall(r"#[0-9a-fA-F]{3,6}\b", html.split("</style>")[0])))   # renk yok
        # alacağımız yönü
        html2 = render_to_string("core/cari_mutabakat_pdf.html", {"cari": self.musteri, "tarih": D(2026, 3, 31), "belge_tarihi": D(2026, 4, 1), "firma": None,
                                                                  "logo_b64": None, "tl": Dc("750.00"), "doviz": [("USD", Dc("-20.00"))]})
        self.assertIn("Bizim alacağımız", html2)
        self.assertIn("20,00", html2)
        self.assertIn("USD", html2)
