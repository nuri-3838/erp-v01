"""ATÖLYE > Personel İş Takibi Excel çıktısı (openpyxl): Geçmiş / Kişi özeti / İstasyon özeti — üç sayfa. Satırlar ``is_takibi`` sözlükleridir."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.tarih import _TR

BASLIK_DOLGU = PatternFill("solid", fgColor="15294D")
GECMIS_KOLON = ("Tarih", "Personel / Ekip", "İstasyon", "Ürün Kodu", "Ürün", "Varyant", "Sağlam", "Fire", "Fire %", "Süre (dk)", "Adam-saat", "Kaynak", "Durum", "Kayıt No")
GECMIS_GEN = (11, 32, 18, 16, 34, 12, 9, 8, 8, 10, 10, 26, 18, 14)
KISI_KOLON = ("Personel", "Katıldığı sağlam", "Pay (adam-saniye oranı)", "Fire", "Fire %", "Oturum", "Süresiz kayıt", "İş süresi (dk)", "Mesai (dk)", "Verimlilik % (oturum ÷ mesai)")
KISI_GEN = (28, 16, 20, 9, 8, 9, 13, 14, 12, 24)
IST_KOLON = ("İstasyon", "Sağlam", "Fire", "Fire %", "Oturum", "Süresiz kayıt", "Oturum süresi (dk)", "Adam-saat", "Adet/saat")
IST_GEN = (26, 10, 8, 8, 9, 13, 16, 10, 10)


def _sayfa(wb, ad, baslik, alt, kolonlar, genislik, ilk=False):
    ws = wb.active if ilk else wb.create_sheet()
    ws.title = ad
    ws["A1"] = baslik
    ws["A1"].font = Font(bold=True, size=13, color="15294D")
    ws["A2"] = alt
    ws["A2"].font = Font(color="5B6678")
    for i, k in enumerate(kolonlar, start=1):
        c = ws.cell(row=4, column=i, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BASLIK_DOLGU
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(genislik, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    return ws


def takip_xlsx(satirlar, kisiler, istasyonlar, bas, bit) -> bytes:
    wb = Workbook()
    aralik = f"{bas:%d.%m.%Y} – {bit:%d.%m.%Y}"
    ws = _sayfa(wb, "Geçmiş", "Personel İş Takibi — Geçmiş", aralik + " · ekip işinde sağlam/fire ekipçe; oturumsuz (ofis/emirsiz) kayıtta süre yok", GECMIS_KOLON, GECMIS_GEN, ilk=True)
    r = 5
    for s in satirlar:
        ws.append([s["tarih"], s["personel_metni"], f"{s['istasyon'].kod} {s['istasyon'].ad}", s["cikti"].kod, s["cikti"].ad, s["varyant"] or "",
                   s["saglam"], s["fire"], s["fire_yuzde"], (s["sure_saniye"] // 60) if s["sure_saniye"] is not None else None,
                   round(s["adam_saniye"] / 3600, 2) if s["adam_saniye"] else None, s["kaynak_metni"],
                   s["durum"] + (" · kayıt geri alındı" if s["geri_alindi"] else ""), s["kayit"].no if s["kayit"] is not None else ""])
        ws.cell(row=r, column=1).number_format = "DD.MM.YYYY"
        for col in (7, 8):
            ws.cell(row=r, column=col).number_format = "#,##0.###"
        ws.cell(row=r, column=9).number_format = "0.0"
        r += 1
    ws = _sayfa(wb, "Kişi özeti", "Personel İş Takibi — Kişi özeti", aralik + " · verimlilik = oturum süresi ÷ mesai süresi", KISI_KOLON, KISI_GEN)
    r = 5
    for k in kisiler:
        ws.append([k["personel"].ad_soyad, k["saglam"], k["pay"], k["fire"], k["fire_yuzde"], k["oturum"], k["suresiz"], k["is_saniye"] // 60, k["mesai_saniye"] // 60, k["verimlilik"]])
        for col in (2, 3, 4):
            ws.cell(row=r, column=col).number_format = "#,##0.###"
        for col in (5, 10):
            ws.cell(row=r, column=col).number_format = "0.0"
        r += 1
    ws = _sayfa(wb, "İstasyon özeti", "Personel İş Takibi — İstasyon özeti", aralik, IST_KOLON, IST_GEN)
    r = 5
    for k in istasyonlar:
        ws.append([f"{k['istasyon'].kod} {k['istasyon'].ad}", k["saglam"], k["fire"], k["fire_yuzde"], k["oturum"], k["suresiz"], k["sure_saniye"] // 60, k["adam_saat"], k["adet_saat"]])
        for col in (2, 3):
            ws.cell(row=r, column=col).number_format = "#,##0.###"
        for col in (4, 9):
            ws.cell(row=r, column=col).number_format = "0.0"
        ws.cell(row=r, column=8).number_format = "0.00"
        r += 1
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
