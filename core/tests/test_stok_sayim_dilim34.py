"""Stok Sayımı — dilim 3 ekranları (onay / geri alma / donmuş değerler) + dilim 4 (Excel indir-yükle, stok kartı "Son sayım" + "Sayım SY-…" kaynağı, raporlar)."""
import io
from decimal import Decimal

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from core.models import EkranYetki, StokHareket, StokSayimGirdisi
from core.services import stok_sayim as ss
from core.services import stok_sayim_excel as sx
from core.tests.test_stok_sayim_onay import BUGUN, OnayTaban

D = Decimal


class Taban(OnayTaban):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.personel = User.objects.create_user("sayici", password="x")
        EkranYetki.objects.create(kullanici=cls.personel, ekran_kod="stok_sayim")
        cls.stokcu = User.objects.create_user("stokcu", password="x")
        EkranYetki.objects.create(kullanici=cls.stokcu, ekran_kod="stoklar")

    def hazir(self, kor=False):
        a, b = self.kart("151", ad="ayak a"), self.kart("152", ad="mamul b")
        self.giris(a, "10", "1000")
        self.giris(b, "5", "1000")
        s = ss.sayim_ac(depo_id=self.depo.pk, tarih=BUGUN, kor=kor, kullanici=self.yon)
        return a, b, s

    def xlsx_doldur(self, icerik, degerler, sayim_no=None):
        """Şablon baytlarını açıp {stok kodu: sayılan} değerlerini 'Sayılan' sütununa yazar."""
        wb = load_workbook(io.BytesIO(icerik))
        ws = wb.active
        if sayim_no:
            ws["B1"] = sayim_no
        baslik = next(r for r in range(1, 12) if ws.cell(row=r, column=1).value == sx.BASLIK_KIMLIK)
        kolon = [c.value for c in ws[baslik]].index(sx.BASLIK_SAYILAN) + 1
        for r in range(baslik + 1, ws.max_row + 1):
            kod = ws.cell(row=r, column=2).value
            if kod in degerler:
                ws.cell(row=r, column=kolon, value=degerler[kod])
        cikti = io.BytesIO()
        wb.save(cikti)
        return cikti.getvalue()

    def yukle(self, s, icerik, ad="sayim.xlsx"):
        return self.client.post(reverse("core:stok_sayim_excel_yukle", args=[s.pk]), {"dosya": SimpleUploadedFile(ad, icerik, content_type=sx_turu())})


def sx_turu():
    return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class OnayEkraniTest(Taban):
    def test_onay_ve_geri_alma_ekrandan_donmus_degerler(self):
        a, b, s = self.hazir()
        for st, m in ((a, "8"), (b, "7")):
            sat = s.satirlar.get(stok=st)
            ss.girdi_ekle(sat, kullanici=self.personel, miktar=m)
            ss.karar_ver(sat, "DUZELT", kullanici=self.yon)
        self.client.force_login(self.yon)
        belge = reverse("core:stok_sayim_belge", args=[s.pk])
        self.assertContains(self.client.get(belge), "Önce sayım girişini kapatın")
        self.client.post(reverse("core:stok_sayim_durum", args=[s.pk]), {"islem": "kapat"})
        r = self.client.get(belge)
        self.assertContains(r, "Sayımı onayla")
        self.assertContains(r, "62x ↔ stok hesabı")
        r = self.client.post(reverse("core:stok_sayim_onay", args=[s.pk]), follow=True)
        self.assertContains(r, "onaylandı")
        s.refresh_from_db()
        self.assertEqual(s.durum, "ONAYLI")
        self.assertContains(r, f"{s.fis.yil}/{s.fis.fis_no}")
        self.assertContains(r, "Onayı geri al")
        self.assertTrue(r.context["donmus"])
        d = {x["stok"].pk: x for x in r.context["satirlar"]}
        self.assertEqual((d[a.pk]["referans"], d[a.pk]["sayilan"], d[a.pk]["fark"], d[a.pk]["tutar"]), (D("10"), D("8"), D("-2"), D("-200.00")))     # canlı fark 0 değil: onay anı değerleri
        self.assertEqual(d[b.pk]["tutar"], D("400.00"))
        self.assertEqual(r.context["ozet"]["duzelt_adet"], 2)
        r = self.client.post(reverse("core:stok_sayim_geri_al", args=[s.pk]), follow=True)
        self.assertContains(r, "geri alındı")
        self.assertContains(r, "yeniden onaylanamaz")
        s.refresh_from_db()
        self.assertEqual(s.durum, "GERI_ALINDI")
        self.assertEqual(StokHareket.objects.filter(kaynak="SAYIM", silindi=False).count(), 0)

    def test_onay_hata_mesaji_ve_yetki(self):
        a, b, s = self.hazir()
        ss.giris_kapat(s, self.yon)
        self.client.force_login(self.yon)
        r = self.client.post(reverse("core:stok_sayim_onay", args=[s.pk]), follow=True)
        self.assertContains(r, "Onaylanacak satır yok")
        for kisi in (self.personel, self.stokcu):
            self.client.force_login(kisi)
            self.assertEqual(self.client.post(reverse("core:stok_sayim_onay", args=[s.pk])).status_code, 403)
            self.assertEqual(self.client.post(reverse("core:stok_sayim_geri_al", args=[s.pk])).status_code, 403)


class ExcelTest(Taban):
    def sablon(self, s, kullanici):
        self.client.force_login(kullanici)
        r = self.client.get(reverse("core:stok_sayim_excel_indir", args=[s.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        return r.content

    def test_sablon_yapisi_ve_kor_belgede_sistem_sutunu_yok(self):
        a, b, s = self.hazir(kor=False)
        wb = load_workbook(io.BytesIO(self.sablon(s, self.personel)))
        ws = wb.active
        self.assertEqual((ws["B1"].value, ws.column_dimensions["A"].hidden), (s.no, True))
        basliklar = [c.value for c in ws[6]]
        self.assertEqual(basliklar, ["Satır Kimliği", "Stok Kodu", "Ürün", "Varyant", "Birim", "Sistem (kesme anı)", "Sayılan"])
        self.assertEqual(ws.cell(row=7, column=6).value, 10.0)
        kor = ss.sayim_ac(depo_id=self.depo.pk, tarih=BUGUN, kor=True, kullanici=self.yon)
        ws2 = load_workbook(io.BytesIO(self.sablon(kor, self.personel))).active
        self.assertEqual([c.value for c in ws2[6]], ["Satır Kimliği", "Stok Kodu", "Ürün", "Varyant", "Birim", "Sayılan"])      # KÖR: sistem sütunu dosyada YOK
        tum = " ".join(str(c.value) for row in ws2.iter_rows() for c in row if c.value is not None)
        self.assertNotIn("Sistem", tum)
        self.assertNotIn("10.0", tum.replace("Satır", ""))                                                                  # sistem miktarı sızmaz

    def test_yetki_kapali_belge_ve_fark_excel_yalniz_yonetici(self):
        a, b, s = self.hazir()
        self.client.force_login(self.stokcu)
        self.assertEqual(self.client.get(reverse("core:stok_sayim_excel_indir", args=[s.pk])).status_code, 403)
        self.client.force_login(self.personel)
        self.assertEqual(self.client.get(reverse("core:stok_sayim_excel_indir", args=[s.pk]), {"tur": "fark"}).status_code, 403)
        ss.girdi_ekle(s.satirlar.get(stok=a), kullanici=self.personel, miktar="9")
        self.client.force_login(self.yon)
        r = self.client.get(reverse("core:stok_sayim_excel_indir", args=[s.pk]), {"tur": "fark"})
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[4]][:7], ["Stok Kodu", "Ürün", "Varyant", "Sistem (kesme)", "Sistem (referans)", "Sayılan", "Fark"])
        satir = next(row for row in ws.iter_rows(min_row=5, values_only=True) if row[0] == a.kod)
        self.assertEqual((satir[3], satir[5], satir[6]), (10.0, 9.0, -1.0))
        ss.giris_kapat(s, self.yon)
        self.client.force_login(self.personel)
        r = self.client.get(reverse("core:stok_sayim_excel_indir", args=[s.pk]), follow=True)
        self.assertContains(r, "sayım girişi kapalı")

    def test_yukleme_onizleme_uygula_toplamsal_ve_tekrar_yok(self):
        a, b, s = self.hazir()
        ss.girdi_ekle(s.satirlar.get(stok=a), kullanici=self.personel, miktar="2")           # mevcut giriş: Excel EKLENİR
        dosya = self.xlsx_doldur(self.sablon(s, self.personel), {a.kod: 6, b.kod: "7,5"})
        self.client.force_login(self.personel)
        r = self.yukle(s, dosya)
        self.assertContains(r, "Excel yükleme önizlemesi")
        self.assertEqual(len(r.context["satirlar"]), 2)
        self.assertEqual(StokSayimGirdisi.objects.count(), 1)                                 # önizleme henüz yazmaz
        token = r.context["token"]
        r = self.client.post(reverse("core:stok_sayim_excel_yukle", args=[s.pk]), {"token": token}, follow=True)
        self.assertContains(r, "2 satıra sayım girişi kaydedildi")
        sayilan = ss.sayilanlar(s)
        self.assertEqual((sayilan[s.satirlar.get(stok=a).pk], sayilan[s.satirlar.get(stok=b).pk]), (D("8"), D("7.5")))     # 2 + 6 ; 7,5
        g = StokSayimGirdisi.objects.filter(kaynak="EXCEL").first()
        self.assertEqual((g.kullanici_id, g.aciklama.startswith("Excel ")), (self.personel.pk, True))
        r = self.yukle(s, dosya)                                                              # aynı dosya: önizleme uyarır, uygulama reddedilir
        self.assertContains(r, "daha önce yüklenmiş")
        r = self.client.post(reverse("core:stok_sayim_excel_yukle", args=[s.pk]), {"token": r.context["token"]}, follow=True)
        self.assertContains(r, "daha önce yüklenmiş; tekrar yüklenmedi")
        self.assertEqual(StokSayimGirdisi.objects.filter(kaynak="EXCEL").count(), 2)

    def test_hatali_dosyalar_ve_satirlar(self):
        a, b, s = self.hazir()
        diger = ss.sayim_ac(depo_id=self.depo.pk, tarih=BUGUN, kullanici=self.yon)
        self.client.force_login(self.personel)
        sablon = self.sablon(s, self.personel)
        r = self.client.post(reverse("core:stok_sayim_excel_yukle", args=[s.pk]), {"dosya": SimpleUploadedFile("y.xlsx", self.xlsx_doldur(sablon, {a.kod: 1}, sayim_no=diger.no))}, follow=True)
        self.assertContains(r, "belgesine ait değil")
        r = self.client.post(reverse("core:stok_sayim_excel_yukle", args=[s.pk]), {"dosya": SimpleUploadedFile("k.xlsx", b"bu bir xlsx degil")}, follow=True)
        self.assertContains(r, "Yalnızca .xlsx")
        r = self.client.post(reverse("core:stok_sayim_excel_yukle", args=[s.pk]), {"dosya": SimpleUploadedFile("b.xlsx", b"PK" + b"x" * (sx.MAKS_BOYUT + 10))}, follow=True)
        self.assertContains(r, "2 MB")
        r = self.client.post(reverse("core:stok_sayim_excel_yukle", args=[s.pk]), {}, follow=True)
        self.assertContains(r, "dosyası seçin")
        # satır hataları: negatif, metin, bilinmeyen kimlik, tekrar
        wb = load_workbook(io.BytesIO(sablon))
        ws = wb.active
        ws.cell(row=7, column=7, value=-3)
        ws.cell(row=8, column=7, value="abc")
        ws.cell(row=9, column=1, value=999999)
        ws.cell(row=9, column=7, value=2)
        cikti = io.BytesIO()
        wb.save(cikti)
        r = self.yukle(s, cikti.getvalue())
        mesajlar = " ".join(m for _n, m in r.context["hatalar"])
        self.assertIn("negatif", mesajlar)
        self.assertIn("geçerli bir sayı değil", mesajlar)
        self.assertIn("belgesinde yok", mesajlar)
        self.assertEqual(len(r.context["satirlar"]), 0)
        self.assertEqual(StokSayimGirdisi.objects.count(), 0)

    def test_formul_enjeksiyonu_hucreye_guvenli_yazilir_ve_yukleme_yetkisi(self):
        kotu = self.kart("151", ad="=HYPERLINK(\"http://x\")")
        self.giris(kotu, "3", "300")
        s = ss.sayim_ac(depo_id=self.depo.pk, tarih=BUGUN, kullanici=self.yon)
        ws = load_workbook(io.BytesIO(self.sablon(s, self.personel))).active
        degerler = [ws.cell(row=r, column=3).value for r in range(7, ws.max_row + 1)]
        self.assertIn("'=HYPERLINK(\"HTTP://X\")", degerler)
        self.client.force_login(self.stokcu)
        self.assertEqual(self.client.post(reverse("core:stok_sayim_excel_yukle", args=[s.pk]), {}).status_code, 403)


class KartVeRaporTest(Taban):
    def onayli(self):
        a, b, s = self.hazir()
        for st, m in ((a, "8"), (b, "7")):
            sat = s.satirlar.get(stok=st)
            ss.girdi_ekle(sat, kullanici=self.personel, miktar=m)
            ss.karar_ver(sat, "DUZELT", kullanici=self.yon)
        ss.giris_kapat(s, self.yon)
        return a, b, ss.onayla(s, self.yon)

    def test_stok_kartinda_son_sayim_ve_sayim_kaynagi(self):
        a, b, _ = self.hazir()
        self.client.force_login(self.yon)
        self.assertContains(self.client.get(reverse("core:stok_detay", args=[a.pk])), "Son sayım: henüz sayılmadı")
        a, b, s = self.onayli()
        r = self.client.get(reverse("core:stok_detay", args=[a.pk]))
        self.assertContains(r, "Son sayım:")
        self.assertContains(r, s.no)
        self.assertContains(r, f"Sayım {s.no}")                                                                     # hareket kaynağı
        self.assertContains(r, reverse("core:stok_sayim_belge", args=[s.pk]))                                       # yönetici: belgeye bağlantı
        self.assertEqual((r.context["son_sayim"]["sayilan"], r.context["son_sayim"]["fark"], r.context["son_sayim"]["duzeltildi"]), (D("8"), D("-2"), True))
        self.client.force_login(self.stokcu)                                                                         # yönetici olmayan: metin, bağlantı yok
        r = self.client.get(reverse("core:stok_detay", args=[a.pk]))
        self.assertContains(r, f"Sayım {s.no}")
        self.assertNotContains(r, reverse("core:stok_sayim_belge", args=[s.pk]))
        ss.geri_al(s, self.yon)                                                                                      # geri alınınca "son sayım" onaylı belge olmadığından düşer
        self.client.force_login(self.yon)
        self.assertContains(self.client.get(reverse("core:stok_detay", args=[a.pk])), "Son sayım: henüz sayılmadı")

    def test_gecmis_ve_stok_bazinda_rapor(self):
        a, b, s = self.onayli()
        acik = ss.sayim_ac(depo_id=self.depo.pk, tarih=BUGUN, kullanici=self.yon)
        self.client.force_login(self.yon)
        r = self.client.get(reverse("core:stok_sayim_raporlari"))
        self.assertContains(r, s.no)
        self.assertContains(r, acik.no)
        satir = next(x for x in r.context["satirlar"] if x["sayim"].pk == s.pk)
        self.assertEqual((satir["duzelt"], satir["fazla_tutar"], satir["eksik_tutar"], satir["fis"].pk), (2, D("400.00"), D("200.00"), s.fis_id))
        self.assertIsNone(next(x for x in r.context["satirlar"] if x["sayim"].pk == acik.pk)["fazla_tutar"])
        self.assertNotContains(self.client.get(reverse("core:stok_sayim_raporlari"), {"durum": "ONAYLI"}), acik.no)
        x = self.client.get(reverse("core:stok_sayim_raporlari"), {"excel": "1"})
        ws = load_workbook(io.BytesIO(x.content)).active
        self.assertEqual([c.value for c in ws[1]][:3], ["Sayım No", "Tarih", "Depo"])
        self.assertIn(s.no, [row[0] for row in ws.iter_rows(min_row=2, values_only=True)])
        # stok bazında
        r = self.client.get(reverse("core:stok_sayim_rapor_stok"), {"ara": a.kod})
        self.assertEqual(r.context["stok"].pk, a.pk)                                                                # tek eşleşme doğrudan açılır
        sat = r.context["satirlar"]
        self.assertEqual(len(sat), 2)                                                                                # onaylı belge + açık belge (devam ediyor)
        onayli_satir = next(x for x in sat if x["sayim"].pk == s.pk)
        self.assertEqual((onayli_satir["fark"], onayli_satir["sayilan"], onayli_satir["tutar"], onayli_satir["devam"]), (D("-2"), D("8"), D("-200.00"), False))
        self.assertTrue(next(x for x in sat if x["sayim"].pk == acik.pk)["devam"])
        x = self.client.get(reverse("core:stok_sayim_rapor_stok"), {"stok": a.pk, "excel": "1"})
        ws = load_workbook(io.BytesIO(x.content)).active
        self.assertEqual(ws["A1"].value, f"{a.kod} — {a.ad}")
        self.assertContains(self.client.get(reverse("core:stok_sayim_rapor_stok"), {"ara": "yokböyleştok"}), "Eşleşen stok bulunamadı")

    def test_raporlar_yalniz_yonetici_ve_menude(self):
        for kisi in (self.personel, self.stokcu):
            self.client.force_login(kisi)
            self.assertEqual(self.client.get(reverse("core:stok_sayim_raporlari")).status_code, 403)
            self.assertEqual(self.client.get(reverse("core:stok_sayim_rapor_stok")).status_code, 403)
        from core.yetki import kullanici_menusu
        for kisi in (self.personel, self.stokcu):
            (stoklar,) = [m for m in kullanici_menusu(kisi) if m.kod == "STOKLAR"]
            self.assertNotIn("stok_sayim_raporlari", [e.kod for e in stoklar.ekranlar])
        (yon_menu,) = [m for m in kullanici_menusu(self.yon) if m.kod == "STOKLAR"]
        self.assertIn("stok_sayim_raporlari", [e.kod for e in yon_menu.ekranlar])


class KurBeklemeTest(Taban):
    def test_kur_yoksa_onay_net_mesaj_ve_ekranda_uyari(self):
        import datetime
        from core.models import Kur, StokSayimi
        from core.services.hareket import eldeki_miktar
        a, b, s = self.hazir()
        sat = s.satirlar.get(stok=a)
        ss.girdi_ekle(sat, kullanici=self.personel, miktar="8")
        ss.karar_ver(sat, "DUZELT", kullanici=self.yon)
        ss.giris_kapat(s, self.yon)
        Kur.objects.filter(tarih=s.tarih).delete()                                                          # bugünün kuru henüz yayınlanmadı (D+1)
        self.assertTrue(ss.kur_bekleniyor_mu(s))
        with self.assertRaisesMessage(ss.SayimHatasi, "Kur yayınlanınca"):
            ss.onayla(s, self.yon)
        self.assertEqual((StokHareket.objects.filter(kaynak="SAYIM").count(), eldeki_miktar(a)), (0, D("10")))
        self.client.force_login(self.yon)
        belge = reverse("core:stok_sayim_belge", args=[s.pk])
        r = self.client.get(belge)
        self.assertContains(r, "Kur yayınlanınca")
        self.assertNotContains(r, "Sayımı onayla")
        r = self.client.post(reverse("core:stok_sayim_onay", args=[s.pk]), follow=True)
        self.assertContains(r, "USD kuru henüz yok")
        Kur.objects.create(tarih=s.tarih, usd_alis=D("30"))                                                  # kur yayınlandı
        self.assertContains(self.client.get(belge), "Sayımı onayla")
        r = self.client.post(reverse("core:stok_sayim_onay", args=[s.pk]), follow=True)
        self.assertContains(r, "onaylandı")
        s.refresh_from_db()
        self.assertEqual(s.durum, "ONAYLI")
