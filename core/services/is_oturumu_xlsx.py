"""ÜRETİM > İş Oturumları listesinin Excel çıktısı (openpyxl). Satırlar ``is_oturumu.ozet`` sözlükleridir; burada yalnız dizilim var."""
from __future__ import annotations

from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_TR = ZoneInfo("Europe/Istanbul")
BASLIK_DOLGU = PatternFill("solid", fgColor="15294D")
KOLONLAR = ("Tarih", "İstasyon", "ÜS", "İstasyon Emri", "Ürün Kodu", "Ürün", "Ekip", "Kişi", "Başlangıç", "Bitiş", "Süre (dk)",
            "Adam-saat", "Durum", "Adet", "Adet/saat", "Adet başı (sn)", "Kayıt No", "Başlatan", "Bitiren")
GENISLIKLER = (11, 18, 14, 14, 16, 34, 30, 6, 10, 10, 10, 10, 18, 9, 10, 13, 14, 20, 20)


def _saat(dt):
    return dt.astimezone(_TR).strftime("%H:%M") if dt else ""


def oturumlar_xlsx(satirlar, bas, bit) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "İş Oturumları"
    ws["A1"] = "İş Oturumları"
    ws["A1"].font = Font(bold=True, size=13, color="15294D")
    ws["A2"] = f"{bas:%d.%m.%Y} – {bit:%d.%m.%Y}"
    ws["A2"].font = Font(color="5B6678")
    for i, k in enumerate(KOLONLAR, start=1):
        c = ws.cell(row=4, column=i, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BASLIK_DOLGU
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(GENISLIKLER, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    r = 5
    for s in satirlar:
        o = s["oturum"]
        ie = o.istasyon_emri
        ws.append([o.is_tarihi, f"{o.istasyon.kod} {o.istasyon.ad}", o.uretim_emri.no, ie.no, ie.operasyon.cikti.kod, ie.operasyon.cikti.ad,
                   s["ekip_metni"], s["kisi"], _saat(o.baslangic_zamani), _saat(o.bitis_zamani), s["sure_saniye"] // 60, s["adam_saat"],
                   o.get_durum_display(), s["adet"], s["adet_saat"], s["adet_basi_saniye"], o.kayit.no if o.kayit_id else "",
                   o.baslatan.ad_soyad, o.bitiren.ad_soyad if o.bitiren_id else ""])
        ws.cell(row=r, column=1).number_format = "DD.MM.YYYY"
        ws.cell(row=r, column=12).number_format = "0.00"
        ws.cell(row=r, column=14).number_format = "#,##0"
        ws.cell(row=r, column=15).number_format = "0.0"
        r += 1
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
