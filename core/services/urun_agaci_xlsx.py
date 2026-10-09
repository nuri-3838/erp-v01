"""ÜRETİM > Ürün Ağacı görünümlerinin Excel çıktıları (openpyxl). Hesap ``urun_agaci`` servisinden gelir; burada yalnız dizilim var.
Sayılar Decimal olarak yazılır (float'a çevrilmez), biçimi hücre number_format'ı belirler."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.services.urun_agaci import KAYNAK_AD

MIKTAR_FMT = "#,##0.0000"
PARA_FMT = "#,##0.00"
YUZDE_FMT = "0.0"
BASLIK_DOLGU = PatternFill("solid", fgColor="15294D")
GRUP_DOLGU = PatternFill("solid", fgColor="EEF1F7")
KIRMIZI = Font(color="B3253F")
GRI = Font(color="8A93A5", italic=True)


def _yeni(baslik, alt_baslik, kolonlar, genislikler):
    wb = Workbook()
    ws = wb.active
    ws.title = baslik[:31]
    ws["A1"] = baslik
    ws["A1"].font = Font(bold=True, size=13, color="15294D")
    ws["A2"] = alt_baslik
    ws["A2"].font = Font(color="5B6678")
    for i, k in enumerate(kolonlar, start=1):
        c = ws.cell(row=4, column=i, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BASLIK_DOLGU
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(genislikler, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    return wb, ws


def _bayt(wb) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _grup_satiri(ws, satir_no, metin, kolon_sayisi):
    for i in range(1, kolon_sayisi + 1):
        ws.cell(row=satir_no, column=i).fill = GRUP_DOLGU
    ws.cell(row=satir_no, column=1, value=metin).font = Font(bold=True)


def _sayi(ws, satir, kolon, deger, fmt):
    c = ws.cell(row=satir, column=kolon, value=deger)
    if deger is not None:
        c.number_format = fmt
    return c


def _ad(stok):
    return stok.ad


# --- Ürün / Ağaç -----------------------------------------------------------------------------------------------------------

def _agac_duz(dugum, derinlik=0):
    yield derinlik, dugum
    for c in dugum["cocuklar"]:
        yield from _agac_duz(c, derinlik + 1)


def agac_xlsx(agac, miktar, baslik_ek="") -> bytes:
    wb, ws = _yeni(f"Ürün Ağacı — {agac['stok'].kod} {agac['stok'].ad}", f"Miktar: {miktar}  {baslik_ek}".strip(),
                   ["Seviye", "Kod", "Ad", "Miktar", "İstasyon", "Not"], [8, 16, 46, 14, 30, 60])
    r = 5
    for derinlik, d in _agac_duz(agac):
        notlar = []
        if d.get("tam"):
            notlar.append(f"tam boy · çalıştırma {d['calistirma'].normalize():f} · üretilecek {d['uretilecek'].normalize():f}")
        for y in d.get("yan_ciktilar", []):
            notlar.append(f"yan çıktı: {y['miktar'].normalize():f} × {y['stok'].kod}")
        if d.get("tekrar"):
            notlar.append("↑ aynı stok yukarıda toplam talep üzerinden hesaplandı")
        ws.cell(row=r, column=1, value=derinlik)
        ws.cell(row=r, column=2, value=("    " * derinlik) + d["stok"].kod)
        ws.cell(row=r, column=3, value=_ad(d["stok"]))
        _sayi(ws, r, 4, d["miktar"], MIKTAR_FMT)
        ws.cell(row=r, column=5, value=(f"{d['istasyon'].kod} · {d['istasyon'].ad}" if d.get("istasyon") else "hazır stok"))
        ws.cell(row=r, column=6, value="; ".join(notlar))
        r += 1
    return _bayt(wb)


# --- Ürün / Malzeme ve Stok ------------------------------------------------------------------------------------------------

def malzeme_xlsx(urun, miktar, sonuc, depo_ad="Tüm depolar") -> bytes:
    wb, ws = _yeni(f"Malzeme ve Stok — {urun.kod} {urun.ad}",
                   f"Miktar: {miktar} · Depo: {depo_ad} · {sonuc['yeterli']} kalem yeterli · {sonuc['eksik']} kalem eksik "
                   "(açık kayıtlara ayrılan stok düşülmez)",
                   ["Kategori", "Kod", "Ad", "Birim", "Gerekli", "Eldeki stok", "Eksik"], [28, 16, 46, 8, 14, 14, 14])
    r = 5
    for g in sonuc["gruplar"]:
        for s in g["satirlar"]:
            ws.cell(row=r, column=1, value=g["kategori"])
            ws.cell(row=r, column=2, value=s["stok"].kod)
            ws.cell(row=r, column=3, value=_ad(s["stok"]))
            ws.cell(row=r, column=4, value=s["birim"])
            _sayi(ws, r, 5, s["gerekli"], MIKTAR_FMT)
            _sayi(ws, r, 6, s["eldeki"], MIKTAR_FMT)
            c = _sayi(ws, r, 7, s["eksik"], MIKTAR_FMT)
            if s["eksik"] > 0:
                c.font = KIRMIZI
                ws.cell(row=r, column=2).font = KIRMIZI
            r += 1
    return _bayt(wb)


# --- Ürün / Maliyet --------------------------------------------------------------------------------------------------------

def maliyet_xlsx(urun, miktar, sonuc) -> bytes:
    """TL ve USD AYRI sütunlarda (birim ve — miktar 1 değilse — seçilen miktar için); kaynak sütunu: Ortalama / Alış fiyatı (kart) / Yok."""
    n = miktar != 1
    ozet = (f"{sonuc['ortalama_sayi']} kalem ortalama · {sonuc['kart_sayi']} kalem alış fiyatı · {sonuc['maliyetsiz']} kalem maliyetsiz"
            + (f" (alış fiyatı payı %{sonuc['kart_pay']:.1f})" if sonuc["kart_pay"] is not None and sonuc["kart_sayi"] else ""))
    uyari = f" · ⚠ {sonuc['maliyetsiz']} kalemin maliyeti yok — toplam eksik" if sonuc["maliyetsiz"] else ""
    kolonlar = ["Kategori", "Kod", "Ad", "Birim", "Kaynak", "Birim tüketim", "Birim maliyet TL", "Birim maliyet USD", "Tutar TL", "Tutar USD", "Pay %"]
    genis = [28, 16, 46, 8, 18, 14, 16, 16, 14, 14, 9]
    if n:
        kolonlar += [f"Tutar TL ({miktar.normalize():f} adet)", f"Tutar USD ({miktar.normalize():f} adet)"]
        genis += [18, 18]
    alt = f"1 adet için kesirli malzeme maliyeti · {ozet}{uyari}" + (f" · {sonuc['kur_notu']}" if sonuc["kur_notu"] else "")
    wb, ws = _yeni(f"Maliyet — {urun.kod} {urun.ad}", alt, kolonlar, genis)
    r = 5

    def toplam_hucreleri(satir, try_, usd, try_n, usd_n):
        _sayi(ws, satir, 9, try_, PARA_FMT).font = Font(bold=True)
        _sayi(ws, satir, 10, usd, PARA_FMT).font = Font(bold=True)
        if n:
            _sayi(ws, satir, 12, try_n, PARA_FMT).font = Font(bold=True)
            _sayi(ws, satir, 13, usd_n, PARA_FMT).font = Font(bold=True)

    for g in sonuc["gruplar"]:
        for s in g["satirlar"]:
            ws.cell(row=r, column=1, value=g["kategori"])
            ws.cell(row=r, column=2, value=s["stok"].kod)
            ws.cell(row=r, column=3, value=_ad(s["stok"]))
            ws.cell(row=r, column=4, value=s["birim"])
            ws.cell(row=r, column=5, value=KAYNAK_AD[s["kaynak"]])
            _sayi(ws, r, 6, s["tuketim"], MIKTAR_FMT)
            if s["maliyet_yok"]:
                ws.cell(row=r, column=7, value="maliyet yok").font = GRI
            else:
                _sayi(ws, r, 7, s["ort_try"], "#,##0.0000")
                _sayi(ws, r, 8, s["ort_usd"], "#,##0.0000")
                _sayi(ws, r, 9, s["tutar_try"], PARA_FMT)
                _sayi(ws, r, 10, s["tutar_usd"], PARA_FMT)
                _sayi(ws, r, 11, s["pay"], YUZDE_FMT)
                if n:
                    _sayi(ws, r, 12, s["tutar_try_n"], PARA_FMT)
                    _sayi(ws, r, 13, s["tutar_usd_n"], PARA_FMT)
            r += 1
        _grup_satiri(ws, r, f"{g['kategori']} ara toplam", len(kolonlar))
        toplam_hucreleri(r, g["toplam_try"], g["toplam_usd"], g["toplam_try_n"], g["toplam_usd_n"])
        r += 1
    _grup_satiri(ws, r, "GENEL TOPLAM", len(kolonlar))
    toplam_hucreleri(r, sonuc["toplam_try"], sonuc["toplam_usd"], sonuc["miktar_toplam_try"], sonuc["miktar_toplam_usd"])
    return _bayt(wb)


# --- Nerede kullanılıyor ---------------------------------------------------------------------------------------------------

def kullanim_xlsx(sonuc) -> bytes:
    stok = sonuc["stok"]
    wb, ws = _yeni(f"Nerede kullanılıyor — {stok.kod} {stok.ad}", f"{sonuc['sayi']} bitmiş ürün · tüketim 1 adet ürün başına ({sonuc['birim']})",
                   ["Seri", "Ürün kodu", "Ürün adı", "Tüketim", "Yol", "Not"], [14, 16, 46, 14, 90, 50])
    r = 5
    for g in sonuc["gruplar"]:
        for s in g["satirlar"]:
            ws.cell(row=r, column=1, value=g["ad"])
            ws.cell(row=r, column=2, value=s["urun"].kod)
            ws.cell(row=r, column=3, value=_ad(s["urun"]))
            _sayi(ws, r, 4, s["tuketim"], MIKTAR_FMT)
            ws.cell(row=r, column=5, value=" → ".join(f"{x.kod} {x.ad}" for x in s["yol"]))
            ws.cell(row=r, column=6, value=("yan çıktı olarak da üretiliyor: " + ", ".join(x.kod for x in s["yan_cikti_ureten"]))
                    if s["yan_cikti_ureten"] else "")
            r += 1
    return _bayt(wb)


# --- Karşılaştır -----------------------------------------------------------------------------------------------------------

def _karsilastir_sayfa(wb, ws, sonuc, ilk, baslik_ek=""):
    urunler = sonuc["urunler"]
    maliyet = sonuc["mod"] == "maliyet"
    pb = sonuc["pb"]
    ws.title = f"Maliyet {pb}" if maliyet else "Miktar"
    ws["A1"] = "Karşılaştırma"
    ws["A1"].font = Font(bold=True, size=13, color="15294D")
    alt = (f"Hücre = 1 adet ürün başına malzeme maliyeti ({pb})" if maliyet else "Hücre = 1 adet ürün başına kesirli tüketim")
    if maliyet and sonuc["kur_notu"]:
        alt += " · " + sonuc["kur_notu"]
    ws["A2"] = (alt + " " + baslik_ek).strip()
    ws["A2"].font = Font(color="5B6678")
    kolonlar = ["Kategori", "Kod", "Ad", "Birim"] + ([("Kaynak")] if maliyet else []) + [f"{u.kod}\n{u.ad}" for u in urunler]
    ofset = 5 if maliyet else 4
    for i, k in enumerate(kolonlar, start=1):
        c = ws.cell(row=4, column=i, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BASLIK_DOLGU
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate([26, 16, 40, 8] + ([18] if maliyet else []) + [16] * len(urunler), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    ws.row_dimensions[4].height = 48
    r = 5
    for g in sonuc["gruplar"]:
        for s in g["satirlar"]:
            ws.cell(row=r, column=1, value=g["kategori"])
            ws.cell(row=r, column=2, value=s["stok"].kod)
            ws.cell(row=r, column=3, value=_ad(s["stok"]))
            ws.cell(row=r, column=4, value=s["birim"])
            if maliyet:
                ws.cell(row=r, column=5, value=KAYNAK_AD[s["kaynak"]])
            for i, h in enumerate(s["hucreler"]):
                if h is None:
                    continue
                if maliyet and h["maliyet_yok"]:
                    ws.cell(row=r, column=ofset + 1 + i, value="maliyet yok").font = GRI
                else:
                    _sayi(ws, r, ofset + 1 + i, h["deger"], PARA_FMT if maliyet else MIKTAR_FMT)
            r += 1
    if maliyet:
        _grup_satiri(ws, r, f"ÜRÜN BAŞINA TOPLAM ({pb})", ofset + len(urunler))
        for i, t in enumerate(sonuc["toplamlar"]):
            _sayi(ws, r, ofset + 1 + i, t["toplam"], PARA_FMT).font = Font(bold=True)
        eksik = [t["maliyetsiz"] for t in sonuc["toplamlar"]]
        if any(eksik):
            r += 1
            ws.cell(row=r, column=1, value=f"{pb} maliyeti olmayan kalem sayısı").font = GRI
            for i, n in enumerate(eksik):
                if n:
                    ws.cell(row=r, column=ofset + 1 + i, value=n).font = KIRMIZI


def karsilastir_xlsx(sonuclar, baslik_ek="") -> bytes:
    """``sonuclar``: tek sayfa (Miktar) ya da maliyette [TL sonucu, USD sonucu] → iki AYRI sayfa."""
    wb = Workbook()
    for i, sonuc in enumerate(sonuclar):
        ws = wb.active if i == 0 else wb.create_sheet()
        _karsilastir_sayfa(wb, ws, sonuc, i, baslik_ek)
    return _bayt(wb)
