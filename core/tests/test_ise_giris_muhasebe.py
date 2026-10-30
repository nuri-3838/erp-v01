"""İşe giriş Dilim 3: muhasebe paketi — hazır mı kontrolü (eksik zorunlu alan + işe girişe kalan gün), Excel/PDF bilgi paketi, ücret yetkisi,
"muhasebeye gönderildi" işareti, GET yazmaz."""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.template.loader import render_to_string
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import IkAyari, IseGirisKaydi
from core.services import ise_giris as ig
from core.services import ise_giris_muhasebe as mu
from core.services.firma import firma_bilgisi_getir
from core.services.personel_ucret import ucret_ekle
from core.tarih import tr_bugun
from core.tests.ik_yardimci import OzelDizinTestTemel, personel_kur
from core.tests.test_ise_giris_paket import dolu_firma, dolu_personel, kullanici

BUGUN = date(2026, 10, 30)


def excel_hucreleri(veri: bytes) -> dict:
    """{etiket: değer} (A sütunu → B sütunu)."""
    ws = load_workbook(BytesIO(veri)).active
    return {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value}


class HazirlikTest(TestCase):
    def test_eksik_zorunlu_alanlar_ve_onerilenler(self):
        bos = personel_kur(ad="a", soyad="b", ise_giris_tarihi=BUGUN + timedelta(days=10))
        h = mu.hazirlik(bos, ig.kayit_oku(bos), bugun=BUGUN)
        self.assertFalse(h["hazir"])
        self.assertEqual(set(h["eksikler"]), {"TC kimlik no", "Doğum tarihi", "Doğum yeri", "Baba adı", "Ana adı", "Adres", "Görev"})
        self.assertEqual(set(h["onerilen_eksikler"]), {"IBAN (maaş ödemesi için)", "Telefon"})
        dolu = dolu_personel(ise_giris_tarihi=BUGUN + timedelta(days=10))
        h = mu.hazirlik(dolu, ig.kayit_oku(dolu), bugun=BUGUN)
        self.assertTrue(h["hazir"])
        self.assertEqual((h["eksikler"], h["onerilen_eksikler"]), ([], []))

    def test_kalan_gun_durumlari(self):
        for fark, durum, parca in ((10, "normal", "10 gün var"), (1, "acil", "1 gün kaldı"), (0, "acil", "bugün"), (-3, "gecmis", "3 gün önce işe başladı")):
            p = personel_kur(ad="k", soyad=f"s{fark}", ise_giris_tarihi=BUGUN + timedelta(days=fark))
            h = mu.hazirlik(p, ig.kayit_oku(p), bugun=BUGUN)
            self.assertEqual((h["durum"], h["kalan_gun"]), (durum, fark))
            self.assertIn(parca, h["uyari"])

    def test_gonderildiyse_uyari_yok(self):
        p = personel_kur(ad="g", soyad="g", ise_giris_tarihi=BUGUN + timedelta(days=1))
        mu.gonderildi_isaretle(p, tarih=BUGUN, bugun=BUGUN)
        h = mu.hazirlik(p, ig.kayit_oku(p), bugun=BUGUN)
        self.assertEqual((h["uyari"], h["gonderildi"]), ("", BUGUN))


class PaketIcerigiTest(OzelDizinTestTemel):
    def setUp(self):
        super().setUp()
        self.p = dolu_personel(cinsiyet="KADIN", medeni_durum="EVLI", telefon="05321112233")
        dolu_firma()
        self.firma = firma_bilgisi_getir()
        self.kayit = ig.kayit_oku(self.p)
        ucret_ekle(self.p, gecerlilik_baslangic=date(2026, 11, 2), tip="NET", net_tutar=Decimal("40000"))

    def test_ucret_yalniz_yetkiliye(self):
        yok = {b for b, _ in mu.bolumler(self.p, self.firma, self.kayit, ucret_gorebilir=False, bugun=BUGUN)}
        var = dict(mu.bolumler(self.p, self.firma, self.kayit, ucret_gorebilir=True, bugun=BUGUN))
        self.assertNotIn("Ücret", yok)
        self.assertEqual(var["Ücret"][0][:2], ("Ücret", "40.000,00 TL"))
        self.assertEqual(var["Ücret"][1][:2], ("Ücret geçerlilik başlangıcı", "02.11.2026"))
        self.assertEqual(var["Ücret"][2][:2], ("Ücret ödeme şekli/günü", "Bir sonraki ayın 10. iş günü"))

    def test_excel_alanlar_ve_ucret_yetkisi(self):
        h = excel_hucreleri(mu.excel(self.p, self.firma, self.kayit, ucret_gorebilir=False, bugun=BUGUN))
        self.assertEqual((h["Adı"], h["Soyadı"], h["T.C. Kimlik No"]), ("AYŞE", "DEMİR", "10000000146"))
        self.assertEqual((h["Doğum tarihi"], h["Doğum yeri"], h["Cinsiyet"], h["Medeni durum"]), ("12.04.1995", "KAYSERİ", "Kadın", "Evli"))
        self.assertEqual((h["Görev"], h["İşe giriş tarihi"], h["Sözleşme türü"], h["Deneme süresi"]), ("KAYNAKÇI", "02.11.2026", "Belirsiz süreli", "2 ay"))
        self.assertEqual((h["Unvan"], h["SGK işyeri sicil no"]), ("SEMTA ALÜMİNYUM A.Ş.", "1.2345.67"))
        self.assertTrue(h["IBAN"].startswith("TR33 0006"))
        self.assertNotIn("Ücret", h)                                                            # ücret yetkisi yok
        h = excel_hucreleri(mu.excel(self.p, self.firma, self.kayit, ucret_gorebilir=True, bugun=BUGUN))
        self.assertEqual(h["Ücret"], "40.000,00 TL")

    def test_excel_eksik_zorunlu_alan_isaretlenir(self):
        bos = personel_kur(ad="Veli", soyad="Kaya")
        h = excel_hucreleri(mu.excel(bos, self.firma, ig.kayit_oku(bos), bugun=BUGUN))
        for etiket in ("T.C. Kimlik No", "Doğum tarihi", "Doğum yeri", "Baba adı", "Ana adı", "Adres", "Görev"):
            self.assertEqual(h[etiket], "EKSİK", etiket)
        self.assertEqual(h["E-posta"], None)                                                    # zorunlu olmayan boş alan boş kalır, EKSİK yazmaz
        self.assertEqual(h["Hesap sahibi"], "VELİ KAYA")

    def test_excel_tc_metin_olarak_saklanir(self):
        ws = load_workbook(BytesIO(mu.excel(self.p, self.firma, self.kayit, bugun=BUGUN))).active
        hucre = next(ws.cell(row=r, column=2) for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "T.C. Kimlik No")
        self.assertEqual((hucre.value, hucre.data_type), ("10000000146", "s"))

    def test_pdf_belgesi_imzasiz_ve_ucret_kurali(self):
        b = mu.pdf_belgesi(self.p, self.firma, self.kayit, ucret_gorebilir=False, bugun=BUGUN)
        html = render_to_string("core/ise_giris_paket_pdf.html", {"firma": self.firma, "logo_b64": None, "belgeler": [b], "belge_tarihi": BUGUN, "personel": self.p})
        self.assertIn("İŞE GİRİŞ BİLGİ FORMU (MUHASEBE İÇİN)", html)
        self.assertIn("10000000146", html)
        self.assertNotIn('class="imza"', html)                                                  # imzasız: işçi/işveren kutuları yok
        self.assertNotIn("40.000", html)
        b = mu.pdf_belgesi(self.p, self.firma, self.kayit, ucret_gorebilir=True, bugun=BUGUN)
        self.assertIn("40.000,00 TL", str([s for _b, s in b["bolumler"]]))

    def test_dosya_adi_turkce_karakter_temiz(self):
        p = personel_kur(ad="özgür", soyad="çağlayan şahin")
        self.assertEqual(mu.dosya_adi(p, "xlsx", BUGUN), "ise_giris_muhasebe_ozgur_caglayan_sahin_20261030.xlsx")


class GonderildiTest(TestCase):
    def test_isaretle_geri_al_ve_diger_alanlara_dokunmaz(self):
        p = personel_kur(ad="g", soyad="h")
        ig.paket_basildi(p)
        once = IseGirisKaydi.objects.get(personel=p).paket_basim_zamani
        k = mu.gonderildi_isaretle(p, tarih=BUGUN - timedelta(days=1), bugun=BUGUN)
        self.assertEqual(k.muhasebe_gonderim_tarihi, BUGUN - timedelta(days=1))
        k = IseGirisKaydi.objects.get(personel=p)
        self.assertEqual((k.muhasebe_gonderim_tarihi, k.paket_basim_zamani), (BUGUN - timedelta(days=1), once))
        mu.gonderildi_isaretle(p, geri_al=True, bugun=BUGUN)
        self.assertIsNone(IseGirisKaydi.objects.get(personel=p).muhasebe_gonderim_tarihi)

    def test_tarih_bos_bugun_ve_gelecek_reddedilir(self):
        p = personel_kur(ad="g", soyad="h")
        self.assertEqual(mu.gonderildi_isaretle(p, bugun=BUGUN).muhasebe_gonderim_tarihi, BUGUN)
        with self.assertRaises(ig.IseGirisHatasi):
            mu.gonderildi_isaretle(p, tarih=BUGUN + timedelta(days=1), bugun=BUGUN)

    def test_kayit_yokken_olusturur(self):
        p = personel_kur(ad="y", soyad="z")
        self.assertEqual(IseGirisKaydi.objects.count(), 0)
        mu.gonderildi_isaretle(p, bugun=BUGUN)
        self.assertEqual(IseGirisKaydi.objects.count(), 1)


class MuhasebeGorunumTest(OzelDizinTestTemel):
    def setUp(self):
        super().setUp()
        self.yon = User.objects.create_superuser("muyon", password="x")
        self.belgeci = kullanici("mubelgeci", "personel", "personel_belgeleri")
        self.ucretli = kullanici("muucretli", "personel", "personel_belgeleri", "personel_ucret")
        self.sade = kullanici("musade", "personel")
        self.p = dolu_personel(ise_giris_tarihi=tr_bugun() + timedelta(days=5))
        dolu_firma()
        self.xlsx = reverse("core:ise_giris_muhasebe_indir", args=[self.p.pk, "xlsx"])
        self.pdf = reverse("core:ise_giris_muhasebe_indir", args=[self.p.pk, "pdf"])
        self.gonder = reverse("core:ise_giris_muhasebe_gonderildi", args=[self.p.pk])
        self.sekme = reverse("core:personel_detay", args=[self.p.pk]) + "?sekme=ise_giris"

    def test_indirmeler_ve_yazmama(self):
        self.client.force_login(self.belgeci)
        r = self.client.get(self.xlsx)
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn("attachment", r["Content-Disposition"])
        self.assertIn("ise_giris_muhasebe_ayse_demir_", r["Content-Disposition"])
        self.assertIn("no-store", r["Cache-Control"])
        r = self.client.get(self.pdf)
        self.assertEqual((r.status_code, r["Content-Type"]), (200, "application/pdf"))
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertEqual(IseGirisKaydi.objects.count(), 0)                                      # indirme gönderim sayılmaz, kayıt oluşturmaz

    def test_gecersiz_bicim_404(self):
        self.client.force_login(self.belgeci)
        self.assertEqual(self.client.get(self.xlsx.replace("xlsx", "doc")).status_code, 404)

    def test_yetki(self):
        self.client.logout()
        self.assertEqual(self.client.get(self.xlsx).status_code, 302)
        self.assertEqual(self.client.post(self.gonder, {}).status_code, 302)
        for kisi in (self.sade, kullanici("muyalnizbelge", "personel_belgeleri")):               # personel + personel_belgeleri ikisi de gerekli
            self.client.force_login(kisi)
            for url in (self.xlsx, self.pdf):
                self.assertEqual(self.client.get(url).status_code, 403)
            self.assertEqual(self.client.post(self.gonder, {}).status_code, 403)
        self.client.force_login(self.yon)
        self.assertEqual(self.client.get(self.xlsx).status_code, 200)

    def test_excel_ucret_yalniz_ucret_yetkilisine(self):
        ucret_ekle(self.p, gecerlilik_baslangic=tr_bugun() + timedelta(days=5), tip="NET", net_tutar=Decimal("40000"))
        self.client.force_login(self.belgeci)
        self.assertNotIn("Ücret", excel_hucreleri(self.client.get(self.xlsx).content))
        self.client.force_login(self.ucretli)
        self.assertEqual(excel_hucreleri(self.client.get(self.xlsx).content)["Ücret"], "40.000,00 TL")

    def test_gonderildi_isaretle_ve_geri_al(self):
        self.client.force_login(self.belgeci)
        self.assertEqual(self.client.get(self.gonder).status_code, 302)                         # GET yazmaz
        self.assertEqual(IseGirisKaydi.objects.count(), 0)
        r = self.client.post(self.gonder, {"tarih": tr_bugun().isoformat()}, follow=True)
        self.assertRedirects(r, self.sekme + "#muhasebe-paketi", fetch_redirect_response=False)
        self.assertEqual(IseGirisKaydi.objects.get(personel=self.p).muhasebe_gonderim_tarihi, tr_bugun())
        html = self.client.get(self.sekme).content.decode()
        self.assertIn("Muhasebeye gönderildi: " + tr_bugun().strftime("%d.%m.%Y"), html)
        self.assertIn("Gönderildi işaretini kaldır", html)
        self.client.post(self.gonder, {"geri_al": "1"})
        self.assertIsNone(IseGirisKaydi.objects.get(personel=self.p).muhasebe_gonderim_tarihi)

    def test_gelecek_ve_gecersiz_tarih_mesaj(self):
        self.client.force_login(self.belgeci)
        r = self.client.post(self.gonder, {"tarih": (tr_bugun() + timedelta(days=2)).isoformat()}, follow=True)
        self.assertContains(r, "Gönderim tarihi gelecekte olamaz")
        r = self.client.post(self.gonder, {"tarih": "yanlis"}, follow=True)
        self.assertContains(r, "Geçersiz tarih")
        self.assertEqual(IseGirisKaydi.objects.count(), 0)

    def test_sekme_bolumu_eksik_ve_kalan_gun(self):
        self.client.force_login(self.belgeci)
        r = self.client.get(self.sekme)
        self.assertContains(r, "Muhasebe Paketi")
        self.assertContains(r, self.xlsx)
        self.assertContains(r, self.pdf)
        self.assertContains(r, "İşe girişe 5 gün var")
        self.assertContains(r, "Hazır")
        self.assertNotContains(r, "Muhasebe için eksik bilgiler")
        self.assertContains(r, "Ücret görme yetkiniz olmadığı için pakete ücret girmez")
        bos = personel_kur(ad="Veli", soyad="Kaya")
        r = self.client.get(reverse("core:personel_detay", args=[bos.pk]) + "?sekme=ise_giris")
        self.assertContains(r, "Muhasebe için eksik bilgiler")
        self.assertContains(r, "Eksik bilgi var")
        self.assertContains(r, "TC kimlik no")
        self.assertContains(r, "gün önce işe başladı")                                          # 2020'de işe girmiş: bildirge gecikmiş uyarısı
