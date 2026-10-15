"""KDV Listeleri Excel çıktısı (yalnız dizilim/biçim — hesaplar ``kdv_listeleri``'nde). 1. sayfa GİB e-beyanname yükleme şablonunun KOLON SIRASIYLA (başlık satırı + veri; başka satır yok);
sonraki sayfalar okunur: iade/ihraç, özet, mutabakat, uyarılar, faturasız hareketler."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.services.kdv_listeleri import AYLAR

BASLIK_DOLGU = PatternFill("solid", fgColor="15294D")
PARA = "#,##0.00"
INDIRILECEK_KOLONLAR = ("Sıra No", "Belge Tarihi", "Belge Seri No", "Belge Sıra No", "Satıcının Adı Soyadı / Unvanı", "Satıcının VKN / TCKN", "Mal / Hizmet Cinsi", "Miktar",
                        "KDV Hariç Tutar", "KDV Tutarı", "KDV'nin İndirildiği Dönem")
SATIS_KOLONLAR = ("Sıra No", "Belge Tarihi", "Belge Seri No", "Belge Sıra No", "Alıcının Adı Soyadı / Unvanı", "Alıcının VKN / TCKN", "KDV Hariç Tutar", "KDV Tutarı")


def _num(x):
    return None if x is None else float(x)


def _baslik(ws, kolonlar, satir=1, genislik=None):
    for i, k in enumerate(kolonlar, start=1):
        c = ws.cell(row=satir, column=i, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BASLIK_DOLGU
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate(genislik or [14] * len(kolonlar), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{satir + 1}"


def _vkn(r):
    return r["vkn"] or ""


def _indirilecek_sayfa(ws, satirlar):
    _baslik(ws, INDIRILECEK_KOLONLAR, genislik=[8, 12, 10, 18, 44, 16, 34, 12, 16, 14, 14])
    for r in satirlar:
        ws.append([r["sira"], r["belge_tarihi"], r["seri"], r["sira_no"], r["satici"], _vkn(r), r["cins"], r["miktar"], _num(r["matrah_tl"]), _num(r["indirilecek_tl"]), r["donem"]])
        n = ws.max_row
        ws.cell(row=n, column=2).number_format = "DD.MM.YYYY"
        for col in (9, 10):
            ws.cell(row=n, column=col).number_format = PARA
        ws.cell(row=n, column=6).number_format = "@"


def _satis_sayfa(ws, satirlar, tur_kolonu=False):
    kolonlar = SATIS_KOLONLAR + (("Tür",) if tur_kolonu else ())
    _baslik(ws, kolonlar, genislik=[8, 12, 10, 18, 44, 16, 16, 14] + ([14] if tur_kolonu else []))
    for r in satirlar:
        ws.append([r["sira"], r["belge_tarihi"], r["seri"], r["sira_no"], r["alici"], _vkn(r), _num(r["matrah_tl"]), _num(r["hesaplanan_tl"])] + ([{"iade": "İade", "ihrac": "İhraç kayıtlı"}.get(r["tur"], "")] if tur_kolonu else []))
        n = ws.max_row
        ws.cell(row=n, column=2).number_format = "DD.MM.YYYY"
        for col in (7, 8):
            ws.cell(row=n, column=col).number_format = PARA
        ws.cell(row=n, column=6).number_format = "@"


def _ozet_sayfa(ws, veri, baslik, hesap_adi, ek_satirlar):
    ws.column_dimensions["A"].width = 54
    for col in "BCDEF":
        ws.column_dimensions[col].width = 20
    ws["A1"] = f"{baslik} — {AYLAR[veri['ay'] - 1]} {veri['yil']}"
    ws["A1"].font = Font(bold=True, size=13, color="15294D")
    r = 3
    for ad, deger in ek_satirlar:
        ws.cell(row=r, column=1, value=ad)
        c = ws.cell(row=r, column=2, value=_num(deger) if hasattr(deger, "quantize") else deger)
        if hasattr(deger, "quantize"):
            c.number_format = PARA
        r += 1
    m = veri["mutabakat"]
    r += 1
    ws.cell(row=r, column=1, value=f"Mutabakat — {hesap_adi}").font = Font(bold=True)
    r += 1
    for ad, deger in ((f"Liste KDV toplamı (faturalardan)", m["liste"]), ("Faturasız hareketler (aynı taraf)", m["faturasiz"]), ("Toplam", m["toplam"]),
                      (f"{hesap_adi} dönem {m['taraf']} hareketi", m["hareket"]), ("FARK", m["fark"]), (f"{hesap_adi} dönem {m['karsi_ad']} hareketi (bilgi: mahsup vb.)", m["karsi_taraf"])):
        ws.cell(row=r, column=1, value=ad)
        c = ws.cell(row=r, column=2, value=_num(deger))
        c.number_format = PARA
        r += 1
    if m["kalemler"]:
        r += 1
        ws.cell(row=r, column=1, value="Fatura bazında fark (hesaplanan ↔ fişteki)").font = Font(bold=True)
        for k in m["kalemler"]:
            r += 1
            ws.cell(row=r, column=1, value=k["ad"])
            for j, key in enumerate(("hesaplanan", "fis", "fark"), start=2):
                ws.cell(row=r, column=j, value=_num(k[key])).number_format = PARA
    r += 2
    u = veri["uyarilar"]
    ws.cell(row=r, column=1, value="Uyarılar").font = Font(bold=True)
    for ad, rows in (("VKN/TCKN eksik", u["eksik_vkn"]), ("Fatura no eksik", u["eksik_no"]), ("Gerçek tarih kullanıldı", u["gercek_tarih"]), ("01.01 girişli, gerçek tarih açıklamada yok", u["tarih_isaretli"])):
        r += 1
        ws.cell(row=r, column=1, value=f"{ad}: {len(rows)}")
        ws.cell(row=r, column=2, value=", ".join(f"{x['belge_no'] or '#' + str(x['fatura'].pk)}" for x in rows)[:200])
    for grup in u["mukerrer"]:
        r += 1
        ws.cell(row=r, column=1, value=f"Mükerrer belge no: {grup[0]['belge_no']} ({len(grup)} kayıt)")
    r += 2
    ws.cell(row=r, column=1, value=f"Faturaya bağlı olmayan {hesap_adi} hareketleri").font = Font(bold=True)
    r += 1
    for j, k in enumerate(("Tarih / fiş", "Borç", "Alacak", "Kaynak", "Açıklama"), start=1):
        c = ws.cell(row=r, column=j, value=k)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BASLIK_DOLGU
    for x in veri["faturasiz"]:
        r += 1
        ws.cell(row=r, column=1, value=f"{x['tarih']:%d.%m.%Y} · {x['fis'].yil}/{x['fis'].fis_no}")
        ws.cell(row=r, column=2, value=_num(x["borc"])).number_format = PARA
        ws.cell(row=r, column=3, value=_num(x["alacak"])).number_format = PARA
        ws.cell(row=r, column=4, value=x["kaynak"])
        ws.cell(row=r, column=5, value=x["aciklama"][:150])


def indirilecek_xlsx(veri) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "İndirilecek KDV Listesi"
    _indirilecek_sayfa(ws, veri["satirlar"])
    ozet = wb.create_sheet("Özet ve Mutabakat")
    t = veri["toplam"]
    _ozet_sayfa(ozet, veri, "İndirilecek KDV Listesi", "191", [("Fatura adedi", t["adet"]), ("KDV hariç toplam (TL)", t["matrah"]), ("İndirilecek KDV toplamı (TL)", t["kdv"]),
                                                              ("  ‣ bunun tevkifat kısmı (sorumlu sıfatıyla beyan, 360)", t["tevkifat"]), ("Faturasız 191 borç hareketi", t["faturasiz_borc"]), ("Faturasız 191 alacak hareketi", t["faturasiz_alacak"])])
    cikti = BytesIO()
    wb.save(cikti)
    return cikti.getvalue()


def satis_xlsx(veri) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Satışlara Ait Fatura Listesi"
    _satis_sayfa(ws, veri["normal"])
    ws2 = wb.create_sheet("İade ve İhraç Kayıtlı")
    _satis_sayfa(ws2, veri["iade"] + veri["ihrac"], tur_kolonu=True)
    ozet = wb.create_sheet("Özet ve Mutabakat")
    t = veri["toplam"]
    _ozet_sayfa(ozet, veri, "Satışlara Ait Fatura Listesi", "391", [("Normal satış: matrah / KDV", f"{t['normal_matrah']:,.2f} / {t['normal_kdv']:,.2f}"), ("İade: matrah / KDV", f"{t['iade_matrah']:,.2f} / {t['iade_kdv']:,.2f}"),
                                                                   ("İhraç kayıtlı: matrah / KDV", f"{t['ihrac_matrah']:,.2f} / {t['ihrac_kdv']:,.2f}"), ("Toplam KDV (391'e yazılan)", t["kdv"]),
                                                                   ("Faturasız 391 borç hareketi", t["faturasiz_borc"]), ("Faturasız 391 alacak hareketi", t["faturasiz_alacak"])])
    cikti = BytesIO()
    wb.save(cikti)
    return cikti.getvalue()
