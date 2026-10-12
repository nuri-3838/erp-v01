"""ATÖLYE iş oturumları (Başla / Bitir — core.services.is_oturumu): ekip seçimi (mesaisi açık + atölye yetkili; meşgul olanlar seçilemez), bir
personel aynı anda tek açık oturumda, ekipteki herkes bitirebilir (kayıt personeli = başlatan, bitiren izi), üretim hatasında oturum açık kalır,
İptal (iş yapılmadı; süre kalır), mesai bitirme (üye ayrılır / tek kişiyse engel), dünden kalan açık iş (bitiş saati sorulur, varsayılan mesai
çıkışı, kayıt tarihi = bitiş tarihi), pano uyarısı (18:00 sonrası / dünden kalan), yönetici izlemede sayaç+ekip, ofis listesi + Excel."""
import datetime
from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.urls import reverse
from django.utils import timezone

from core.models import IsOturumu, IsOturumuUyesi, MesaiKaydi, OperasyonKaydi
from core.services import atolye, is_oturumu
from core.services.mesai import acik_kayit
from core.services.mesai_hesap import hesap_olustur
from core.tarih import _TR, tr_bugun
from core.tests.ik_yardimci import personel_kur
from core.tests.test_atolye import IZINLI, AtolyeTaban

D = Decimal


class OturumTaban(AtolyeTaban):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        atolye.operator_ayarla(cls.p2, True)                                                          # ayşe de operatör
        cls.p3 = personel_kur(ad="mehmet", soyad="usta", ise_giris_tarihi=date(2020, 1, 1))
        cls.u3 = hesap_olustur(cls.p3, kullanici_adi="mehmetu", sifre="cokGizli!2026", ilk_giris_sifre=False)
        atolye.operator_ayarla(cls.p3, True)

    def mesai_bitir(self, user):
        self.client.force_login(user)
        return self.client.post(reverse("core:atolye_mesai"), {"eylem": "bitir"}, HTTP_X_REAL_IP=IZINLI, follow=True)


class BaslaBitirTest(OturumTaban):
    def test_ekip_adaylari_ve_tek_oturum_kurali(self):
        self.baslat(self.u2)                                                                          # ayşe mesaide
        self.baslat()                                                                                 # errahman mesaide (mehmet mesaide değil)
        self.assertEqual([(a["personel"].pk, a["mesgul"]) for a in atolye.ekip_adaylari(self.u)], [(self.p2.pk, None)])
        r = self.client.get(reverse("core:atolye_istasyon", args=[self.lazer.pk]))
        self.assertContains(r, 'name="ekip" value="%d"' % self.p2.pk)
        self.assertNotContains(r, 'value="%d"' % self.p3.pk)
        r = self.basla(self.ie_kesim, ekip=[self.p2.pk])
        self.assertTrue(any("iş başladı" in m and "AYSE MESAİ" in m for m in self.mesajlar(r)), self.mesajlar(r))
        o = self.acik_oturum(self.ie_kesim)
        self.assertEqual([(u.personel_id, u.baslatan, u.mesai_kaydi_id) for u in o.uyeler.order_by("-baslatan")],
                         [(self.p.pk, True, acik_kayit(self.p).pk), (self.p2.pk, False, acik_kayit(self.p2).pk)])
        self.assertEqual((o.istasyon_id, o.uretim_emri_id, o.is_tarihi, o.baslatan_id, o.durum, o.baslangic_ip), (self.lazer.pk, self.emir.pk, tr_bugun(), self.p.pk, "ACIK", IZINLI))
        self.assertContains(r, "Çalışıyor")
        self.assertContains(r, "ERRAHMAN OPERATOR, AYSE MESAİ")
        self.assertContains(r, "BİTİR")
        self.assertNotContains(r, "▶ BAŞLA")                                                          # açık işi varken yeni iş başlatamaz
        self.assertContains(r, "Açık işiniz var")
        r = self.basla(self.ie_bukum)
        self.assertTrue(any("şu an başka bir işte: LAZER" in m for m in self.mesajlar(r)), self.mesajlar(r))
        # ayşe (üye) de aynı oturumu görür ve bitirebilir; başka iş başlatamaz; mehmet'e ayşe/errahman meşgul görünür
        self.client.force_login(self.u2)
        r = self.client.get(reverse("core:atolye_istasyon", args=[self.lazer.pk]))
        self.assertContains(r, "BİTİR")
        r = self.basla(self.ie_bukum)
        self.assertTrue(any("başka bir işte" in m for m in self.mesajlar(r)))
        self.baslat(self.u3)
        adaylar = {a["personel"].pk: a["mesgul"] for a in atolye.ekip_adaylari(self.u3)}
        self.assertEqual((adaylar[self.p.pk].pk, adaylar[self.p2.pk].pk), (o.pk, o.pk))
        r = self.basla(self.ie_bukum, ekip=[self.p2.pk])
        self.assertTrue(any("AYSE MESAİ şu an başka bir işte" in m for m in self.mesajlar(r)), self.mesajlar(r))
        self.assertEqual(IsOturumu.objects.count(), 1)
        r = self.client.get(reverse("core:atolye_istasyon", args=[self.bukum.pk]))
        self.assertContains(r, "(LAZER")                                                               # aday listesinde meşgul notu
        self.assertContains(r, "disabled")

    def test_herkes_bitirebilir_kayit_baslatana_bitiren_izi(self):
        self.baslat(self.u2)
        self.baslat()
        self.basla(self.ie_kesim, ekip=[self.p2.pk])
        o = self.acik_oturum(self.ie_kesim)
        IsOturumu.objects.filter(pk=o.pk).update(baslangic_zamani=timezone.now() - datetime.timedelta(minutes=30))
        IsOturumuUyesi.objects.filter(oturum=o).update(katilim_zamani=timezone.now() - datetime.timedelta(minutes=30))
        self.client.force_login(self.u2)                                                              # ÜYE bitiriyor
        r = self.bitir(o, "4")
        self.assertTrue(any("4 adet yazıldı" in m for m in self.mesajlar(r)), self.mesajlar(r))
        o.refresh_from_db()
        k = o.kayit
        self.assertEqual((o.durum, o.bitiren_id, o.miktar, k.personel_id, k.mesai_kaydi_id, k.created_by_id), ("BITTI", self.p2.pk, D("4"), self.p.pk, acik_kayit(self.p).pk, self.u2.pk))
        self.assertTrue(all(u.ayrilma_zamani == o.bitis_zamani for u in o.uyeler.all()))
        oz = is_oturumu.ozet(IsOturumu.objects.prefetch_related("uyeler__personel").get(pk=o.pk))
        self.assertEqual((oz["kisi"], oz["adet"]), (2, D("4")))
        self.assertTrue(1790 <= oz["sure_saniye"] <= 1900, oz["sure_saniye"])
        self.assertTrue(2 * 1790 <= oz["adam_saniye"] <= 2 * 1900)
        self.assertEqual(oz["adet_basi_saniye"], oz["adam_saniye"] // 4)
        self.assertEqual(oz["sure_metni"], "30 dk")
        self.assertTrue(D("7.5") <= oz["adet_saat"] <= D("8.1"), oz["adet_saat"])
        self.assertIsNone(is_oturumu.aktif_uyelik(self.p))
        r = self.bitir(o, "4")                                                                        # ikinci kez bitirilemez
        self.assertTrue(any("zaten kapatılmış" in m for m in self.mesajlar(r)))

    def test_uretim_hatasinda_oturum_acik_kalir_sonra_duzelir(self):
        self.baslat()
        self.basla(self.ie_bukum)                                                                      # bükümün girdisi (kesilmiş) henüz yok
        o = self.acik_oturum(self.ie_bukum)
        r = self.bitir(o, "3")
        self.assertTrue(any("açık kaldı" in m for m in self.mesajlar(r)), self.mesajlar(r))
        o.refresh_from_db()
        self.assertEqual((o.durum, o.kayit_id), ("ACIK", None))
        self.assertFalse(OperasyonKaydi.objects.filter(istasyon_emri=self.ie_bukum).exists())
        self.assertIsNotNone(is_oturumu.aktif_uyelik(self.p))
        r = self.bitir(o, "0")
        self.assertTrue(any("0 adetle bitirilemez" in m for m in self.mesajlar(r)), self.mesajlar(r))   # "0" → servis kuralı (boş → "Adet girin")
        with self.assertRaisesMessage(is_oturumu.OturumHatasi, "0 adetle bitirilemez"):
            is_oturumu.bitir(self.p, o, D("0"))
        # kesilmiş üretilince büküm bitirilir — aynı oturum (süre korunur)
        from core.services import uretim
        uretim.istasyon_emri_uret(self.ie_kesim, miktar=D("10"), tarih=tr_bugun(), kullanici=self.yon)
        r = self.bitir(o, "3")
        self.assertTrue(any("3 adet yazıldı" in m for m in self.mesajlar(r)), self.mesajlar(r))
        o.refresh_from_db()
        self.assertEqual((o.durum, o.miktar), ("BITTI", D("3")))

    def test_iptal_is_yapilmadi(self):
        self.baslat()
        self.basla(self.ie_kesim)
        o = self.acik_oturum(self.ie_kesim)
        r = self.client.post(reverse("core:atolye_oturum_iptal", args=[o.pk]), HTTP_X_REAL_IP=IZINLI, follow=True)
        self.assertTrue(any("iptal edildi" in m for m in self.mesajlar(r)), self.mesajlar(r))
        o.refresh_from_db()
        self.assertEqual((o.durum, o.kayit_id, o.miktar, o.bitiren_id), ("IPTAL", None, None, self.p.pk))
        self.assertIsNotNone(o.bitis_zamani)
        self.assertIsNotNone(o.uyeler.get().ayrilma_zamani)
        self.assertFalse(OperasyonKaydi.objects.filter(istasyon_emri=self.ie_kesim).exists())
        self.assertIsNone(is_oturumu.aktif_uyelik(self.p))
        r = self.client.get(reverse("core:atolye_istasyon", args=[self.lazer.pk]))
        self.assertContains(r, "▶ BAŞLA")                                                              # yeniden başlatabilir

    def test_mesai_bitirme_uye_ayrilir_tek_kisi_engel(self):
        self.baslat(self.u2)
        self.baslat()
        self.basla(self.ie_kesim, ekip=[self.p2.pk])
        o = self.acik_oturum(self.ie_kesim)
        r = self.mesai_bitir(self.u2)                                                                 # üye: ayrılır, oturum sürer
        self.assertTrue(any("Mesai bitirildi" in m for m in self.mesajlar(r)), self.mesajlar(r))
        self.assertIsNone(acik_kayit(self.p2))
        u2 = IsOturumuUyesi.objects.get(oturum=o, personel=self.p2)
        self.assertIsNotNone(u2.ayrilma_zamani)
        o.refresh_from_db()
        self.assertEqual(o.durum, "ACIK")
        r = self.mesai_bitir(self.u)                                                                  # tek kişi kaldı: engel
        self.assertTrue(any("Önce açık işinizi bitirin: LAZER" in m for m in self.mesajlar(r)), self.mesajlar(r))
        self.assertIsNotNone(acik_kayit(self.p))
        self.client.force_login(self.u)
        self.bitir(o, "4")
        r = self.mesai_bitir(self.u)
        self.assertTrue(any("Mesai bitirildi" in m for m in self.mesajlar(r)))
        oz = is_oturumu.ozet(IsOturumu.objects.prefetch_related("uyeler__personel").get(pk=o.pk))
        self.assertEqual(oz["kisi"], 2)
        self.assertTrue(oz["adam_saniye"] <= 2 * oz["sure_saniye"])                                   # ayrılan üyenin süresi kısa

    def test_dunden_kalan_acik_is(self):
        simdi = timezone.now()
        dun = tr_bugun() - datetime.timedelta(days=1)
        m_dun = MesaiKaydi.objects.create(personel=self.p, is_tarihi=dun, giris_zamani=simdi - datetime.timedelta(hours=22),
                                          cikis_zamani=simdi - datetime.timedelta(hours=15))
        o = IsOturumu.objects.create(istasyon_emri=self.ie_kesim, uretim_emri=self.emir, istasyon=self.lazer, baslatan=self.p, is_tarihi=dun,
                                     baslangic_zamani=simdi - datetime.timedelta(hours=20))
        IsOturumuUyesi.objects.create(oturum=o, personel=self.p, mesai_kaydi=m_dun, baslatan=True, katilim_zamani=o.baslangic_zamani)
        # pano: dünden kalan açık iş saatten bağımsız uyarı
        self.client.force_login(self.yon)
        r = self.client.get(reverse("core:pano"))
        self.assertContains(r, "açık atölye iş oturumu var")
        self.assertContains(r, "ERRAHMAN OPERATOR")
        # operatör ekranı: en üstte DÜNKÜ AÇIK İŞ, bitiş saati sorulur (varsayılan = dünkü mesai çıkışı)
        self.baslat()
        r = self.client.get(reverse("core:atolye"))
        self.assertContains(r, "DÜNKÜ AÇIK İŞ")
        self.assertContains(r, "Bu işe git")
        r = self.client.get(reverse("core:atolye_istasyon", args=[self.lazer.pk]))
        self.assertContains(r, 'name="bitis"')
        varsayilan = m_dun.cikis_zamani.astimezone(_TR).strftime("%Y-%m-%dT%H:%M")
        self.assertContains(r, f'value="{varsayilan}"')
        r = self.bitir(o, "4")
        self.assertTrue(any("bitiş saatini girin" in m for m in self.mesajlar(r)), self.mesajlar(r))
        r = self.bitir(o, "4", bitis=(simdi - datetime.timedelta(hours=21)).astimezone(_TR).strftime("%Y-%m-%dT%H:%M"))
        self.assertTrue(any("başlangıçtan sonra" in m for m in self.mesajlar(r)), self.mesajlar(r))
        r = self.bitir(o, "4", bitis=(simdi + datetime.timedelta(hours=3)).astimezone(_TR).strftime("%Y-%m-%dT%H:%M"))
        self.assertTrue(any("ileri bir zaman" in m for m in self.mesajlar(r)), self.mesajlar(r))
        r = self.bitir(o, "4", bitis=varsayilan)
        self.assertTrue(any("4 adet yazıldı" in m for m in self.mesajlar(r)), self.mesajlar(r))
        o.refresh_from_db()
        beklenen_bitis = datetime.datetime.strptime(varsayilan, "%Y-%m-%dT%H:%M").replace(tzinfo=_TR)
        self.assertEqual((o.durum, o.bitis_zamani, o.kayit.tarih, o.kayit.mesai_kaydi_id), ("BITTI", beklenen_bitis, beklenen_bitis.date(), m_dun.pk))
        self.assertTrue(5 * 3600 - 60 < o.sure_saniye <= 5 * 3600, o.sure_saniye)                    # datetime-local dakika hassasiyeti
        self.assertNotContains(self.client.get(reverse("core:atolye")), "DÜNKÜ AÇIK İŞ")

    def test_pano_uyarisi_bugunku_18_sonrasi(self):
        self.baslat()
        self.basla(self.ie_kesim)
        gun = tr_bugun()
        erken = datetime.datetime.combine(gun, datetime.time(17, 59), tzinfo=_TR)
        gec = datetime.datetime.combine(gun, datetime.time(18, 0), tzinfo=_TR)
        self.assertEqual(len(is_oturumu.unutulanlar(erken)), 0)
        self.assertEqual(len(is_oturumu.unutulanlar(gec)), 1)
        self.client.force_login(self.yon)
        with patch("core.services.is_oturumu.timezone.now", return_value=erken):
            self.assertNotContains(self.client.get(reverse("core:pano")), "açık atölye iş oturumu")
        with patch("core.services.is_oturumu.timezone.now", return_value=gec):
            r = self.client.get(reverse("core:pano"))
        self.assertContains(r, "<strong>1</strong> açık atölye iş oturumu")
        self.assertContains(r, "?durum=ACIK")

    def test_yonetici_izlemede_sayac_ve_ekip_gorunur(self):
        self.baslat(self.u2)
        self.baslat()
        self.basla(self.ie_kesim, ekip=[self.p2.pk])
        self.client.force_login(self.yon)
        r = self.client.get(reverse("core:atolye_istasyon", args=[self.lazer.pk]))
        self.assertContains(r, "Çalışıyor")
        self.assertContains(r, "ERRAHMAN OPERATOR, AYSE MESAİ")
        self.assertContains(r, 'class="at-sayac" data-baslangic=')
        self.assertNotContains(r, "BİTİR")
        self.assertNotContains(r, "▶ BAŞLA")


class OfisListesiTest(OturumTaban):
    def test_liste_filtre_ve_excel(self):
        self.baslat(self.u2)
        self.baslat()
        self.uret(self.ie_kesim, "4")                                                                 # ekipsiz, bitti
        self.basla(self.ie_kesim, ekip=[self.p2.pk])
        o2 = self.acik_oturum(self.ie_kesim)
        self.client.post(reverse("core:atolye_oturum_iptal", args=[o2.pk]), HTTP_X_REAL_IP=IZINLI)
        self.basla(self.ie_bukum)                                                                      # açık kalan
        self.client.force_login(self.yon)
        r = self.client.get(reverse("core:is_oturumlari"))
        self.assertContains(r, "3 iş oturumu")
        self.assertContains(r, self.emir.no)
        self.assertContains(r, "ERRAHMAN OPERATOR, AYSE MESAİ")
        self.assertContains(r, "İptal (iş yapılmadı)")
        self.assertContains(r, "açık")
        self.assertContains(r, OperasyonKaydi.objects.get(istasyon_emri=self.ie_kesim).no)
        r = self.client.get(reverse("core:is_oturumlari") + "?durum=IPTAL")
        self.assertContains(r, "1 iş oturumu")
        r = self.client.get(reverse("core:is_oturumlari") + f"?personel={self.p2.pk}")
        self.assertContains(r, "1 iş oturumu")
        r = self.client.get(reverse("core:is_oturumlari") + f"?istasyon={self.bukum.pk}")
        self.assertContains(r, "1 iş oturumu")
        r = self.client.get(reverse("core:is_oturumlari") + "?ara=IK-KESILMIS")
        self.assertContains(r, "2 iş oturumu")
        r = self.client.get(reverse("core:is_oturumlari") + "?xlsx=1")
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        from openpyxl import load_workbook
        ws = load_workbook(BytesIO(r.content)).active
        satirlar = list(ws.iter_rows(min_row=4, values_only=True))
        self.assertEqual(satirlar[0][:3], ("Tarih", "İstasyon", "ÜS"))
        self.assertEqual(len(satirlar), 4)
        self.assertEqual(sorted(s[12] for s in satirlar[1:]), ["Açık", "Bitti", "İptal (iş yapılmadı)"])
        bitti = next(s for s in satirlar[1:] if s[12] == "Bitti")
        self.assertEqual((bitti[2], bitti[7], bitti[13]), (self.emir.no, 1, D("4")))
        # operatör ofis listesini göremez
        self.client.force_login(self.u)
        self.assertEqual(self.client.get(reverse("core:is_oturumlari")).status_code, 403)
