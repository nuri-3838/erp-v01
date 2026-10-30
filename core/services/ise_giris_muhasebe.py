"""İNSAN KAYNAKLARI > İşe giriş > MUHASEBE PAKETİ (Dilim 3).

Resmî işe girişi (SGK bildirgesi) muhasebe/mali müşavir yapar; ERP resmî kanala girmez. Bu servis muhasebeye iletilecek bilgileri personel
kartından toplar: "hazır mı?" kontrolü (eksik zorunlu alanlar + işe girişe kalan gün uyarısı), Excel/PDF bilgi paketi ve "muhasebeye
gönderildi" işareti. Hiçbir değer saklanmaz; her indirmede canlı üretilir (yalnız gönderim tarihi ``IseGirisKaydi``'ne yazılır).
Ücret yalnız ``ucret_gorebilir`` (personel_ucret yetkisi) olana dahil edilir.
"""
from __future__ import annotations

from io import BytesIO

from django.db import transaction
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.page import PageMargins

from core.models import IkAyari, Personel
from core.services import ise_giris as ig
from core.tarih import tr_bugun

# (alan, görünen ad) — bunlar yoksa SGK işe giriş bildirgesi hazırlanamaz; "hazır" bunlara bağlıdır.
ZORUNLU_ALANLAR = (
    ("tc_kimlik_no", "TC kimlik no"), ("dogum_tarihi", "Doğum tarihi"), ("dogum_yeri", "Doğum yeri"), ("baba_adi", "Baba adı"),
    ("ana_adi", "Ana adı"), ("adres", "Adres"), ("gorev", "Görev"), ("ise_giris_tarihi", "İşe giriş tarihi"),
)
# Önerilen (eksikse paket hazır SAYILIR, yalnız uyarı çıkar): maaş ödemesi / iletişim.
ONERILEN_ALANLAR = (("iban", "IBAN (maaş ödemesi için)"), ("telefon", "Telefon"))

_TR_ASCII = str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")
_BASLIK_DOLGU = PatternFill("solid", fgColor="21385F")
_BOLUM_DOLGU = PatternFill("solid", fgColor="DDE3EE")
_EKSIK_DOLGU = PatternFill("solid", fgColor="FFCDD2")


def _bos(deger) -> bool:
    return deger is None or str(deger).strip() == ""


def hazirlik(personel: Personel, kayit, *, bugun=None) -> dict:
    """Muhasebe için hazır mı? {"hazir", "eksikler": [ad], "onerilen_eksikler": [ad], "kalan_gun", "durum": "gecmis|acil|normal|belirsiz", "uyari",
    "gonderildi": date|None}. Eksik bilgi paketi indirmeyi ENGELLEMEZ (eksikler Excel'de işaretlenir)."""
    bugun = bugun or tr_bugun()
    eksikler = [ad for alan, ad in ZORUNLU_ALANLAR if _bos(getattr(personel, alan))]
    onerilen = [ad for alan, ad in ONERILEN_ALANLAR if _bos(getattr(personel, alan))]
    kalan = (personel.ise_giris_tarihi - bugun).days if personel.ise_giris_tarihi else None
    gonderildi = kayit.muhasebe_gonderim_tarihi
    if kalan is None:
        durum, uyari = "belirsiz", ""
    elif kalan < 0:
        durum, uyari = "gecmis", f"Personel {-kalan} gün önce işe başladı. İşe giriş bildirgesi işe başlamadan önce verilmeli; muhasebeye hemen iletin."
    elif kalan <= 1:
        durum, uyari = "acil", "İşe giriş bildirgesi işe başlamadan önce verilmeli; muhasebeye bugün iletin." if kalan == 0 else \
            "İşe başlamaya 1 gün kaldı. İşe giriş bildirgesi işe başlamadan önce verilmeli; muhasebeye bugün iletin."
    else:
        durum, uyari = "normal", f"İşe girişe {kalan} gün var. İşe giriş bildirgesi işe başlamadan önce verilmeli."
    if gonderildi:                                                      # gönderilmişse süre baskısı gösterilmez
        uyari = ""
    return {"hazir": not eksikler, "eksikler": eksikler, "onerilen_eksikler": onerilen, "kalan_gun": kalan, "durum": durum, "uyari": uyari,
            "gonderildi": gonderildi}


def bolumler(personel: Personel, firma, kayit, *, ucret_gorebilir=False, bugun=None, ayar=None) -> list:
    """Paket içeriği: [(bölüm adı, [(etiket, değer, zorunlu_mu)])]. Boş değer "" kalır (Excel'de işaretlenir, PDF'te noktalı çizgi)."""
    from core.dogrulama import iban_bicimli
    bugun = bugun or tr_bugun()
    gun = lambda d: d.strftime("%d.%m.%Y") if d else ""
    secenek = lambda alan: getattr(personel, f"get_{alan}_display")() if getattr(personel, alan) else ""
    zorunlu = {alan for alan, _ in ZORUNLU_ALANLAR}
    e = ig.etkin_sozlesme(kayit, ayar or IkAyari.get())
    saat = e["haftalik_calisma_saati"]
    ucret_satirlari = []
    if ucret_gorebilir:
        from core.services import personel_ucret as ucret_servis
        tarih = max(bugun, personel.ise_giris_tarihi) if personel.ise_giris_tarihi else bugun
        u = ucret_servis.gecerli_ucret(personel, tarih)
        ucret_satirlari = [("Ücret", ig.ucret_metni(personel, bugun=bugun), False),
                           ("Ücret geçerlilik başlangıcı", gun(u.gecerlilik_baslangic) if u else "", False),
                           ("Ücret ödeme şekli/günü", e["ucret_odeme"], False)]
    z = lambda alan: alan in zorunlu
    out = [
        ("İşveren", [("Unvan", firma.unvan, False), ("Vergi dairesi", firma.vergi_dairesi, False), ("Vergi no", firma.vergi_no, False),
                     ("SGK işyeri sicil no", firma.sgk_isyeri_sicil_no, False)]),
        ("Kimlik Bilgileri", [
            ("Adı", personel.ad, False), ("Soyadı", personel.soyad, False), ("T.C. Kimlik No", personel.tc_kimlik_no, z("tc_kimlik_no")),
            ("Doğum tarihi", gun(personel.dogum_tarihi), z("dogum_tarihi")), ("Doğum yeri", personel.dogum_yeri, z("dogum_yeri")),
            ("Baba adı", personel.baba_adi, z("baba_adi")), ("Ana adı", personel.ana_adi, z("ana_adi")),
            ("Cinsiyet", secenek("cinsiyet"), False), ("Medeni durum", secenek("medeni_durum"), False),
            ("Çocuk sayısı", "" if personel.cocuk_sayisi is None else str(personel.cocuk_sayisi), False),
            ("Öğrenim durumu", secenek("ogrenim_durumu"), False)]),
        ("İletişim", [("Adres", " ".join((personel.adres or "").split()), z("adres")), ("Telefon", personel.telefon, False),
                      ("E-posta", personel.eposta, False)]),
        ("İşe Giriş", [
            ("İşe giriş tarihi", gun(personel.ise_giris_tarihi), z("ise_giris_tarihi")), ("Departman", personel.departman, False),
            ("Görev", personel.gorev, z("gorev")), ("Sözleşme türü", "Belirsiz süreli", False),
            ("Deneme süresi", f"{e['deneme_suresi_ay']} ay" if e["deneme_suresi_ay"] else "", False),
            ("Haftalık çalışma saati", f"{saat:g}" if saat else "", False), ("Çalışma günleri", e["calisma_gunleri"], False)]),
    ]
    if ucret_satirlari:
        out.append(("Ücret", ucret_satirlari))
    out.append(("Banka", [("IBAN", iban_bicimli(personel.iban) if personel.iban else "", False), ("Banka", personel.banka_adi, False),
                          ("Hesap sahibi", personel.hesap_sahibi or personel.ad_soyad, False)]))
    return out


def dosya_adi(personel: Personel, uzanti: str, bugun=None) -> str:
    bugun = bugun or tr_bugun()
    temiz = "_".join("".join(c if c.isalnum() else " " for c in personel.ad_soyad.translate(_TR_ASCII)).split()).lower()
    return f"ise_giris_muhasebe_{temiz or personel.pk}_{bugun:%Y%m%d}.{uzanti}"


def excel(personel: Personel, firma, kayit, *, ucret_gorebilir=False, bugun=None, ayar=None) -> bytes:
    """Tek sayfa, iki sütun (Alan / Değer); zorunlu olup boş alanlar kırmızı "EKSİK". A4 dikey, yazdırılabilir."""
    bugun = bugun or tr_bugun()
    ws_kitap = Workbook()
    ws = ws_kitap.active
    ws.title = "İşe Giriş Bilgileri"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 60
    ws.cell(row=1, column=1, value=f"İŞE GİRİŞ BİLGİLERİ — {personel.ad_soyad}").font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value=f"Hazırlanma tarihi: {bugun:%d.%m.%Y}").font = Font(italic=True)
    satir = 4
    for baslik, alanlar in bolumler(personel, firma, kayit, ucret_gorebilir=ucret_gorebilir, bugun=bugun, ayar=ayar):
        for sutun in (1, 2):
            h = ws.cell(row=satir, column=sutun, value=baslik.upper() if sutun == 1 else None)
            h.font, h.fill = Font(bold=True, color="FFFFFF"), _BASLIK_DOLGU
        satir += 1
        for etiket, deger, zorunlu in alanlar:
            a = ws.cell(row=satir, column=1, value=etiket)
            a.font, a.fill = Font(bold=True), _BOLUM_DOLGU
            b = ws.cell(row=satir, column=2, value=deger if not _bos(deger) else ("EKSİK" if zorunlu else ""))
            b.number_format = "@"                                       # TC/IBAN/tarih metin kalsın (baştaki sıfır, bilimsel gösterim yok)
            b.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
            if zorunlu and _bos(deger):
                b.fill, b.font = _EKSIK_DOLGU, Font(bold=True, color="B71C1C")
            satir += 1
        satir += 1
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)
    buf = BytesIO()
    ws_kitap.save(buf)
    return buf.getvalue()


def pdf_belgesi(personel: Personel, firma, kayit, *, ucret_gorebilir=False, bugun=None, ayar=None) -> dict:
    """Evrak paketi PDF şablonu için tek belge (imzasız): ``ise_giris.paket_hazirla`` çıktısıyla aynı sözlük biçimi."""
    from core.services.belge_sablon import BOS
    ps = [(b, [(e, d if not _bos(d) else BOS) for e, d, _z in a])
          for b, a in bolumler(personel, firma, kayit, ucret_gorebilir=ucret_gorebilir, bugun=bugun, ayar=ayar)]
    return {"kod": "MUHASEBE_BILGI", "baslik": "İŞE GİRİŞ BİLGİ FORMU (MUHASEBE İÇİN)", "bolumler": ps,
            "paragraflar": ["Bu form personel kartından otomatik hazırlanmıştır; SGK işe giriş bildirgesinin hazırlanması için muhasebeye iletilir."],
            "imza": "yok", "imza_notu": ""}


@transaction.atomic
def gonderildi_isaretle(personel: Personel, *, tarih=None, geri_al=False, kullanici=None, bugun=None):
    """"Muhasebeye gönderildi" tarihini yazar (yalnız o alan); ``geri_al`` ise temizler. Gelecek tarih reddedilir."""
    bugun = bugun or tr_bugun()
    if personel.silindi:
        raise ig.IseGirisHatasi("Silinmiş personel için işlem yapılamaz.")
    if not geri_al:
        tarih = tarih or bugun
        if tarih > bugun:
            raise ig.IseGirisHatasi("Gönderim tarihi gelecekte olamaz.")
    kayit = ig.kayit_oku(personel)
    kayit.muhasebe_gonderim_tarihi = None if geri_al else tarih
    kayit.updated_by = kullanici
    if kayit.pk is None:
        kayit.created_by = kullanici
        kayit.save()
    else:
        kayit.save(update_fields=["muhasebe_gonderim_tarihi", "updated_by", "updated_at"])
    return kayit
